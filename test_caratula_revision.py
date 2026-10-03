# -*- coding: utf-8 -*-
"""LA CARÁTULA DEL AMPARO EN REVISIÓN, COMO LA ESCRIBE EL CORPUS (28-sep-2026,
AR 631/2025), SIN LLAMAR A NINGÚN MODELO.

Medido sobre la primera página de 500 amparos en revisión del Tercer Tribunal
Colegiado en Materias Administrativa y Civil del Vigésimo Segundo Circuito
(478 carátulas legibles):
  · ningún rubro lleva el órgano recurrido ni el juzgado (0 de 478);
  · si recurre la quejosa: «QUEJOSA Y RECURRENTE» (162 de 226);
  · si recurre otra parte: el recurrente con su carácter («TERCERA INTERESADA
    Y RECURRENTE», «AUTORIDAD RESPONSABLE Y RECURRENTE») y, en la forma que se
    adopta (AR 60/2020), la quejosa en su propio renglón encima.

Nombres sintéticos.

    .venv/bin/python test_caratula_revision.py
"""
import sys

from docx import Document

import documento_generado as dg
import fase_partes as fp
import tipos_asunto as ta

FALLOS = []


def ok(cond, que):
    print(f"   {'PASA ' if cond else 'FALLA'}  {que}")
    if not cond:
        FALLOS.append(que)


QUEJOSA = "Unión Ejemplo, A.C."
TERCERA = "Inmobiliaria Ejemplo, S.A. de C.V."

print("1 · EL CATÁLOGO: SIN ÓRGANO RECURRIDO EN EL RUBRO")
_cat = ta.caratula_de("amparo_revision")
ok(not any("RECURRIDO" in e or c == "responsable" for e, c, _o in _cat),
   f"la revisión no rotula al juzgado ni a la responsable: {_cat}")
# C3 y C4 (3-oct-2026): la queja y la revisión fiscal ya NO llevan el renglón
# del órgano ni el de la Sala (David: «hay que quitar autoridades recurrentes y
# sala»; 0 de 8 quejas y 28 de 28 revisiones fiscales del banco). Antes esta
# comprobación exigía que lo conservaran.
ok(not any(c == "responsable" for _e, c, _o in ta.caratula_de("queja"))
   and not any(c == "responsable" for _e, c, _o in ta.caratula_de("revision_fiscal"))
   and ta.caratula_de("amparo_directo")[-1][0] == "AUTORIDAD RESPONSABLE",
   "sólo el amparo directo conserva el renglón de la autoridad; la queja y la revisión fiscal, no (C3, C4)")
ok(ta.etiqueta_de_figura("amparo_revision", "responsable").startswith("ÓRGANO RECURRIDO")
   and ta.etiqueta_de_figura("amparo_directo", "responsable") == "AUTORIDAD RESPONSABLE",
   "la ficha de partes sigue sabiendo cómo se llama el juzgado, aunque no vaya en el rubro")
_ficha = fp.Partes(quejoso=QUEJOSA, tercero_interesado=TERCERA,
                   autoridad_responsable="Juzgado Séptimo de Distrito",
                   tipo_asunto="amparo_revision").bloque()
ok("ÓRGANO RECURRIDO (Juzgado de Distrito que dictó la sentencia recurrida): Juzgado Séptimo" in _ficha and "AUTORIDAD RESPONSABLE: Juzgado" not in _ficha,
   "la ficha no llama «autoridad responsable» al juzgado de la recurrida")

print("\n2 · QUIÉN RECURRE DECIDE LOS RENGLONES")
_f = ta.filas_caratula("amparo_revision", {"quejoso": QUEJOSA})
ok(_f == [("QUEJOSA Y RECURRENTE", "quejoso", QUEJOSA)],
   f"recurre la quejosa: una figura, concordada (A.C. es femenino): {_f}")
_f = ta.filas_caratula("amparo_revision", {
    "quejoso": QUEJOSA, "recurrente": TERCERA, "papel_recurrente": "tercero",
    "quejoso_moral": True, "organo_recurrido": "Juzgado Séptimo de Distrito",
    "responsable": "Magistrada de la Sala Civil"})
ok(_f == [("QUEJOSA", "quejoso", QUEJOSA),
          ("TERCERA INTERESADA Y RECURRENTE", "recurrente", TERCERA)],
   f"el 631: quejosa y tercera recurrente en renglones distintos, sin juzgado ni Magistrada: {_f}")
_f = ta.filas_caratula("amparo_revision", {
    "quejoso": "Juan Pérez", "recurrente": "Gobernador del Estado",
    "papel_recurrente": "autoridad", "quejoso_moral": True})
ok(_f == [("PARTE QUEJOSA", "quejoso", "Juan Pérez"),
          ("AUTORIDAD RESPONSABLE Y RECURRENTE", "recurrente", "Gobernador del Estado")],
   f"recurre la autoridad; «quejoso_moral» del formulario no vale para otra persona: {_f}")
_f = ta.filas_caratula("amparo_revision", {"quejoso": "Juan Pérez", "recurrente": "Pedro Gómez"})
ok(_f[-1] == ("RECURRENTE", "recurrente", "Pedro Gómez"),
   f"sin carácter conocido, «RECURRENTE» a secas: {_f}")
_f = ta.filas_caratula("amparo_revision", {"quejoso": QUEJOSA, "adherente": "Gobernador"})
ok(_f[-1] == ("RECURRENTE ADHESIVO", "adherente", "Gobernador"), f"el adhesivo cuando consta: {_f}")
# C3: el órgano de la queja ya no va en el rubro (antes: «el órgano leído sigue
# mandando en su renglón»); la ficha de partes lo sigue rotulando.
_f = ta.filas_caratula("queja", {"quejoso": "X", "responsable": "Sala", "organo_recurrido": "Juzgado Segundo"})
ok(not any(c == "responsable" or "Juzgado" in str(v) for _e, c, v in _f)
   and ta.etiqueta_de_figura("queja", "responsable") == "ÓRGANO QUE DICTÓ EL AUTO RECURRIDO",
   f"en la queja el órgano no va en el rubro y la ficha de partes lo sigue nombrando: {_f}")

print("\n3 · EL .docx Y LA HOJA DEL PROMPT DICEN LO MISMO")
_datos = {"encabezado": "AMPARO EN REVISIÓN CIVIL: 631/2025", "tipo_asunto": "amparo_revision",
          "quejoso": QUEJOSA, "recurrente": TERCERA, "papel_recurrente": "tercero",
          "organo_recurrido": "Juzgado Séptimo de Distrito en el Estado de Ejemplo",
          "responsable": "Magistrada de la Sala Civil", "magistrado": "M", "secretario": "S"}
_doc = Document()
dg._caratula(_doc, _datos, "amparo_revision")
_cab = " ".join(p.text for p in _doc.paragraphs).upper()
ok("QUEJOSA: UNIÓN EJEMPLO, A.C." in _cab and "TERCERA INTERESADA Y RECURRENTE: INMOBILIARIA EJEMPLO" in _cab
   and "QUEJOSA Y RECURRENTE" not in _cab and "RECURRIDO" not in _cab
   and "JUZGADO" not in _cab and "MAGISTRADA" not in _cab,
   f"la carátula del .docx: {_cab[:260]}")
_p = dg.prompt_estructura(_datos)
ok("TERCERA INTERESADA Y RECURRENTE (quien interpuso el recurso): Inmobiliaria Ejemplo" in _p
   and "ÓRGANO RECURRIDO:" not in _p
   and "sentencia recurrida: Juzgado Séptimo" in _p,
   "la hoja del prompt: el recurrente con su carácter y el juzgado como dato, no como rótulo")

print("\n4 · CON NOMBRES COMO LOS DE LA FILA DEL 631 (revisión adversarial de la fase E)")
# Misma forma que los nombres reales de la fila 462 —sustantivo de cabeza sin
# forma jurídica, artículo de la prosa en minúscula, «C.T.M.» al final—, con
# datos de prueba. Con el nombre sintético («…, A.C.») el defecto no se veía.
import tipos_asunto as _ta4
_Q4 = ("la Unión de Trabajadores de la Construcción, Transportistas y Similares del Estado de "
       "Ejemplo, C.T.M.")
_R4 = "Impulsora de Desarrollos Ejemplo VV"
_f4 = _ta4.filas_caratula("amparo_revision", {"quejoso": _Q4, "recurrente": _R4,
                                               "papel_recurrente": "tercero"})
ok([e for e, _c, _v in _f4 if _c in ("quejoso", "recurrente")] == ["QUEJOSA", "TERCERA INTERESADA Y RECURRENTE"],
   f"los rótulos concuerdan con «Unión…, C.T.M.» e «Impulsora…» sin forma jurídica: {[e for e, _c, _v in _f4]}")
ok(all(not str(_v).startswith(("la ", "el ")) for _e, _c, _v in _f4),
   "el renglón lleva el nombre, sin el artículo de la prosa")
ok(_ta4.genero_de("Sindicato Único de Trabajadores de Ejemplo, C.T.M.") == "o"
   and _ta4.sin_articulo_de_prosa("La Costeña, S.A. de C.V.") == "La Costeña, S.A. de C.V.",
   "el sindicato es masculino; el «La» de un nombre propio no se quita")

print("\n5 · LA QUEJA Y LA REVISIÓN FISCAL SIN EL ÓRGANO EN EL RUBRO (C3 y C4, 3-oct-2026)")
# David: «hay que quitar autoridades recurrentes y sala» (RF; banco: «REVISIÓN
# FISCAL: {toca} / RECURRENTE: {recurrente}» 28 de 28) y, en la queja, lo más
# práctico: 0 de 8 engroses llevan el órgano en el rubro; ya va en el V I S T O
# y en la competencia.
import contexto_taller as _ct5

ok(ta.caratula_de("queja") == [("{QUEJOSO_A} Y RECURRENTE", "quejoso", True)],
   f"queja: una sola fila en el catálogo, sin «ÓRGANO QUE DICTÓ EL AUTO RECURRIDO»: {ta.caratula_de('queja')}")
ok(ta.caratula_de("revision_fiscal") == [("RECURRENTE", "quejoso", True),
                                         ("RECURRENTE ADHESIVO", "adherente", False),
                                         ("PARTE ACTORA", "tercero", False)],
   f"revisión fiscal: exactamente las tres filas del contrato: {ta.caratula_de('revision_fiscal')}")
_f = ta.filas_caratula("queja", {"quejoso": "Juan Pérez López", "responsable": "Juzgado Cuarto de Distrito",
                                 "organo_recurrido": "Juzgado Cuarto de Distrito en el Estado de Ejemplo"})
ok(_f == [("PARTE QUEJOSA Y RECURRENTE", "quejoso", "Juan Pérez López")],
   f"queja que recurre el quejoso: un renglón y ningún juzgado, aunque se haya leído: {_f}")
_f = ta.filas_caratula("queja", {"quejoso": QUEJOSA, "recurrente": TERCERA, "papel_recurrente": "tercero",
                                 "organo_recurrido": "Juzgado Cuarto de Distrito"})
ok(_f == [("QUEJOSA", "quejoso", QUEJOSA), ("TERCERA INTERESADA Y RECURRENTE", "recurrente", TERCERA)],
   f"queja que recurre la tercera: sigue partiendo en dos renglones, sin el órgano: {_f}")
_f = ta.filas_caratula("revision_fiscal", {
    "quejoso": "la Titular de la Unidad Jurídica de la Delegación Estatal en Ejemplo del Instituto de Ejemplo",
    "adherente": "Ana Ruiz Gómez", "tercero": "Ana Ruiz Gómez",
    "responsable": "Sala Regional de Ejemplo del Tribunal Federal de Justicia Administrativa",
    "organo_recurrido": "Sala Regional de Ejemplo"})
_et = [e for e, _c, _v in _f]
# LA ACTORA QUE ES LA ADHESIVA VA UNA VEZ (revisión RF, 3-oct-2026; RF 4/2025:
# «RECURRENTE ADHESIVO: ***** (ACTORA)»). Antes se comprobaban los tres renglones
# con el mismo nombre en el del adhesivo y en el de la actora.
ok(_et == ["RECURRENTE", "RECURRENTE ADHESIVO"]
   and _f[1][2] == "Ana Ruiz Gómez (ACTORA)"
   and _f[0][2].startswith("Titular de la Unidad Jurídica")
   and not any("Sala" in str(v) for _e, _c, v in _f),
   f"revisión fiscal: «RECURRENTE» y la actora adhesiva en un renglón; ni «AUTORIDAD RECURRENTE» ni la Sala: {_f}")
_f = ta.filas_caratula("revision_fiscal", {
    "quejoso": "la Titular de la Unidad Jurídica de la Delegación Estatal en Ejemplo del Instituto de Ejemplo",
    "adherente": "Ana Ruiz Gómez", "tercero": "Pedro Gil Mora"})
ok([e for e, _c, _v in _f] == ["RECURRENTE", "RECURRENTE ADHESIVO", "PARTE ACTORA"] and _f[2][2] == "Pedro Gil Mora",
   "si el adhesivo no es la actora, los tres renglones como C4")
_f = ta.filas_caratula("revision_fiscal", {
    "quejoso": "Secretario de Hacienda y Crédito Público, Jefe del Servicio de Administración Tributaria y Titular "
               "de la Administración de Operación de Padrones “2”",
    "tercero": "Juan Pérez López, María Gómez Ruiz, Pedro Sánchez Díaz y Ana Torres Vega"})
ok(_f[0][0] == "RECURRENTES" and _f[1] == ("PARTE ACTORA", "tercero", "Juan Pérez López Y OTROS"),
   f"varias autoridades recurren: «RECURRENTES»; más de dos actores: el primero «Y OTROS» (RF 33/2024, 42/2023): {_f}")
_f = ta.filas_caratula("revision_fiscal", {"quejoso": "Titular de la Unidad Jurídica del Instituto de Ejemplo",
                                           "tercero": "Juan Pérez López y María Gómez Ruiz"})
ok(_f[0][0] == "RECURRENTE" and _f[1][2] == "Juan Pérez López y María Gómez Ruiz",
   "una autoridad, «RECURRENTE»; dos actores, los dos")
ok(ta.etiqueta_de_figura("revision_fiscal", "responsable").startswith("SALA QUE DICTÓ LA SENTENCIA RECURRIDA")
   and ta.etiqueta_de_figura("queja", "responsable") == "ÓRGANO QUE DICTÓ EL AUTO RECURRIDO",
   "la ficha de partes sigue sabiendo cómo se llaman el órgano de la queja y la Sala, aunque no vayan en el rubro")
_bl = fp.Partes(quejoso="Titular de la Unidad Jurídica", tercero_interesado="Ana Ruiz Gómez",
                autoridad_responsable="Sala Regional de Ejemplo", tipo_asunto="revision_fiscal").bloque()
ok("SALA QUE DICTÓ LA SENTENCIA RECURRIDA" in _bl and "Sala Regional de Ejemplo" in _bl
   and "(no existe en este tipo" not in _bl,
   "y la ficha de partes de la revisión fiscal la nombra (sin el rubro, `fase_partes` la daba por inexistente)")
try:
    _doc5 = Document()
    dg._caratula(_doc5, {"encabezado": "REVISIÓN FISCAL: 33/2024", "tipo_asunto": "revision_fiscal",
                         "quejoso": "Titular de la Unidad Jurídica", "tercero": "Ana Ruiz Gómez",
                         "responsable": "Sala Regional de Ejemplo", "magistrado": "M", "secretario": "S"},
                 "revision_fiscal")
    _cab5 = " ".join(p.text for p in _doc5.paragraphs).upper()
    ok("RECURRENTE: TITULAR DE LA UNIDAD JURÍDICA" in _cab5 and "AUTORIDAD RECURRENTE" not in _cab5
       and "SALA RESPONSABLE" not in _cab5 and "SALA REGIONAL" not in _cab5,
       f"el .docx de la revisión fiscal: {_cab5[:200]}")
except Exception as _e5:  # noqa: BLE001 — documento_generado puede estar a medio editar por otro
    ok(False, f"el .docx de la revisión fiscal reventó: {type(_e5).__name__}: {_e5}")
# SIN LA BANDERA, EL RUBRO DE ANTES: C3 y C4 no son [siempre].
_ct5.poner(True, {"banderas": {"procedencia_por_tipo": False}}, pruebas=True)
try:
    ok(ta.caratula_de("queja") == [("RECURRENTE", "quejoso", True),
                                   ("ÓRGANO QUE DICTÓ EL AUTO RECURRIDO", "responsable", True)]
       and [e for e, _c, _o in ta.caratula_de("revision_fiscal")]
       == ["AUTORIDAD RECURRENTE", "PARTE ACTORA", "SALA RESPONSABLE"],
       "sin la bandera (PROCEDENCIA_POR_TIPO=0) vuelven el órgano de la queja y la Sala de la revisión fiscal")
finally:
    _ct5.poner(True, {"banderas": {"procedencia_por_tipo": True}}, pruebas=True)

print()
if FALLOS:
    print(f"FALLAN {len(FALLOS)}: " + " · ".join(FALLOS))
    sys.exit(1)
print("RESULTADO: TODAS LAS COMPROBACIONES PASAN")
