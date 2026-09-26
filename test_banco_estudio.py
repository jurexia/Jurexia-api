"""El banco del estudio — métricas, arnés y comparador, sin red. 26-sep-2026.

    .venv/bin/python test_banco_estudio.py

Todo corre en local: el servidor se sustituye por un transporte simulado de
httpx que habla el mismo SSE que /taller/resolver/stream. Las comprobaciones
contra los engroses Kingston corren sólo si el corpus está en este Mac.
"""
import asyncio
import base64
import io
import json
import re
import tempfile
from pathlib import Path

import banco_estudio as be
import comparar_estudio as ce
import metricas_estudio as me

FALLOS = []


def ok(cond, que):
    print(f"   {'PASA ' if cond else 'FALLA'}  {que}")
    if not cond:
        FALLOS.append(que)


# ═══════════════════════════════════════════════════════════════════════════
# Textos de prueba: escritos a mano para ejercitar cada medida, NO copiados de
# ningún expediente.
# ═══════════════════════════════════════════════════════════════════════════
ESTANDAR = """R E S U L T A N D O
SEXTO. Celebración de la sesión vía remota. El presente asunto se listó el *********.
C O N S I D E R A N D O
SEXTO. Estudio. Los conceptos de violación son infundados.
Sentencia reclamada
La Sala confirmó la sentencia de primera instancia por las razones que se resumen.
Conceptos de violación
En el primer concepto de violación la parte quejosa alega que se valoró mal la confesional.
Solución
Por cuestión de método, se analizarán conjuntamente los conceptos de violación primero y segundo.
Sobre los conceptos de violación primero y segundo, en los que la parte quejosa sostiene que la Sala valoró mal la confesional. Se considera infundado.
Lo anterior, porque el artículo 16 de la Constitución Política de los Estados Unidos Mexicanos establece que todo acto de molestia debe estar fundado y motivado, lo que aquí se cumplió.
Sirve de apoyo la jurisprudencia de la Primera Sala, de registro digital 176546, de rubro «FUNDAMENTACIÓN Y MOTIVACIÓN DE LAS RESOLUCIONES JURISDICCIONALES, DEBEN ANALIZARSE A LA LUZ DE LOS ARTÍCULOS 14 Y 16».
No pasa inadvertido que la quejosa invoca la cesión verbal; sin embargo, no la probó en el juicio natural.
El crédito por doscientos setenta y nueve mil doscientos diez pesos se notificó el trece de agosto de dos mil veinticinco, conforme al artículo 17 de la Ley Federal de Procedimiento Contencioso Administrativo.
Sobre el tercer concepto de violación, en el que la parte quejosa sostiene que se le dejó en indefensión. Resulta inoperante.
El artículo 16 de la Constitución Política de los Estados Unidos Mexicanos dispone de nuevo que los actos deben motivarse, y eso rige también para este tramo del asunto.
Así lo ilustra la jurisprudencia de rubro «FUNDAMENTACIÓN Y MOTIVACIÓN DE LAS RESOLUCIONES JURISDICCIONALES, DEBEN ANALIZARSE A LA LUZ DE LOS ARTÍCULOS 14 Y 16», ya citada.
El tercer concepto de violación es inoperante. El tercer concepto es inoperante porque no combate la consideración toral.
Por las razones expuestas, el tercer concepto es inoperante.
Tampoco asiste razón a la quejosa cuando pide que se conceda para que la Sala deje insubsistente la sentencia.
En suma, los conceptos de violación son infundados e inoperantes.
SÉPTIMO. Efectos. 1. Deje insubsistente la sentencia reclamada.
2. Dicte otra en la que reitere lo que no fue materia de la concesión.
Por lo expuesto y fundado, se:
R E S U E L V E
ÚNICO. La Justicia de la Unión no ampara ni protege.
SÍNTESIS
TEMA: resumen del asunto.
"""

MODERNA = """SEXTO. Estudio. Los conceptos de violación son fundados.
Sentencia reclamada
La Sala tuvo por precluido el derecho de ampliar.
Conceptos de violación
En el primero alega la preclusión; en el segundo, los alegatos.
Solución
1. ¿La Sala debió admitir la ampliación de demanda?
Sí. El primer concepto de violación es fundado, porque el plazo corría desde la notificación.
Del artículo 17 de la Ley Federal de Procedimiento Contencioso Administrativo deriva el derecho de ampliar la demanda dentro del plazo.
2. ¿La Sala debía pronunciarse sobre los alegatos?
Es innecesario. El segundo concepto de violación queda sin materia porque se repone el procedimiento.
SÉPTIMO. Efectos. 1. Deje insubsistente la sentencia definitiva.
Por lo expuesto y fundado, se:
R E S U E L V E
"""

# El escrito: un encabezado (datos del propio asunto, que NO son anclas) de
# más de 1,500 caracteres y un cuerpo con cinco anclas.
ESCRITO = ("EXPEDIENTE: 765/25-09-01-8-ST. QUEJOSO: EMPRESA DE PRUEBA, S.A. DE C.V. "
           "ASUNTO: DEMANDA DE AMPARO DIRECTO CONTRA LA SENTENCIA DEFINITIVA DE 28 DE "
           "NOVIEMBRE DE 2025. " + "Texto de relleno del encabezado. " * 50 + "\n"
           "CONCEPTOS DE VIOLACIÓN\n"
           "PRIMERO. La Sala violó el artículo 17 de la Ley Federal de\nProcedimiento "
           "Contencioso Administrativo al negar la ampliación notificada el 13 de agosto "
           "de 2025. Es aplicable la tesis de registro 2010224. El crédito es por "
           "$279,210.00 y se relaciona con el amparo directo 312/2023. " * 1 + "\n")


print("\n1 · EL CORTE POR SUBTÍTULOS")
s = me.secciones(ESTANDAR)
ok(s["solucion_rotulada"] and s["solucion"].startswith("Por cuestión de método"),
   "la Solución empieza después de su rótulo")
ok("En suma" in s["solucion"] and "SÉPTIMO" not in s["solucion"],
   "la Solución termina donde empiezan los Efectos")
ok(s["efectos"].startswith("SÉPTIMO. Efectos") and "Dicte otra" in s["efectos"],
   "los Efectos se aíslan como considerando propio")
ok(s["sintesis"].startswith("TEMA"), "la SÍNTESIS queda aparte")
ok("valoró mal la confesional" in s["conceptos"] and "confirmó" in s["sentencia_reclamada"],
   "los dos resúmenes quedan en su sitio y fuera de la Solución")
ok(me.secciones("Los conceptos de violación son ineficaces.\n\nPrimer párrafo.")["solucion"]
   .startswith("Los conceptos"), "un oro sin rótulos (ADC 722) se mide entero")
ok(not me.secciones("SEXTO. Estudio.\nTexto.\n")["solucion_rotulada"],
   "sin «Solución» la fila lo dice")

print("\n2 · NÚMEROS, FECHAS Y CIFRAS ESCRITOS CON LETRA")
ok(me.numero_de_palabras("doscientos setenta y nueve mil doscientos diez") == 279210,
   "doscientos setenta y nueve mil doscientos diez = 279210")
ok(me.numero_de_palabras("dos mil veinticinco") == 2025 and
   me.numero_de_palabras("treinta y uno") == 31, "años y días")
ok(me.numero_de_palabras("un millón doscientos mil") == 1200000, "millones")
f1 = {k for _, _, k in me.fechas("el 28 de noviembre de 2025")}
f2 = {k for _, _, k in me.fechas("de veintiocho de noviembre de dos mil veinticinco")}
f3 = {k for _, _, k in me.fechas("a los veintiocho días del mes de noviembre de dos mil veinticinco")}
ok(f1 == f2 == f3 == {"fecha 2025-11-28"}, "la misma fecha en cifras, en letra y en fórmula notarial")
ok({k for _, _, k in me.fechas("el primero de julio de dos mil veinticinco")} == {"fecha 2025-07-01"},
   "«primero de julio»")
c1 = {k for _, _, k in me.cifras("por $279,210.00 (doscientos setenta y nueve mil doscientos diez pesos 00/100 M.N.)")}
ok(c1 == {"cifra 279210"}, "la cifra en número y en letra es la misma ancla")

print("\n3 · ARTÍCULO CON SU LEY")
P = lambda t: [k for _, _, k in me.preceptos(t)]
ok(P("los artículos 14 y 16 de la Constitución Política de los Estados Unidos Mexicanos")
   == ["art 14 cpeum", "art 16 cpeum"], "lista de artículos con una ley")
ok(P("el artículo 16 constitucional") == ["art 16 cpeum"], "«constitucional»")
ok(P("los artículos 8, numeral 1, y 25, numeral 1, de la Convención Americana") ==
   ["art 8 cadh", "art 25 cadh"], "el «numeral 1» no es un artículo")
ok(P("artículo 79, fracción IV, inciso b), de la Ley de Amparo. Y el artículo 80 del mismo "
     "ordenamiento.") == ["art 79 ley_amparo", "art 80 ley_amparo"], "anáfora «del mismo ordenamiento»")
ok(P("artículos 1o. y 17 constitucionales") == ["art 1 cpeum", "art 17 cpeum"],
   "«1o.» no corta la frase")
ok(P("el artículo 36 de la ley de la materia y el artículo 17 de la Ley de Amparo") ==
   ["art 17 ley_amparo"], "«la ley de la materia» no se atribuye a la última ley")
ok(P("artículo 210-A del Código Fiscal de la Federación") == ["art 210-a cff"], "sufijo «-A»")
ok(P("artículo 47 de la Ley Federal del Procedimiento Contencioso Administrativo") ==
   ["art 47 lfpca"], "«del Procedimiento» con artículo")
ok(P("el artículo 17 establece que toda persona tiene derecho a que se le administre "
     "justicia por tribunales expeditos para impartirla en los plazos y términos que fijen "
     "las leyes, en términos de la Ley Orgánica") == [],
   "una ley a más de 90 caracteres ya no es la del artículo")
ok(P("artículo 5 de la Ley de Ingresos de la Federación y el artículo 20 de la Ley Agraria")
   == ["art 5 ley ingresos federacion", "art 20 ley_agraria"],
   "el nombre de una ley desconocida no se come el artículo siguiente")
ok(P("artículo 17 de la Ley Federal de\nProcedimiento Contencioso Administrativo") ==
   ["art 17 ley federal"], "con el salto del PDF la ley sale truncada: por eso anclas() une líneas")
ok({k for _, _, k in me.anclas("artículo 17 de la Ley Federal de\nProcedimiento Contencioso "
                               "Administrativo")} == {"art 17 lfpca"},
   "anclas() une el salto de línea del PDF")

print("\n4 · EXPEDIENTES Y REGISTROS")
A = lambda t: {k for _, _, k in me.anclas(t)}
ok(A("amparo directo 312/2023 y juicio 1114/2017") == {"exp 312/2023", "exp 1114/2017"},
   "número/año es expediente")
ok(A("la jurisprudencia 1a./J. 104/2008 (9a.)") == set(), "la clave de una tesis no")
ok(A("Acuerdo General 6/2026 del Pleno") == set(), "el número de un acuerdo general no")
ok(A("el 28/11/2025") == {"fecha 2025-11-28"}, "una fecha con barras es fecha, no expediente")
ok(A("Registro No. 196451") == {"reg 196451"} and A("registro digital: 2,014,643") ==
   {"reg 2014643"}, "registros con sus grafías")

print("\n5 · LAS MEDIDAS SOBRE UNA SOLUCIÓN ESTÁNDAR")
m = me.medir(ESTANDAR, n=3)
d = m["detalle"]
ok(m["apartados"] == 2 and d["ordinales_por_apartado"] == [[1, 2], [3]],
   "dos apartados: el anuncio de método no abre uno; el tercero sí")
ok(all(x <= 30 for x in d["aperturas"]), f"apertura hasta la calificación ({d['aperturas']})")
ok(m["metodo_molde"] == 1, "el molde «Por cuestión de método»")
ok(m["citas"] == 2 and m["citas_repetidas"] == 1 and d["citas_repetidas"] == ["reg 176546"],
   "la tesis traída otra vez sólo por el rubro es la misma (gen_642)")
ok(d["preceptos_reexpuestos"] == ["art 16 cpeum"],
   "el art. 16 CPEUM expuesto como premisa en dos apartados")
ok(m["calif_max_concepto"] == 3 and m["conceptos_calif_mas_2"] == 1,
   "el tercer concepto calificado tres veces (V8)")
ok(m["objeciones"] == 2, "«No pasa inadvertido» y «Tampoco» son objeciones")
ok(m["remisiones"] == 1 and m["remisiones_en_blanco"] == 1,
   "«Por las razones expuestas, el tercer concepto…» es remisión en blanco: el ordinal es "
   "del concepto que se califica, no el destino")
R = me.remisiones
ok(R("Como se precisó al examinar el primer concepto, la identidad quedó acreditada.") ==
   {"remisiones": 1, "con_destino": 1, "en_blanco": 0}, "remisión con destino")
ok(R("Como se dijo, la identidad quedó acreditada con la confesión del quejoso de poseer "
     "dentro de la parcela de la actora desde hace muchos años.") ==
   {"remisiones": 1, "con_destino": 0, "en_blanco": 0},
   "sin destino pero con la proposición que decide: no es en blanco")
ok(R("Por las razones expuestas en el apartado anterior, es infundado.")["con_destino"] == 1,
   "«en el apartado anterior» es destino")
ok(m["ordenes_fuera_efectos"] == 1, "«para que la Sala deje insubsistente» en la Solución; "
   "los Efectos no cuentan")
ok(m["recapitulacion_final"] == 1 and m["recapitulaciones"] == 1, "«En suma» al final")
_rec = next(x[4] for x in me.METRICAS if x[0] == "recapitulacion_final")
ok(_rec(1, m), "cierre con dos apartados: acusa (Decisión 3 b)")
ok(not _rec(1, {"apartados": 3, "resultados_distintos": True}) and
   _rec(1, {"apartados": 3, "resultados_distintos": False}),
   "con tres apartados, sólo si sus resultados son distintos (Decisión 3 b, las dos condiciones)")
ok(me.resultados_distintos([{"resultado": "infundado"}] * 3) is False and
   me.resultados_distintos([{"resultado": "infundado"}, {"resultado": "inoperante"}]) and
   me.resultados_distintos([{"resultado": "infundado"}, {"resultado": None}]),
   "resultados distintos por familia; el ilegible no acusa")
ok(me.familia_calificacion("Infundados") == "infundado" and
   me.familia_calificacion("sin materia") == "innecesario", "familias de calificación")
ok(m["cobertura_ordinal"] == 1.0 and me.medir(ESTANDAR, n=4)["cobertura_ordinal"] == 0.75,
   "cobertura por ordinal con la n del contador")
ok(m["n_planteamientos"] == 3, "la n forzada se respeta")

print("\n6 · COBERTURA CONTRA LA DEMANDA")
esc = me.anclas_del_escrito(ESCRITO)
ok("fecha 2025-11-28" not in esc and "exp 765/25-09-01-8-st" not in esc,
   "los datos del encabezado del escrito no son anclas")
ok(esc == {"art 17 lfpca", "fecha 2025-08-13", "reg 2010224", "cifra 279210", "exp 312/2023"},
   f"las cinco anclas del cuerpo ({sorted(esc)})")
m2 = me.medir(ESTANDAR, ESCRITO, n=3)
ok(m2["cobertura_demanda"] == 0.6 and m2["detalle"]["anclas_faltan"] ==
   ["exp 312/2023", "reg 2010224"], "3 de 5: el artículo, la cifra en letra y la fecha en letra")

print("\n7 · LA MODERNA")
mm = me.medir(MODERNA, n=2)
aps = me.apartados(me.secciones(MODERNA)["solucion"])
ok(len(aps) == 2 and [a["tipo"] for a in aps] == ["pregunta", "pregunta"],
   "dos preguntas, dos apartados; la respuesta no abre otro")
ok([sorted(a["ordinales"]) for a in aps] == [[1], [2]],
   "«Sí. El primer concepto…»: el concepto se lee en la segunda frase")
ok(mm["ordenes_fuera_efectos"] == 0 and mm["cobertura_ordinal"] == 1.0, "moderna limpia")
ok(me.medir("")["palabras_solucion"] == 0, "un texto vacío no revienta")

print("\n8 · FALSOS ORDINALES")
ok(me.ordinales("depósitos realizados por el tercero interesado por concepto de alimentos")
   == set(), "«tercero interesado» y «por concepto de» no nombran el tercer concepto (ADC 526)")
ok(me.ordinales("En primer lugar, el concepto de violación es infundado") == set(),
   "«en primer lugar» no es el primer concepto")
ok(me.ordinales("el único concepto de violación") == {1}, "«único» es el primero")
ok(not me.califica("la sentencia está debidamente fundada y motivada"),
   "«fundada y motivada» no es una calificación")

# ═══════════════════════════════════════════════════════════════════════════
print("\n9 · CALIBRACIÓN CONTRA LOS ENGROSES (si el corpus está)")
if me.CASOS_KINGSTON.exists():
    oros = {f["caso"]: f for f in me.filas_oro()}
    ok(len(oros) == 24, f"los 24 de oro sólido ({len(oros)})")
    o103 = me.medir(oros["ADA 103-2025"]["texto"], oros["ADA 103-2025"]["escrito"])
    ok(o103["apartados"] == 3 and o103["apertura_mediana"] <= 30,
       f"oro 103: tres apartados y apertura corta ({o103['apertura_mediana']}; w2 midió 27)")
    ok(o103["recapitulacion_final"] == 1 and o103["resultados_distintos"] and
       not _rec(o103["recapitulacion_final"], o103),
       "oro 103: cierra con infundado/inoperante/ineficaz en tres apartados: no se acusa")
    k702 = next(k for k in oros if k.startswith("ADA 702"))
    ok(me.medir(oros[k702]["texto"])["ordenes_fuera_efectos"] == 0,
       "oro 702: «impide que una decisión… deje insubsistente» no es una orden")
    k43 = next(k for k in oros if k.startswith("ADC 43"))
    ok(me.medir(oros[k43]["texto"])["citas_repetidas"] <= 1,
       "oro 43: discutir la tesis que aplicó la responsable no es citarla cuatro veces")
    k526 = next(k for k in oros if k.startswith("ADC 526"))
    ok(3 not in {o for a in me.apartados(me.secciones(oros[k526]["texto"])["solucion"])
                 for o in a["ordinales"]}, "oro 526: no hay apartado del «tercer concepto»")
    malos = []
    for clave, _, _, _, acusa in me.METRICAS:
        if acusa in (None,) or clave in ("apertura_mediana", "remisiones_con_contenido_pct"):
            continue
        ms = [me.medir(f["texto"], f["escrito"]) for f in oros.values()]
        ev = [x for x in ms if x.get(clave) is not None]
        if ev and sum(1 for x in ev if acusa(x[clave], x)) / len(ev) > me.TOPE_ACUSADOS:
            malos.append(clave)
    ok(not malos, f"las alarmas que se proponen usar acusan ≤ 10 % de los buenos ({malos})")
else:
    print("   (sin corpus Kingston en este equipo: se salta)")

# ═══════════════════════════════════════════════════════════════════════════
print("\n10 · SÓLO CUENTAS DE CASA")
ok(be.es_de_casa("Administracion@iurexia.com ") and be.es_de_casa("soporte@iurexia.com"),
   "las dos cuentas del banco")
for ajena in ("jdm.juridico@gmail.com", "otro@iurexia.com", ""):
    try:
        be.exigir_casa(ajena)
        ok(False, f"«{ajena}» debió rechazarse")
    except be.CuentaAjena:
        ok(True, f"«{ajena or 'vacío'}» se rechaza (la lista es cerrada)")
try:
    be.formulario(be.Caso("1/2026", "jdm.juridico@gmail.com", "acervo"), "estandar",
                  be.Variante("v1", "v1"))
    ok(False, "el formulario de una cuenta ajena debió rechazarse")
except be.CuentaAjena:
    ok(True, "ni siquiera se arma el formulario de una cuenta ajena")
try:
    be.plan([be.Caso("1/2026", "alguien@gmail.com", "acervo")], be.variantes("v1"), 1, "x",
            Path(tempfile.mkdtemp()))
    ok(False, "el plan con una cuenta ajena debió rechazarse")
except be.CuentaAjena:
    ok(True, "tampoco se planea")

print("\n11 · VARIANTES Y FORMULARIO")
vs = be.variantes("v1,v2,v1@r,prod,v1")
ok([(v.etiqueta, v.servidor) for v in vs] ==
   [("v1", "v1"), ("v2", "v2"), ("v1@r", "v1"), ("prod", None)],
   "«v1@r» pide v1 y se guarda aparte; «prod» no pide nada; sin duplicados")
f = be.formulario(be.POR_NUMERO["103/2025"], "moderna", be.Variante("v2", "v2"))
ok(f == {"numero": "103/2025", "user_email": "administracion@iurexia.com",
         "formato": "moderna", "modo_decision": "acervo", "variante_estudio": "v2"},
   "acervo: sólo número, cuenta, forma, modo y variante")
ok("variante_estudio" not in be.formulario(be.POR_NUMERO["103/2025"], "estandar",
                                           be.Variante("prod", None)), "«prod» no manda el campo")
tmp = Path(tempfile.mkdtemp())
(tmp / "crit.json").write_text(json.dumps([{"problema": "¿P?", "sentido": "fundado"}]))
(tmp / "glob.json").write_text(json.dumps({"sentido": "fundado",
                                           "contexto": {"resolvio": "Declaró la nulidad."}}))
(tmp / "ctx.txt").write_text("contexto del secretario")
c93 = be.Caso("93/2026", be.SOPORTE, "por_problema", criterios=str(tmp / "crit.json"),
              global_=str(tmp / "glob.json"), contexto=str(tmp / "ctx.txt"))
raiz = Path(tempfile.mkdtemp())
f = be.formulario(c93, "estandar", be.Variante("v1", "v1"), raiz)
ok(f["modo_decision"] == "por_problema" and f["resolvio_declarado"] == "Declaró la nulidad."
   and json.loads(f["criterios_json"])[0]["sentido"] == "fundado"
   and f["contexto"] == "contexto del secretario",
   "93: criterio, global, contexto y resolvio_declarado como en generar_formatos.sh")
ok((raiz / "_entradas" / "93-2026" / "crit.json").exists(),
   "las entradas del 93 se congelan en bancos/estudio/_entradas")
(tmp / "crit.json").unlink()
f2 = be.formulario(c93, "estandar", be.Variante("v1", "v1"), raiz)
ok(json.loads(f2["criterios_json"])[0]["problema"] == "¿P?",
   "si el scratchpad desaparece, se lee la copia congelada")
ok(be.elegir_casos("103-2025, 93/2026")[1].numero == "93/2026" and
   len(be.elegir_casos("")) == 14, "casos por número en cualquier grafía; 14 por omisión")

print("\n12 · EL FLUJO SSE Y LA CORRIDA")
evs = be.eventos_sse([": latido", "", 'data: {"tipo": "texto", "dato": "SEXTO."}', "",
                      'data: {"tipo": "texto",', 'data: "dato": " Estudio."}', "",
                      'data: {"tipo": "componiendo"}', ""])
ok([e["tipo"] for e in evs] == ["texto", "texto", "componiendo"],
   "los latidos se saltan; un evento de dos líneas se junta")


def _docx(parrafos):
    import docx
    d = docx.Document()
    for p in parrafos:
        d.add_paragraph(p)
    d.add_paragraph("")
    b = io.BytesIO()
    d.save(b)
    return base64.b64encode(b.getvalue()).decode()


DOC = _docx(["SEXTO. Estudio. Los conceptos son fundados.", "Solución",
             "Sobre el primer concepto de violación, en el que sostiene X. Es fundado."])


def _acc(listo):
    a = be.Acumulador(t0=0)
    a.meter({"tipo": "texto", "dato": "SEXTO. "}, ahora=5)
    a.meter({"tipo": "texto", "dato": "Estudio."}, ahora=6)
    a.meter({"tipo": "componiendo"}, ahora=90)
    if listo is not None:
        a.meter(listo, ahora=120)
    return a


c = be.POR_NUMERO["103/2025"]
fila, docx = be.resultado(_acc({"tipo": "listo", "docx_b64": DOC, "variante": "v2",
                                "commit": "abc1234", "version": 7, "avisos": ["a"],
                                "tokens": 123}),
                          c, be.Variante("v2", "v2"), 1, "estandar", "e")
ok(fila["ok"] and not fila["descartada"] and fila["estudio_crudo"] == "SEXTO. Estudio."
   and fila["texto"].splitlines()[1] == "Solución" and docx,
   "estudio crudo de los «texto» y el texto del .docx, párrafos no vacíos")
ok(fila["commit"] == "abc1234" and fila["listo"]["tokens"] == 123 and
   "docx_b64" not in fila["listo"] and fila["t_primer_texto"] == 5,
   "variante, commit y lo que añada el servidor se guardan; el base64 no")
f_mal, _ = be.resultado(_acc({"tipo": "listo", "docx_b64": DOC, "variante": "v1"}),
                        c, be.Variante("v2", "v2"), 1, "estandar", "e")
ok(f_mal["descartada"] and "corrió «v1»" in f_mal["descartada"], "variante que no casa: descartada")
f_sin, _ = be.resultado(_acc({"tipo": "listo", "docx_b64": DOC}), c,
                        be.Variante("v2", "v2"), 1, "estandar", "e")
ok(bool(f_sin["descartada"]), "servidor que no dice la variante: descartada (no se sabe qué corrió)")
f_prod, _ = be.resultado(_acc({"tipo": "listo", "docx_b64": DOC}), c,
                         be.Variante("prod", None), 1, "estandar", "e")
ok(f_prod["ok"] and not f_prod["descartada"], "«prod» acepta lo que haya")
f_corta, _ = be.resultado(_acc(None), c, be.Variante("v1", "v1"), 1, "estandar", "e")
ok(not f_corta["ok"] and "sin evento «listo»" in f_corta["error"], "sin «listo» es error")

# Los flujos grabados de producción (25-sep, diagnóstico): si siguen en el
# scratchpad, el lector tiene que reconstruir exactamente lo que se guardó.
_G = be.DIR93.parent
_grabados = [(_G / "diag" / "gen_103-2025_stream.txt", _G / "diag" / "gen_103-2025_estudio_crudo.txt",
              _G / "diag" / "gen_103-2025.txt"),
             (_G / "93" / "f_estandar2_stream.txt", None, _G / "93" / "estandar2.txt")]
for st, crudo, txt in _grabados:
    if not (st.exists() and txt.exists()):
        print(f"   (sin {st.name}: se salta)")
        continue
    a = be.Acumulador(t0=0)
    for e in be.eventos_sse(st.read_text(encoding="utf-8").split("\n")):
        a.meter(e, ahora=1)
    fg, _ = be.resultado(a, c, be.Variante("prod", None), 1, "estandar", "e")
    ok(fg["ok"] and fg["texto"] == txt.read_text(encoding="utf-8") and
       (crudo is None or fg["estudio_crudo"] == crudo.read_text(encoding="utf-8")),
       f"flujo grabado {st.name}: .docx y estudio crudo idénticos a lo guardado")

print("\n13 · REANUDABLE Y EN ORDEN")
raiz = Path(tempfile.mkdtemp())
vs = be.variantes("v1,v2")
p = be.plan([c], vs, 2, "e", raiz)
ok([f"{v.etiqueta}#{k}" for v, k, _, _ in p["103/2025"]] == ["v1#1", "v2#1", "v1#2", "v2#2"],
   "las variantes se intercalan corrida a corrida")
be.guardar(be.ruta_corrida(raiz, "e", c, vs[0], 1), dict(fila, variante="v1"), b"docx")
be.guardar(be.ruta_corrida(raiz, "e", c, vs[1], 1), f_mal)
p = be.plan([c], vs, 2, "e", raiz)
ok([ya for _, _, _, ya in p["103/2025"]] == [True, False, False, False],
   "lo bien guardado no se repite; lo descartado sí")
ok(be.ruta_corrida(raiz, "e", c, vs[0], 1).with_suffix(".docx").exists(), "el .docx al lado")
r = be.en_seco([c], vs, 2, "estandar", "e", raiz, imprimir=lambda *a: None)
ok(r["pendientes"] == 3 and not r["problemas"], "en seco: 3 por hacer, sin red")


# ── el servidor simulado ─────────────────────────────────────────────────
def _servidor(activos, maximos, lento=0.05, falla=(), cuelga=(), inicios=None):
    import httpx
    import time as _t

    async def manejar(req):
        cuerpo = req.content.decode("utf-8", "replace")
        num = re.search(r'name="numero"\r\n\r\n([^\r]+)', cuerpo).group(1)
        var = re.search(r'name="variante_estudio"\r\n\r\n([^\r]+)', cuerpo)
        if inicios is not None:
            inicios.setdefault(num, []).append(_t.monotonic())
        if num in falla:
            return httpx.Response(409, text="No hay propuesta en este proceso.")
        activos[num] = activos.get(num, 0) + 1
        activos["_"] = activos.get("_", 0) + 1
        maximos[num] = max(maximos.get(num, 0), activos[num])
        maximos["_"] = max(maximos.get("_", 0), activos["_"])
        try:
            await asyncio.sleep(10 if num in cuelga else lento)
        finally:
            # La corrida cortada por el tope se cancela aquí dentro: sin el
            # finally el simulador seguiría contándola como activa.
            activos[num] -= 1
            activos["_"] -= 1
        listo = {"tipo": "listo", "docx_b64": DOC, "palabras": 12, "version": 3,
                 "variante": var.group(1) if var else None, "commit": "c0ffee1"}
        sse = (": latido\n\n" + 'data: {"tipo": "texto", "dato": "SEXTO. Estudio."}\n\n'
               + 'data: {"tipo": "componiendo"}\n\n'
               + "data: " + json.dumps(listo) + "\n\n")
        return httpx.Response(200, content=sse.encode(),
                              headers={"content-type": "text/event-stream"})
    return httpx.AsyncClient(transport=httpx.MockTransport(manejar))


casos3 = [be.POR_NUMERO[n] for n in ("103/2025", "642/2024", "174/2026")]
activos, maximos = {}, {}
raiz = Path(tempfile.mkdtemp())
hechas = asyncio.run(be.correr(casos3, vs, 2, "estandar", "e", paralelo=3, base="http://x",
                               raiz=raiz, cliente=_servidor(activos, maximos),
                               imprimir=lambda *a: None, espera_tras_corte=0))
ok(len(hechas) == 12 and all(h["ok"] and not h["descartada"] for h in hechas),
   "12 corridas buenas (3 casos × 2 variantes × 2)")
ok(all(maximos[n] == 1 for n in ("103/2025", "642/2024", "174/2026")),
   "nunca dos corridas a la vez sobre el mismo expediente")
ok(maximos["_"] >= 2, f"sí en paralelo entre expedientes (máx. {maximos['_']})")
ok(len(list((raiz / "e").glob("*/*.json"))) == 12 and (raiz / "e" / "manifiesto.json").exists(),
   "cada corrida en bancos/estudio/<etiqueta>/<caso>/<variante>_<k>.json")
otra = asyncio.run(be.correr(casos3, vs, 2, "estandar", "e", paralelo=3, base="http://x",
                             raiz=raiz, cliente=_servidor({}, {}), imprimir=lambda *a: None))
ok(otra == [], "reanudar sin nada pendiente no repite nada")
try:
    asyncio.run(be.correr(casos3, vs, 2, "moderna", "e", paralelo=3, base="http://x",
                          raiz=raiz, cliente=_servidor({}, {}), imprimir=lambda *a: None))
    ok(False, "una etiqueta estándar no admite corridas modernas")
except SystemExit as ex:
    ok("usa otra etiqueta" in str(ex),
       "una etiqueta, una forma: la moderna no se da por hecha con corridas estándar")

activos, maximos, inicios = {}, {}, {}
raiz = Path(tempfile.mkdtemp())
msgs_corte = []
hechas = asyncio.run(be.correr(casos3, vs, 2, "estandar", "e", paralelo=2, base="http://x",
                               raiz=raiz, cliente=_servidor(activos, maximos,
                                                            falla={"642/2024"},
                                                            cuelga={"174/2026"},
                                                            inicios=inicios),
                               timeout=0.3, imprimir=msgs_corte.append,
                               espera_tras_corte=0.4))
por = {}
for h in hechas:
    por.setdefault(h["caso"], []).append(h)
ok(len(por["642/2024"]) == 1 and por["642/2024"][0]["http"] == 409,
   "un 409 de la sesión para el caso a la primera")
ok(all("sin evento «listo» a los 0 s" in h["error"] or "0 s" in h["error"]
       for h in por["174/2026"]) and len(por["174/2026"]) == 4,
   "el tope de tiempo corta la corrida colgada y la marca como error")
ok(maximos["_"] <= 2, "el paralelo pedido se respeta")
# Revisión 26-sep: la corrida cortada sigue viva en el servidor (tarea propia
# que no se cancela) y acabará escribiendo la fila; la siguiente del MISMO
# expediente espera. Entre dos arranques del 174: 0.3 s de tope + 0.4 de espera.
_ini = inicios.get("174/2026", [])
ok(len(_ini) == 4 and all(b - a >= 0.65 for a, b in zip(_ini, _ini[1:])),
   f"tras una corrida cortada el expediente espera antes de la siguiente "
   f"({[round(b - a, 2) for a, b in zip(_ini, _ini[1:])]})")
ok(any("se cortó sin «listo»" in m for m in msgs_corte), "y lo dice")
ok(len(inicios.get("103/2025", [])) == 4 and max(inicios["103/2025"]) < _ini[-1],
   "la espera de un expediente no frena a los demás (el 103 acabó antes)")
F = lambda **k: dict({"ok": False, "eventos": {}, "error": "x", "http": 200}, **k)
ok(be.quedo_corriendo(F()) and be.quedo_corriendo(F(http=None)),
   "flujo cortado o sin respuesta: puede seguir vivo en el servidor")
ok(not any(be.quedo_corriendo(x) for x in (
    F(ok=True), F(listo={}), F(eventos={"error": 1}), F(http=409), F(http=502),
    F(error="formulario: FileNotFoundError", http=None))),
   "con «listo», con «error», con 4xx/5xx o sin mandar nada, no se espera")

# ── --congelar: la huella de la sesión (Supabase simulado) ──────────────
def _con_huella(huella_de_sesion):
    import httpx
    import os as _os
    _os.environ["SUPABASE_URL"] = "http://supa"
    _os.environ["SUPABASE_SERVICE_KEY"] = "clave-de-prueba"
    sse_srv = _servidor({}, {})

    async def manejar(req):
        if req.url.path.startswith("/rest/v1/taller_sesiones"):
            ok(req.url.params.get("email") == "eq.administracion@iurexia.com"
               and req.method == "GET",
               "la huella se pide sólo para la cuenta de casa del caso, y sólo se lee")
            sel = req.url.params.get("select") or ""
            ok("material:estado->material" in sel and "estado->fases," not in sel
               and "fuentes" not in sel,
               "la huella pide el acervo y las piezas del criterio, no `fuentes` (1.2 MB)")
            return httpx.Response(200, json=[{"problemas": [huella_de_sesion],
                                              "propuestas": []}])
        return await sse_srv._transport.handle_async_request(req)
    return httpx.AsyncClient(transport=httpx.MockTransport(manejar))


raiz = Path(tempfile.mkdtemp())
msgs = []
h1 = asyncio.run(be.correr([c], vs, 1, "estandar", "e", base="http://x", raiz=raiz,
                           congelar=True, cliente=_con_huella("¿P1?"), imprimir=msgs.append))
ok(len(h1) == 2 and (raiz / "e" / "103-2025" / "huella.txt").exists() and
   all(h["huella_sesion"] for h in h1), "la primera corrida fija la huella de la sesión")
h2 = asyncio.run(be.correr([c], vs, 2, "estandar", "e", base="http://x", raiz=raiz,
                           congelar=True, cliente=_con_huella("¿P1 retocada?"),
                           imprimir=msgs.append))
ok(h2 == [] and any("LA SESIÓN CAMBIÓ" in m for m in msgs),
   "si la sesión cambió (otro criterio), el caso se para y se dice")
ok(be.huella_de({"problemas": ["¿P?"], "material": {"tesis": [1]}}) !=
   be.huella_de({"problemas": ["¿P?"], "material": {"tesis": [2]}}),
   "otro acervo, otra huella: una consulta a media tanda cambia las citas")

# ── --sesiones: sólo lee y dice qué fallaría o ensuciaría ──
def _supa_sesiones(filas):
    import httpx

    async def manejar(req):
        ok(req.method == "GET", "revisar las sesiones sólo lee")
        num = req.url.params.get("expediente", "")[3:]
        return httpx.Response(200, json=[filas[num]] if num in filas else [])
    return httpx.AsyncClient(transport=httpx.MockTransport(manejar))


rev = asyncio.run(be.revisar_sesiones(
    [be.POR_NUMERO[n] for n in ("103/2025", "642/2024", "174/2026", "93/2026")],
    cliente=_supa_sesiones({
        "103/2025": {"consultado": True, "propuestas": [{"sentido": "fundado"}],
                     "material_completo": None, "contraste": None},
        "642/2024": {"consultado": True, "propuestas": [], "material_completo": True,
                     "contraste": "listo"},
        "93/2026": {"consultado": True, "propuestas": [], "material_completo": True,
                    "contraste": "listo"}}),
    imprimir=lambda *a: None))
ok([g for g, _ in rev["103/2025"]] == [False, False],
   "acervo viejo y sin contraste: avisa, no para (ensucia la medida, no la tumba)")
ok(any(g and "sin propuestas" in t for g, t in rev["642/2024"]),
   "modo acervo sin propuestas: la corrida daría 409")
ok(any(g and "404" in t for g, t in rev["174/2026"]), "sin fila: 404")
ok(rev["93/2026"] == [], "el 93 por problema no necesita propuestas guardadas")
import os as _os
del _os.environ["SUPABASE_URL"], _os.environ["SUPABASE_SERVICE_KEY"]

# ═══════════════════════════════════════════════════════════════════════════
print("\n14 · EL COMPARADOR")
ok(abs(ce.signo(12, 16) - 0.0384) < 0.001, "prueba de signo: 12 de 16 → p ≈ 0.04 (w2 §6.6)")
ok(ce.signo(5, 16) > 0.5,
   "5 mejoras y 11 «≈» no es significativo: los empates cuentan como no ganados")
ok(ce.veredicto(-3, 1, "menos") == "mejora" and ce.veredicto(3, 1, "menos") == "EMPEORA"
   and ce.veredicto(0.5, 1, "menos") == "≈" and ce.veredicto(1, None, None) == "—",
   "veredicto con dirección y banda de ruido")

# v1: la Solución nombra las cinco anclas; v2 pierde el registro y el
# expediente en el caso A (bloquea) y repite menos (menos objeciones).
SOL_V1 = ESTANDAR.replace(
    "En suma,", "Consta el registro 2010224 y el amparo directo 312/2023. En suma,")
SOL_V2 = ESTANDAR.replace("No pasa inadvertido que", "Además,")
raiz = Path(tempfile.mkdtemp())
for caso_n, textos in (("103/2025", {"v1": [SOL_V1] * 3, "v2": [SOL_V2] * 3}),
                       ("642/2024", {"v1": [SOL_V1] * 3, "v2": [SOL_V1] * 3})):
    cc = be.POR_NUMERO[caso_n]
    for v, ts in textos.items():
        for k, t in enumerate(ts, 1):
            be.guardar(be.ruta_corrida(raiz, "e", cc, be.Variante(v, v), k),
                       {"caso": caso_n, "variante": v, "k": k, "ok": True, "descartada": None,
                        "texto": t, "commit": "c0ffee1", "version": 3, "t_total": 150,
                        "palabras": 3000})
be.guardar(be.ruta_corrida(raiz, "e", be.POR_NUMERO["103/2025"], be.Variante("v2", "v2"), 4),
           {"caso": "103/2025", "variante": "v2", "k": 4, "ok": True,
            "descartada": "se pidió «v2» y corrió «v1»", "texto": SOL_V1})
datos = ce.cargar(raiz, "e")
ok(len(datos["103/2025"]["v2"]) == 4 and len(ce.buenas(datos["103/2025"]["v2"])) == 3,
   "las descartadas se cargan pero no se comparan")
med = ce.Medidor(raiz, escritos={"103/2025": ESCRITO, "642/2024": ESCRITO},
                 oros={"103/2025": None, "642/2024": SOL_V1})
par = ce.comparar_par(datos, med, "v1", "v2")
bl = {f["caso"]: f["regresion"] for f in par["bloquea"]}
ok(list(bl) == ["103/2025"] and bl["103/2025"]["anclas"] == ["exp 312/2023", "reg 2010224"],
   "BLOQUEA el caso que pierde anclas que v1 nombraba en todas sus corridas")
ok(par["resumen"]["objeciones_x1000"]["mejora"] == 1 and
   par["resumen"]["objeciones_x1000"]["ruido"] == 1,
   "menos objeciones: mejora en un caso, igual en el otro")
fila642 = next(f for f in par["filas"] if f["caso"] == "642/2024")
ok(fila642["oro"] is not None and fila642["metricas"]["citas_repetidas"]["ruido"] == 0,
   "el oro del caso se mide al lado; tres corridas iguales dan ruido cero")
md = ce.informe("e", [par], ce.operacion(datos), {"palabras_solucion": {"mediana": 2982,
                                                                          "p90": 5425}})
ok("BLOQUEA en 1 caso" in md and "#### Palabras de la Solución" in md and "2982" in md
   and "c0ffee1" in md, "el informe trae el bloqueo, las tablas por métrica, el oro y el commit")
ok("1× se pidió «v2» y corrió «v1»" in md, "y dice por qué se descartó lo descartado")

# ── el bloqueo, calibrado contra el ruido (revisión 26-sep) ──────────────
# ESTANDAR nombra 3 de las 5 anclas (0.6); SOL_V1 las 5 (1.0).
def _tanda(raiz, caso_n, variante, textos, **extra):
    cc = be.POR_NUMERO[caso_n]
    for k, t in enumerate(textos, 1):
        f = {"caso": caso_n, "variante": variante, "k": k, "ok": True, "descartada": None,
             "texto": t, "commit": "c0ffee1", "version": 3, "t_total": 150, "palabras": 3000}
        f.update(extra)
        be.guardar(be.ruta_corrida(raiz, "e", cc, be.Variante(variante, variante), k), f)


raiz = Path(tempfile.mkdtemp())
# 103: un ancla que sale 2 de 3 en la base y 1 de 3 en la variante es azar.
_tanda(raiz, "103/2025", "v1", [SOL_V1, SOL_V1, ESTANDAR])
_tanda(raiz, "103/2025", "v2", [SOL_V1, ESTANDAR, ESTANDAR])
# 642: la variante trunca.
_tanda(raiz, "642/2024", "v1", [SOL_V1] * 3)
_tanda(raiz, "642/2024", "v2", [SOL_V1] * 3, listo={"finish_reason": "length"})
# 174: la variante no entrega ninguna corrida buena.
_tanda(raiz, "174/2026", "v1", [SOL_V1] * 3)
for k in (1, 2, 3):
    be.guardar(be.ruta_corrida(raiz, "e", be.POR_NUMERO["174/2026"], be.Variante("v2", "v2"), k),
               {"caso": "174/2026", "variante": "v2", "k": k, "ok": False,
                "error": "HTTP 500: Internal Server Error"})
datos = ce.cargar(raiz, "e")
med = ce.Medidor(raiz, escritos={c: ESCRITO for c in ("103/2025", "642/2024", "174/2026")},
                 oros={c: None for c in ("103/2025", "642/2024", "174/2026")})
par = ce.comparar_par(datos, med, "v1", "v2")
f103 = next(f for f in par["filas"] if f["caso"] == "103/2025")
ok(not f103["regresion"]["anclas"] and f103["regresion"]["anclas_revisar"] ==
   ["exp 312/2023", "reg 2010224"] and "103/2025" not in [f["caso"] for f in par["bloquea"]],
   "un ancla 2/3 → 1/3 no bloquea (salta por azar con el mismo prompt): se lista a revisar")
ok([f["caso"] for f in par["bloquea"]] == ["642/2024"] and par["bloquea"][0]["truncadas"] == 3,
   "una corrida truncada (finish_reason=length) bloquea")
ok([f["caso"] for f in par["incompletos"]] == ["174/2026"] and ce.detiene(par),
   "la variante sin corridas buenas en un caso no se certifica: detiene")
md = ce.informe("e", [par], ce.operacion(datos), {})
ok("NO SE PUEDE CERTIFICAR en 1 caso" in md and "A revisar (no bloquea)" in md
   and "el bloqueo no está calibrado" in md and "| 3 |" in md,
   "el informe dice lo incompleto, lo que se revisa, que sin réplica no está calibrado "
   "y las truncadas")

# M1b con banda, y la réplica que calibra el bloqueo.
raiz = Path(tempfile.mkdtemp())
_tanda(raiz, "103/2025", "v1", [SOL_V1] * 3)
_tanda(raiz, "103/2025", "v1@r", [SOL_V1, SOL_V1, ESTANDAR])
_tanda(raiz, "103/2025", "v2", [ESTANDAR] * 3)
datos = ce.cargar(raiz, "e")
med = ce.Medidor(raiz, escritos={"103/2025": ESCRITO}, oros={"103/2025": None})
par = ce.comparar_par(datos, med, "v1", "v2", "v1@r")
r = par["filas"][0]["regresion"]
ok(not r["anclas"] and r["anclas_revisar"] == ["exp 312/2023", "reg 2010224"],
   "con réplica, el ancla que la réplica no nombró siempre no es estable: no bloquea")
ok(r["peor_corrida"] is None and r["peor_corrida_bruta"] == (1.0, 0.6),
   "la peor corrida más baja dentro de la banda (réplica 1.0-0.6) no bloquea")
ok(par["calibrado"] and [f["caso"] for f in par["falsos_bloqueos"]] == ["103/2025"],
   "la base contra su réplica bloquearía aquí: el informe lo dice como falta de calibración")
md = ce.informe("e", [par], ce.operacion(datos), {})
ok("El bloqueo NO está calibrado" in md, "y lo escribe")
par0 = ce.comparar_par(datos, med, "v1", "v2")
ok(par0["filas"][0]["regresion"]["anclas"] == ["exp 312/2023", "reg 2010224"],
   "sin réplica, el ancla que la base nombró SIEMPRE y la variante NUNCA sí bloquea")

print()
if FALLOS:
    print(f"FALLAN {len(FALLOS)}: " + " · ".join(FALLOS))
    raise SystemExit(1)
print("RESULTADO: TODAS LAS COMPROBACIONES PASAN")
