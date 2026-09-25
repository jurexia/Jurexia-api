"""El calendario que se explica solo, su leyenda y el mapa del plazo — 25-sep-2026.

    .venv/bin/python test_calendario_computo.py
"""
import datetime as d

import docx

import documento_generado as dg
import fase0_oportunidad as f0

FALLOS = []


def ok(cond, que):
    print(f"   {'PASA ' if cond else 'FALLA'}  {que}")
    if not cond:
        FALLOS.append(que)


def texto_tablas(doc):
    return "\n".join(c.text for t in doc.tables for f in t.rows for c in f.cells)


C93 = f0.computar(d.date(2025, 12, 5), d.date(2026, 1, 19), regla="lfpca_boletin", plazo=15,
                  responsable="Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa",
                  tipo_asunto="amparo_directo")
CEX = f0.computar(d.date(2026, 3, 2), d.date(2026, 3, 30), regla="personal", plazo=15,
                  responsable="Primera Sala Civil del Tribunal Superior de Justicia del Estado de Querétaro",
                  tipo_asunto="amparo_directo")
CRV = f0.computar(d.date(2026, 6, 22), d.date(2026, 7, 2), regla="personal", plazo=10,
                  responsable="Juzgado Primero de Distrito", tipo_asunto="amparo_revision")

print("\n1 · EL PÁRRAFO DICE POR QUÉ SE CUENTA ASÍ")
p = f0.parrafo_oportunidad(C93, "", "amparo_directo")
ok("artículo 65 de la Ley Federal de Procedimiento Contencioso Administrativo" in p, "93/2026: el boletín surte al tercer día, con su artículo 65")
ok("en términos del artículo 18 de la Ley de Amparo" in p, "y el plazo corre desde el día siguiente, artículo 18")
ok("ni del dieciséis de diciembre de dos mil veinticinco al uno de enero de dos mil veintiséis" in p,
   "los trece inhábiles son UN tramo, no trece fechas")
ok("último día del plazo" in p, "dice que se presentó el último día")
ok(not p.startswith("Igualmente, la presentación"), "ya no la versión corta sin cómputo")
pe = f0.parrafo_oportunidad(CEX, "", "amparo_directo")
ok("conforme a la ley del acto" in pe and "artículo 31" not in pe, "personal en amparo directo: la ley del acto, sin inventar artículo")
ok("del cuatro al veinticinco de marzo de dos mil veintiséis" in pe, "el rango no repite mes y año")
ok(f0.aviso_fundamento(CEX, "amparo_directo").startswith("EL SURTIMIENTO VA SIN PRECEPTO"), "y avisa que falta el precepto")
ok(f0.aviso_fundamento(C93, "amparo_directo") == "", "con precepto, no avisa")
pr = f0.parrafo_oportunidad(CRV, "", "amparo_revision")
ok("artículo 31, fracción II, de la Ley de Amparo" in pr and "artículo 22 de la Ley de Amparo" in pr,
   "revisión: surte por el 31, fracción II, y corre por el 22")
pf = f0.parrafo_oportunidad(f0.computar(d.date(2026, 3, 2), d.date(2026, 3, 10), regla="lfpca", plazo=15,
                                        responsable="Sala Regional", tipo_asunto="revision_fiscal"), "", "revision_fiscal")
ok("artículo 18" not in pf and "artículo 22" not in pf, "revisión fiscal: su 63 ya dice desde cuándo; no se le cuelga el 18")
ok("según la ley del acto" in str(f0.reglas_para("amparo_directo", "Tribunal de Justicia Administrativa del Estado de Jalisco")),
   "el desplegable ya no dice «art. 31, fr. I» para la personal ante un tribunal estatal")

print("\n2 · CADA DÍA DICE QUÉ ES")
dd = dg._dias_del_computo(C93)
ok(dd[d.date(2025, 12, 5)] == ("notificacion", ["notificación"]), "el 5 de diciembre: notificación")
ok(dd[d.date(2025, 12, 10)] == ("surte", ["surte efectos"]), "el 10: surte efectos")
ok(dd[d.date(2025, 12, 11)] == ("plazo", ["día 1"]), "el 11: día 1")
ok(dd[d.date(2025, 12, 17)] == ("inhabil", ["inhábil"]), "el 17 (vacaciones): inhábil, dicho")
ok(d.date(2025, 12, 20) not in dd, "el sábado no lleva palabra: se ve")
ok(dd[d.date(2026, 1, 19)] == ("presentacion", ["presentación", "día 15"]), "el 19 de enero: presentación, día 15")
de = dg._dias_del_computo(CEX)
ok(de[d.date(2026, 3, 25)] == ("plazo", ["día 15 · vence"]), "el vencimiento se rotula")
ok(de[d.date(2026, 3, 26)] == ("fuera", ["fuera +1"]) and de[d.date(2026, 3, 27)] == ("fuera", ["fuera +2"]),
   "los hábiles después del vencimiento, contados")
ok(de[d.date(2026, 3, 30)] == ("presentacion", ["presentación", "fuera +3"]), "la presentación tardía dice cuánto")
ok(dg._habiles_de_retraso(CEX) == 3 and dg._habiles_de_retraso(C93) == 0, "tres hábiles de retraso; cero si fue en tiempo")

print("\n3 · LA LEYENDA Y EL MAPA")
doc = docx.Document(); dg._pagina(doc)
dg.calendario_computo(doc, C93, "amparo_directo"); dg.mapa_computo(doc, C93, f0.fecha_en_letra, "amparo_directo")
t = texto_tablas(doc)
ok("Notificación y día en que surtió efectos." in t and "(art. 65 LFPCA)" in t, "la leyenda dice qué es el ámbar y por qué")
ok("Días hábiles del plazo, numerados del 1 al 15." in t and "(art. 18 LA)" in t, "qué son los días en gris azulado")
ok("Presentación de la demanda de amparo." in t and "El último día del plazo." in t, "qué es el negro")
ok("Sin color: días que no corren." in t and "16 dic 2025 a 1 ene 2026 (art. 19 LA)" in t, "y por qué los blancos no corren")
ok("NOTIFICACIÓN" in t and "SURTE EFECTOS" in t and "PLAZO" in t and "PRESENTACIÓN" in t, "las cuatro tarjetas del mapa")
ok("15 días hábiles · arts. 17 y 18 LA" in t, "la tarjeta del plazo junta sus fundamentos")
ok("EN TIEMPO · PRESENTADA EL ÚLTIMO DÍA DEL PLAZO" in t, "el veredicto del 93/2026")
ok("CÓMPUTO DEL PLAZO" not in t, "ya no la tabla de dos columnas")
doc2 = docx.Document(); dg._pagina(doc2)
dg.calendario_computo(doc2, CEX, "amparo_directo"); dg.mapa_computo(doc2, CEX, f0.fecha_en_letra, "amparo_directo")
t2 = texto_tablas(doc2)
ok("FUERA DE PLAZO · PRESENTADA 3 DÍAS HÁBILES DESPUÉS DEL VENCIMIENTO" in t2, "el veredicto tardío dice la distancia")
ok("3 días hábiles después del vencimiento" in t2 and "día(s)" not in t2, "sin «día(s) hábil(es)»")
ok("Días hábiles transcurridos después del vencimiento." in t2, "la leyenda explica el rosa sólo cuando lo hay")
ok("Días hábiles transcurridos" not in t, "y no lo explica cuando no lo hay")
barra = [tb for tb in doc2.tables if len(tb.columns) == 18]
ok(bool(barra), "la barra tardía tiene 15 del plazo + 2 fuera + la presentación")

print()
if FALLOS:
    print(f"FALLAN {len(FALLOS)}: " + " · ".join(FALLOS))
    raise SystemExit(1)
print("RESULTADO: TODAS LAS COMPROBACIONES PASAN")
