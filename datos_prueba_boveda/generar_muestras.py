"""Contratos FICTICIOS para probar la Bóveda. Nada de usuarios.

Genera N contratos (PDF con texto, Word y PDF escaneado = sólo imagen) y un
verdad.json con lo que la tabla de revisión debería encontrar en cada uno.
Uso: python generar_muestras.py <carpeta_salida> [n]
"""
import json
import random
import sys
from pathlib import Path

import docx
import fitz

random.seed(20261012)

EMPRESAS = [
    "Constructora Peña Alta, S.A. de C.V.", "Servicios Logísticos del Bajío, S.A. de C.V.",
    "Agroindustrias Sierra Gorda, S.P.R. de R.L.", "Textiles Cuauhtémoc, S.A.",
    "Inmobiliaria Los Arcos de Celaya, S.A. de C.V.", "Tecnologías Nopal, S.A.P.I. de C.V.",
    "Distribuidora Farmacéutica Lerma, S.A. de C.V.", "Transportes Huasteca Norte, S.A. de C.V.",
    "Alimentos Mesa Central, S. de R.L. de C.V.", "Consultores Jurídicos Corregidora, S.C.",
    "Metalúrgica El Pueblito, S.A. de C.V.", "Hotelera Bernal, S.A. de C.V.",
]
PERSONAS = [
    "Mariana Olvera Ruiz", "Joaquín Ledesma Prieto", "Rosalba Cuéllar Mendoza", "Ernesto Villagrán Soto",
    "Itzel Barrera Campos", "Humberto Zamudio Lara", "Patricia Anaya Rincón", "Gilberto Nieves Ortega",
    "Fernanda Quiroz Salas", "Rodrigo Esparza Medina", "Lucía Montes de Oca Ibarra", "Octavio Rangel Pliego",
]
CIUDADES = [
    ("Santiago de Querétaro, Querétaro", "los tribunales competentes de la ciudad de Santiago de Querétaro, Querétaro"),
    ("Ciudad de México", "los tribunales competentes de la Ciudad de México"),
    ("Guadalajara, Jalisco", "los tribunales del Primer Partido Judicial del Estado de Jalisco, con sede en Guadalajara"),
    ("Monterrey, Nuevo León", "los tribunales competentes de la ciudad de Monterrey, Nuevo León"),
    ("León, Guanajuato", "los juzgados del Partido Judicial de León, Guanajuato"),
    ("San Luis Potosí, San Luis Potosí", "los tribunales del fuero común de la capital del Estado de San Luis Potosí"),
]
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
TIPOS = ["arrendamiento", "prestación de servicios", "compraventa", "confidencialidad", "suministro"]

RELLENO = [
    "Las partes reconocen que en la celebración del presente contrato no existe error, dolo, mala fe, lesión ni cualquier otro vicio del consentimiento que pudiera invalidarlo.",
    "Ninguna de las partes podrá ceder ni transmitir, total o parcialmente, los derechos y obligaciones derivados de este instrumento sin el consentimiento previo y por escrito de la otra.",
    "Todas las notificaciones relacionadas con este contrato se harán por escrito en los domicilios señalados en las declaraciones, y surtirán efectos el día hábil siguiente a su recepción.",
    "El presente contrato constituye el acuerdo total entre las partes y deja sin efecto cualquier negociación, propuesta o convenio verbal o escrito anterior respecto de su objeto.",
    "Si alguna disposición de este contrato fuere declarada nula o inválida, las demás conservarán su plena vigencia y obligatoriedad, y las partes negociarán de buena fe una disposición sustituta.",
    "Ninguna de las partes será responsable por el incumplimiento derivado de caso fortuito o fuerza mayor, siempre que lo notifique a la otra dentro de los cinco días hábiles siguientes.",
    "Las partes se obligan a guardar estricta confidencialidad respecto de la información técnica, comercial y financiera que reciban con motivo del presente contrato.",
    "Los encabezados de las cláusulas se utilizan únicamente por conveniencia y no afectarán la interpretación de su contenido.",
]


def fecha(d, m, a):
    return f"{d} de {MESES[m - 1]} de {a}"


def contrato(i):
    tipo = TIPOS[i % len(TIPOS)]
    a, b = random.sample(EMPRESAS, 2)
    rep_a, rep_b = random.sample(PERSONAS, 2)
    ciudad, foro = random.choice(CIUDADES)
    d, m, anio = random.randint(1, 28), random.randint(1, 12), random.choice([2024, 2025, 2026])
    meses = random.choice([6, 12, 18, 24, 36])
    monto = random.choice([18500, 42000, 75000, 128000, 250000, 980000])
    pena_pct = random.choice([5, 10, 15, 20, 25])
    con_pena = random.random() > 0.25
    hojas = random.choice([2, 3, 4, 6, 9])  # páginas aproximadas
    renovacion = random.random() > 0.5

    roles = {
        "arrendamiento": ("EL ARRENDADOR", "EL ARRENDATARIO", f"el uso y goce temporal del inmueble ubicado en calle Andador {random.randint(1, 99)}, colonia Centro, {ciudad}"),
        "prestación de servicios": ("EL CLIENTE", "EL PRESTADOR", "servicios de consultoría, mantenimiento y soporte técnico"),
        "compraventa": ("EL VENDEDOR", "EL COMPRADOR", f"{random.randint(20, 400)} toneladas de materia prima conforme al anexo de especificaciones"),
        "confidencialidad": ("LA PARTE REVELADORA", "LA PARTE RECEPTORA", "la protección de la información confidencial intercambiada para evaluar una posible alianza comercial"),
        "suministro": ("EL PROVEEDOR", "EL DISTRIBUIDOR", "el suministro periódico de insumos para la línea de producción"),
    }[tipo]
    rol_a, rol_b, objeto = roles

    clausulas = [
        ("PRIMERA. OBJETO.", f"{rol_a} y {rol_b} convienen en que el objeto del presente contrato es {objeto}."),
        ("SEGUNDA. CONTRAPRESTACIÓN.", f"Como contraprestación, se pagará la cantidad de ${monto:,} M.N. (más el Impuesto al Valor Agregado), pagadera en mensualidades vencidas dentro de los primeros diez días naturales de cada mes."),
        ("TERCERA. VIGENCIA.", f"El presente contrato tendrá una vigencia de {meses} meses contados a partir del {fecha(d, m, anio)}."
         + (" Se renovará automáticamente por periodos iguales salvo aviso por escrito con treinta días de anticipación." if renovacion else " Concluida la vigencia, no se renovará de manera automática.")),
    ]
    if con_pena:
        clausulas.append(("CUARTA. PENA CONVENCIONAL.", f"En caso de incumplimiento de cualquiera de las obligaciones a su cargo, la parte incumplida pagará a la otra, por concepto de pena convencional, una cantidad equivalente al {pena_pct}% (por ciento) del monto total del contrato, sin perjuicio de exigir el cumplimiento forzoso."))
    else:
        clausulas.append(("CUARTA. INCUMPLIMIENTO.", "En caso de incumplimiento, la parte afectada podrá exigir el cumplimiento o la rescisión del contrato, con el pago de daños y perjuicios que acredite en juicio."))
    n = 5
    ordinales = ["QUINTA", "SEXTA", "SÉPTIMA", "OCTAVA", "NOVENA", "DÉCIMA", "DÉCIMA PRIMERA", "DÉCIMA SEGUNDA", "DÉCIMA TERCERA", "DÉCIMA CUARTA", "DÉCIMA QUINTA", "DÉCIMA SEXTA", "DÉCIMA SÉPTIMA", "DÉCIMA OCTAVA", "DÉCIMA NOVENA", "VIGÉSIMA", "VIGÉSIMA PRIMERA", "VIGÉSIMA SEGUNDA"]
    k = 0
    for _ in range(hojas * 2):
        texto = " ".join(random.sample(RELLENO, 3))
        clausulas.append((f"{ordinales[k]}. DISPOSICIONES GENERALES.", texto))
        k += 1
        if k >= len(ordinales) - 1:
            break
    clausulas.append((f"{ordinales[k]}. JURISDICCIÓN.", f"Para la interpretación y cumplimiento del presente contrato, las partes se someten expresamente a la jurisdicción de {foro}, renunciando a cualquier otro fuero que pudiera corresponderles por razón de sus domicilios presentes o futuros."))

    titulo = f"CONTRATO DE {tipo.upper()}"
    proemio = (f"{titulo} QUE CELEBRAN, POR UNA PARTE, {a.upper()}, REPRESENTADA EN ESTE ACTO POR {rep_a.upper()}, "
               f"A QUIEN EN LO SUCESIVO SE LE DENOMINARÁ «{rol_a}», Y POR LA OTRA, {b.upper()}, REPRESENTADA POR {rep_b.upper()}, "
               f"A QUIEN SE LE DENOMINARÁ «{rol_b}», AL TENOR DE LAS SIGUIENTES DECLARACIONES Y CLÁUSULAS.")
    cierre = f"Leído que fue el presente contrato y enteradas las partes de su contenido y alcance legal, lo firman por duplicado en {ciudad}, el {fecha(d, m, anio)}."
    verdad = {
        "tipo": tipo, "partes": [a, b], "representantes": [rep_a, rep_b],
        "fecha_firma": fecha(d, m, anio), "vigencia_meses": meses, "renovacion_automatica": renovacion,
        "monto": monto, "pena_convencional_pct": pena_pct if con_pena else None,
        "jurisdiccion": foro, "ciudad": ciudad,
    }
    return titulo, proemio, clausulas, cierre, verdad


def a_pdf(ruta, titulo, proemio, clausulas, cierre):
    doc = fitz.open()
    bloques = [("titulo", titulo), ("p", proemio)] + [(None, (c, t)) for c, t in clausulas] + [("p", cierre)]
    pagina = doc.new_page(width=612, height=792)
    y = 72

    def escribir(texto, tam=10.5, negrita=False):
        nonlocal pagina, y
        fuente = "hebo" if negrita else "helv"
        if y > 730:
            pagina = doc.new_page(width=612, height=792)
            y = 72
        rect = fitz.Rect(72, y, 540, 760)
        sobra = pagina.insert_textbox(rect, texto, fontsize=tam, fontname=fuente, align=3)
        if sobra < 0:  # no cupo: página nueva
            pagina = doc.new_page(width=612, height=792)
            y = 72
            rect = fitz.Rect(72, y, 540, 760)
            sobra = pagina.insert_textbox(rect, texto, fontsize=tam, fontname=fuente, align=3)
        y = 760 - sobra + 10

    for kind, cont in bloques:
        if kind == "titulo":
            escribir(cont, 13, True)
        elif kind == "p":
            escribir(cont)
        else:
            c, t = cont
            escribir(c, 10.5, True)
            escribir(t)
    doc.save(ruta)
    n = doc.page_count
    doc.close()
    return n


def a_docx(ruta, titulo, proemio, clausulas, cierre):
    d = docx.Document()
    d.add_heading(titulo, level=1)
    d.add_paragraph(proemio)
    for c, t in clausulas:
        p = d.add_paragraph()
        p.add_run(c + " ").bold = True
        p.add_run(t)
    d.add_paragraph(cierre)
    d.save(ruta)


def escanear(pdf_texto, ruta):
    """El mismo PDF pero sin capa de texto: cada página como imagen (fuerza OCR)."""
    src = fitz.open(pdf_texto)
    out = fitz.open()
    for p in src:
        pix = p.get_pixmap(dpi=150, colorspace=fitz.csGRAY)
        pag = out.new_page(width=p.rect.width, height=p.rect.height)
        pag.insert_image(pag.rect, pixmap=pix)
    out.save(ruta, deflate=True)
    out.close()
    src.close()


def main():
    salida = Path(sys.argv[1])
    n = int(sys.argv[2]) if len(sys.argv) > 2 else 60
    salida.mkdir(parents=True, exist_ok=True)
    verdades = {}
    for i in range(n):
        titulo, proemio, clausulas, cierre, verdad = contrato(i)
        base = f"{i + 1:03d}_{verdad['tipo'].replace(' ', '_').replace('ó', 'o')}"
        if i % 10 == 9:          # uno de cada diez, escaneado
            tmp = salida / f"_tmp_{base}.pdf"
            a_pdf(tmp, titulo, proemio, clausulas, cierre)
            nombre = f"{base}_escaneado.pdf"
            escanear(tmp, salida / nombre)
            tmp.unlink()
        elif i % 4 == 3:         # uno de cada cuatro, Word
            nombre = f"{base}.docx"
            a_docx(salida / nombre, titulo, proemio, clausulas, cierre)
        else:
            nombre = f"{base}.pdf"
            a_pdf(salida / nombre, titulo, proemio, clausulas, cierre)
        verdades[nombre] = verdad
    (salida / "verdad.json").write_text(json.dumps(verdades, ensure_ascii=False, indent=1))
    print(f"{n} contratos en {salida}")


if __name__ == "__main__":
    main()
