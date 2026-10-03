# -*- coding: utf-8 -*-
"""LOS CONSIDERANDOS Y EL RESOLUTIVO PROCESAL, POR TIPO (3-oct-2026), sin red.

    .venv/bin/python test_considerandos_tipo.py

David: «que los resultandos y considerandos de procedencia vayan impecables,
disminuyendo el margen de error si el secretario introduce el auto de admisión
y los datos correctos… aplicar, por ahora en toda la república, el código
federal de procedimientos civiles de manera supletoria a la ley de amparo,
salvo en los proyectos de ciudad de mexico donde ya opera el CNPCF».

Dos clases de arreglos, y se prueban por separado:
  · [siempre] —correcciones, también sin la bandera—: el supletorio por sede,
    el tribunal en versales en la competencia, «por el Jueza», «localizado» con
    la Sala, «de el auto», «por analogía» fuera de la revisión fiscal, la
    legitimación de la autoridad, el surtimiento del oficio (art. 31, fr. I),
    la procedencia del amparo en revisión (C1: ya sin su existencia), la
    competencia del auto que sobresee fuera de audiencia y la fecha del
    recurrido con una sola candidata.
  · [bandera `procedencia_por_tipo`] —diseño nuevo—: lo que trae
    `datos["procesal"]` manda sobre lo que se leía de la prosa (expediente,
    toca, fecha, inciso del 97, responsable…), el adhesivo con su considerando
    y su punto, los resolutivos de la revisión (firmeza, «en contra del», el
    ordinal), y la verja procesal al final con sus avisos primero.

SEGUNDA RONDA (3-oct-2026), secciones 11-19: (A) la legitimación de la persona
física en neutro y la fracción del 5o. en la queja; (B) «a la parte quejosa»;
(C) el resolutivo del amparo directo con los datos del V I S T O; (D) cada
inhábil con su fundamento; (E) el precepto del surtimiento; (F) la procedencia
de la revisión fiscal por la fracción de la ficha; (G) «el Titular» y las
versales; (H) el cierre y la dispensa de la queja; (I) la responsable
originaria del AR; (J) la antesala de los resolutivos en la verja.

CUARTA RONDA (3-oct-2026), secciones 32-39, contra la verificación final de las
32 sentencias reales: E1 un solo convertidor de nombres a prosa; E3 el cargo del
ponente; E4 «resolución definitiva» vs «que puso fin al juicio»; E5 el adhesivo
sin auto; E8 la unidad de la revisión fiscal con su artículo y sin valores
cortados; E10 el órgano descartado en competencia y existencia; E11 un hecho,
un aviso; E2 el plural del compositor en la legitimación y la carátula.

QUINTA RONDA (3-oct-2026), secciones 40-46, contra la segunda verificación
final: el plural de la carátula del AR es el de quien recurre (AR 208,
regresión); F3 la queja sin fracción del 97 no supone la I ni «indirecto»; los
efectos nombran lo reclamado (AD 349); F6 la figura del resolutivo por la puerta
única; F1 la forma de notificación que no consta no se afirma; la carátula de la
queja sin «con residencia en esta ciudad» (Q 337); F5 el artículo del papel en
los cargos de dos géneros.

SEXTA RONDA (3-oct-2026), secciones 47-52, lo que David respondió con su visto
bueno para todos los usuarios: C1 el amparo en revisión sin existencia y con su
procedencia (81, fr. I, con el inciso de lo recurrido); C2 la fracción del 35
en el banco; C5 el punto que confirma nombra acto y autoridad; C6 los asuntos
relacionados que marca el secretario (rubro y conexidad o hecho notorio); C3/C4
las carátulas de la queja y de la revisión fiscal; C7 el surtimiento agrario
por la sede. Las comprobaciones de §8, §9, §26, §37 y §45 que el contrato volvió
falsas (la existencia del AR, el renglón del órgano en la queja) se reescribieron
y dicen qué comprobaban antes; la de §14 (el agrario sin sede) también.

Las piezas de otros agentes (`resultandos_por_tipo`, `verja_procesal`) se
sustituyen por dobles con la firma del §6 de la especificación: aquí se prueba
el CABLEADO de `documento_generado`, no esas piezas. Al final, si están, una
pasada con las de verdad.
"""
import datetime as dt
import os
import sys
import tempfile
import types

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document

import banco
import contexto_taller as ct
import documento_generado as dg
import fase0_oportunidad as f0
import fase_origen as fo
import tipos_asunto as ta

FALLOS = []
H = dg.HUECO


def ok(c, q):
    print(f"   {'PASA ' if c else 'FALLA'}  {q}")
    if not c:
        FALLOS.append(q)


XXII = "Tercer Tribunal Colegiado en Materias Administrativa y Civil del Vigésimo Segundo Circuito"
YUC = "Primer Tribunal Colegiado en Materias Civil y Administrativa del Décimo Cuarto Circuito"
CDMX = "Décimo Tribunal Colegiado en Materia Civil del Primer Circuito"
SALA_QRO = "Primera Sala Civil del Tribunal Superior de Justicia del Estado de Querétaro"
_TMP = tempfile.mkdtemp(prefix="considerandos_tipo_")


def bandera(encendida: bool):
    ct.poner(True, {"banderas": {"procedencia_por_tipo": bool(encendida)}}, pruebas=True)


def componer(tipo, datos_extra, resultandos=None, calif=("infundado",), regla="personal",
             plazo=None, estudio=("El estudio de fondo.",), nombre="x"):
    """(texto del .docx, estructura) — los mismos datos base para todos."""
    plazo = plazo or {"amparo_directo": 15, "amparo_revision": 10, "queja": 5,
                      "revision_fiscal": 15}[tipo]
    pres = dt.date(2025, 3, 27) if tipo in ("queja", "amparo_revision") else dt.date(2025, 4, 8)
    datos = {"tipo_asunto": tipo, "numero": "1/2026", "magistrado": "Magistrado Ejemplo",
             "secretario": "Secretario Ejemplo", "tribunal": XXII, "ciudad": "Querétaro, Querétaro",
             "acto": "texto del acto", "antecedentes": "", "quejoso": "Juan Pérez López",
             "encabezado": "ASUNTO: 1/2026"}
    datos.update(datos_extra)
    c = f0.computar(dt.date(2025, 3, 20), pres, regla=regla, plazo=plazo,
                    responsable=datos.get("responsable"), tipo_asunto=tipo)
    est = dg.Estructura(apertura="", visto="para resolver el asunto 1/2026; y,",
                        resultandos=[dict(r) for r in (resultandos or [
                            {"titulo": "Trámite", "texto": "Texto del resultando."}])],
                        competencia="", existencia="", procedencia="")
    ruta = os.path.join(_TMP, f"{nombre}.docx")
    dg.componer(datos, est, c, f0.fecha_en_letra, ruta, estudio=list(estudio),
                calificaciones=list(calif), tipo_asunto=tipo, antecedentes=["Antecedente uno."])
    return "\n".join(p.text for p in Document(ruta).paragraphs if p.text.strip()), est


def entre(txt, desde, hasta):
    i = txt.find(desde)
    j = txt.find(hasta, i + 1) if i >= 0 else -1
    return txt[i:j] if (i >= 0 and j > i) else (txt[i:] if i >= 0 else "")


# ═══════════════════════════════════════════════════════════════════════════
print("\n1 · EL SUPLETORIO DE LA LEY DE AMPARO, POR SEDE [siempre]")
_q = ta.supletorio(XXII, "Querétaro, Querétaro")
_y = ta.supletorio(YUC, "Mérida, Yucatán")
_c = ta.supletorio(CDMX, "Ciudad de México")
ok(_q["documentales"] == "los artículos 129 y 202 del Código Federal de Procedimientos Civiles, de "
   "aplicación supletoria a la Ley de Amparo" and _y["documentales"] == _q["documentales"],
   "fuera de la Ciudad de México: la fórmula medida del CFPC (banco_formulas_medidas.json:64), sin el 2o.")
ok(_c["cdmx"] and "312, fracciones II y VIII, y 344 del Código Nacional" in _c["documentales"]
   and _c["documentales"].endswith("conforme a su artículo 2o."), "en la Ciudad de México: el CNPCF, con su 2o.")
ok(ta.es_cdmx("", "México, D.F.") and ta.es_cdmx("", "CDMX") and ta.es_cdmx(CDMX, "")
   and ta.es_cdmx(XXII, "") is False and ta.es_cdmx("", "") is None,
   "la sede se lee de la ciudad (CDMX, D.F.) o, sin ella, del circuito (Primer Circuito)")

print("\n2 · EL BANCO: VARIANTES Y EL CIRCUITO ESCRITO DE TODAS LAS MANERAS [siempre]")
_v, _f = banco.texto_de("amparo_directo", "existencia",
                        {"toca": "toca 374/2024", "expediente": "905/2023", "juzgado": "la Primera Sala",
                         "supletorio_documentales": _q["documentales"]}, variante="ad-c2-civil")
ok("los autos del toca 374/2024 y del expediente 905/2023" in _v and not _f and "la Sala responsable" not in _v,
   "la variante ad-c2-civil se lee por su id: toca y expediente por separado, el órgano por su nombre")
ok(banco.texto_de("amparo_directo", "existencia", {}, variante="no-existe")[0] == "",
   "una variante que no existe no se inventa: vacío")
ok(banco.fraccion_del_acuerdo("Tercer Tribunal Colegiado del XXII Circuito") == "XXII"
   and banco.fraccion_del_acuerdo("Tercer Tribunal Colegiado del 22o. Circuito") == "XXII"
   and banco.fraccion_del_acuerdo(XXII.upper()) == "XXII" and banco.fraccion_del_acuerdo(CDMX) == "I",
   "la fracción del AG 3/2013 con «XXII Circuito», «22o. Circuito», en versales y en letra")
ok(banco.fraccion_del_acuerdo("Primer Tribunal Colegiado del Centro Auxiliar de la Primera Región") == "",
   "un Centro Auxiliar no tiene fracción por circuito: hueco, no una fracción ajena")
_qc = banco.apartado("queja", "competencia")["plantilla"]
ok("por {juzgado}, {concordancia}" in _qc and "por el {responsable}" not in _qc,
   "la competencia de la queja ya no dice «por el {responsable}, localizado»")

print("\n3 · LA LEGITIMACIÓN [siempre]")
_l = ta.legitimacion_de("amparo_revision", "Director de Ingresos del Municipio de Querétaro", "", H, papel="autoridad")
ok("interpuesto por el Director de Ingresos" in _l and "legitimado para" in _l and "legitimada" not in _l,
   "la autoridad con su artículo y concordada por su cargo, no por el indicador de persona moral")
_l = ta.legitimacion_de("revision_fiscal", "Administración Desconcentrada Jurídica de Querétaro «1»",
                        autoridad_demandada="Administración Desconcentrada de Auditoría Fiscal de Querétaro «1»")
ok("por parte legítima" in _l and "63, párrafo primero" in _l
   and "lo hizo valer la Administración Desconcentrada Jurídica" in _l
   and "defensa jurídica de la Administración Desconcentrada de Auditoría Fiscal de Querétaro «1», autoridad "
       "demandada en el juicio de nulidad" in _l and "autoridad facultada" not in _l,
   "revisión fiscal: la unidad de defensa jurídica DE la autoridad demandada (art. 63, párrafo primero, LFPCA)")
ok("defensa jurídica de la autoridad demandada en el juicio de nulidad" in
   ta.legitimacion_de("revision_fiscal", "Administración Desconcentrada Jurídica de Yucatán «1»"),
   "sin la autoridad demandada en la ficha: «la autoridad demandada», sin inventar cuál")
_l = ta.legitimacion_de("amparo_directo", "Juan Pérez López", "María Ruiz", figura="apoderada")
ok("la parte quejosa está legitimada" in _l and "legitimado" not in _l,
   "con representante, el adjetivo concuerda con «la parte quejosa», no con el nombre")
ok(ta.legitimacion_de("amparo_directo", "Comercializadora Ejemplo, S.A. de C.V.").startswith(
   "La demanda de amparo fue promovida por Comercializadora"), "una sociedad no lleva artículo")

print("\n4 · LA COMPETENCIA DE LA REVISIÓN, POR LO QUE SE RECURRE [siempre]")
_cr, _ = banco.texto_de("amparo_revision", "competencia",
                        {"materia": "administrativa", "juez_distrito": "el Juzgado Quinto", "concordancia": "localizado"})
_d, _a = ta.competencia_revision(_cr, "Querétaro, a diez de marzo. Visto el estado de los autos, se sobresee "
                                      "en el juicio fuera de la audiencia constitucional.")
ok("81, fracción I, inciso d), y 84" in _d and "35, fracción II y 210" in _d and "fuera de la audiencia" in _d and _a,
   "el auto que sobresee fuera de la audiencia: 81-I-d) y 35-II (no el e) de la sentencia de audiencia)")
_d, _ = ta.competencia_revision(_cr, "VISTOS para resolver los autos del juicio de amparo 55/2026; audiencia "
                                     "constitucional… se sobresee")
ok("inciso e) y 84" in _d, "la sentencia de audiencia que sobresee sigue en el inciso e)")
_d, _ = ta.competencia_revision(_cr, "", clase="auto_sobreseimiento")
ok("inciso d)" in _d, "la clase que da la ficha manda sobre el proemio")

print("\n5 · EL SURTIMIENTO CUANDO RECURRE UNA AUTORIDAD (art. 31) [siempre]")
_ofi = f0.computar(dt.date(2025, 3, 20), dt.date(2025, 3, 27), regla="oficio", plazo=10,
                   tipo_asunto="amparo_revision")
_per = f0.computar(dt.date(2025, 3, 20), dt.date(2025, 3, 27), regla="personal", plazo=10,
                   tipo_asunto="amparo_revision")
ok(_ofi.surtio == dt.date(2025, 3, 20) and _per.surtio == dt.date(2025, 3, 24),
   "el oficio surte el mismo día (fr. I); la personal, al día hábil siguiente (fr. II)")
ok(f0.fundamento_de_surtimiento(_ofi.regla, "amparo_revision") == "artículo 31, fracción I, de la Ley de Amparo"
   and f0.fundamento_de_surtimiento(f0.REGLAS_SURTE["electronica"], "queja") == "artículo 31, fracción III, de la Ley de Amparo"
   and f0.fundamento_de_surtimiento(_per.regla, "amparo_revision") == "artículo 31, fracción II, de la Ley de Amparo"
   and f0.fundamento_de_surtimiento(_per.regla, "amparo_revision", "autoridad") == "",
   "la fracción que corresponde: I oficio, II particulares, III electrónica; la II nunca para la autoridad")
_p = f0.parrafo_oportunidad(_ofi, "", "amparo_revision", papel="autoridad")
ok("por oficio y surtió efectos el mismo día, conforme al artículo 31, fracción I" in _p,
   "el considerando lo dice con la regla de la autoridad")
_p = f0.parrafo_oportunidad(_per, "", "amparo_revision", papel="autoridad")
ok("conforme al *********" in _p and "fracción II" not in _p
   and "RECURRE UNA AUTORIDAD" in f0.aviso_fundamento(_per, "amparo_revision", "autoridad"),
   "autoridad con la regla de los particulares: el precepto en hueco y un aviso que lo nombra")
_r = f0.reglas_para("amparo_revision", "Juzgado Quinto de Distrito", papel="autoridad")
ok(_r["por_omision"] == "oficio" and {"oficio", "electronica", "personal", "lista", "otra"} <= {x["clave"] for x in _r["reglas"]}
   and f0.reglas_para("queja", "Juzgado Quinto de Distrito")["por_omision"] == "personal"
   and f0.reglas_para("amparo_directo", "Sala Regional del TFJA")["por_omision"] == "lfpca_boletin",
   "el desplegable: en los recursos, las tres fracciones del 31 y, si recurre la autoridad, el oficio por omisión")

print("\n6 · LA FECHA DE LO RECURRIDO, CON UNA SOLA CANDIDATA [siempre]")
ok(fo.fecha_de("Seguido el juicio, dictó sentencia el trece de marzo de dos mil veinticinco. Por auto de "
               "Presidencia de veintidós de abril de dos mil veinticinco se admitió el recurso.") == "",
   "RF_yucatan: la única fecha era la del auto de Presidencia y ya no se toma por la recurrida")
ok(fo.fecha_de("contra el auto de seis de marzo de dos mil veintiséis, dictado por la Jueza")
   == "seis de marzo de dos mil veintiséis", "con la marca de lo recurrido, se toma")

print("\n7 · EL 97, FRACCIÓN II")
ok(ta.cola_97("b", "II").startswith("en el que la autoridad responsable proveyó sobre la suspensión")
   and ta.cola_97("b") == ta.COLA_97["b"] and ta.cola_97("b)", "I") == ta.COLA_97["b"],
   "la cola de la fracción II es la suya; la de la I, la de siempre")

# ═══════════════════════════════════════════════════════════════════════════
print("\n8 · EL CAMINO DE SIEMPRE, CON SUS ARREGLOS [siempre] (bandera apagada)")
bandera(False)
_llamadas = {"verja": 0}
_verja_doble = types.ModuleType("verja_procesal")
_textos_verja = [""]
_verja_doble.revisar = lambda texto, ficha, datos, tipo: (_llamadas.__setitem__("verja", _llamadas["verja"] + 1)
                                                          or _textos_verja.append(texto)
                                                          or ["AVISO DE LA VERJA"])
_verja_doble.arreglar = lambda t: (t, [])
_reales = {k: sys.modules.get(k) for k in ("verja_procesal", "resultandos_por_tipo")}
sys.modules["verja_procesal"] = _verja_doble

_ad, _ = componer("amparo_directo", {"tribunal": XXII.upper(), "responsable": "la Jueza Primero de Oralidad "
                  "Mercantil del Distrito Judicial de Querétaro", "materia": "mercantil",
                  "magistrado": "Magistrada Ejemplo", "procesal": {"expediente": "999/2099"}},
                  resultandos=[{"titulo": "Presentación", "texto": "dictada en el toca civil 374/2024, que "
                                "confirmó la del expediente 905/2023."}], nombre="ad_viejo")
_comp = entre(_ad, "PRIMERO. Competencia.", "SEGUNDO.")
ok("Este Tercer Tribunal Colegiado en Materias Administrativa y Civil del Vigésimo Segundo Circuito es" in _comp,
   "la competencia nombra al tribunal como nombre, no en versales (AD_tribunal_versales)")
ok("artículos 129 y 202 del Código Federal de Procedimientos Civiles" in _ad and "Código Nacional" not in _ad,
   "XXII Circuito: la existencia con el CFPC, y el CNPCF en ninguna parte")
ok("MAGISTRADA PONENTE: MAGISTRADA EJEMPLO" in _ad, "la carátula concuerda con lo que el secretario escribió")
ok("Sustenta esa consideración la jurisprudencia 2a./J. 58/2010" in _ad and "por analogía" not in _ad,
   "amparo directo: la 2a./J. 58/2010 se aplica directamente, sin «por analogía»")
ok("999/2099" not in _ad and _llamadas["verja"] == 0,
   "sin la bandera, `procesal` no manda y la verja no corre: el camino de hoy")
_ad_cdmx, _ = componer("amparo_directo", {"tribunal": CDMX, "ciudad": "Ciudad de México",
                       "responsable": "Tercera Sala Civil del Tribunal Superior de Justicia de la Ciudad de México"},
                       nombre="ad_cdmx")
ok("312, fracciones II y VIII, y 344 del Código Nacional de Procedimientos Civiles y Familiares" in _ad_cdmx
   and "Código Federal de Procedimientos Civiles" not in _ad_cdmx,
   "Ciudad de México: la existencia con el CNPCF, y el CFPC en ninguna parte")

_qj, _ = componer("queja", {"responsable": "Jueza Tercero de Distrito en el Estado de Querétaro"},
                  resultandos=[{"titulo": "Interposición", "texto": "contra el auto de trece de marzo de dos "
                                "mil veinticinco, dictado por la Jueza Tercero de Distrito en el juicio de "
                                "amparo 742/2025-II, mediante el cual desechó de plano la demanda de amparo."}],
                  nombre="q_viejo")
ok("por la Jueza Tercero de Distrito en el Estado de Querétaro, localizada en la circunscripción" in _qj
   and "por el Jueza" not in _qj, "queja: «por la Jueza…, localizada» (no «por el Jueza…, localizado»)")
ok("algún apartado del auto recurrido" in _qj and "de el auto" not in _qj, "dispensa: «del auto recurrido»")

_rf, _ = componer("revision_fiscal", {"quejoso": "Administración Desconcentrada Jurídica de Querétaro «1»",
                  "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"},
                  regla="lfpca_boletin", nombre="rf_viejo")
ok("Tribunal Federal de Justicia Administrativa, localizada en la circunscripción" in _rf,
   "revisión fiscal: «la Sala Regional…, localizada»")
ok("por analogía, la jurisprudencia 2a./J. 58/2010" in _rf, "revisión fiscal: aquí sí «por analogía»")
ok("lo hizo valer la Administración Desconcentrada Jurídica de Querétaro «1», en su carácter de unidad "
   "administrativa encargada de la defensa jurídica" in _rf, "revisión fiscal: la legitimación del corpus")

_ar, _ = componer("amparo_revision", {"quejoso": "Comercializadora Ejemplo, S.A. de C.V.",
                  "recurrente": "Director de Ingresos del Municipio de Querétaro", "papel_recurrente": "autoridad",
                  "responsable": "Director de Ingresos del Municipio de Querétaro",
                  "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                  "materia": "administrativa", "resolvio_a_quo": "concede"},
                  resultandos=[{"titulo": "Trámite", "texto": "el Juzgado Quinto de Distrito registró la "
                                "demanda bajo el número 950/2024 y concedió el amparo."}],
                  regla="oficio", nombre="ar_viejo")
# C1 (3-oct-2026) [siempre]: la revisión ya no reproduce la existencia (David:
# «ya viene en la sentencia recurrida»). Antes se comprobaba aquí «SEGUNDO.
# Existencia de la resolución recurrida.» con los autos que remitió el juzgado
# (art. 89); ahora, en ese sitio, la procedencia del recurso.
_ex = entre(_ar, "SEGUNDO. Procedencia.", "TERCERO.")
ok(_ex and "El presente recurso de revisión es procedente, de conformidad con el artículo 81, fracción I, "
   "inciso e), de la Ley de Amparo, en razón de que se impugna una sentencia dictada en la audiencia "
   "constitucional." in _ex and "Existencia" not in _ar and "informe justificado" not in _ar,
   "amparo en revisión (sin bandera): «SEGUNDO. Procedencia.» con el 81-I-e), y sin existencia (C1)")
ok("interpuesto por el Director de Ingresos" in _ar and "por oficio y surtió efectos el mismo día, conforme al "
   "artículo 31, fracción I" in _ar, "amparo en revisión: la autoridad con su artículo y su regla de surtimiento")

_pe = dg.prompt_estructura({"tipo_asunto": "amparo_directo"})
ok('{{"titulo"' not in _pe and '{"titulo"' in _pe and "omitió formular pedimento" not in _pe
   and "ficialía" not in _pe and "obran en autos»" not in _pe,
   "el prompt de la estructura (camino viejo): sin llaves dobles, sin el pedimento por defecto ni la oficialía")

# ═══════════════════════════════════════════════════════════════════════════
print("\n9 · CON LA BANDERA: LA FICHA MANDA SOBRE LA PROSA (con dobles del §6)")
bandera(True)
_adh_llamadas = []
_rpt_doble = types.ModuleType("resultandos_por_tipo")
_rpt_doble.considerando_adhesivo = lambda tipo, ficha, datos: (
    "Legitimación y oportunidad del amparo adhesivo.", "TEXTO DEL CONSIDERANDO ADHESIVO.", ["AVISO DEL ADHESIVO"])


def _res_adh(tipo, prospera, ficha):
    _adh_llamadas.append(prospera)
    return ("Se declara sin materia el amparo adhesivo promovido por María Gómez Ruiz.", "")


_rpt_doble.resolutivo_adhesivo = _res_adh
_rpt_doble.componer = lambda tipo, ficha, datos: {"visto": "", "resultandos": [], "avisos": [], "datos_extra": {}}
sys.modules["resultandos_por_tipo"] = _rpt_doble
_cambios_verja = {"n": 0}


def _arreglar_doble(t):
    _cambios_verja["n"] += 1
    return (t.replace("TOCA_MAL", "toca"), ["x"] if "TOCA_MAL" in t else [])


_verja_doble.arreglar = _arreglar_doble
_llamadas["verja"] = 0

_PROSA_ENGAÑOSA = [{"titulo": "Presentación de la demanda de amparo", "texto":
                    "contra la sentencia de primero de enero de dos mil veinte, dictada en el expediente "
                    "111/2020."}]
_pro_ad = {"expediente": "905/2023", "toca": "374/2024", "toca_en_prosa": "toca 374/2024",
           "fecha_acto": "trece de marzo de dos mil veinticinco", "fecha_acto_iso": "2025-03-13",
           "responsable": SALA_QRO, "organo_acto": SALA_QRO, "materia": "civil",
           "adhesivo": {"quien": "María Gómez Ruiz", "admision": "2025-05-05"}}
_ficha_ad = {"tipo": "amparo_directo", "adhesivo": dict(_pro_ad["adhesivo"]), "responsable": SALA_QRO}
_ad, _est = componer("amparo_directo", {"responsable": "Primera Sala Civil del Estado",
                     "procesal": _pro_ad, "tramite": _ficha_ad}, resultandos=_PROSA_ENGAÑOSA, nombre="ad_ficha")
_exi = entre(_ad, "SEGUNDO. Existencia del acto reclamado.", "TERCERO.")
ok("con los autos del toca 374/2024 y del expediente 905/2023" in _exi and "111/2020" not in _ad.split("C O N S I D E R A N D O")[1],
   "existencia AD: toca y expediente de la ficha (variante ad-c2-civil), nunca el número de la prosa")
ok("AUTORIDAD RESPONSABLE: " + SALA_QRO.upper() in _ad and f"dictada por la {SALA_QRO}, localizada" in _ad
   and f"rendido por la {SALA_QRO}" in _exi
   and f"contra la sentencia dictada el trece de marzo de dos mil veinticinco, por la {SALA_QRO}, en el "
       f"toca 374/2024, derivado del expediente 905/2023." in _ad
   and "Primera Sala Civil del Estado" not in _ad,
   "la responsable de la ficha, la misma en carátula, competencia, existencia y resolutivo")
ok("precisada en el primer resultando" not in _ad.split("R E S U E L V E")[1],
   "(C) el resolutivo del amparo directo identifica la sentencia con los datos del V I S T O, no «precisada…»")
ok("Por lo expuesto y fundado" in _textos_verja[-1] and "R E S U E L V E" in _textos_verja[-1],
   "(J) a la verja se le pasa la antesala de los resolutivos («Por lo expuesto y fundado…»)")
_cons = _ad.split("C O N S I D E R A N D O")[1]
ok(_cons.find("Legitimación y oportunidad del amparo adhesivo.") < _cons.find("Acto reclamado y conceptos de violación.")
   and "TEXTO DEL CONSIDERANDO ADHESIVO." in _cons and "CUARTO. Legitimación y oportunidad del amparo adhesivo." in _cons,
   "el considerando del adhesivo va antes de la dispensa, con su ordinal")
ok("PRIMERO. La Justicia de la Unión no ampara ni protege" in _ad
   and "SEGUNDO. Se declara sin materia el amparo adhesivo promovido por María Gómez Ruiz." in _ad
   and "ÚNICO." not in _ad and _adh_llamadas[-1] is False,
   "el resolutivo: PRIMERO el amparo y SEGUNDO el adhesivo (se le dice que el principal no prospera)")
ok(_est.avisos[:1] == ["AVISO DE LA VERJA"] and _llamadas["verja"] == 1 and _cambios_verja["n"] > 0
   and "AVISO DEL ADHESIVO" in _est.avisos,
   "la verja corre una vez, su arreglo pasa por el bloque procesal y sus avisos van PRIMERO")
# D1 (tercera ronda): sin el precepto, la cláusula se OMITE —como en los engroses de la
# ponencia de David— y un aviso lo RECOMIENDA; ni hueco ni «conforme a la ley del acto».
ok("surtió efectos al día hábil siguiente, es decir" in _ad and "conforme al *********" not in _ad
   and "ley del acto" not in _ad and any("PRECEPTO DEL SURTIMIENTO" in a for a in _est.avisos),
   "sin el artículo de la ley del acto (D1): la cláusula se omite y un aviso recomienda el precepto")

_ad2, _est2 = componer("amparo_directo", {"responsable": SALA_QRO,
                       "procesal": {"fecha_acto": "trece de marzo de dos mil veinticinco", "materia": "civil",
                                    "responsable": SALA_QRO}}, resultandos=_PROSA_ENGAÑOSA, nombre="ad_sin_exp")
ok("los autos del expediente *********" in _ad2 and "111/2020" not in _ad2.split("C O N S I D E R A N D O")[1]
   and any("EXISTENCIA" in a and "número del expediente" in a for a in _est2.avisos),
   "la ficha sin expediente: hueco con el aviso que nombra el dato, y la prosa no se lee")
_sin, _ = componer("amparo_directo", {"responsable": SALA_QRO}, resultandos=_PROSA_ENGAÑOSA, nombre="ad_sin_pro")
bandera(False)
_sin_b, _ = componer("amparo_directo", {"responsable": SALA_QRO}, resultandos=_PROSA_ENGAÑOSA, nombre="ad_sin_pro_b")
bandera(True)
_norm = lambda t: t.replace(", conforme a la ley del acto", "")
ok(_sin == _norm(_sin_b), "con la bandera y SIN `procesal`, el mismo documento de hoy (salvo la cláusula del "
                         "precepto, D1)")

_rf, _est = componer("revision_fiscal", {"quejoso": "Administración Desconcentrada Jurídica de Querétaro «1»",
                     "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa",
                     "fecha_origen": "uno de enero de dos mil veinte",
                     "procesal": {"expediente_tfja": "695/25-09-01-7-OT", "sala": "Sala Regional del Centro II del "
                                  "Tribunal Federal de Justicia Administrativa",
                                  "fecha_acto": "trece de marzo de dos mil veinticinco", "fecha_acto_iso": "2025-03-13",
                                  "autoridad_demandada": "Administración Desconcentrada de Auditoría Fiscal de Querétaro «1»"}},
                     resultandos=[{"titulo": "Trámite", "texto": "Por auto de Presidencia de veintidós de abril de "
                                   "dos mil veinticinco se admitió el recurso."}], regla="lfpca_boletin", nombre="rf_ficha")
ok("Se confirma la sentencia de trece de marzo de dos mil veinticinco, dictada en el expediente 695/25-09-01-7-OT"
   in _rf and "veintidós de abril" not in _rf.split("R E S U E L V E")[1],
   "revisión fiscal: fecha y expediente del resolutivo, de la ficha (no del auto de Presidencia ni del PDF)")
ok("defensa jurídica de la Administración Desconcentrada de Auditoría Fiscal de Querétaro «1»" in _rf,
   "revisión fiscal: la autoridad demandada de la ficha en la legitimación")
_rf2, _est2 = componer("revision_fiscal", {"responsable": "Sala Regional Peninsular del Tribunal Federal de Justicia "
                       "Administrativa", "procesal": {"sala": "Sala Regional Peninsular del Tribunal Federal de "
                                                       "Justicia Administrativa"}},
                       regla="lfpca_boletin", nombre="rf_sin_fecha")
ok("Se confirma la sentencia de *********, dictada en el expediente *********" in _rf2
   and any("falta su FECHA" in a and "EXPEDIENTE del juicio de nulidad" in a for a in _est2.avisos),
   "revisión fiscal sin fecha ni expediente en la ficha: huecos y el aviso que los nombra")

_qb, _estq = componer("queja", {"responsable": "Juzgado Décimo de Distrito en Materia Civil en la Ciudad de México",
                      "tribunal": CDMX, "ciudad": "Ciudad de México",
                      "procesal": {"fraccion_97": "I", "inciso_97": "b", "cola_97": "en el que se negó la suspensión provisional",
                                   "juicio_amparo": "742/2025-II", "fecha_acto": "trece de marzo de dos mil veinticinco",
                                   "juzgado": "Juzgado Décimo de Distrito en Materia Civil en la Ciudad de México",
                                   "descripcion_acto": "del auto que negó la suspensión provisional"}},
                      resultandos=[{"titulo": "Interposición", "texto": "contra el auto que desechó la demanda."}],
                      nombre="q_b")
ok("97, fracción I, inciso b), de la Ley de Amparo; 35" in _qb and "inciso b) de la Ley" not in _qb
   and "97, fracción I, inciso b), de la Ley de Amparo, en razón" in _qb
   and "en el que se negó la suspensión provisional." in _qb and "inciso a)" not in _qb,
   "queja: el inciso y la cola de la ficha, una sola vez para competencia y procedencia (la prosa decía «desechó»)")
ok(any("su plazo es de DOS" in a for a in _estq.avisos), "inciso b) de la fracción I con plazo de cinco días: aviso")
_q2, _estq2 = componer("queja", {"responsable": SALA_QRO,
                       "procesal": {"fraccion_97": "II", "inciso_97": "b", "juicio_amparo": "380/2025",
                                    "fecha_acto": "trece de marzo de dos mil veinticinco", "juzgado": SALA_QRO,
                                    "cola_97": "en el que negó la suspensión del acto reclamado",
                                    "descripcion_acto": "del auto que negó la suspensión del acto reclamado"}},
                       calif=("fundado",), nombre="q_ii")
ok("97, fracción II, inciso b), de la Ley de Amparo; 35" in _q2 and "autoridad responsable en el juicio de amparo "
   "directo 380/2025, radicado en este órgano colegiado" in _q2 and "97, fracción II, inciso b), de la Ley" in _q2
   and "fracción I," not in entre(_q2, "PRIMERO. Competencia.", "TERCERO."),
   "queja de la fracción II: su propia cadena en la competencia y en la procedencia")
_q3, _estq3 = componer("queja", {"responsable": "Jueza Tercero de Distrito en el Estado de Querétaro",
                       "procesal": {"fraccion_97": "I", "juicio_amparo": "742/2025-II"}}, nombre="q_sin_inciso")
ok("inciso *********)" in _q3 and any("FRACCIÓN I, SALE EN HUECO" in a for a in _estq3.avisos),
   "queja sin inciso en la ficha: hueco y aviso, no la deducción de la prosa")

_RES_JUZ = ("La Justicia de la Unión ampara y protege a Comercializadora Ejemplo, S.A. de C.V., en contra el "
            "acto reclamado al Director de Ingresos del Municipio de Querétaro, para los efectos precisados en "
            "el último considerando de esta resolución.")
_ar1, _esta = componer("amparo_revision", {"quejoso": "Comercializadora Ejemplo, S.A. de C.V.",
                       "recurrente": "Director de Ingresos del Municipio de Querétaro", "papel_recurrente": "autoridad",
                       "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                       "resolvio_a_quo": "sobresee_concede", "resolutivo_recurrida": _RES_JUZ,
                       "reasuncion": {"quien_recurre": "autoridad", "sobresee_ademas": True, "adhesiva": False},
                       "procesal": {"juicio_amparo": "950/2024", "juzgado": "Juzgado Quinto de Distrito en el Estado "
                                    "de Querétaro", "materia": "administrativa", "clase_recurrida": "sentencia"}},
                       regla="oficio", nombre="ar_firme")
_rs = _ar1.split("R E S U E L V E")[1]
ok("PRIMERO. Queda firme el sobreseimiento decretado en la sentencia recurrida." in _rs
   and "SEGUNDO. En la materia de la revisión, se confirma la sentencia recurrida." in _rs
   and "Se sobresee" not in _rs,
   "revisión: queda firme lo que nadie recurrió y la confirmación se acota a la materia de la revisión")
_ar2, _ = componer("amparo_revision", {"quejoso": "Comercializadora Ejemplo, S.A. de C.V.",
                   "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                   "resolvio_a_quo": "concede", "resolutivo_recurrida": _RES_JUZ,
                   "procesal": {"juicio_amparo": "950/2024", "juzgado": "Juzgado Quinto de Distrito en el "
                                "Estado de Querétaro"}}, nombre="ar_contra")
_rs2 = _ar2.split("R E S U E L V E")[1]
ok("en contra del acto reclamado" in _rs2 and "en contra el" not in _rs2
   and "el último considerando de esta resolución" in _rs2,
   "revisión: «en contra el acto» del resolutivo reproducido se contrae, y su remisión (la del juzgado) no se toca")
_ar3, _ = componer("amparo_revision", {"quejoso": "Juan Pérez López",
                   "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                   "resolvio_a_quo": "sobresee", "procesal": {"juicio_amparo": "950/2024"}},
                   calif=("fundado",), estudio=("Es fundado el agravio; procede revocar el sobreseimiento y "
                                                "conceder el amparo.",), nombre="ar_ordinal")
_rs3 = _ar3.split("R E S U E L V E")[1]
_ult = [l for l in _ar3.splitlines() if l.startswith(("QUINTO. Estudio", "SEXTO. Estudio", "SÉPTIMO. Estudio"))]
ok(_ult and f"el considerando {_ult[0].split('.')[0].lower()} de esta ejecutoria" in _rs3
   and "último considerando de esta ejecutoria" not in _rs3,
   f"revisión: los puntos remiten al considerando de esta ejecutoria por su ordinal ({_ult[0][:14] if _ult else '—'})")
_ar4, _ = componer("amparo_revision", {"quejoso": "Juan Pérez López",
                   "organo_recurrido": "Jueza Primera de Distrito en el Estado de Yucatán", "resolvio_a_quo": "niega",
                   "procesal": {"juicio_amparo": "77/2025", "juzgado": "Jueza Primera de Distrito en el Estado de Yucatán",
                                "clase_recurrida": "auto_sobreseimiento"}}, nombre="ar_auto")
# C1 (3-oct-2026): la clase ya no se comprueba en la existencia —que la
# revisión no reproduce— sino en la procedencia, con el inciso d) del 81.
ok("de conformidad con el artículo 81, fracción I, inciso d), de la Ley de Amparo, en razón de que se impugna "
   "el auto que sobreseyó en el juicio fuera de la audiencia constitucional." in _ar4
   and "81, fracción I, inciso d), y 84" in _ar4 and "Existencia" not in _ar4
   and "por la Jueza Primera de Distrito en el Estado de Yucatán, localizada" in _ar4,
   "revisión de un auto de sobreseimiento: la clase de la ficha en la procedencia y en la competencia (inciso d)")
ok("Se confirma el auto recurrido." in _ar4 and "el auto recurrido se notificó" in _ar4
   and "Se confirma la sentencia recurrida" not in _ar4,
   "y el auto se llama auto en la oportunidad y en el resolutivo, no «sentencia»")
_ar5, _ = componer("amparo_revision", {"quejoso": "Juan Pérez López",
                   "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro", "resolvio_a_quo": "niega",
                   "procesal": {"juicio_amparo": "950/2024", "juzgado": "Juzgado Quinto de Distrito en el Estado "
                                "de Querétaro", "clase_recurrida": "interlocutoria_suspension"}}, nombre="ar_incidente")
_rs5 = _ar5.split("R E S U E L V E")[1]
ok("PRIMERO. Se confirma la interlocutoria recurrida." in _rs5
   and "SEGUNDO. Se niega la suspensión definitiva solicitada por Juan Pérez López" in _rs5
   and "ampara" not in _rs5 and "81, fracción I, inciso a), de la Ley de Amparo, en razón de que se impugna la "
   "interlocutoria que resolvió sobre la suspensión definitiva." in _ar5
   and "81, fracción I, inciso a), y 84" in _ar5 and "Existencia" not in _ar5,
   "incidente en revisión: la interlocutoria, la suspensión definitiva (no el amparo) y el 81-I-a) en la "
   "competencia y en la procedencia (C1: sin existencia)")

for k, v in _reales.items():
    if v is None:
        sys.modules.pop(k, None)
    else:
        sys.modules[k] = v

# ═══════════════════════════════════════════════════════════════════════════
print("\n10 · LA PASADA CON LAS PIEZAS DE VERDAD (si ya están)")
try:
    import importlib
    rpt = importlib.import_module("resultandos_por_tipo")
    importlib.import_module("verja_procesal")
    _hay = hasattr(rpt, "componer") and hasattr(rpt, "considerando_adhesivo")
except Exception:
    _hay = False
if not _hay:
    print("   (resultandos_por_tipo / verja_procesal no están: se omite)")
else:
    _ficha = {"formato": 1, "tipo": "amparo_directo", "numero": "1/2026", "materia": "administrativa",
              "sede": {"tribunal": YUC, "ciudad": "Mérida, Yucatán", "circuito": "", "cdmx": False},
              "admision": {"fecha": "2025-04-22"}, "turno": {"fecha": "2025-05-06", "ponente": "Magistrado Ejemplo"},
              "ministerio_publico": "", "promovente": "Juan Pérez López", "caracter": "quejoso",
              "presentacion": "2025-04-08", "via_presentacion": "responsable", "notificacion": "2025-03-20",
              "acto": {"clase": "sentencia", "fecha": "2025-03-13",
                       "organo": "Sala Regional Peninsular del Tribunal Federal de Justicia Administrativa",
                       "expediente": "1234/24-16-01-3"},
              "responsable": "Sala Regional Peninsular del Tribunal Federal de Justicia Administrativa",
              "terceros": ["Administración Desconcentrada Jurídica de Yucatán «1»"], "derechos": ["14", "16"],
              "avisos": [], "fuentes": {}, "fechas_imposibles": []}
    _datos = {"tipo_asunto": "amparo_directo", "numero": "1/2026", "quejoso": "Juan Pérez López",
              "responsable": "Sala Regional Peninsular del Tribunal Federal de Justicia Administrativa",
              "materia": "administrativa", "tribunal": YUC, "ciudad": "Mérida, Yucatán",
              "magistrado": "Magistrado Ejemplo", "secretario": "S", "encabezado": "AMPARO DIRECTO ADMINISTRATIVO: 1/2026"}
    _comp = rpt.componer("amparo_directo", _ficha, _datos)
    _datos.update({"procesal": _comp["datos_extra"], "tramite": _ficha})
    _c = f0.computar(dt.date(2025, 3, 20), dt.date(2025, 4, 8), regla="lfpca_boletin", plazo=15,
                     responsable=_datos["responsable"], tipo_asunto="amparo_directo")
    _e = dg.Estructura(visto=_comp["visto"], resultandos=_comp["resultandos"])
    _r = os.path.join(_TMP, "ad_real.docx")
    dg.componer(_datos, _e, _c, f0.fecha_en_letra, _r, estudio=["El estudio."], calificaciones=["infundado"],
                tipo_asunto="amparo_directo")
    _t = "\n".join(p.text for p in Document(_r).paragraphs if p.text.strip())
    _proc = _t.split("Antecedentes.")[0] if "Antecedentes." in _t else _t.split("Estudio.")[0]
    _proc_y_res = _proc + _t.split("R E S U E L V E")[1]
    _sin_sesion = _proc_y_res.replace(f"listó el {H}", "").replace(f"sesión ordinaria de {H}", "") \
                             .replace(f"correspondiente a la sesión de {H}", "")
    ok(H not in _sin_sesion,
       "amparo directo con datos completos y las piezas reales: CERO huecos fuera de la sesión")
    ok("autos del expediente 1234/24-16-01-3" in _t and "artículos 129 y 202 del Código Federal" in _t,
       "y la existencia con el expediente de la ficha y el CFPC (Yucatán)")
    ok(_t.count("Sala Regional Peninsular del Tribunal Federal de Justicia Administrativa") >= 4,
       "la responsable de la ficha escrita igual en visto, resultandos, competencia, existencia y resolutivo")

# ═══════════════════════════════════════════════════════════════════════════
# SEGUNDA RONDA (3-oct-2026): A-J
# ═══════════════════════════════════════════════════════════════════════════
print("\n11 · (A) LA LEGITIMACIÓN DE LA PERSONA FÍSICA, SIN GÉNERO [siempre]")
bandera(False)
_l = ta.legitimacion_de("amparo_directo", "Ana López Ruiz")
ok(_l.startswith("La demanda de amparo fue promovida por Ana López Ruiz; la parte quejosa está legitimada "
                 "para ello conforme al artículo 5o., fracción I,") and "quien está" not in _l,
   "AD, persona física: «; la parte quejosa está legitimada», nunca «quien está legitimado/a»")
ok("; la parte quejosa está legitimada" in ta.legitimacion_de("amparo_directo", "Isadora Méndez Ruiz"),
   "el nombre de pila no decide nada (Isadora acaba como «Comercializadora»)")
_l = ta.legitimacion_de("amparo_revision", "Juan Pérez López")
ok("Juan Pérez López; la parte recurrente se encuentra legitimada" in _l and "artículo 6o. de la Ley" in _l
   and "6º" not in _l, "AR, persona física: «; la parte recurrente se encuentra legitimada», con el 6o.")
_l = ta.legitimacion_de("amparo_revision", "Ana López Ruiz", papel="tercero")
ok("Ana López Ruiz; la parte recurrente se encuentra legitimada" in _l and "5o., fracción III" in _l,
   "AR, tercera interesada persona física: neutra y con el 5o., fr. III")
ok("Comercializadora Ejemplo, S.A. de C.V., quien está legitimada" in
   ta.legitimacion_de("amparo_directo", "Comercializadora Ejemplo, S.A. de C.V.")
   and "el Director de Ingresos del Municipio de Querétaro, quien se encuentra legitimado" in
   ta.legitimacion_de("amparo_revision", "Director de Ingresos del Municipio de Querétaro", papel="autoridad"),
   "las morales y los órganos, como estaban: concuerdan por su artículo")
ok(all(x not in ta.legitimacion_de(t, "Juan Pérez López", "María Ruiz", figura="apoderada")
       for t in ("amparo_directo", "amparo_revision", "queja") for x in ("5º", "6º", "5°", "6°"))
   and "6o. y 11" in ta.legitimacion_de("amparo_directo", "Juan Pérez López", "María Ruiz"),
   "«5º»/«6º» → «5o.»/«6o.» en los moldes")
_fr = {p_: ta.legitimacion_de("queja", "Juan Pérez López", papel=p_)
       for p_ in ("quejoso", "autoridad", "tercero", "ministerio_publico")}
ok("5o., fracción I," in _fr["quejoso"] and "5o., fracción II," in _fr["autoridad"]
   and "5o., fracción III," in _fr["tercero"] and "5o., fracción IV," in _fr["ministerio_publico"],
   "queja: el 5o. CON su fracción por el papel (I quejoso, II autoridad, III tercero, IV MP)")
ok("5o., fracción I," in ta.legitimacion_de("queja", "Juan Pérez López", recurre_el_quejoso=True)
   and f"5o., fracción {H}," in ta.legitimacion_de("queja", "Juan Pérez López", recurre_el_quejoso=False),
   "queja sin papel: la del quejoso si recurre el quejoso; si no se sabe, hueco")
_qx, _estx = componer("queja", {"responsable": "Jueza Tercero de Distrito en el Estado de Querétaro",
                      "recurrente": "María Gómez Ruiz", "papel_recurrente": "tercero"}, nombre="q_tercera")
ok("El recurso fue interpuesto por María Gómez Ruiz; la parte recurrente está legitimada para hacerlo en "
   "términos del artículo 5o., fracción III," in _qx, "queja de la tercera: el párrafo es suyo, con la fr. III")
_qy, _esty = componer("queja", {"responsable": "Jueza Tercero de Distrito en el Estado de Querétaro",
                      "recurrente": "María Gómez Ruiz"}, nombre="q_sin_papel")
ok(f"5o., fracción {H}," in _qy and any("FRACCIÓN DEL ARTÍCULO 5o." in a for a in _esty.avisos),
   "queja de otra parte sin su carácter: la fracción en hueco y el aviso que la pide")
_qz, _ = componer("queja", {"responsable": "Jueza Tercero de Distrito en el Estado de Querétaro"}, nombre="q_quejoso")
ok("Juan Pérez López; la parte recurrente está legitimada para hacerlo en términos del artículo 5o., fracción I,"
   in _qz, "queja del propio quejoso (recurrente vacío = el mismo): fr. I")

print("\n12 · (B) A QUIÉN SE NOTIFICÓ, SIN GÉNERO [siempre]")
_c_ad = f0.computar(dt.date(2025, 3, 20), dt.date(2025, 4, 8), regla="personal", plazo=15)
_c_ar = f0.computar(dt.date(2025, 3, 20), dt.date(2025, 3, 27), regla="personal", plazo=10,
                    tipo_asunto="amparo_revision")
_c_of = f0.computar(dt.date(2025, 3, 20), dt.date(2025, 3, 27), regla="oficio", plazo=10,
                    tipo_asunto="amparo_revision")
_c_rf = f0.computar(dt.date(2025, 3, 20), dt.date(2025, 4, 8), regla="lfpca_boletin", plazo=15,
                    tipo_asunto="revision_fiscal")
_c_ot = f0.computar(dt.date(2025, 3, 20), dt.date(2025, 4, 8), regla="otra", plazo=15,
                    surtio_manual=dt.date(2025, 3, 21))
_ps = [f0.parrafo_oportunidad(_c_ad, "", "amparo_directo"), f0.parrafo_oportunidad(_c_ar, "", "queja"),
       f0.parrafo_oportunidad(_c_of, "", "amparo_revision", papel="autoridad"),
       f0.parrafo_oportunidad(_c_rf, "", "revision_fiscal"), f0.parrafo_oportunidad(_c_ot, "", "amparo_directo")]
ok("se notificó a la parte quejosa" in _ps[0] and "se notificó a la parte recurrente" in _ps[1]
   and "se notificó a la autoridad recurrente" in _ps[2] and "se notificó a la autoridad recurrente" in _ps[3],
   "«a la parte quejosa», «a la parte recurrente», «a la autoridad recurrente» (y en revisión fiscal)")
ok(not any("se notificó al " in x for x in _ps) and "según lo manifestado por la parte quejosa" in _ps[4],
   "nunca «se notificó al quejoso/recurrente»; la regla «otra» dice «por la parte quejosa»")

print("\n13 · (D) CADA INHÁBIL CON SU FUNDAMENTO [siempre]")
_k = f0.clave_del_inhabil
ok(_k(dt.date(2025, 5, 2)) == "c1_2025" and _k(dt.date(2025, 4, 17)) == "c1_2025"
   and _k(dt.date(2025, 3, 17)) == "lft_iii" and _k(dt.date(2026, 2, 2)) == "lft_ii"
   and _k(dt.date(2026, 11, 16)) == "lft_vi" and _k(dt.date(2026, 11, 2)) == "c3_2026"
   and _k(dt.date(2025, 9, 15)) == "c3_2025" and _k(dt.date(2025, 5, 1)) == "art19"
   and _k(dt.date(2025, 12, 25)) == "vac" and _k(dt.date(2026, 1, 1)) == "art19"
   and _k(dt.date(2024, 12, 20)) == "art19",
   "la nota de la página del OAJ para cada día: circular, lunes de la LFT, vacaciones (desde 2025) o el 19")
_cm = f0.computar(dt.date(2025, 4, 10), dt.date(2025, 5, 6), regla="personal", plazo=15, tipo_asunto="queja")
_pm = f0.parrafo_oportunidad(_cm, "", "queja")
# D6 (tercera ronda): los grupos en orden cronológico, cada uno con su fundamento.
_i_ss, _i_my = _pm.find("ni del dieciséis al dieciocho de abril de dos mil veinticinco"), \
    _pm.find("ni el uno y el cinco de mayo de dos mil veinticinco")
ok(0 <= _i_my < _i_ss and "el dos de mayo de dos mil veinticinco, por ser inhábiles conforme a la Circular "
   "1/2025 de la Secretaría Ejecutiva del Pleno del otrora Consejo de la Judicatura Federal" in _pm
   and "ni el uno y el cinco de mayo de dos mil veinticinco, por ser inhábiles en términos del artículo 19 de la "
       "Ley de Amparo" in _pm and "del uno al cinco de mayo" not in _pm,
   "el 2 de mayo de 2025 y la Semana Santa, con la Circular 1/2025; el 1 y el 5, con el 19 (el grupo del 19 primero, una vez)")
_cl = f0.computar(dt.date(2025, 3, 6), dt.date(2025, 3, 28), regla="personal", plazo=15)
ok("ni el diecisiete de marzo de dos mil veinticinco, por ser inhábil conforme al artículo 74, fracción III, de la "
   "Ley Federal del Trabajo, en relación con el artículo 6, fracción III, del Acuerdo General del Pleno del otrora "
   "Consejo de la Judicatura Federal" in f0.parrafo_oportunidad(_cl, "", "amparo_directo"),
   "el tercer lunes de marzo, con el 74 de la LFT y el acuerdo que lo hace inhábil (no el 19)")
_cv = f0.computar(dt.date(2025, 12, 5), dt.date(2026, 1, 8), regla="personal", plazo=15)
_pv = f0.parrafo_oportunidad(_cv, "", "amparo_directo")
ok("ni del dieciséis al treinta y uno de diciembre de dos mil veinticinco, por corresponder al periodo vacacional "
   "del Poder Judicial de la Federación, conforme al artículo 226 de la Ley Orgánica del Poder Judicial de la "
   "Federación" in _pv and _pv.count("artículo 19") == 1 and "ni el uno de enero de dos mil veintiséis, por ser "
   "inhábiles en términos del artículo 19" in _pv,
   "las vacaciones con el 226 de la LOPJF; el 1o. de enero, con el 19 y los sábados y domingos (una vez)")
_tr = f0.tramos_inhabiles(_cm.inhabiles_en_medio, _cm.cal_amparo)
ok(isinstance(_tr, list) and all(len(x) == 2 for x in _tr)
   and f0.tramos_en_letra(_tr).endswith("ni el uno y el cinco de mayo de dos mil veinticinco")
   and "(inhábiles conforme a la Circular 1/2025" in f0.tramos_en_letra(_tr),
   "para quien escribe «ni {tramos}, por ser inhábiles en términos del 19» (el adhesivo): los ajenos con su "
   "fundamento entre paréntesis y los del 19 al final; sigue siendo una lista de pares")
_crf = f0.computar(dt.date(2025, 4, 10), dt.date(2025, 5, 6), regla="lfpca_boletin", plazo=15,
                   tipo_asunto="revision_fiscal")
ok("Circular" not in f0.parrafo_oportunidad(_crf, "", "revision_fiscal")
   and "Acuerdo SS/1/2025" in f0.parrafo_oportunidad(_crf, "", "revision_fiscal"),
   "la revisión fiscal sigue con el calendario del TFJA y su acuerdo (no se le aplican las notas del OAJ)")
_dl = Document()
dg._pagina(_dl)
dg.calendario_computo(_dl, _cv, "amparo_directo")
_tl = "\n".join(c.text for t in _dl.tables for f in t.rows for c in f.cells)
ok("Sábados y domingos (art. 19 LA)" in _tl and "16 dic 2025 a 31 dic 2025 (vacaciones, art. 226 LOPJF)" in _tl
   and "1 ene 2026 (art. 19 LA)" in _tl, "la leyenda de la tabla, cada tramo con su fundamento")

print("\n14 · (E) EL PRECEPTO DEL SURTIMIENTO (bandera)")
ok(f0.conforme_al_precepto("el artículo 126 del Código de Procedimientos Civiles del Estado de Querétaro")
   == "conforme al artículo 126 del Código de Procedimientos Civiles del Estado de Querétaro"
   and f0.conforme_al_precepto("conforme a los artículos 3 y 4 de la Ley X.") == "conforme a los artículos 3 y 4 de la Ley X"
   and f0.conforme_al_precepto(H) == "", "lo que escribe el secretario, en la forma que pide la frase")
_p126 = f0.parrafo_oportunidad(_c_ad, "", "amparo_directo", sin_precepto_en_hueco=True,
                               fundamento_surtimiento="el artículo 126 del Código de Procedimientos Civiles del "
                                                      "Estado de Querétaro")
ok("surtió efectos al día hábil siguiente, conforme al artículo 126 del Código de Procedimientos Civiles del Estado "
   "de Querétaro, es decir," in _p126 and H not in _p126
   and f0.aviso_fundamento(_c_ad, "amparo_directo", en_hueco=True, fundamento_surtimiento="el artículo 126 X") == "",
   "con el precepto declarado: «conforme a {ese texto}» en vez del hueco, y sin aviso")
ok(f0.surtimiento_nacional("amparo_directo", "mercantil", "Juez Primero Civil", "personal")[0]
   == "el artículo 1075 del Código de Comercio"
   and "artículos 65 y 70 de la Ley Federal de Procedimiento Contencioso" in f0.surtimiento_nacional(
       "amparo_directo", "administrativa", "Sala Regional Peninsular del Tribunal Federal de Justicia "
       "Administrativa", "personal")[0]
   # C7 (3-oct-2026): lo agrario se decide por la sede. Antes: («», «») sin
   # sede, porque el 167 de la Ley Agraria ya remite al CNPCF; ahora el 321 del
   # CFPC fuera de la Ciudad de México (o sin sede), con el aviso del
   # transitorio, y nada en ella (la regla del Código Nacional trae su precepto).
   and f0.surtimiento_nacional("amparo_directo", "agraria", "Tribunal Unitario Agrario del Distrito 42",
                               "personal")[0].startswith("el artículo 321 del Código Federal de Procedimientos Civiles")
   and f0.surtimiento_nacional("amparo_directo", "agraria", "Tribunal Unitario Agrario del Distrito 42",
                               "personal", tribunal="Primer Tribunal Colegiado en Materia Administrativa del "
                                                    "Primer Circuito") == ("", "")
   and f0.surtimiento_nacional("queja", "mercantil", "", "personal") == ("", "")
   and f0.surtimiento_nacional("amparo_directo", "mercantil", "", "lfpca_boletin") == ("", ""),
   "la omisión nacional sólo donde el texto local la confirma (Código de Comercio 1075, LFPCA 65 y 70); "
   "el agrario por la sede (C7: el 321 del CFPC fuera de la Ciudad de México o sin sede; en ella, nada, "
   "porque la regla del Código Nacional trae su precepto); ni los recursos, ni la regla que ya trae su precepto")
bandera(True)
_pro_e = {"responsable": SALA_QRO, "fecha_acto": "trece de marzo de dos mil veinticinco", "materia": "civil",
          "expediente": "905/2023", "toca": "374/2024", "toca_en_prosa": "toca 374/2024",
          "expediente_en_prosa": "expediente 905/2023",
          "fundamento_surtimiento": "el artículo 126 del Código de Procedimientos Civiles del Estado de Querétaro"}
_ade, _este = componer("amparo_directo", {"responsable": SALA_QRO, "procesal": _pro_e}, nombre="ad_surte")
ok("conforme al artículo 126 del Código de Procedimientos Civiles del Estado de Querétaro, es decir" in _ade
   and "conforme al *********" not in _ade
   and not any("EL SURTIMIENTO VA SIN PRECEPTO" in a for a in _este.avisos),
   "con la bandera, `procesal.fundamento_surtimiento` llena el hueco y el aviso desaparece")
_adm, _estm = componer("amparo_directo", {"responsable": "Juez Primero de lo Mercantil del Distrito Judicial de "
                       "Querétaro", "materia": "mercantil", "procesal": {"materia": "mercantil",
                       "responsable": "Juez Primero de lo Mercantil del Distrito Judicial de Querétaro"}},
                       nombre="ad_mercantil")
ok("conforme al artículo 1075 del Código de Comercio, es decir" in _adm
   and any("1075 DEL CÓDIGO DE COMERCIO" in a and "COMPRUÉBALO" in a for a in _estm.avisos),
   "mercantil sin precepto declarado: el 1075 del Código de Comercio, con aviso «compruébalo»")
bandera(False)
_adn, _ = componer("amparo_directo", {"responsable": SALA_QRO, "procesal": _pro_e}, nombre="ad_surte_sin")
ok("conforme a la ley del acto" in _adn and "126 del Código" not in _adn,
   "sin la bandera, el camino de siempre")

print("\n15 · (C) EL RESOLUTIVO DEL AMPARO DIRECTO, INDIVIDUALIZADO (bandera)")
bandera(True)
_adc, _estc = componer("amparo_directo", {"responsable": SALA_QRO, "procesal": _pro_e}, calif=("fundado",),
                       estudio=("El concepto es fundado; procede conceder el amparo para que la Sala deje "
                                "insubsistente la sentencia.",), nombre="ad_concede")
_rc = _adc.split("R E S U E L V E")[1]
_ef = [l_ for l_ in _adc.splitlines() if ". Efectos." in l_[:22]]
ok("La Justicia de la Unión ampara y protege a Juan Pérez López, contra la sentencia dictada el trece de marzo de "
   f"dos mil veinticinco, por la {SALA_QRO}, en el toca 374/2024, derivado del expediente 905/2023" in _rc
   and (not _ef or f"para los efectos precisados en el considerando {_ef[0].split('.')[0].lower()} de esta "
                   "ejecutoria." in _rc),
   "concede: la sentencia con su fecha, su órgano y su toca, y los efectos por el ordinal de esta ejecutoria")
_adh_, _esth_ = componer("amparo_directo", {"responsable": SALA_QRO, "procesal": {"responsable": SALA_QRO,
                         "materia": "civil"}}, nombre="ad_res_hueco")
ok(f"contra la sentencia dictada el {H}, por la {SALA_QRO}, en el {H}." in _adh_.split("R E S U E L V E")[1]
   and any("EL RESOLUTIVO DEL AMPARO DIRECTO NO IDENTIFICA EL ACTO" in a for a in _esth_.avisos),
   "sin fecha ni toca en la ficha: hueco y el aviso que nombra los datos")
_lau, _ = componer("amparo_directo", {"responsable": "Junta Especial Número Cincuenta de la Federal de "
                   "Conciliación y Arbitraje", "procesal": {"responsable": "Junta Especial Número Cincuenta de la "
                   "Federal de Conciliación y Arbitraje", "fecha_acto": "trece de marzo de dos mil veinticinco",
                   "expediente": "77/2024", "expediente_en_prosa": "expediente 77/2024",
                   "descripcion_acto": "del laudo reclamado"}}, nombre="ad_laudo")
ok("contra el laudo dictado el trece de marzo de dos mil veinticinco, por la Junta Especial Número Cincuenta" in _lau,
   "el laudo concuerda: «el laudo dictado»")

print("\n16 · (F) LA PROCEDENCIA DE LA REVISIÓN FISCAL CON LA FICHA (bandera)")
_rf_base = {"quejoso": "Administración Desconcentrada Jurídica de Querétaro «1»",
            "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"}
_pro_rf = {"sala": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa",
           "expediente_tfja": "695/25-09-01-7-OT", "fecha_acto": "trece de marzo de dos mil veinticinco",
           "fecha_acto_iso": "2025-03-13"}
_r1, _ = componer("revision_fiscal", dict(_rf_base, procesal=dict(_pro_rf, fraccion_63="I", cuantia="$847,738.77")),
                  regla="lfpca_boletin", nombre="rf_fr1")
ok("artículo 63, fracción I, de la Ley Federal de Procedimiento Contencioso Administrativo, toda vez que el "
   "asunto es de una cuantía de $847,738.77" in _r1 and "$113.14 en 2025" in _r1,
   "fracción I con la cuantía de la ficha contra la UMA de la fecha de la sentencia")
_r2, _e2 = componer("revision_fiscal", dict(_rf_base, procesal=dict(_pro_rf, fraccion_63="II", cuantia="$100,000")),
                    regla="lfpca_boletin", nombre="rf_fr2")
ok("fracción II" in _r2 and "razonó su importancia y trascendencia" in _r2
   and any("IMPORTANCIA Y TRASCENDENCIA" in a for a in _e2.avisos),
   "fracción II: importancia y trascendencia razonada por la autoridad, con aviso para comprobarlo")
_r3, _ = componer("revision_fiscal", dict(_rf_base, procesal=dict(_pro_rf, fraccion_63="III",
                  autoridad_demandada="Administración Desconcentrada de Auditoría Fiscal de Querétaro «1»")),
                  regla="lfpca_boletin", nombre="rf_fr3")
ok("fracción III" in _r3 and "la dictó una autoridad del Servicio de Administración Tributaria" in _r3
   and f"supuesto del inciso {H})" in _r3, "fracción III: la autoridad hacendaria y el inciso en hueco")
_r5, _e5 = componer("revision_fiscal", dict(_rf_base, procesal=dict(_pro_rf, fraccion_63="V")),
                    regla="lfpca_boletin", nombre="rf_fr5")
ok("fracción V, de la Ley Federal de Procedimiento Contencioso Administrativo, que lo admite cuando se trata de "
   f"una resolución dictada en materia de comercio exterior, supuesto que se actualiza porque {H}." in _r5
   and any("FRACCIÓN V DEL ARTÍCULO 63" in a for a in _e5.avisos),
   "las demás fracciones: la fórmula del artículo con el porqué en hueco y aviso")
bandera(False)
_r0, _ = componer("revision_fiscal", dict(_rf_base, procesal=dict(_pro_rf, fraccion_63="V")),
                  regla="lfpca_boletin", nombre="rf_fr_sin")
ok("comercio exterior" not in _r0, "sin la bandera, la ficha no entra a la procedencia")

print("\n17 · (G) ARTÍCULOS Y VERSALES [siempre]")
ok(dg._con_articulo("Titular de la Jefatura de Servicios Jurídicos") == "el Titular de la Jefatura de Servicios Jurídicos"
   and dg._con_articulo("Jueza Tercero de Distrito").startswith("la "),
   "«el Titular» (corpus 59 contra 12); la Jueza sigue con «la»")
ok(dg._nombre_de_organo("SALA REGIONAL DEL CENTRO II DEL TRIBUNAL FEDERAL DE JUSTICIA ADMINISTRATIVA")
   == "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"
   and dg._nombre_de_organo("TERCER TRIBUNAL COLEGIADO DEL XXII CIRCUITO") == "Tercer Tribunal Colegiado del XXII Circuito"
   and dg._nombre_de_organo("DELEGACIÓN ESTATAL DEL IMSS EN QUERÉTARO") == "Delegación Estatal del IMSS en Querétaro"
   and all(x in dg._nombre_de_organo("SAT TFJA ISSSTE INFONAVIT CONAGUA SHCP")
           for x in ("SAT", "TFJA", "ISSSTE", "INFONAVIT", "CONAGUA", "SHCP"))
   and dg._nombre_de_organo("JUZGADO CIVIL") == "Juzgado Civil",
   "_nombre_de_organo no rompe romanos (II, XXII) ni siglas (IMSS, SAT, TFJA, ISSSTE, INFONAVIT, CONAGUA, SHCP)")

print("\n18 · (H) LA QUEJA: SU CIERRE Y SU DISPENSA")
_qd, _ = componer("queja", {"responsable": "Jueza Tercero de Distrito en el Estado de Querétaro"}, nombre="q_disp")
ok("examinar las cuestiones efectivamente planteadas" in _qd and "previamente planteadas" not in _qd,
   "[siempre] la dispensa de la queja dice «efectivamente planteadas»")
bandera(True)
_q2b, _ = componer("queja", {"responsable": SALA_QRO, "procesal": {"fraccion_97": "II", "inciso_97": "b",
                   "juicio_amparo": "380/2025", "juzgado": SALA_QRO,
                   "fecha_acto": "trece de marzo de dos mil veinticinco",
                   "cola_97": "en el que negó la suspensión del acto reclamado"}}, calif=("fundado",),
                   nombre="q_ii_cierre")
_q1b, _ = componer("queja", {"responsable": "Jueza Tercero de Distrito en el Estado de Querétaro",
                   "procesal": {"fraccion_97": "I", "inciso_97": "a", "juicio_amparo": "742/2025-II"}},
                   nombre="q_i_cierre")
ok("envíese testimonio de esta resolución a la autoridad responsable" in _q2b
   and "al juzgado de origen" not in _q2b and "envíese testimonio de esta resolución al juzgado de origen" in _q1b,
   "[bandera] fracción II: el testimonio va a la autoridad responsable; la fracción I, al juzgado de origen")

print("\n19 · (I) LA AUTORIDAD RESPONSABLE ORIGINARIA DE LA FICHA (bandera)")
_ari, _esti = componer("amparo_revision", {"quejoso": "Juan Pérez López",
                       "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                       "resolvio_a_quo": "sobresee",
                       "procesal": {"juicio_amparo": "950/2024", "juzgado": "Juzgado Quinto de Distrito en el "
                                    "Estado de Querétaro", "autoridades": ["Director de Ingresos del Municipio de "
                                                                           "Querétaro"],
                                    "responsable": "Director de Ingresos del Municipio de Querétaro"}},
                       calif=("fundado",), estudio=("Es fundado el agravio; procede revocar el sobreseimiento y "
                                                    "conceder el amparo.",), nombre="ar_originaria")
_rsi = _ari.split("R E S U E L V E")[1]
ok("Director de Ingresos del Municipio de Querétaro" in _rsi and H not in _rsi
   and not any("NO SE PUDO LEER LA AUTORIDAD RESPONSABLE ORIGINARIA" in a for a in _esti.avisos),
   "el punto del amparo nombra a la responsable de la ficha, sin comodín ni aviso")
_arj, _estj = componer("amparo_revision", {"quejoso": "Juan Pérez López",
                       "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                       "resolvio_a_quo": "sobresee", "procesal": {"juicio_amparo": "950/2024"}},
                       calif=("fundado",), estudio=("Es fundado el agravio; procede revocar el sobreseimiento y "
                                                    "conceder el amparo.",), nombre="ar_originaria_sin")
ok(any("NO SE PUDO LEER LA AUTORIDAD RESPONSABLE ORIGINARIA" in a for a in _estj.avisos),
   "sin autoridades en la ficha, el comodín y su aviso, como antes")


# ═══════════════════════════════════════════════════════════════════════════
# TERCERA RONDA (3-oct-2026): el banco de oráculo (32 sentencias reales) y la
# prueba de punta a punta (AD 274/2025, AR 631/2025). Dueño A de FIXES_R3.
# ═══════════════════════════════════════════════════════════════════════════
def componer3(tipo, datos_extra, ficha=None, procesal=None, resultandos=None, calif=("infundado",),
              regla="personal", notif=dt.date(2025, 3, 20), pres=None, existencia="", procedencia="",
              avisos_previos=(), estudio=("El estudio de fondo.",), nombre="x3"):
    """(texto, estructura, tablas, ruta) — como `componer`, con la ficha, el cómputo y la
    estructura a la medida de cada caso."""
    plazo = {"amparo_directo": 15, "amparo_revision": 10, "queja": 5, "revision_fiscal": 15}[tipo]
    pres = pres or (dt.date(2025, 3, 27) if tipo in ("queja", "amparo_revision") else dt.date(2025, 4, 8))
    datos = {"tipo_asunto": tipo, "numero": "1/2026", "magistrado": "Magistrado Ejemplo",
             "secretario": "Secretario Ejemplo", "tribunal": XXII, "ciudad": "Querétaro, Querétaro",
             "acto": "texto del acto", "antecedentes": "", "quejoso": "Juan Pérez López",
             "encabezado": "ASUNTO: 1/2026"}
    datos.update(datos_extra)
    if procesal is not None:
        datos["procesal"] = procesal
    if ficha is not None:
        datos["tramite"] = ficha
    c = f0.computar(notif, pres, regla=regla, plazo=plazo, responsable=datos.get("responsable"),
                    tipo_asunto=tipo)
    est = dg.Estructura(apertura="", visto="para resolver el asunto 1/2026; y,",
                        resultandos=[dict(r) for r in (resultandos or [
                            {"titulo": "Trámite", "texto": "Texto del resultando."}])],
                        competencia="", existencia=existencia, procedencia=procedencia,
                        avisos=list(avisos_previos))
    ruta = os.path.join(_TMP, f"{nombre}.docx")
    dg.componer(datos, est, c, f0.fecha_en_letra, ruta, estudio=list(estudio),
                calificaciones=list(calif), tipo_asunto=tipo, antecedentes=["Antecedente uno."])
    d_ = Document(ruta)
    texto = "\n".join(p_.text for p_ in d_.paragraphs if p_.text.strip())
    tablas = "\n".join(c_.text for t_ in d_.tables for f_ in t_.rows for c_ in f_.cells)
    return texto, est, tablas, ruta


def considerandos(t):
    return t.split("C O N S I D E R A N D O")[1].split("R E S U E L V E")[0] if "C O N S I D E R A N D O" in t else ""


print("\n20 · LOS ARTÍCULOS DE LOS CARGOS Y LOS PREFIJOS TEMPORALES [siempre] (AD 128, RF 2/7/26/49, AR 201)")
bandera(False)
_ca = dg._con_articulo
ok(_ca("Coordinadora de Recursos Humanos del Municipio de Cadereyta de Montes").startswith("la Coordinadora")
   and _ca("Jefa de la Unidad Jurídica").startswith("la Jefa") and _ca("JEFA DE LA UNIDAD JURÍDICA").startswith("la ")
   and _ca("Subdirectora de Afiliación").startswith("la ") and _ca("Subdelegada de Prestaciones").startswith("la ")
   and _ca("Procuradora Fiscal de la Federación").startswith("la ") and _ca("Contralora Interna").startswith("la ")
   and _ca("Tesorera Municipal").startswith("la ") and _ca("Presidenta Municipal").startswith("la "),
   "los cargos en femenino concuerdan con «la» (Coordinadora, Jefa, Subdirectora, Procuradora, Contralora…)")
ok(all(_ca(x).startswith("la ") for x in ("Legislatura del Estado de Querétaro", "Cámara de Diputados",
                                          "Jefatura de Gobierno", "Subjefatura de Servicios", "Oficina Recaudadora",
                                          "Sindicatura Municipal", "Gubernatura", "Regiduría de Hacienda",
                                          "Alcaldía Benito Juárez")),
   "«la Legislatura», «la Cámara», «la Jefatura», «la Oficina», «la Sindicatura» (no «el Legislatura»)")
ok(_ca("Actual Sala Regional en Querétaro del Tribunal Federal de Justicia Administrativa")
   == "la actual Sala Regional en Querétaro del Tribunal Federal de Justicia Administrativa"
   and _ca("entonces Director de Ingresos") == "el entonces Director de Ingresos"
   and _ca("hoy Jueza Tercero de Distrito") == "la hoy Jueza Tercero de Distrito",
   "el prefijo temporal en minúscula y detrás del artículo, que lo decide el sustantivo («la actual Sala»)")
ok(_ca("Titular de la Unidad Jurídica").startswith("el Titular")
   and _ca("la titular del Juzgado Cuarto de Distrito") == "la titular del Juzgado Cuarto de Distrito"
   and not _ca("Isadora Méndez Ruiz").startswith("la ") and not _ca("Pastora de la Cruz").startswith("la "),
   "«el Titular» salvo que el papel diga «la titular»; un nombre de pila en -dora no decide el artículo")
_rfa, _, _, _ = componer3("revision_fiscal", {"quejoso": "Administración Desconcentrada Jurídica de Querétaro «1»",
                          "responsable": "Actual Sala Regional en Querétaro del Tribunal Federal de Justicia "
                          "Administrativa"}, regla="lfpca_boletin", nombre="r3_actual")
ok("tramitado ante la actual Sala Regional en Querétaro del Tribunal Federal de Justicia Administrativa, "
   "localizada" in _rfa and "el Actual" not in _rfa, "RF 49/2025: «ante la actual Sala…, localizada», en el documento")

print("\n21 · LO RECLAMADO SE LLAMA POR SU CLASE: RESOLUCIÓN O LAUDO (bandera; AD 349 y 552)")
bandera(True)
_pro_res = {"responsable": SALA_QRO, "materia": "civil", "fecha_acto": "trece de marzo de dos mil veinticinco",
            "toca": "374/2024", "toca_en_prosa": "toca 374/2024", "expediente": "905/2023",
            "expediente_en_prosa": "expediente 905/2023", "clase": "resolucion",
            "acto_reclamado": "la resolución reclamada"}
_adr, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO}, procesal=_pro_res,
                          ficha={"tipo": "amparo_directo", "acto": {"clase": "resolucion"}}, nombre="r3_res")
_cr = considerandos(_adr)
# E4 (cuarta ronda): «resolución» sin más es «una resolución definitiva»; «que puso fin al juicio»
# sólo si el sentido lo dice (sección 34).
ok("en contra de una resolución definitiva en materia civil, dictada por la " + SALA_QRO in _cr
   and "La existencia de la resolución reclamada" in _cr and "le causa la resolución reclamada" in _cr
   and "la resolución reclamada se notificó" in _cr and "el contenido de la resolución reclamada" in _cr
   and "algún apartado de la resolución reclamada" in _cr
   and "sentencia reclamada" not in _cr and "sentencia definitiva" not in _cr,
   "resolución: competencia, existencia, legitimación, oportunidad y dispensa la llaman resolución")
_pro_lau = {"responsable": "Junta Especial Número Cincuenta de la Federal de Conciliación y Arbitraje",
            "materia": "laboral", "fecha_acto": "trece de marzo de dos mil veinticinco", "expediente": "77/2024",
            "expediente_en_prosa": "expediente 77/2024", "clase": "laudo"}
_adl, _, _, _ = componer3("amparo_directo", {"responsable": _pro_lau["responsable"], "materia": "laboral"},
                          procesal=_pro_lau, ficha={"tipo": "amparo_directo", "acto": {"clase": "laudo"}},
                          nombre="r3_laudo")
_cl = considerandos(_adl)
ok("en contra de un laudo en materia laboral, dictado por la Junta Especial" in _cl
   and "el laudo reclamado se notificó" in _cl and "le causa el laudo reclamado" in _cl
   and "el contenido del laudo reclamado" in _cl and "sentencia reclamada" not in _cl and "de el laudo" not in _cl,
   "laudo: «un laudo…, dictado por», «el laudo reclamado», «del laudo reclamado»")
_ads, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO}, procesal=dict(_pro_res, clase="sentencia",
                          acto_reclamado="la sentencia reclamada"), nombre="r3_sent")
ok("en contra de una sentencia definitiva en materia civil" in _ads and "la sentencia reclamada se notificó" in _ads,
   "con la sentencia, como siempre")

print("\n22 · LA CARÁTULA CON LA FICHA: PONENTE, ENCABEZADO, ADHERENTE Y RÓTULOS (bandera)")
_f_ret = {"tipo": "amparo_directo", "turno": {"fecha": "2025-05-23", "ponente": "J. Guadalupe Tafoya Hernández"},
          "returno": {"fecha": "2025-11-25", "ponente": "Luis Armando Pérez Topete"}}
_t1, _e1, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO,
                           "magistrado": "Magistrado J. Guadalupe Tafoya Hernández"},
                           procesal={"responsable": SALA_QRO, "materia": "civil"}, ficha=_f_ret, nombre="r3_pon1")
ok("PONENTE: LUIS ARMANDO PÉREZ TOPETE." in _t1 and "TAFOYA" not in _t1.split("Querétaro, Querétaro.")[0]
   and any("SE RETURNARON A Luis Armando Pérez Topete" in a for a in _e1.avisos),
   "AD 274: con returno, la carátula nombra al ponente del ÚLTIMO returno y avisa del cambio")
_t2, _e2, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "magistrado": ""},
                           procesal={"responsable": SALA_QRO, "materia": "civil"},
                           ficha=dict(_f_ret, returno=[{"fecha": "2025-08-01", "ponente": "Otra Persona Ruiz"},
                                                       {"fecha": "2025-11-25",
                                                        "ponente": "Magistrada Jenica Campos Juárez"}]),
                           nombre="r3_pon2")
ok("MAGISTRADA PONENTE: JENICA CAMPOS JUÁREZ." in _t2 and "MAGISTRADA JENICA" not in _t2,
   "sin ponente en el formulario: el del último returno de la lista, con el cargo en el rótulo y no en el valor")
_t3, _e3, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "magistrado": "Magistrada Jenica Campos Juárez"},
                           procesal={"responsable": SALA_QRO, "materia": "civil"},
                           ficha={"tipo": "amparo_directo", "turno": {"ponente": "Jenica Campos Juárez"}},
                           nombre="r3_pon3")
ok("MAGISTRADA PONENTE: JENICA CAMPOS JUÁREZ." in _t3 and not any("PONENTE" in a for a in _e3.avisos),
   "RF 26: «MAGISTRADO PONENTE: MAGISTRADO…» ya no; el cargo del formulario decide el rótulo")
_t4, _e4, _, _ = componer3("amparo_revision", {"magistrado": "Bertha Martínez Vega, Secretaria en funciones de "
                           "Magistrada", "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                           procesal={"juicio_amparo": "950/2024"},
                           ficha={"tipo": "amparo_revision", "turno": {"ponente": "Bertha Martínez Vega"}},
                           nombre="r3_pon4")
ok("\nPONENTE: BERTHA MARTÍNEZ VEGA." in _t4 and "MAGISTRADO PONENTE" not in _t4,
   "AR 222: la secretaria en funciones se rotula «PONENTE», no «MAGISTRADO PONENTE»")
_t5, _e5, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "magistrado": "Otro Magistrado Distinto"},
                           procesal={"responsable": SALA_QRO, "materia": "civil"}, ficha=_f_ret, nombre="r3_pon5")
ok("OTRO MAGISTRADO DISTINTO" in _t5 and any("NO ES EL DEL ÚLTIMO RETURNO" in a for a in _e5.avisos),
   "si el formulario dice otra persona, se respeta lo escrito y se avisa")
# E3 (cuarta ronda): el aviso del rótulo por omisión sale SIEMPRE que nadie diga el cargo, venga el
# nombre de la ficha o del formulario (sección 33).
ok(any("POR OMISIÓN" in a for a in _e1.avisos) and not any("POR OMISIÓN" in a for a in _e2.avisos)
   and any("POR OMISIÓN" in a for a in _e5.avisos),
   "el rótulo masculino sin cargo va con aviso, lo haya puesto la ficha (Q 342) o el secretario (E3); con el "
   "cargo en el papel, sin aviso")
_t6, _e6, _, _ = componer3("amparo_revision", {"encabezado": "", "materia": "civil",
                           "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                           procesal={"juicio_amparo": "950/2024", "materia": "civil"},
                           ficha={"tipo": "amparo_revision", "numero": "201/2025", "materia": "civil",
                                  "adhesivo": {"quien": "MARÍA GÓMEZ RUIZ (TERCERA INTERESADA)"}},
                           nombre="r3_enc")
ok(_t6.startswith("AMPARO EN REVISIÓN CIVIL: 201/2025.") and "RECURRENTE ADHESIVO: MARÍA GÓMEZ RUIZ." in _t6,
   "encabezado vacío: se compone con la ficha; el adherente de la ficha a su renglón, sin la etiqueta de rol")
_t6b, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO},
                          procesal={"responsable": SALA_QRO, "materia": "civil"},
                          ficha={"tipo": "amparo_directo", "adhesivo": {"quien": "Comercializadora Ejemplo, S.A. de C.V.",
                                                                        "admision": "2025-05-05"}}, nombre="r3_adh_ad")
_t6c, _, _, _ = componer3("revision_fiscal", {"quejoso": "Administración Desconcentrada Jurídica de Querétaro «1»",
                          "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"},
                          procesal={"sala": "Sala Regional del Centro II del Tribunal Federal de Justicia "
                                    "Administrativa"},
                          ficha={"tipo": "revision_fiscal", "adhesivo": {"quien": "Juan Pérez López"}},
                          regla="lfpca_boletin", nombre="r3_adh_rf")
ok("ADHERENTE: COMERCIALIZADORA EJEMPLO, S.A. DE C.V." in _t6b.split("Querétaro, Querétaro.")[0]
   and "RECURRENTE ADHESIVO: JUAN PÉREZ LÓPEZ." in _t6c.split("Querétaro, Querétaro.")[0],
   "AD 279 y RF 4: el adherente de la ficha también en la carátula del amparo directo y de la revisión fiscal")
ok("R E S U L T A N D O:" in _t6 and "C O N S I D E R A N D O:" in _t6,
   "D7: «R E S U L T A N D O:» y «C O N S I D E R A N D O:» con dos puntos (8 de 8 engroses)")
_t7, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO}, nombre="r3_sin_ficha")
ok("R E S U L T A N D O\n" in _t7 and "R E S U L T A N D O:" not in _t7
   and "MAGISTRADO PONENTE: MAGISTRADO EJEMPLO" in _t7,
   "sin la ficha, la carátula y los rótulos de siempre")

print("\n23 · UN DATO QUE FALTA, UN AVISO (bandera; Q 335: ocho avisos para dos datos)")
_esp = ["FALTA EL ÓRGANO QUE DICTÓ EL AUTO RECURRIDO (acto.organo): está en el auto recurrido. Va en hueco en "
        "«V I S T O».",
        "FALTA LA FECHA DEL AUTO (informe_101.fecha): está en el auto. Va en hueco en «Trámite del recurso».",
        "SE RECOMIENDA CITAR EL PRECEPTO DEL SURTIMIENTO: …",
        "EL CONSIDERANDO DE COMPETENCIA LLEVA HUECO (*********): falta el órgano."]
_det = [{"regla": "a", "aviso": "HUECO EN EL V I S T O: falta el órgano — «dictado por *********, en el» . Sólo…"},
        {"regla": "a", "aviso": "HUECO EN EL RESULTANDO «TRÁMITE DEL RECURSO»: falta la fecha del auto — «por "
                                "auto de *********» . Sólo…"},
        {"regla": "a", "aviso": "HUECO EN EL CONSIDERANDO «COMPETENCIA»: falta el órgano — «por *********, "
                                "localizado» . Sólo…"},
        {"regla": "a", "aviso": "HUECO EN EL CONSIDERANDO «LEGITIMACIÓN Y OPORTUNIDAD»: falta el precepto que rige "
                                "el surtimiento de la notificación — «…» . Sólo…"},
        {"regla": "a", "aviso": "HUECO EN EL RESULTANDO «TURNO»: falta el ponente — «ponencia de *********» . Sólo…"},
        {"regla": "a", "aviso": "HUECO EN LOS RESOLUTIVOS: falta el ponente — «*********» . Sólo…"},
        {"regla": "a", "campo": "presentacion", "aviso": "HUECO EN EL RESULTANDO «INTERPOSICIÓN»: falta la fecha de "
                                                         "presentación — «el *********» . Sólo…"},
        {"regla": "h", "aviso": "NORMA ABROGADA: …"}]
_fu = getattr(dg, "_fundir_avisos_de_hueco", None)
_r23 = _fu(_det, _esp + ["FALTA LA FECHA DE PRESENTACIÓN DEL RECURSO (presentacion): … Va en hueco en «Otra»."]) \
    if callable(_fu) else []
ok(callable(_fu) and len(_r23) == 2
   and _r23[0].startswith("HUECO EN EL RESULTANDO «TURNO» Y LOS RESOLUTIVOS: falta el ponente")
   and _r23[1] == "NORMA ABROGADA: …",
   "la verja no repite el hueco que otra pieza ya avisó (por apartado o por campo) y junta el mismo dato en uno")
_v23 = types.ModuleType("verja_procesal")
_v23.revisar_detalle = lambda t, f, d, tp: [dict(x) for x in _det[:1]] + [{"regla": "b", "aviso": "EVASIVA X"}]
_v23.revisar = lambda t, f, d, tp: [x["aviso"] for x in _v23.revisar_detalle(t, f, d, tp)]
_v23.arreglar = lambda t: (t, [])
_real_v = sys.modules.get("verja_procesal")
sys.modules["verja_procesal"] = _v23
try:
    _t23, _e23, _, _ = componer3("queja", {"responsable": "Juzgado Tercero de Distrito en el Estado de Querétaro"},
                                 procesal={"fraccion_97": "I", "inciso_97": "a", "juicio_amparo": "742/2025"},
                                 ficha={"tipo": "queja", "promovente": "Juan Pérez López", "caracter": "quejoso"},
                                 avisos_previos=_esp[:1], nombre="r3_avisos")
finally:
    if _real_v is None:
        sys.modules.pop("verja_procesal", None)
    else:
        sys.modules["verja_procesal"] = _real_v
ok(_e23.avisos[:1] == ["EVASIVA X"] and not any(a.startswith("HUECO EN EL V I S T O") for a in _e23.avisos)
   and _esp[0] in _e23.avisos, "en el documento: el genérico de la verja cae y el del compositor se queda")
_t23b, _e23b, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO},
                               procesal={"responsable": SALA_QRO, "materia": "civil",
                                         "fecha_acto": "trece de marzo de dos mil veinticinco"},
                               avisos_previos=["FALTA EL EXPEDIENTE DE ORIGEN (acto.expediente): está en la "
                                               "sentencia reclamada. Va en hueco en «V I S T O»."], nombre="r3_exp")
ok("los autos del expediente *********" in _t23b
   and not any("LLEVA HUECO" in a and "expediente" in a for a in _e23b.avisos),
   "el «LLEVA HUECO» del considerando no repite el campo que el compositor ya pidió (acto.expediente)")
_t23c, _e23c, _, _ = componer3("amparo_directo", {"responsable": ""},
                               procesal={"materia": "civil", "expediente": "905/2023",
                                         "fecha_acto": "trece de marzo de dos mil veinticinco"}, nombre="r3_agrupa")
_lh = [a for a in _e23c.avisos if "LLEVAN HUECO" in a or "LLEVA HUECO" in a]
ok("dictada por *********, localizado" in _t23c and "rendido por *********," in _t23c
   and len([a for a in _lh if "el órgano que dictó" in a]) == 1
   and any(a.startswith("LOS CONSIDERANDOS DE COMPETENCIA Y DE EXISTENCIA LLEVAN HUECO") and "el órgano que dictó" in a
           for a in _lh),
   "un dato que falta en dos considerandos (el órgano, en competencia y existencia): un aviso para los dos")

print("\n24 · LA TABLA DEL CÓMPUTO NO CONTRADICE AL PÁRRAFO [siempre] (regresión sin bandera)")
bandera(False)
_t24, _, _tab24, _ = componer3("amparo_revision", {"quejoso": "Comercializadora Ejemplo, S.A. de C.V.",
                               "recurrente": "Director de Ingresos del Municipio de Querétaro",
                               "papel_recurrente": "autoridad",
                               "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                               regla="personal", nombre="r3_tabla")
ok("conforme al *********" in _t24 and "fr. II" not in _tab24 and "fracción II" not in _tab24
   and "ley del acto" not in _tab24,
   "AR de la autoridad con la regla de los particulares: ni la leyenda ni la tarjeta citan la fr. II del 31")
_t24b, _, _tab24b, _ = componer3("amparo_revision", {"organo_recurrido": "Juzgado Quinto de Distrito en el Estado "
                                 "de Querétaro"}, regla="personal", nombre="r3_tabla_q")
ok("art. 31, fr. II LA" in _tab24b, "la quejosa sigue con la fr. II en su tabla")

print("\n25 · «PROCEDENCIA.» SIN BANDERA, CON SU CONDICIÓN DE SIEMPRE (D5)")
_t25, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO}, existencia="",
                          procedencia="El juicio es procedente.", nombre="r3_proc")
ok("Procedencia. El juicio es procedente." in _t25,
   "sin bandera, si el modelo dejó vacía la existencia, «Procedencia.» vuelve a salir (como en main)")

print("\n26 · LA PROCEDENCIA DE LA REVISIÓN POR LO RECURRIDO, SIN BANDERA [siempre] (antes: el 90 en la existencia)")
_t26, _, _, _ = componer3("amparo_revision", {"acto": "VISTOS para resolver el incidente de suspensión relativo al "
                          "juicio de amparo 55/2026, y RESULTANDO… se niega la suspensión definitiva.",
                          "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                          resultandos=[{"titulo": "Trámite", "texto": "registró el juicio de amparo 55/2026."}],
                          nombre="r3_art90")
# C1 (3-oct-2026): la existencia que citaba el 90 (incidente) o el 89
# (sentencia) ya no se escribe; el proemio de la recurrida decide ahora el
# inciso de la procedencia, también sin la bandera.
ok("SEGUNDO. Procedencia. El presente recurso de revisión es procedente, de conformidad con el artículo 81, "
   "fracción I, inciso a), de la Ley de Amparo, en razón de que se impugna la interlocutoria que resolvió sobre "
   "la suspensión definitiva." in _t26 and "artículo 90 de la Ley" not in _t26 and "Existencia" not in _t26,
   "la interlocutoria de suspensión (por su proemio, sin bandera): procedencia con el 81-I-a), sin existencia")
_t26b, _, _, _ = componer3("amparo_revision", {"organo_recurrido": "Juzgado Quinto de Distrito en el Estado de "
                           "Querétaro"}, resultandos=[{"titulo": "Trámite", "texto": "registró el juicio de amparo "
                           "950/2024 y concedió el amparo."}], nombre="r3_art89")
ok("81, fracción I, inciso e), de la Ley de Amparo, en razón de que se impugna una sentencia dictada en la "
   "audiencia constitucional." in _t26b and "artículo 89 de la Ley" not in _t26b,
   "la sentencia, con el 81-I-e) (y ya no el 89 de la existencia)")

print("\n27 · LA QUEJA: DÓNDE SE DICTÓ EL AUTO, LA COMA DEL INCISO Y EL INCISO e) (bandera; Q 337)")
bandera(True)
_pro_q = {"fraccion_97": "I", "inciso_97": "b", "juicio_amparo": "1180/2026", "incidente": True,
          "via_amparo": "indirecto", "cola_97": "en el que se concedió la suspensión provisional",
          "fecha_acto": "trece de julio de dos mil veintiséis",
          "juzgado": "Juzgado Quinto de Distrito en el Estado de Querétaro"}
_t27, _, _, _ = componer3("queja", {"responsable": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                          procesal=_pro_q, nombre="r3_q_inc")
ok("en el incidente de suspensión relativo al juicio de amparo indirecto 1180/2026, en el que se concedió" in _t27
   and "97, fracción I, inciso b), de la Ley de Amparo; 35, fracción III" in _t27,
   "el incidente y la vía en la procedencia; la coma tras el inciso en la competencia")
_t27b, _, _, _ = componer3("queja", {"responsable": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                           procesal=dict(_pro_q, incidente=False, inciso_97="a",
                                         cola_97="por el que desechó la demanda"), nombre="r3_q_juicio")
ok("en el juicio de amparo indirecto 1180/2026, por el que desechó" in _t27b,
   "sin incidente: «en el juicio de amparo indirecto»")
_t27c, _e27c, _, _ = componer3("queja", {"responsable": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                               procesal={"fraccion_97": "I", "inciso_97": "e", "juicio_amparo": "300/2026",
                                         "fecha_acto": "dos de marzo de dos mil veintiséis"}, nombre="r3_q_e")
ok("posterioridad a la sentencia" not in _t27c
   and "en el juicio de amparo indirecto 300/2026, *********." in _t27c
   and any("QUÉ RESOLVIÓ EL AUTO RECURRIDO SALE EN HUECO" in a for a in _e27c.avisos),
   "inciso e) sin el sentido del auto: hueco y aviso, no «posterior a la sentencia» (que no siempre es cierto)")
bandera(False)
_t27d, _, _, _ = componer3("queja", {"responsable": "Jueza Tercero de Distrito en el Estado de Querétaro"},
                           resultandos=[{"titulo": "Interposición", "texto": "contra el auto de trece de marzo de dos "
                                         "mil veinticinco, dictado en el juicio de amparo 742/2025-II, mediante el cual "
                                         "desechó de plano la demanda de amparo."}], nombre="r3_q_viejo")
ok("en el juicio de amparo 742/2025-II," in _t27d and "inciso a) de la Ley de Amparo; 35" in _t27d,
   "sin bandera, la procedencia y la competencia de siempre")

print("\n28 · LA DELEGACIÓN DE LA SCJN Y EL AVISO DE «TODAS LAS AUTORIDADES» (bandera; AR 201, 208, 222, 72)")
bandera(True)
_t28, _e28, _, _ = componer3("amparo_revision", {"quejoso": "Juan Pérez López",
                             "recurrente": "Gobernador del Estado de Querétaro", "papel_recurrente": "autoridad",
                             "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                             "resolvio_a_quo": "concede"},
                             procesal={"juicio_amparo": "950/2024", "materia": "administrativa",
                                       "autoridades": ["Legislatura del Estado de Querétaro",
                                                       "Gobernador del Estado de Querétaro"]},
                             ficha={"tipo": "amparo_revision", "caracter": "autoridad",
                                    "promovente": "Gobernador del Estado de Querétaro",
                                    "acto": {"resolvio": "concede"},
                                    "demanda": {"actos": ["La Ley de Hacienda del Estado de Querétaro, "
                                                          "artículos 90 y 97"]}},
                             regla="oficio", nombre="r3_deleg")
ok(any("DELEGACIÓN de la SCJN" in a and "73, segundo párrafo" in a and "MUY PROBABLEMENTE" in a
       for a in _e28.avisos), "concedido contra una ley y recurre la autoridad: aviso de delegación y del art. 73")
_t28b, _e28b, _, _ = componer3("amparo_revision", {"quejoso": "Juan Pérez López",
                               "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                               "resolvio_a_quo": "niega"},
                               procesal={"juicio_amparo": "950/2024", "autoridades": [
                                   "Director de Ingresos del Municipio de Querétaro", "Tesorera Municipal de Querétaro"]},
                               ficha={"tipo": "amparo_revision", "caracter": "quejoso", "promovente": "Juan Pérez López",
                                      "demanda": {"actos": ["El oficio DI/123/2024 de diez de enero"]}},
                               nombre="r3_todas_no")
ok(not any("DELEGACIÓN" in a for a in _e28b.avisos)
   and not any("NOMBRA A TODAS LAS AUTORIDADES" in a for a in _e28b.avisos),
   "sin ley impugnada, sin aviso de delegación; y «nombra a todas» no sale si el punto no las nombra")
_t28c, _e28c, _, _ = componer3("amparo_revision", {"quejoso": "Juan Pérez López",
                               "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro",
                               "resolvio_a_quo": "sobresee"},
                               procesal={"juicio_amparo": "950/2024", "autoridades": [
                                   "Director de Ingresos del Municipio de Querétaro", "Tesorera Municipal de Querétaro"]},
                               ficha={"tipo": "amparo_revision", "caracter": "quejoso", "promovente": "Juan Pérez López"},
                               calif=("fundado",), estudio=("Es fundado el agravio; procede revocar el sobreseimiento y "
                                                            "conceder el amparo.",), nombre="r3_todas_si")
ok(any("NOMBRA A TODAS LAS AUTORIDADES" in a and "la Tesorera Municipal" in a for a in _e28c.avisos),
   "cuando el punto sí las nombra, el aviso sale (y «la Tesorera», no «el Tesorera»)")

print("\n29 · LA LEGITIMACIÓN CON LA FICHA (bandera; AR 208/72/222, Q 229, RF 2/7/21/6-2026)")
_t29, _, _, _ = componer3("amparo_revision", {"quejoso": "Juan Pérez López",
                          "recurrente": "GOBERNADOR DEL ESTADO DE QUERÉTARO (AUTORIDAD RESPONSABLE)",
                          "papel_recurrente": "autoridad",
                          "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                          procesal={"juicio_amparo": "950/2024"},
                          ficha={"tipo": "amparo_revision", "caracter": "autoridad",
                                 "promovente": "GOBERNADOR DEL ESTADO DE QUERÉTARO (AUTORIDAD RESPONSABLE)",
                                 "representante": "Juan Ruiz Soto", "figura_representante": "delegado"},
                          regla="oficio", nombre="r3_leg_ar")
_l29 = entre(_t29, "Legitimación", "Por cuanto hace")
ok("interpuesto por el Gobernador del Estado de Querétaro, por conducto de su delegado Juan Ruiz Soto" in _l29
   and "AUTORIDAD RESPONSABLE)" not in _l29 and "GOBERNADOR" not in _l29,
   "AR: la autoridad de la ficha, en prosa, sin la etiqueta de rol y con su delegado")
_t29q, _, _, _ = componer3("queja", {"responsable": "Juzgado Tercero de Distrito en el Estado de Querétaro",
                           "quejoso": "Ana López Ruiz y Rosa López Ruiz", "recurrente": "ANA LÓPEZ RUIZ Y OTRA"},
                           procesal={"fraccion_97": "I", "inciso_97": "a", "juicio_amparo": "742/2025"},
                           ficha={"tipo": "queja", "caracter": "quejoso",
                                  "promovente": "Ana López Ruiz y Rosa López Ruiz", "representante": "Pedro Gil Mora",
                                  "figura_representante": "autorizado en términos amplios"}, nombre="r3_leg_q")
_l29q = entre(_t29q, "Legitimación", "Por cuanto hace")
# F6 (quinta ronda): una regla de separador en el resultando y la legitimación —coma si la figura cita
# un artículo o pasa de tres palabras— (Q 172 y 300/2025: «…amplios, Representante,» frente a
# «…amplios Representante,» en el mismo documento).
ok("por Ana López Ruiz y Rosa López Ruiz, por conducto de su autorizado en términos amplios, Pedro Gil Mora" in _l29q
   and "5o., fracción I," in _l29q and "Y OTRA" not in _l29q,
   "Q 229: quejoso y recurrente escritos distinto ya no pierden la representación que la ficha trae")
_PROM_RF = ("TITULAR DE LA UNIDAD JURÍDICA DE LA DELEGACIÓN ESTATAL DEL ISSSTE, EN REPRESENTACIÓN DEL SUBDELEGADO DE "
            "PRESTACIONES (DEMANDADA)")
_t29r, _, _, _ = componer3("revision_fiscal", {"quejoso": _PROM_RF,
                           "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"},
                           procesal={"sala": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa",
                                     "expediente_tfja": "695/25-09-01-7-OT"},
                           ficha={"tipo": "revision_fiscal", "promovente": _PROM_RF}, regla="lfpca_boletin",
                           nombre="r3_leg_rf")
_l29r = entre(_t29r, "Legitimación", "Por cuanto hace")
ok("63, párrafo primero, en relación con el 5o., cuarto párrafo, de la Ley Federal de Procedimiento Contencioso" in _l29r
   and "lo hizo valer el Titular de la Unidad Jurídica de la Delegación Estatal del ISSSTE, en su carácter de unidad "
       "administrativa encargada de la defensa jurídica del Subdelegado de Prestaciones" in _l29r
   and "EN REPRESENTACIÓN" not in _l29r and "DEMANDADA)" not in _l29r,
   "RF 2/7: la unidad y la autoridad demandada partidas, en prosa, sin la etiqueta de rol y con el 5o. LFPCA")
_t29s, _e29s, _, _ = componer3("revision_fiscal", {"quejoso": "Jefe del Departamento de Pensiones",
                               "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"},
                               procesal={"recurrente_unidad": "Jefe del Departamento de Pensiones",
                                         "autoridad_demandada": "Jefe del Departamento de Pensiones de la Delegación"},
                               ficha={"tipo": "revision_fiscal", "promovente": "Jefe del Departamento de Pensiones",
                                      "figura_representante": "Titular de la Unidad de Asuntos Jurídicos"},
                               regla="lfpca_boletin", nombre="r3_leg_rf2")
_l29s = entre(_t29s, "Legitimación", "Por cuanto hace")
ok(("por conducto del Titular de la Unidad de Asuntos Jurídicos" in _l29s
    or "lo hizo valer el Titular de la Unidad de Asuntos Jurídicos" in _l29s)
   and "Pensiones, en su carácter de unidad administrativa encargada de la defensa jurídica del Jefe" not in _l29s,
   "RF 21 y 6/2026: recurre la propia demandada; la legitimación se funda en la unidad de su figura, "
   "no en «X, unidad de defensa jurídica de X»")
_t29t, _e29t, _, _ = componer3("revision_fiscal", {"quejoso": "Jefe del Departamento de Pensiones",
                               "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"},
                               procesal={"recurrente_unidad": "Jefe del Departamento de Pensiones",
                                         "autoridad_demandada": "Jefe del Departamento de Pensiones"},
                               ficha={"tipo": "revision_fiscal", "promovente": "Jefe del Departamento de Pensiones"},
                               regla="lfpca_boletin", nombre="r3_leg_rf3")
_l29t = entre(_t29t, "Legitimación", "Por cuanto hace")
ok((f"lo hizo valer {H}," in _l29t or f"por conducto de {H}" in _l29t) and f"el {H}" not in _l29t
   and "Jefe del Departamento de Pensiones, en su carácter" not in _l29t
   and any("FALTA LA UNIDAD JURÍDICA" in a for a in _e29t.avisos),
   "sin la unidad que firmó el oficio: hueco y aviso, no «X, unidad de defensa jurídica de X»")
_t29u, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "quejoso": "GABRIEL REYES ALAMO"},
                           procesal={"responsable": SALA_QRO, "materia": "civil", "fecha_acto": "veinte de febrero de "
                                     "dos mil veinticinco", "expediente": "479/2024"}, nombre="r3_leg_ad")
ok("promovida por Gabriel Reyes Alamo;" in _t29u and "protege a Gabriel Reyes Alamo" in _t29u.split("R E S U E L V E")[1]
   and "GABRIEL REYES ALAMO" in _t29u.split("Querétaro, Querétaro.")[0],
   "AD 274 de punta a punta: el quejoso en prosa en la legitimación y el resolutivo; en versales sólo en la carátula")

print("\n30 · REVISIÓN FISCAL EXTEMPORÁNEA: SIN AVISOS DE UNA PROCEDENCIA QUE NO SE ESCRIBE (bandera; RF 6/2025)")
_t30, _e30, _, _ = componer3("revision_fiscal", {"quejoso": "Administración Desconcentrada Jurídica de Querétaro «1»",
                             "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"},
                             procesal={"sala": "Sala Regional del Centro II del Tribunal Federal de Justicia "
                                       "Administrativa", "fraccion_63": "V", "expediente_tfja": "695/25-09-01-7-OT"},
                             regla="lfpca_boletin", pres=dt.date(2025, 7, 30), nombre="r3_rf_extemp")
ok("Se desecha por extemporáneo" in _t30 and "Procedencia." not in _t30
   and not any("ARTÍCULO 63" in a and "FRACCIÓN V" in a for a in _e30.avisos),
   "se desecha por extemporáneo: no hay considerando de procedencia ni sus avisos")

print("\n31 · LA REGLA (g) DEL SUPLETORIO TAMBIÉN EN ANTECEDENTES Y ESTUDIO (bandera; revisión de fundamentos)")
try:
    import importlib
    _vp_real = importlib.import_module("verja_procesal")
    _hay_vp = callable(getattr(_vp_real, "revisar", None))
except Exception:
    _hay_vp = False
if not _hay_vp:
    print("   (verja_procesal no está: se omite)")
else:
    _t31, _e31, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO},
                                 procesal={"responsable": SALA_QRO, "materia": "civil",
                                           "fecha_acto": "trece de marzo de dos mil veinticinco",
                                           "expediente": "905/2023"},
                                 estudio=("Es un hecho notorio, en términos del artículo 269 del Código Nacional de "
                                          "Procedimientos Civiles y Familiares, de aplicación supletoria a la Ley de "
                                          "Amparo, que el expediente existe.",), nombre="r3_supl")
    ok(any("SUPLETORIEDAD QUE NO CORRESPONDE A LA SEDE" in a and "ESTUDIO" in a for a in _e31.avisos),
       "Querétaro con el CNPCF en el Estudio: la verja lo acusa (antes sólo miraba el bloque procesal)")

# ═══════════════════════════════════════════════════════════════════════════
# CUARTA RONDA (3-oct-2026): lo que dejó abierto la verificación final contra
# las 32 sentencias reales (rev/verif_final.txt). Dueño A de FIXES_R4: E1, E2,
# E3, E4, E5, E8, E10 y E11. Las piezas de B y D que cambian AHORA se prueban
# con dobles de su firma (el cableado es lo de aquí) y, si ya están, de verdad.
# ═══════════════════════════════════════════════════════════════════════════
import resultandos_por_tipo as rpt4

bandera(True)
_np_real = rpt4.nombre_en_prosa
_NP_MAPA = {
    "Sucesión a Bienes de JOSÉ GARCÍA RUIZ": "Sucesión a Bienes de José García Ruiz",
    ("SOCIEDAD DE PRODUCCIÓN RURAL DE RESPONSABILIDAD ILIMITADA, POR CONDUCTO DE SU CONSEJO DE "
     "ADMINISTRACIÓN, Y JUAN PÉREZ LÓPEZ, POR PROPIO DERECHO"):
        ("Sociedad de Producción Rural de Responsabilidad Ilimitada, por conducto de su consejo de "
         "administración, y Juan Pérez López, por propio derecho"),
}
_np_llamadas = []


def _np_doble(nombre, autoridad=False, con_articulo=False):
    """La firma de E1: `nombre_en_prosa(nombre, autoridad=False, con_articulo=False)`."""
    _np_llamadas.append((nombre, autoridad, con_articulo))
    n = _NP_MAPA.get(nombre)
    if n is None:
        return _np_real(nombre, autoridad=autoridad)
    if con_articulo and n.lower().startswith("sucesión"):
        return "la " + n
    return n


def _np_doble_viejo(nombre, autoridad=False):
    """La firma vieja: siempre pone el artículo de la sucesión."""
    n = _NP_MAPA.get(nombre) or _np_real(nombre, autoridad=autoridad)
    return ("la " + n) if n.lower().startswith("sucesión") else n


print("\n32 · E1: UN SOLO CONVERTIDOR DE NOMBRES A PROSA (bandera; AD 335, AD 552, Q 342, AR 208)")
rpt4.nombre_en_prosa = _np_doble
try:
    ok(dg._en_prosa("Sucesión a Bienes de JOSÉ GARCÍA RUIZ") == "Sucesión a Bienes de José García Ruiz"
       and dg._en_prosa_con_articulo("Sucesión a Bienes de JOSÉ GARCÍA RUIZ")
       == "la Sucesión a Bienes de José García Ruiz",
       "_en_prosa pasa por resultandos_por_tipo.nombre_en_prosa (nombre mezclado: la racha en versales a prosa)")
    _t32, _e32, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO,
                                 "quejoso": "Sucesión a Bienes de JOSÉ GARCÍA RUIZ"},
                                 procesal={"responsable": SALA_QRO, "materia": "civil", "fecha_acto": "veinte de "
                                           "febrero de dos mil veinticinco", "expediente": "479/2024"},
                                 nombre="r4_sucesion")
    _r32 = _t32.split("R E S U E L V E")[1] if "R E S U E L V E" in _t32 else ""
    ok("ni protege a la Sucesión a Bienes de José García Ruiz," in _r32 and "JOSÉ GARCÍA" not in _r32
       and "Sucesión a Bienes de José García Ruiz" in entre(_t32, "Legitimación", "Por cuanto hace"),
       "AD 335: el resolutivo dice «la Sucesión a Bienes de José García Ruiz», como el V I S T O")
    _Q552 = ("SOCIEDAD DE PRODUCCIÓN RURAL DE RESPONSABILIDAD ILIMITADA, POR CONDUCTO DE SU CONSEJO DE "
             "ADMINISTRACIÓN, Y JUAN PÉREZ LÓPEZ, POR PROPIO DERECHO")
    _t32b, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "quejoso": _Q552},
                               procesal={"responsable": SALA_QRO, "materia": "civil", "fecha_acto": "veinte de "
                                         "febrero de dos mil veinticinco", "expediente": "479/2024"},
                               nombre="r4_spr")
    _l32b = entre(_t32b, "Legitimación", "Por cuanto hace")
    _r32b = _t32b.split("R E S U E L V E")[1] if "R E S U E L V E" in _t32b else ""
    ok("por conducto de su consejo de administración" in _l32b and "por conducto de su consejo de "
       "administración" in _r32b and not any(x in _l32b + _r32b for x in ("por Conducto", "de Su ", "Por Propio",
                                                                          "por Propio")),
       "AD 552: «por conducto de su consejo de administración» y «por propio derecho» en minúscula en la "
       "legitimación y el resolutivo")
    rpt4.nombre_en_prosa = _np_doble_viejo
    ok(dg._en_prosa("Sucesión a Bienes de JOSÉ GARCÍA RUIZ") == "Sucesión a Bienes de José García Ruiz"
       and dg._en_prosa_con_articulo("Sucesión a Bienes de JOSÉ GARCÍA RUIZ")
       == "la Sucesión a Bienes de José García Ruiz" and dg._en_prosa("la Titular de la Unidad Jurídica")
       == "la Titular de la Unidad Jurídica",
       "con la firma vieja (sin con_articulo), _en_prosa quita el artículo que el papel no traía y respeta el suyo")
finally:
    rpt4.nombre_en_prosa = _np_real
_t32c, _, _, _ = componer3("amparo_revision", {"quejoso": "Juan Pérez López",
                           "recurrente": "gobernador del Estado de Querétaro", "papel_recurrente": "autoridad",
                           "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                           procesal={"juicio_amparo": "950/2024"},
                           ficha={"tipo": "amparo_revision", "caracter": "autoridad",
                                  "promovente": "gobernador del Estado de Querétaro"},
                           regla="oficio", nombre="r4_gob")
_l32c = entre(_t32c, "Legitimación", "Por cuanto hace")
ok("interpuesto por el Gobernador del Estado de Querétaro" in _l32c and "el gobernador" not in _l32c,
   "AR 208: la autoridad que recurre, como órgano y con su artículo («el Gobernador», no «el gobernador»)")
try:
    import inspect as _insp4
    _np_nueva = "con_articulo" in _insp4.signature(_np_real).parameters
except Exception:
    _np_nueva = False
if not _np_nueva:
    print("   (resultandos_por_tipo.nombre_en_prosa aún sin con_articulo: se omite la pasada de verdad)")
else:
    ok(dg._en_prosa("Sucesión a Bienes de JOSÉ GARCÍA RUIZ") == "Sucesión a Bienes de José García Ruiz"
       and "por conducto de su consejo de administración" in dg._en_prosa(_Q552)
       and "por propio derecho" in dg._en_prosa(_Q552),
       "con la pieza de verdad: el nombre mezclado y las fórmulas de representación en minúscula")

print("\n33 · E3: EL CARGO DEL PONENTE (bandera; Q 342, AD 128/279/552, RF 4/49, AR 239)")
_t33, _e33, _, _ = componer3("queja", {"responsable": "Juzgado Tercero de Distrito en el Estado de Querétaro",
                             "magistrado": ""},
                             procesal={"fraccion_97": "I", "inciso_97": "a", "juicio_amparo": "742/2025"},
                             ficha={"tipo": "queja", "turno": {"fecha": "2025-10-01", "ponente": "Jenica Campos Juárez",
                                                               "titulo": "Magistrada"}}, nombre="r4_pon1")
ok("MAGISTRADA PONENTE: JENICA CAMPOS JUÁREZ." in _t33 and not any("POR OMISIÓN" in a for a in _e33.avisos),
   "Q 342: el cargo del auto de turno (turno.titulo) decide «MAGISTRADA PONENTE» y no hay aviso falso")
_t33b, _e33b, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "magistrado": "Jenica Campos Juárez"},
                               procesal={"responsable": SALA_QRO, "materia": "civil"},
                               ficha={"tipo": "amparo_directo",
                                      "turno": {"fecha": "2025-05-23", "ponente": "Otra Persona Ruiz"},
                                      "returno": [{"fecha": "2025-11-25", "ponente": "Jenica Campos Juárez",
                                                   "titulo": "Magistrada"}]}, nombre="r4_pon2")
ok("MAGISTRADA PONENTE: JENICA CAMPOS JUÁREZ." in _t33b and not any("POR OMISIÓN" in a for a in _e33b.avisos),
   "el nombre tecleado sin cargo toma el del returno (returno.titulo) si es la misma persona")
_t33c, _e33c, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "magistrado": "Jenica Campos Juárez"},
                               procesal={"responsable": SALA_QRO, "materia": "civil"},
                               ficha={"tipo": "amparo_directo", "turno": {"fecha": "2025-05-23",
                                                                          "ponente": "Jenica Campos Juárez"}},
                               nombre="r4_pon3")
ok("MAGISTRADO PONENTE: JENICA CAMPOS JUÁREZ." in _t33c
   and any("«MAGISTRADO PONENTE» VA POR OMISIÓN" in a for a in _e33c.avisos),
   "E3: tecleado sin cargo y sin cargo en el auto: rótulo por omisión y SIEMPRE con aviso")
_t33d, _e33d, _, _ = componer3("amparo_revision", {"magistrado": "",
                               "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                               procesal={"juicio_amparo": "950/2024"},
                               ficha={"tipo": "amparo_revision", "turno": {
                                   "fecha": "2025-05-23",
                                   "ponente": "licenciada Bertha Martínez Vega, secretaria en funciones de Magistrada"}},
                               nombre="r4_pon4")
ok("\nPONENTE: BERTHA MARTÍNEZ VEGA." in _t33d and "LICENCIADA" not in _t33d.split("Querétaro, Querétaro.")[0]
   and dg._ponente_sin_cargo("Magistrada Licenciada Jenica Campos Juárez") == ("Jenica Campos Juárez", "a")
   and dg._cargo_del_titulo("Secretaria en funciones de Magistrada") == "funciones",
   "AR 239: el tratamiento («licenciada») sale del nombre de la carátula; no decide el cargo")
_t33e, _e33e, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "magistrado": ""},
                               procesal={"responsable": SALA_QRO, "materia": "civil",
                                         "ponente": "Jenica Campos Juárez", "ponente_titulo": "Magistrada"},
                               ficha={"tipo": "amparo_directo"}, nombre="r4_pon5")
ok("MAGISTRADA PONENTE: JENICA CAMPOS JUÁREZ." in _t33e and not any("POR OMISIÓN" in a for a in _e33e.avisos),
   "el cargo que expone el compositor (procesal.ponente_titulo) también decide el rótulo")
ok(dg._con_articulo("licenciada Bertha Martínez Vega") == "la licenciada Bertha Martínez Vega"
   and dg._con_articulo("maestra Ana Ruiz") == "la maestra Ana Ruiz"
   and dg._con_articulo("doctora Ana Ruiz") == "la doctora Ana Ruiz",
   "[siempre] «la licenciada», «la maestra», «la doctora» (no «el licenciada»): el artículo del tratamiento")

print("\n34 · E4: «UNA RESOLUCIÓN QUE PUSO FIN AL JUICIO» SÓLO SI LO CONCLUYÓ SIN EL FONDO (bandera; AD 552, 349)")
_t34, _e34, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO}, procesal=_pro_res,
                             ficha={"tipo": "amparo_directo", "acto": {"clase": "resolucion",
                                                                       "sentido": "confirmó la sentencia de primera "
                                                                                  "instancia en el juicio sumario "
                                                                                  "hipotecario"}}, nombre="r4_res_def")
ok("en contra de una resolución definitiva en materia civil, dictada por" in _t34 and "puso fin" not in _t34
   and not any("PUSO FIN AL JUICIO" in a for a in _e34.avisos),
   "AD 552: la apelación resuelta en el fondo es «una resolución definitiva en materia civil», sin aviso")
_t34b, _e34b, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO}, procesal=_pro_res,
                               ficha={"tipo": "amparo_directo", "acto": {"clase": "resolucion",
                                                                         "sentido": "confirmó el auto que desechó la "
                                                                                    "demanda"}}, nombre="r4_res_fin")
ok("en contra de una resolución que puso fin al juicio en materia civil, dictada por" in _t34b
   and any("PUSO FIN AL JUICIO" in a and "170, fracción I" in a and "desechó" in a for a in _e34b.avisos),
   "AD 349: confirmó un desechamiento: «que puso fin al juicio», con aviso de comprobarlo (art. 170, fr. I)")

print("\n35 · E5: ADHESIVO SIN NINGÚN AUTO: NI CONSIDERANDO NI RESOLUTIVO (bandera; AD 274)")
_t35, _e35, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO},
                             procesal={"responsable": SALA_QRO, "materia": "civil", "fecha_acto": "veinte de "
                                       "febrero de dos mil veinticinco", "expediente": "479/2024"},
                             ficha={"tipo": "amparo_directo", "adhesivo": {"quien": "María Gómez Ruiz"}},
                             nombre="r4_adh_sin_auto")
_c35 = considerandos(_t35)
_r35 = _t35.split("R E S U E L V E")[1] if "R E S U E L V E" in _t35 else ""
ok("adhesivo" not in _c35.lower() and "adhesivo" not in _r35.lower() and "ÚNICO." in _r35
   and any("¿HUBO AMPARO ADHESIVO?" in a for a in _e35.avisos),
   "sin admisión ni presentación: sin considerando ni punto del adhesivo, y un aviso que pregunta")
_t35b, _e35b, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO},
                               procesal={"responsable": SALA_QRO, "materia": "civil", "fecha_acto": "veinte de "
                                         "febrero de dos mil veinticinco", "expediente": "479/2024"},
                               ficha={"tipo": "amparo_directo", "adhesivo": {"quien": "María Gómez Ruiz"}},
                               avisos_previos=["¿HUBO AMPARO ADHESIVO? No consta el auto que lo admite."],
                               nombre="r4_adh_sin_auto2")
ok(len([a for a in _e35b.avisos if "¿HUBO AMPARO ADHESIVO?" in a]) == 1,
   "si el compositor ya preguntó «¿HUBO AMPARO ADHESIVO?», no se pregunta dos veces")
_ADH_V = ("EL AMPARO ADHESIVO de MARÍA GÓMEZ RUIZ CONSTA Y NO SE TRATA en los considerandos, los resolutivos: "
          "lleva su resultando, su considerando de legitimación y oportunidad, y su punto resolutivo.")
_ua35 = getattr(dg, "_un_aviso_por_hecho", None)
ok(callable(_ua35) and _ua35([_ADH_V, "¿HUBO AMPARO ADHESIVO? No consta el auto que lo admite."],
                            de_composicion=[_ADH_V]) == ["¿HUBO AMPARO ADHESIVO? No consta el auto que lo admite."],
   "AD 274: la verja ya no pide el considerando y el punto de un adhesivo sin auto; queda la pregunta")

print("\n36 · E8: LA UNIDAD QUE RECURRE EN LA RF, CON SU ARTÍCULO Y SIN VALORES CORTADOS (bandera; RF 4, 21, 49, 2)")
_U36 = "Titular de la Unidad Jurídica de la Oficina de Representación del IMSS en Querétaro"
_SALA = "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa"
_t36, _e36, _, _ = componer3("revision_fiscal", {"quejoso": "la " + _U36, "responsable": _SALA},
                             procesal={"sala": _SALA, "recurrente_unidad": _U36,
                                       "autoridad_demandada": "Titular de la Subdelegación Querétaro del IMSS"},
                             ficha={"tipo": "revision_fiscal", "promovente": "la " + _U36}, regla="lfpca_boletin",
                             nombre="r4_rf_la_titular")
_l36 = entre(_t36, "Legitimación", "Por cuanto hace")
ok("lo hizo valer la Titular de la Unidad Jurídica" in _l36 and "el Titular de la Unidad Jurídica" not in _l36
   and not any("TITULAR» VA POR OMISIÓN" in a for a in _e36.avisos),
   "RF 4/6/49: «la Titular» del papel se conserva en la legitimación (antes volvía «el Titular»)")
_t36b, _e36b, _, _ = componer3("revision_fiscal", {"quejoso": _U36, "responsable": _SALA},
                               procesal={"sala": _SALA, "recurrente_unidad": _U36,
                                         "autoridad_demandada": "Titular de la Subdelegación Querétaro del IMSS"},
                               ficha={"tipo": "revision_fiscal", "promovente": _U36}, regla="lfpca_boletin",
                               nombre="r4_rf_titular")
ok("lo hizo valer el Titular de la Unidad Jurídica" in entre(_t36b, "Legitimación", "Por cuanto hace")
   and len([a for a in _e36b.avisos if "«EL TITULAR» VA POR OMISIÓN" in a]) == 1,
   "«Titular» sin artículo en el papel: «el Titular» por omisión, con UN aviso")
_t36d, _e36d, _, _ = componer3("revision_fiscal", {"quejoso": _U36, "responsable": _SALA},
                               procesal={"sala": _SALA, "recurrente_unidad": _U36,
                                         "recurrente_unidad_con_articulo": "la " + _U36,
                                         "recurrente_la_titular": True,
                                         "autoridad_demandada": "Titular de la Subdelegación Querétaro del IMSS"},
                               ficha={"tipo": "revision_fiscal", "promovente": _U36}, regla="lfpca_boletin",
                               nombre="r4_rf_con_art")
ok("lo hizo valer la Titular de la Unidad Jurídica" in entre(_t36d, "Legitimación", "Por cuanto hace")
   and not any("TITULAR» VA POR OMISIÓN" in a for a in _e36d.avisos),
   "E8: la unidad con el artículo que expone el compositor (recurrente_unidad_con_articulo, recurrente_la_titular)")
_t36c, _e36c, _, _ = componer3("revision_fiscal", {"quejoso": "Jefa de la Unidad Jurídica", "responsable": _SALA},
                           procesal={"sala": _SALA, "recurrente_unidad": "Jefa de la Unidad Jurídica",
                                     "autoridad_demandada": "Subdelegado de Prestaciones"},
                           ficha={"tipo": "revision_fiscal", "promovente": "Jefa de la Unidad Jurídica",
                                  "representante": "Ana Ruiz So…",
                                  "figura_representante": "Subdirectora de lo Contencioso de la Coordinación…"},
                           regla="lfpca_boletin", nombre="r4_rf_trunc")
_l36c = entre(_t36c, "Legitimación", "Por cuanto hace")
ok("…" not in _l36c and "Ana Ruiz So" not in _l36c and "lo hizo valer la Jefa de la Unidad Jurídica" in _l36c
   and len([a for a in _e36c.avisos if a.startswith("DATO TRUNCADO (representante")]) == 1
   and callable(getattr(dg, "_sin_truncar", None)) and dg._sin_truncar("Titular de la Unidad Jurídica...") == ""
   and dg._sin_truncar("Ana Ruiz Soto") == "Ana Ruiz Soto",
   "RF 2 y 6/2026: el representante o la figura cortados con «…» no se firman en la legitimación")

print("\n37 · E10: EL ÓRGANO QUE EL COMPOSITOR DESCARTÓ NO VUELVE EN COMPETENCIA NI EXISTENCIA (bandera; AR 307/448/60)")
_SF = "Sala Familiar del Tribunal Superior de Justicia del Estado de Querétaro"
_AV37 = ("FALTA EL JUZGADO DE DISTRITO QUE DICTÓ LA RESOLUCIÓN RECURRIDA (acto.organo): «Sala Familiar…» es una "
         "autoridad responsable. Va en hueco en «V I S T O», «Trámite del juicio de amparo indirecto», "
         "«Competencia» y «Existencia».")
_t37, _e37, _, _ = componer3("amparo_revision", {"organo_recurrido": _SF, "responsable": _SF},
                             procesal={"juicio_amparo": "950/2024", "juzgado": H, "organo_acto": ""},
                             ficha={"tipo": "amparo_revision", "acto": {"organo": _SF}},
                             avisos_previos=[_AV37], nombre="r4_ar_desc")
_c37 = considerandos(_t37)
# C1 (3-oct-2026): la existencia ya no se escribe en la revisión; antes se
# exigía aquí «que remitió ********* en términos del artículo 89».
ok("Sala Familiar" not in _c37 and f"por {H}, localizad" in _c37 and "Existencia" not in _c37
   and f"el {H}" not in _c37 and f"la {H}" not in _c37,
   "AR 307: la competencia en hueco (sin artículo), no con el órgano descartado; y sin existencia (C1)")
ok(not any("EXISTENCIA SALE CON HUECO" in a and "juzgado" in a for a in _e37.avisos)
   and not any("LLEVA" in a and "HUECO" in a and "órgano que dictó" in a for a in _e37.avisos)
   and _AV37 in _e37.avisos,
   "un aviso para el órgano: el del compositor (acto.organo), sin el de existencia ni el «LLEVA HUECO»")
_od = getattr(dg, "_organo_descartado", None)
ok(callable(_od) and _od({"juzgado_descartado": _SF}) and _od({"organo_recurrido": H})
   and not _od({"juzgado": "Juzgado Quinto de Distrito"}) and dg._con_articulo(H) == H,
   "la marca del compositor (juzgado = HUECO, organo_recurrido = HUECO o juzgado_descartado); al hueco, sin artículo")

print("\n38 · E11: UN HECHO, UN AVISO, VENGA DE DONDE VENGA (bandera; AR 448, Q 229/335)")
_FI_V = ("FECHA IMPOSIBLE: lo reclamado/recurrido es de quince de julio de dos mil veinticinco y el escrito se "
         "presentó el trece de mayo de dos mil veinticinco, antes de que existiera.")
_FI_F = ("FECHA IMPOSIBLE: el recurso se presentó el 13/05/2025 y la resolución recurrida es del 15/07/2025. "
         "Nadie combate una resolución que aún no existe.")
_FI_A = ("FECHA IMPOSIBLE: el recurso de revisión se presentó el 13/05/2025, antes de la audiencia "
         "constitucional (15/07/2025).")
_ua = getattr(dg, "_un_aviso_por_hecho", None)
ok(callable(_ua) and _ua([_FI_V, _FI_F, _FI_A, "OTRO"], de_composicion=[_FI_V]) == [_FI_F, "OTRO"]
   and _ua(["FALTA X (acto.fecha): a", "FALTA Y (acto.fecha): b", "FALTA Z (presentacion): c"])
   == ["FALTA X (acto.fecha): a", "FALTA Z (presentacion): c"]
   and len(_ua(["«EL TITULAR» VA POR OMISIÓN (promovente = «Titular de la X»): a",
                "«EL TITULAR» VA POR OMISIÓN (promovente = «Titular de la X»), en la legitimación: b",
                "«EL TITULAR» VA POR OMISIÓN (figura_representante = «Titular de la Y»): c"])) == 2,
   "la misma fecha imposible en dos formatos, una vez (se queda la de fuera); un «FALTA» y un «EL TITULAR» "
   "por campo")
_v38 = types.ModuleType("verja_procesal")
_v38.revisar_detalle = lambda t, f, d, tp, **k: [{"regla": "j", "aviso": _FI_V}]
_v38.revisar = lambda t, f, d, tp: [_FI_V]
_v38.arreglar = lambda t: (t, [])
_real_v38 = sys.modules.get("verja_procesal")
sys.modules["verja_procesal"] = _v38
try:
    _t38, _e38, _, _ = componer3("amparo_revision", {"organo_recurrido": "Juzgado Quinto de Distrito en el Estado "
                                 "de Querétaro"}, procesal={"juicio_amparo": "950/2024"},
                                 ficha={"tipo": "amparo_revision"}, avisos_previos=[_FI_F], nombre="r4_dup")
finally:
    if _real_v38 is None:
        sys.modules.pop("verja_procesal", None)
    else:
        sys.modules["verja_procesal"] = _real_v38
ok(len([a for a in _e38.avisos if a.startswith("FECHA IMPOSIBLE")]) == 1 and _FI_F in _e38.avisos,
   "AR 448: validar y la verja decían la misma fecha imposible; en el documento, un aviso")

print("\n39 · E2: EL PLURAL DEL COMPOSITOR LLEGA A LA LEGITIMACIÓN Y A LA CARÁTULA (bandera; AD 552, RF 7)")
_leg_real = ta.legitimacion_de
_fil_real = ta.filas_caratula
_kw39, _fil39 = [], []


def _leg_doble(tipo, parte="", representante="", hueco=H, figura="", moral=None, papel="",
               autoridad_demandada="", recurre_el_quejoso=None, avisos=None, plural=None, genero="",
               hay_norma=None):
    """La firma de B en la cuarta ronda: `plural`, `genero` y `hay_norma`."""
    _kw39.append({"tipo": tipo, "plural": plural, "hay_norma": hay_norma, "parte": parte, "genero": genero})
    return _leg_real(tipo, parte, representante, hueco, figura=figura, moral=moral, papel=papel,
                     autoridad_demandada=autoridad_demandada, recurre_el_quejoso=recurre_el_quejoso,
                     avisos=avisos)


def _fil_doble(tipo, datos):
    _fil39.append((datos or {}).get("plural"))
    return _fil_real(tipo, datos)


ta.legitimacion_de, ta.filas_caratula = _leg_doble, _fil_doble
try:
    componer3("amparo_directo", {"responsable": SALA_QRO, "quejoso": "JUAN PÉREZ LÓPEZ Y PEDRO PÉREZ LÓPEZ"},
              procesal={"responsable": SALA_QRO, "materia": "civil", "plural": True}, nombre="r4_plural")
    _kw_ad = [k for k in _kw39 if k["tipo"] == "amparo_directo"]
    ok(bool(_kw_ad) and _kw_ad[-1]["plural"] is True and True in _fil39,
       "AD 552: procesal.plural va a legitimacion_de(plural=) y a filas_caratula (datos['plural'])")
    _kw39.clear()
    componer3("amparo_revision", {"quejoso": "Juan Pérez López", "recurrente": "Gobernador del Estado de Querétaro",
                                  "papel_recurrente": "autoridad",
                                  "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
              procesal={"juicio_amparo": "950/2024", "plural": False, "plural_quejoso": True},
              ficha={"tipo": "amparo_revision", "caracter": "autoridad",
                     "promovente": "Gobernador del Estado de Querétaro",
                     "demanda": {"actos": ["La Ley de Hacienda del Estado de Querétaro, artículos 90 y 97"]}},
              regla="oficio", nombre="r4_norma")
    _kw_ar = [k for k in _kw39 if k["tipo"] == "amparo_revision" and k["hay_norma"] is not None]
    ok(bool(_kw_ar) and _kw_ar[-1]["hay_norma"] is True and _kw_ar[-1]["plural"] is False,
       "AR 208/72 (E12): hay_norma viaja a la legitimación de la autoridad; el número es el de quien recurre "
       "(procesal.plural), no el de la quejosa (plural_quejoso)")
    _kw39.clear()
    componer3("amparo_revision", {"quejoso": "Juan Pérez López", "recurrente": "gobernador del Estado de Querétaro",
                                  "papel_recurrente": "autoridad",
                                  "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
              procesal={"juicio_amparo": "950/2024"},
              ficha={"tipo": "amparo_revision", "caracter": "autoridad",
                     "promovente": "gobernador del Estado de Querétaro"}, regla="oficio", nombre="r4_gob2")
    ok(any(k["parte"] == "el Gobernador del Estado de Querétaro" for k in _kw39),
       "AR 208 (E1): la autoridad llega a legitimacion_de por la puerta única, como órgano y con su artículo")
finally:
    ta.legitimacion_de, ta.filas_caratula = _leg_real, _fil_real


# ═══════════════════════════════════════════════════════════════════════════
# QUINTA RONDA (3-oct-2026): lo que dejó la segunda verificación final
# (rev/verif_final2.txt). Dueño A de FIXES_R5: F1, F3, F6, el plural de la
# carátula del AR/queja, los efectos con lo reclamado y la carátula de la queja.
# ═══════════════════════════════════════════════════════════════════════════
print("\n40 · LA CARÁTULA DEL AR: EL PLURAL DE QUIEN RECURRE NO ES EL DE LA QUEJOSA (bandera; AR 208, regresión)")
bandera(True)
_t40, _, _, _ = componer3("amparo_revision", {"quejoso": "JUAN y PEDRO, ambos de apellidos PÉREZ LÓPEZ",
                          "recurrente": "Gobernador del Estado de Querétaro", "papel_recurrente": "autoridad",
                          "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                          procesal={"juicio_amparo": "950/2024", "plural": False, "plural_quejoso": True},
                          ficha={"tipo": "amparo_revision", "caracter": "autoridad",
                                 "promovente": "Gobernador del Estado de Querétaro"}, regla="oficio", nombre="r5_car1")
_car40 = _t40.split("MAGISTRAD")[0]
ok("QUEJOSOS: JUAN Y PEDRO, AMBOS DE APELLIDOS PÉREZ LÓPEZ." in _car40
   and "AUTORIDAD RESPONSABLE Y RECURRENTE: GOBERNADOR DEL ESTADO DE QUERÉTARO." in _car40
   and "AUTORIDADES RESPONSABLES" not in _car40,
   "AR 208: dos quejosos y un Gobernador que recurre → «QUEJOSOS:» y «AUTORIDAD RESPONSABLE Y RECURRENTE:»")
_t40b, _, _, _ = componer3("amparo_revision", {"quejoso": "María López Díaz",
                           "recurrente": "Gobernador del Estado de Querétaro y Secretario de Finanzas del Estado "
                                         "de Querétaro", "papel_recurrente": "autoridad",
                           "organo_recurrido": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                           procesal={"juicio_amparo": "950/2024", "plural": True, "plural_quejoso": False},
                           ficha={"tipo": "amparo_revision", "caracter": "autoridad",
                                  "promovente": "Gobernador del Estado de Querétaro y Secretario de Finanzas del "
                                                "Estado de Querétaro"}, regla="oficio", nombre="r5_car2")
ok("AUTORIDADES RESPONSABLES" in _t40b.split("MAGISTRAD")[0],
   "y al revés: una quejosa y dos autoridades que recurren → el renglón del recurrente en plural")
_d40 = dg._con_lo_procesal({"plural": None}, {"plural": False, "plural_quejoso": True}, "queja")
ok(_d40.get("plural") is False and _d40.get("plural_quejoso") is True
   and dg._con_lo_procesal({}, {"plural": True, "plural_quejoso": False}, "amparo_directo").get("plural") is False,
   "queja y AR: cada número en su clave; el AD, como antes (plural = el de la quejosa)")

print("\n41 · F3: LA QUEJA SIN FRACCIÓN DEL 97 NO SUPONE LA I NI «INDIRECTO» (bandera; Q 335)")
_ORG41 = "Juzgado Segundo de lo Civil del Distrito Judicial de Querétaro"
_PRO41 = {"fraccion_97": "", "via_amparo": H, "inciso_97": "b", "juicio_amparo": "1180/2026", "juzgado": _ORG41,
          "cola_97": "en el que dejó sin efectos la suspensión del acto reclamado",
          "fecha_acto": "seis de octubre de dos mil veinticinco"}
_FI41 = {"tipo": "queja", "caracter": "quejoso", "promovente": "Juan Pérez López",
         "acto": {"organo": _ORG41, "expediente": "1180/2026"}}
_t41, _e41, _, _ = componer3("queja", {"responsable": _ORG41}, procesal=_PRO41, ficha=_FI41, calif=("fundado",),
                             nombre="r5_q_sinfr")
_c41 = entre(_t41, "Competencia.", "SEGUNDO.")
_p41 = entre(_t41, "Procedencia.", "TERCERO.")
ok(f"dictado en un juicio de amparo {H}, por el Juzgado Segundo" in _c41
   and f"97, fracción {H}, inciso b), de la Ley de Amparo; 35" in _c41
   and "indirecto" not in _c41 and "fracción I," not in _c41,
   "competencia: la fracción y la vía en hueco, sin «fracción I» ni «amparo indirecto»")
ok(f"artículo 97, fracción {H}, inciso b), de la Ley de Amparo" in _p41
   and f"en el juicio de amparo {H} 1180/2026, en el que dejó sin efectos" in _p41 and "indirecto" not in _p41,
   "procedencia: «en el juicio de amparo ********* 1180/2026», sin la vía supuesta")
ok(not any("DOS (artículo 98" in a or "SUSPENSIÓN PROVISIONAL O DE PLANO" in a for a in _e41.avisos)
   and len([a for a in _e41.avisos if "(fraccion_97)" in a]) == 1
   and any("(fraccion_97)" in a and "COMPETENCIA" in a and "PROCEDENCIA" in a for a in _e41.avisos),
   "sin el aviso del plazo de dos días; un aviso del dato que nombra los dos considerandos")
_R41 = _t41.split("R E S U E L V E")[1]
ok("al juzgado de origen" not in _R41 and "al órgano que dictó el auto recurrido" in _R41,
   "el testimonio va a quien dictó el auto, sin afirmar que es un juzgado")
_AV_FR = ("FALTA LA FRACCIÓN DEL ARTÍCULO 97 (fraccion_97): sale del auto recurrido y de la vía. La vía del "
          "juicio va en hueco en el V I S T O y en «Interposición del recurso de queja».")
_AV_INC = "FALTA EL INCISO DEL ARTÍCULO 97 (inciso_97): sale de qué proveyó el auto recurrido."
_t41b, _e41b, _, _ = componer3("queja", {"responsable": _ORG41}, procesal=dict(_PRO41, inciso_97="", cola_97=""),
                               ficha=_FI41, avisos_previos=[_AV_FR, _AV_INC], nombre="r5_q_sinfr2")
ok(len([a for a in _e41b.avisos if "(fraccion_97)" in a]) == 1 and _AV_FR in _e41b.avisos
   and len([a for a in _e41b.avisos if "INCISO DEL ARTÍCULO 97" in a]) == 1 and _AV_INC in _e41b.avisos
   and not any("FRACCIÓN I," in a for a in _e41b.avisos),
   "E11: con los avisos del compositor (fracción e inciso), ninguno repetido y ninguno afirma «FRACCIÓN I»")
ok(dg._fraccion_97_de({}, {"acto": {"via": "indirecto"}}) == "I"
   and dg._fraccion_97_de({"via_amparo": "directo"}, {}) == "II"
   and dg._fraccion_97_de({"fraccion_97": "Fracción II"}, {}) == "II"
   and dg._fraccion_97_de({"fraccion_97": "", "via_amparo": H}, {"acto": {"organo": _ORG41}}) == "",
   "la fracción consta por la ficha o por la vía; el nombre del órgano no la decide aquí")
_t41c, _, _, _ = componer3("queja", {"responsable": "Juzgado Quinto de Distrito en el Estado de Querétaro"},
                           procesal=dict(_PRO41, fraccion_97="I", via_amparo="indirecto",
                                         juzgado="Juzgado Quinto de Distrito en el Estado de Querétaro"),
                           ficha=_FI41, nombre="r5_q_fr1")
ok("97, fracción I, inciso b), de la Ley de Amparo; 35" in _t41c
   and "en el juicio de amparo indirecto 1180/2026" in _t41c,
   "con la fracción que consta, la cadena de siempre")

print("\n42 · LOS EFECTOS NOMBRAN LO RECLAMADO (bandera; AD 349, AD 128 con laudo)")
for _cl42, _ar42, _seg42 in (("resolucion", "la resolución reclamada", "2. Emitir una nueva en la que"),
                             ("laudo", "el laudo reclamado", "2. Emitir otro en el que")):
    _t42, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO},
                              procesal={"responsable": SALA_QRO, "materia": "civil", "clase": _cl42,
                                        "acto_reclamado": _ar42},
                              ficha={"tipo": "amparo_directo", "acto": {"clase": _cl42}}, calif=("fundado",),
                              estudio=("Es fundado el concepto de violación; procede conceder el amparo.",),
                              nombre=f"r5_ef_{_cl42}")
    _ef42 = entre(_t42, "Efectos.", "R E S U E L V E")
    ok(f"1. Dejar insubsistente {_ar42}." in _ef42 and _seg42 in _ef42 and "sentencia reclamada" not in _ef42,
       f"efectos sin escribir: «Dejar insubsistente {_ar42}» y «{_seg42[3:]}»")
bandera(False)
_t42v, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO}, calif=("fundado",),
                           estudio=("Es fundado el concepto de violación; procede conceder el amparo.",),
                           nombre="r5_ef_viejo")
ok("1. Dejar insubsistente la sentencia reclamada." in _t42v
   and "2. Emitir una nueva en la que reitere" in _t42v, "sin la bandera, el esqueleto de siempre")
bandera(True)

print("\n43 · F6: LA FIGURA DEL RESOLUTIVO POR LA PUERTA ÚNICA (bandera; AD 274 y 456 con nombres)")
import promovente as pv
ok(pv.por_conducto("Gabriel Reyes Alamo", "Pedro Gil Mora", "autorizado en términos amplios del artículo 12")
   == "Gabriel Reyes Alamo, por conducto de su autorizado en términos amplios del artículo 12 de la Ley de "
      "Amparo, Pedro Gil Mora"
   and pv.por_conducto("Constructora del Bajío, S.A. de C.V.", "Juan Ruiz Gómez", "Apoderado Legal")
   == "Constructora del Bajío, S.A. de C.V., por conducto de su apoderado legal Juan Ruiz Gómez"
   and pv.por_conducto("GDE Trading Company, S.A. de C.V.", "Alondra Zúñiga Gutiérrez", "")
   == "GDE Trading Company, S.A. de C.V., por conducto de su representante legal Alondra Zúñiga Gutiérrez",
   "por_conducto: «artículo 12» con su ley y la coma; la figura genérica en minúscula; sin figura, la de siempre")
_t43, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "quejoso": "Gabriel Reyes Alamo",
                          "representante": "Pedro Gil Mora",
                          "figura_representante": "autorizado en términos amplios del artículo 12"},
                          procesal={"responsable": SALA_QRO, "materia": "civil"}, ficha={"tipo": "amparo_directo"},
                          nombre="r5_res_aut")
_R43 = _t43.split("R E S U E L V E")[1]
ok("protege a Gabriel Reyes Alamo, por conducto de su autorizado en términos amplios del artículo 12 de la Ley "
   "de Amparo, Pedro Gil Mora, contra" in _R43 and "artículo 12 Pedro" not in _R43,
   "AD 274: el resolutivo cita el 12 con su ley y separa el nombre con coma, como el resultando")
_t43b, _, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO, "quejoso": "Constructora del Bajío, S.A. de C.V.",
                           "representante": "Juan Ruiz Gómez"},
                           procesal={"responsable": SALA_QRO, "materia": "civil"},
                           ficha={"tipo": "amparo_directo", "figura_representante": "Apoderado Legal"},
                           nombre="r5_res_fig")
ok("Constructora del Bajío, S.A. de C.V., por conducto de su apoderado legal Juan Ruiz Gómez, contra"
   in _t43b.split("R E S U E L V E")[1],
   "AD 456: sin figura en el encargo, la de la ficha, la misma que la legitimación (no «representante legal»)")

print("\n44 · F1: LA FORMA DE NOTIFICACIÓN QUE NO CONSTA NO SE AFIRMA (bandera; Q 335/229/261/342, AD 274/335)")
_FI44 = {"tipo": "queja", "caracter": "quejoso", "promovente": "Juan Pérez López"}
_PR44 = {"fraccion_97": "I", "inciso_97": "a", "juicio_amparo": "742/2025"}
_RQ44 = {"responsable": "Juzgado Quinto de Distrito en el Estado de Querétaro"}
_t44, _e44, _tb44, _ = componer3("queja", _RQ44, procesal=_PR44, ficha=dict(_FI44, forma_notificacion=""),
                                 nombre="r5_f1_sin")
_o44 = entre(_t44, "Por cuanto hace", "\n")
ok("se notificó a la parte recurrente el veinte de marzo de dos mil veinticinco y surtió efectos al día hábil "
   "siguiente, conforme al artículo 31, fracción II" in _o44 and "de manera personal" not in _o44
   and any(a.startswith("LA FORMA DE NOTIFICACIÓN NO CONSTA (forma_notificacion)") for a in _e44.avisos)
   and "Forma no consta" in _tb44 and "de manera personal" not in _tb44,
   "Q 335: sin forma en la ficha, ni «de manera personal» en el párrafo ni en la tabla, y el aviso")
_t44b, _e44b, _, _ = componer3("queja", _RQ44, procesal=_PR44,
                               ficha=dict(_FI44, forma_notificacion="personal",
                                          fuentes={"forma_notificacion": "omision"}), nombre="r5_f1_omision")
ok("de manera personal" not in entre(_t44b, "Por cuanto hace", "\n")
   and len([a for a in _e44b.avisos if "FORMA DE NOTIFICACIÓN NO CONSTA" in a]) == 1,
   "la «personal» que sólo viene del formulario (fuente «omision») tampoco se afirma; un aviso")
_t44c, _e44c, _, _ = componer3("queja", _RQ44, procesal=_PR44,
                               ficha=dict(_FI44, forma_notificacion="personal",
                                          fuentes={"forma_notificacion": "secretario"}), nombre="r5_f1_secr")
ok("veinticinco de manera personal y surtió efectos" in entre(_t44c, "Por cuanto hace", "\n")
   and not any("FORMA DE NOTIFICACIÓN NO CONSTA" in a for a in _e44c.avisos),
   "la que declaró el secretario se escribe, sin aviso")
_t44d, _e44d, _, _ = componer3("queja", dict(_RQ44, papel_recurrente="autoridad"), procesal=_PR44,
                               ficha={"tipo": "queja", "caracter": "autoridad",
                                      "promovente": "Juez Quinto de Distrito en el Estado de Querétaro",
                                      "forma_notificacion": "oficio",
                                      "fuentes": {"forma_notificacion": "omision_autoridad"}},
                               regla="oficio", nombre="r5_f1_aut")
_o44d = entre(_t44d, "Por cuanto hace", "\n")
ok("el veinte de marzo de dos mil veinticinco y surtió efectos el mismo día, conforme al artículo 31, fracción I"
   in _o44d and "por oficio" not in _o44d
   and any("FORMA DE NOTIFICACIÓN NO CONSTA" in a for a in _e44d.avisos),
   "D2 sin constancia: «surtió efectos el mismo día» con el 31, fr. I, sin «por oficio»")
_t44e, _e44e, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO},
                               procesal={"responsable": SALA_QRO, "materia": "civil"},
                               ficha={"tipo": "amparo_directo", "forma_notificacion": ""}, nombre="r5_f1_ad")
ok("la sentencia reclamada se notificó a la parte quejosa el veinte de marzo de dos mil veinticinco y surtió "
   "efectos" in _t44e and "de manera personal" not in entre(_t44e, "Por cuanto hace", "\n")
   and not any("de manera personal" in a for a in _e44e.avisos),
   "AD 274/335: «se notificó a la parte quejosa el X y surtió efectos…», como los engroses; tampoco en los avisos")
ok(dg._sin_forma_de_aviso("el considerando dice que la notificación de manera personal surtió efectos al día",
                          f0.computar(dt.date(2025, 3, 20), dt.date(2025, 4, 8), regla="personal", plazo=15,
                                      tipo_asunto="amparo_directo"))
   == "el considerando dice que la notificación surtió efectos al día",
   "y el aviso del precepto («SE RECOMIENDA CITAR…») no repite la forma, aunque su pieza aún no sepa omitirla")
ok(dg._sin_forma_de_notificacion("el seis de octubre de manera personal y surtió efectos al día",
                                 f0.computar(dt.date(2025, 3, 20), dt.date(2025, 3, 27), regla="personal", plazo=5,
                                             tipo_asunto="queja"))
   == "el seis de octubre y surtió efectos al día",
   "la red propia quita la forma si la pieza del párrafo aún no sabe hacerlo")
bandera(False)
_t44v, _, _, _ = componer3("queja", _RQ44, procesal=_PR44, ficha=dict(_FI44, forma_notificacion=""),
                           nombre="r5_f1_viejo")
ok("veinticinco de manera personal y surtió efectos" in _t44v, "sin la bandera, el párrafo de siempre")
bandera(True)

print("\n45 · LA CARÁTULA DE LA QUEJA YA NO LLEVA EL ÓRGANO (C3, 3-oct-2026; antes: Q 337, sin «con residencia en esta ciudad»)")
_ORG45 = ("Juzgado Quinto de Distrito en Materia de Amparo Civil, Administrativo y de Trabajo y de Juicios "
          "Federales en el Estado de Querétaro, con residencia en esta ciudad")
_t45, _, _, _ = componer3("queja", {"responsable": _ORG45, "organo_recurrido": _ORG45},
                          procesal=dict(_PR44, juzgado=_ORG45), ficha=_FI44, nombre="r5_resid")
_car45 = _t45.split("MAGISTRAD")[0]
# C3: David, «lo más práctico»; 0 de 8 quejas del banco llevan el órgano en el
# rubro (va en el V I S T O y en la competencia). Antes se comprobaba que el
# renglón del órgano saliera sin «con residencia en esta ciudad».
ok("JUICIOS FEDERALES EN EL ESTADO DE QUERÉTARO" not in _car45 and "RESIDENCIA" not in _car45
   and "ÓRGANO QUE DICTÓ" not in _car45,
   "C3: la carátula de la queja ya no lleva el renglón del órgano (ni, con él, la residencia)")
_ORG45b = "Juzgado Cuarto de Distrito en el Estado de Sonora, con residencia en Ciudad Obregón"
_t45b, _, _, _ = componer3("queja", {"responsable": _ORG45b, "organo_recurrido": _ORG45b},
                           procesal=dict(_PR44, juzgado=_ORG45b), ficha=_FI44, nombre="r5_resid2")
ok("CIUDAD OBREGÓN" not in _t45b.split("MAGISTRAD")[0].upper(),
   "C3: tampoco el juzgado foráneo va en el rubro de la queja")

print("\n46 · F5: EL ARTÍCULO DEL PAPEL EN LOS CARGOS DE DOS GÉNEROS (bandera; AD 128, RF 4/49)")
ok(dg._con_la_titular("Oficial Mayor del Municipio de Cadereyta", "la Oficial Mayor del Municipio de Cadereyta")
   == "la Oficial Mayor del Municipio de Cadereyta"
   and dg._con_la_titular("Titular de la Unidad Jurídica", "la titular de la Unidad Jurídica")
   == "la Titular de la Unidad Jurídica"
   and dg._con_la_titular("Oficial Mayor del Municipio", "Oficial Mayor del Municipio") == "Oficial Mayor del Municipio",
   "«la Oficial Mayor» del papel se conserva como «la Titular»; sin artículo en el papel, como viene")

# ═══════════════════════════════════════════════════════════════════════════
# SEXTA RONDA (3-oct-2026): lo que David respondió hoy, con su visto bueno para
# todos los usuarios. C1 la revisión sin existencia y con su procedencia; C2 la
# fracción del 35; C3/C4 las carátulas de la queja y de la revisión fiscal; C5
# el punto del amparo que nombra acto y autoridad; C6 los asuntos relacionados
# que marca el secretario; C7 el surtimiento agrario por la sede.
# ═══════════════════════════════════════════════════════════════════════════
print("\n47 · C1: EL AMPARO EN REVISIÓN SIN EXISTENCIA, CON «SEGUNDO. Procedencia.» [siempre]")
_JZ47 = "Juzgado Quinto de Distrito en el Estado de Querétaro"


def _rotulos(t):
    """Los rótulos de los considerandos, en orden: «PRIMERO. Competencia.»…"""
    import re as _re
    return [m.group(0) for m in _re.finditer(
        r"^(?:PRIMERO|SEGUNDO|TERCERO|CUARTO|QUINTO|SEXTO|SÉPTIMO|OCTAVO|NOVENO|DÉCIMO)\. [^.]{2,90}\.",
        considerandos(t), _re.M)]


for _band47 in (True, False):
    bandera(_band47)
    _q47 = "con bandera" if _band47 else "sin bandera"
    _t47, _e47, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega"},
                                 procesal={"juicio_amparo": "950/2024", "juzgado": _JZ47,
                                           "clase_recurrida": "sentencia"} if _band47 else None,
                                 ficha={"tipo": "amparo_revision"} if _band47 else None,
                                 nombre=f"r6_ar_sent_{_band47}")
    _r47 = _rotulos(_t47)
    ok(_r47[:3] == ["PRIMERO. Competencia.", "SEGUNDO. Procedencia.",
                    "TERCERO. Legitimación y oportunidad para interponer el recurso de revisión."]
       and "SEGUNDO. Procedencia. El presente recurso de revisión es procedente, de conformidad con el artículo "
           "81, fracción I, inciso e), de la Ley de Amparo, en razón de que se impugna una sentencia dictada en la "
           "audiencia constitucional." in _t47,
       f"AR sentencia ({_q47}): Competencia, Procedencia (81-I-e), Legitimación, en ese orden: {_r47[:3]}")
    ok("Existencia" not in _t47 and "existencia de la" not in _t47
       and not any("EXISTENCIA" in a for a in _e47.avisos),
       f"AR sentencia ({_q47}): ni el considerando de existencia ni su aviso de hueco")
    ok("35, fracción V y 210" in _t47 and "35, fracción II" not in _t47,
       f"C2 ({_q47}): la sentencia de audiencia con la fracción V del 35 de la Ley Orgánica")
bandera(True)
_t47s, _e47s, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega",
                               "acto": "VISTOS para resolver el incidente de suspensión relativo al juicio de "
                                       "amparo 55/2026; se resuelve sobre la suspensión definitiva."},
                               procesal={"juicio_amparo": "55/2026", "juzgado": _JZ47,
                                         "clase_recurrida": "interlocutoria_suspension"},
                               ficha={"tipo": "amparo_revision"}, nombre="r6_ar_susp_a")
ok("SEGUNDO. Procedencia. El presente recurso de revisión es procedente, de conformidad con el artículo 81, "
   "fracción I, inciso a), de la Ley de Amparo, en razón de que se impugna la interlocutoria que resolvió sobre "
   "la suspensión definitiva." in _t47s and "81, fracción I, inciso a), y 84" in _t47s
   and "35, fracción II" in _t47s and "Existencia" not in _t47s,
   "AR suspensión: la procedencia con el inciso a), el mismo que la competencia (y la fracción II del 35)")
_t47b, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega",
                           "acto": "VISTOS para resolver el incidente de suspensión relativo al juicio de amparo "
                                   "55/2026, en el que se modifica la interlocutoria que negó la suspensión "
                                   "definitiva."},
                           procesal={"juicio_amparo": "55/2026", "juzgado": _JZ47,
                                     "clase_recurrida": "interlocutoria_suspension"},
                           ficha={"tipo": "amparo_revision"}, nombre="r6_ar_susp_b")
# LOS DOS SUPUESTOS DEL INCISO b) (revisión de normas, 3-oct-2026): antes se
# comprobaba «la resolución que modificó o revocó el acuerdo…», falso cuando lo
# recurrido NEGÓ modificarlo.
ok("81, fracción I, inciso b), de la Ley de Amparo, en razón de que se impugna la resolución que se pronunció "
   "sobre la modificación o revocación del acuerdo en que se concedió o negó la suspensión definitiva." in _t47b
   and "81, fracción I, inciso b), y 84" in _t47b,
   "AR suspensión que modificó lo resuelto: inciso b) en la procedencia y en la competencia")
_t47d, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "sobresee"},
                           procesal={"juicio_amparo": "55/2026", "juzgado": _JZ47,
                                     "clase_recurrida": "auto_sobreseimiento"},
                           ficha={"tipo": "amparo_revision"}, nombre="r6_ar_auto")
ok("SEGUNDO. Procedencia. El presente recurso de revisión es procedente, de conformidad con el artículo 81, "
   "fracción I, inciso d), de la Ley de Amparo, en razón de que se impugna el auto que sobreseyó en el juicio "
   "fuera de la audiencia constitucional." in _t47d and "35, fracción II" in _t47d
   and "35, fracción V" not in _t47d and "Existencia" not in _t47d,
   "AR auto de sobreseimiento: procedencia con el inciso d) y competencia con la fracción II del 35 (C2)")
_t47c, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega"},
                           procesal={"juicio_amparo": "55/2026", "juzgado": _JZ47,
                                     "clase_recurrida": "reposicion_constancias"},
                           ficha={"tipo": "amparo_revision"}, nombre="r6_ar_repos")
ok("inciso c), de la Ley de Amparo, en razón de que se impugna la resolución que decidió el incidente de "
   "reposición de constancias de autos." in _t47c,
   "AR reposición de constancias: inciso c)")
_t47x, _e47x, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega"},
                               procesal={"juicio_amparo": "950/2024", "juzgado": _JZ47},
                               ficha={"tipo": "amparo_revision"}, pres=dt.date(2025, 5, 27),
                               nombre="r6_ar_extemp")
ok("Procedencia." not in _t47x and "Existencia" not in _t47x and "Se desecha el recurso de revisión" in _t47x,
   "AR extemporáneo: no se declara procedente lo que se desecha")

print("\n48 · C2: NINGUNA FÓRMULA DEL BANCO FUNDA LA SENTENCIA DE AUDIENCIA EN LA FRACCIÓN II DEL 35")
import json as _json
import re as _re48
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "banco_plantillas.json"),
          encoding="utf-8") as _f48:
    _bp = _json.load(_f48)
_rv = _bp["revision"]["considerandos"][0]
_textos48 = dict([("plantilla", _rv["plantilla"])] + [(v["id"], v["texto"]) for v in _rv["variantes"]])
_mal48 = [i for i, t in _textos48.items()
          if _re48.search(r"inciso\s+e\)", t) and _re48.search(r"(?<!\d)35,\s*fracci[óo]n\s+II\b", t)]
ok(not _mal48, f"inciso e) del 81 nunca con «35, fracción II» (rv-c1-constitucionalidad lo traía): {_mal48}")
ok("35, fracción V y 210" in _textos48["rv-c1-constitucionalidad"]
   and "35, fracción V," in _textos48["rv-c1-antigua"] and "35, fracción V y 210" in _textos48["plantilla"],
   "la delegación de la SCJN (lo «remitido por la Suprema Corte») y la fórmula antigua, con la V")
ok(all("35, fracción II" in _textos48[k] for k in ("rv-c1-suspension", "rv-c1-suspension-definitiva")),
   "las de la suspensión (incisos a y b) se quedan con la II")
_otras48 = {k: _re48.findall(r"(?<!\d)35,\s*fracci[óo]n\s+([IVX]+)", _bp[k]["considerandos"][0].get("plantilla", ""))
            for k in ("amparo-directo", "queja", "revision-fiscal")}
ok(_otras48 == {"amparo-directo": ["I"], "queja": ["III"], "revision-fiscal": ["VI"]},
   f"amparo directo I, queja III, revisión fiscal VI: {_otras48}")

print("\n49 · C5: EL PUNTO DEL AMPARO NOMBRA ACTO Y AUTORIDAD (bandera)")
import copy as _copy
_ramas_antes = _copy.deepcopy(ta.RAMAS_REVISION)
_AUT49 = ["Segunda Sala Civil del Tribunal Superior de Justicia del Estado de Querétaro",
          "Juzgado Mixto de Primera Instancia de Tolimán"]
_t49, _e49, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega"},
                             procesal={"juicio_amparo": "950/2024", "juzgado": _JZ47, "clase_recurrida": "sentencia",
                                       "actos": ["La sentencia de trece de marzo de dos mil veinticinco, dictada en "
                                                 "el toca civil 374/2024."],
                                       "autoridades": _AUT49, "resultando_demanda": "primero"},
                             ficha={"tipo": "amparo_revision"}, nombre="r6_pt_uno")
_rs49 = _t49.split("R E S U E L V E")[1]
ok("SEGUNDO. La Justicia de la Unión no ampara ni protege a Juan Pérez López, respecto del acto que reclamó de la "
   "Segunda Sala Civil del Tribunal Superior de Justicia del Estado de Querétaro y del Juzgado Mixto de Primera "
   "Instancia de Tolimán, consistente en la sentencia de trece de marzo de dos mil veinticinco, dictada en el toca "
   "civil 374/2024, por las razones expuestas en el último considerando de la resolución recurrida." in _rs49
   and "{actos_del_amparo}" not in _t49,
   "confirma y niega: un acto corto, «respecto del acto que reclamó de… y del…, consistente en…»")
ok(any("RESOLUTIVO DE REVISIÓN, rama «confirma_niega»" in a for a in _e49.avisos),
   "el aviso que pide comprobar el resolutivo se conserva")
_t49b, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "concede"},
                           procesal={"juicio_amparo": "950/2024", "juzgado": _JZ47, "clase_recurrida": "sentencia",
                                     "actos": ["La orden de clausura.", "Su ejecución."],
                                     "autoridades": ["Director de Ingresos del Municipio de Querétaro"],
                                     "resultando_demanda": "primero", "plural_quejoso": True},
                           ficha={"tipo": "amparo_revision"}, nombre="r6_pt_varios")
_rs49b = _t49b.split("R E S U E L V E")[1]
ok("respecto de los actos que reclamaron del Director de Ingresos del Municipio de Querétaro, precisados en el "
   "resultando primero de esta ejecutoria" in _rs49b,
   f"confirma y concede: varios actos remiten al resultando primero; «reclamaron» con quejosos plurales: {_rs49b[:300]}")
_t49c, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "sobresee"},
                           procesal={"juicio_amparo": "950/2024", "juzgado": _JZ47, "clase_recurrida": "sentencia",
                                     "actos": ["La orden de clausura."],
                                     "autoridades": ["Director de Ingresos del Municipio de Querétaro"]},
                           ficha={"tipo": "amparo_revision"}, nombre="r6_pt_sob")
_rs49c = _t49c.split("R E S U E L V E")[1]
ok("SEGUNDO. Se sobresee en el presente juicio de amparo, promovido por Juan Pérez López, contra el acto que "
   "reclamó del Director de Ingresos del Municipio de Querétaro, consistente en la orden de clausura, por las "
   "razones expuestas en la resolución recurrida." in _rs49c,
   "confirma el sobreseimiento: con «contra» y sin punto final dentro del acto")
_t49d, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega"},
                           procesal={"juicio_amparo": "950/2024", "juzgado": _JZ47, "clase_recurrida": "sentencia",
                                     "actos": [H], "autoridades": _AUT49},
                           ficha={"tipo": "amparo_revision"}, nombre="r6_pt_hueco")
ok("respecto de los actos precisados en la resolución recurrida, por las razones" in _t49d.split("R E S U E L V E")[1],
   "con los actos en hueco: la remisión genérica, no un punto con «*********»")
_t49e, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "sobresee"},
                           nombre="r6_pt_sin_ficha")
# SIN ACTOS, LA COLA DE SIEMPRE (revisión AR, 3-oct-2026): antes se comprobaba
# «…en la resolución recurrida, por las razones expuestas en la resolución
# recurrida.», con «resolución recurrida» dos veces.
ok("contra los actos precisados en la resolución recurrida y por las razones expuestas en la misma."
   in _t49e.split("R E S U E L V E")[1] and "{actos_del_amparo}" not in _t49e
   and "en la resolución recurrida, por las razones expuestas en la resolución recurrida" not in _t49e,
   "sin ficha: la remisión genérica de siempre («contra» en el sobreseimiento), sin repetir «resolución recurrida»")
_t49s, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega"},
                           procesal={"juicio_amparo": "55/2026", "juzgado": _JZ47,
                                     "clase_recurrida": "interlocutoria_suspension",
                                     "actos": ["La orden de clausura."],
                                     "autoridades": ["Director de Ingresos del Municipio de Querétaro"]},
                           ficha={"tipo": "amparo_revision"}, nombre="r6_pt_susp")
ok("SEGUNDO. Se niega la suspensión definitiva solicitada por Juan Pérez López, respecto del acto que reclamó del "
   "Director de Ingresos del Municipio de Querétaro, consistente en la orden de clausura, por las razones"
   in _t49s.split("R E S U E L V E")[1],
   "en el incidente, la suspensión definitiva (no el amparo) con el acto nombrado")
bandera(False)
_t49f, _, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega"},
                           procesal={"juicio_amparo": "950/2024", "actos": ["La orden de clausura."],
                                     "autoridades": ["Director de Ingresos del Municipio de Querétaro"]},
                           ficha={"tipo": "amparo_revision"}, nombre="r6_pt_sin_bandera")
ok("respecto de los actos precisados en la resolución recurrida" in _t49f.split("R E S U E L V E")[1]
   and "{actos_del_amparo}" not in _t49f,
   "sin la bandera: la remisión genérica, nunca el marcador")
bandera(True)
_t49g, _e49g, _, _ = componer3("amparo_revision", {"quejoso": "Comercializadora Ejemplo, S.A. de C.V.",
                               "recurrente": "Director de Ingresos del Municipio de Querétaro",
                               "papel_recurrente": "autoridad", "organo_recurrido": _JZ47,
                               "resolvio_a_quo": "niega",
                               "reasuncion": {"quien_recurre": "autoridad", "sobresee_ademas": True,
                                              "adhesiva": False}},
                               procesal={"juicio_amparo": "950/2024", "juzgado": _JZ47, "clase_recurrida": "sentencia",
                                         "actos": ["La orden de clausura.", "El acta de visita."],
                                         "autoridades": ["Director de Ingresos del Municipio de Querétaro"]},
                               ficha={"tipo": "amparo_revision"}, regla="oficio", nombre="r6_pt_firme")
_rs49g = _t49g.split("R E S U E L V E")[1]
ok("PRIMERO. Queda firme el sobreseimiento" in _rs49g
   and "TERCERO. La Justicia de la Unión no ampara ni protege a Comercializadora Ejemplo, S.A. de C.V., respecto "
       "de los actos que reclamó del Director de Ingresos" in _rs49g
   and any(a.startswith("EL PUNTO DEL AMPARO NOMBRA LOS ACTOS") for a in _e49g.avisos),
   f"con un sobreseimiento firme y los actos de la demanda en el punto, el aviso pide quitar lo sobreseído: "
   f"{_rs49g[:400]}")
ok(ta.RAMAS_REVISION == _ramas_antes,
   "llenar el punto no toca el catálogo (RAMAS_REVISION sigue con su {actos_del_amparo} para el asunto siguiente)")

print("\n50 · C6: ASUNTOS RELACIONADOS QUE MARCA EL SECRETARIO (bandera)")
_PRO50 = {"expediente": "905/2023", "toca": "374/2024", "toca_en_prosa": "toca 374/2024",
          "fecha_acto": "trece de marzo de dos mil veinticinco", "fecha_acto_iso": "2025-03-13",
          "responsable": SALA_QRO, "organo_acto": SALA_QRO, "materia": "civil"}
_ENC50 = {"responsable": SALA_QRO, "encabezado": "AMPARO DIRECTO CIVIL: 456/2025"}
_t50a, _, _, _ = componer3("amparo_directo", _ENC50, procesal=dict(_PRO50),
                           ficha={"tipo": "amparo_directo", "numero": "456/2025", "materia": "civil"},
                           nombre="r6_rel_sin")
_REL50 = [{"tipo": "revision_fiscal", "numero": "33/2024", "estado": "misma_sesion"}]
_t50, _e50, _, _ = componer3("amparo_directo", _ENC50, procesal=dict(_PRO50, relacionados=_REL50),
                             ficha={"tipo": "amparo_directo", "numero": "456/2025", "materia": "civil",
                                    "relacionados": _REL50}, nombre="r6_rel_ad_rf")
_car50 = _t50.split("MAGISTRAD")[0].splitlines()
ok(_car50[:2] == ["AMPARO DIRECTO CIVIL: 456/2025.", "RELACIONADO CON EL RECURSO DE REVISIÓN FISCAL 33/2024."],
   f"carátula: «RELACIONADO CON…» justo debajo del encabezado: {_car50[:3]}")
_d50 = Document(os.path.join(_TMP, "r6_rel_ad_rf.docx"))
_p50 = next((p_ for p_ in _d50.paragraphs if p_.text.startswith("RELACIONADO CON")), None)
ok(_p50 is not None and _p50.runs[0].text == "RELACIONADO CON " and _p50.runs[0].bold
   and not _p50.runs[1].bold,
   "el renglón pesa como rótulo («RELACIONADO CON» en negrita) y el asunto va en redonda")
# EL ENCABEZADO DE CADA PÁGINA TAMBIÉN LO DICE (3-oct-2026; AD 456, AD 552 y Q 24 del banco).
_h50 = [p_.text for s_ in _d50.sections for h_ in (s_.header, s_.even_page_header) for p_ in h_.paragraphs]
_h50a = [p_.text for s_ in Document(os.path.join(_TMP, "r6_rel_sin.docx")).sections
         for h_ in (s_.header, s_.even_page_header) for p_ in h_.paragraphs]
ok(_h50 and all(h_ == "AMPARO DIRECTO CIVIL: 456/2025 RELACIONADO CON EL RECURSO DE REVISIÓN FISCAL 33/2024"
                for h_ in _h50 if h_.strip())
   and not any("RELACIONADO" in h_ for h_ in _h50a),
   f"el encabezado de página, par e impar, lleva «RELACIONADO CON…»; sin relacionados, no: {_h50[:2]}")
ok("RELACIONADO CON" not in _t50a and "Conexidad" not in _t50a and "Hecho notorio" not in _t50a,
   "sin relacionados marcados: ni renglón ni considerando (nunca conexidad en automático)")
_r50, _r50a = _rotulos(_t50), _rotulos(_t50a)
ok("CUARTO. Conexidad." in _r50 and _r50.index("CUARTO. Conexidad.") == 3
   and _r50[2].startswith("TERCERO. Legitimación"),
   f"el considerando va justo después de la legitimación y antes de la dispensa: {_r50[:5]}")
_ADH50 = {"quien": "María Gómez Ruiz", "presentacion": "2025-04-30", "admision": "2025-05-05"}
_t50h, _, _, _ = componer3("amparo_directo", _ENC50,
                           procesal=dict(_PRO50, relacionados=_REL50, adhesivo=dict(_ADH50)),
                           ficha={"tipo": "amparo_directo", "numero": "456/2025", "materia": "civil",
                                  "relacionados": _REL50, "adhesivo": dict(_ADH50)}, nombre="r6_rel_adh")
_r50h = _rotulos(_t50h)
_i50c = next((i for i, r in enumerate(_r50h) if r.endswith("Conexidad.")), -1)
ok(_i50c > 0 and "adhesivo" in _r50h[_i50c - 1].lower()
   and _r50h[_i50c + 1].endswith("Acto reclamado y conceptos de violación."),
   f"con amparo adhesivo, la conexidad va después de su legitimación y justo antes de la dispensa: {_r50h[2:6]}")
ok("CUARTO. Conexidad. Con vista en la conexión que guarda el presente juicio de amparo directo civil 456/2025 "
   "con el recurso de revisión fiscal 33/2024, ambos del índice de este Tribunal Colegiado de Circuito, se "
   "resuelven en la misma sesión, con fundamento en el artículo 64 de la Ley Federal de Procedimiento Contencioso "
   "Administrativo, porque en ambos se impugna la misma sentencia." in _t50,
   "AD ↔ RF en la misma sesión: el artículo 64 de la LFPCA (corpus AD 469/2024)")
_est50 = [r for r in _r50 if r.endswith("Estudio.")]
_est50a = [r for r in _r50a if r.endswith("Estudio.")]
ok(_est50 and _est50a and len(_r50) == len(_r50a) + 1
   and dg._ORDINALES.index(_est50[0].split(".")[0]) == dg._ORDINALES.index(_est50a[0].split(".")[0]) + 1,
   f"el ordinal del estudio se corre solo: {_est50a[:1]} → {_est50[:1]}")
_REL50q = [{"tipo": "amparo_revision", "numero": "298/2025", "estado": "resuelto"}]
_t50q, _e50q, _, _ = componer3("queja", dict(_RQ44, encabezado="RECURSO DE QUEJA CIVIL: 24/2026"),
                               procesal=dict(_PR44, relacionados=_REL50q),
                               ficha=dict(_FI44, numero="24/2026", materia="civil", relacionados=_REL50q),
                               nombre="r6_rel_q_ar")
ok("RELACIONADO CON EL AMPARO EN REVISIÓN CIVIL 298/2025." in _t50q.split("MAGISTRAD")[0]
   and "Hecho notorio. Se invoca como hecho notorio, en términos del artículo 88 del Código Federal de "
       "Procedimientos Civiles, de aplicación supletoria a la Ley de Amparo, la ejecutoria dictada por este "
       "Tribunal Colegiado de Circuito en el amparo en revisión civil 298/2025, relacionado con el presente "
       "asunto." in _t50q,
   "Q con un AR ya resuelto: el rubro y el hecho notorio con el CFPC 88 (Querétaro)")
_r50q = _rotulos(_t50q)
ok([r.split(". ", 1)[1] for r in _r50q[:5]]
   == ["Competencia.", "Procedencia.", "Legitimación y oportunidad.", "Hecho notorio.",
       "Trascripción innecesaria del auto recurrido y agravios."],
   f"en la queja, después de la procedencia y la legitimación: {_r50q[:5]}")
_t50c, _, _, _ = componer3("queja", dict(_RQ44, encabezado="RECURSO DE QUEJA CIVIL: 24/2026",
                                         tribunal=CDMX, ciudad="Ciudad de México"),
                           procesal=dict(_PR44, relacionados=_REL50q),
                           ficha=dict(_FI44, numero="24/2026", materia="civil", relacionados=_REL50q),
                           nombre="r6_rel_q_cdmx")
ok("en términos del artículo 269 del Código Nacional de Procedimientos Civiles y Familiares" in _t50c
   and "artículo 88 del Código Federal" not in _t50c,
   "en la Ciudad de México, el hecho notorio con el 269 del CNPCF")
_REL50m = [{"tipo": "amparo_directo", "numero": "452/2025", "estado": "misma_sesion"},
           {"tipo": "queja", "numero": "12/2026", "estado": "resuelto"},
           {"tipo": "amparo_directo", "numero": "456/2025", "estado": "misma_sesion"}]
_t50m, _, _, _ = componer3("amparo_directo", _ENC50, procesal=dict(_PRO50),
                           ficha={"tipo": "amparo_directo", "numero": "456/2025", "materia": "civil",
                                  "relacionados": _REL50m}, nombre="r6_rel_mixto")
_cm50 = entre(_t50m, "Asuntos relacionados.", "QUINTO.")
_rub50m = _t50m.split("MAGISTRAD")[0]
ok("RELACIONADO CON EL AMPARO DIRECTO CIVIL 452/2025 Y CON EL RECURSO DE QUEJA CIVIL 12/2026." in _rub50m
   and _rub50m.count("456/2025") == 1
   and "a fin de evitar el dictado de resoluciones contradictorias" in _cm50
   and "Se invoca como hecho notorio" in _cm50 and "artículo 64" not in _cm50,
   "de los dos estados: «Asuntos relacionados.», sin el 64 (no hay RF), y el propio asunto fuera del rubro")
bandera(False)
_t50f, _, _, _ = componer3("amparo_directo", _ENC50, procesal=dict(_PRO50, relacionados=_REL50),
                           ficha={"tipo": "amparo_directo", "relacionados": _REL50}, nombre="r6_rel_sin_bandera")
ok("RELACIONADO CON" not in _t50f and "Conexidad" not in _t50f, "sin la bandera, nada de relacionados")
bandera(True)

print("\n51 · C3/C4: LAS CARÁTULAS DE LA QUEJA Y DE LA REVISIÓN FISCAL (bandera)")
_t51q, _, _, _ = componer3("queja", dict(_RQ44, organo_recurrido=_RQ44["responsable"]), procesal=_PR44,
                           ficha=_FI44, nombre="r6_car_q")
_c51q = _t51q.split("MAGISTRAD")[0]
ok("Y RECURRENTE: JUAN PÉREZ LÓPEZ." in _c51q and "ÓRGANO" not in _c51q and "JUZGADO QUINTO" not in _c51q,
   f"queja: «… Y RECURRENTE», sin el renglón del órgano: {_c51q.splitlines()[:3]}")
_t51r, _, _, _ = componer3("revision_fiscal", {"quejoso": "Administración Desconcentrada Jurídica de Querétaro «1»",
                           "responsable": "Sala Regional del Centro II del Tribunal Federal de Justicia Administrativa",
                           "tercero": "Comercializadora Ejemplo, S.A. de C.V."},
                           regla="lfpca_boletin", nombre="r6_car_rf")
_c51r = _t51r.split("MAGISTRAD")[0]
ok("RECURRENTE: ADMINISTRACIÓN DESCONCENTRADA JURÍDICA DE QUERÉTARO «1»." in _c51r
   and "AUTORIDAD RECURRENTE" not in _c51r and "SALA RESPONSABLE" not in _c51r and "SALA REGIONAL" not in _c51r
   and "PARTE ACTORA: COMERCIALIZADORA EJEMPLO, S.A. DE C.V." in _c51r,
   f"revisión fiscal: «RECURRENTE:» y «PARTE ACTORA:», sin «SALA RESPONSABLE»: {_c51r.splitlines()[:4]}")
_pe51 = dg.prompt_estructura({"tipo_asunto": "queja", "quejoso": "Juan Pérez López",
                              "responsable": "Juzgado Quinto de Distrito en el Estado de Querétaro"})
ok("órgano que dictó el auto recurrido: Juzgado Quinto de Distrito en el Estado de Querétaro" in _pe51,
   "la hoja de datos del prompt sigue dando el órgano de la queja como dato (no va en el rubro)")

print("\n52 · C7: EL SURTIMIENTO AGRARIO SEGÚN LA SEDE (bandera)")
_TUA = "Tribunal Unitario Agrario del Distrito 42"
_PRO52 = {"materia": "agraria", "responsable": _TUA, "organo_acto": _TUA}
_t52q, _e52q, _, _ = componer3("amparo_directo", {"responsable": _TUA, "materia": "agraria"},
                               procesal=dict(_PRO52), ficha={"tipo": "amparo_directo", "materia": "agraria"},
                               nombre="r6_agr_qro")
_o52q = entre(_t52q, "Por cuanto hace", "\n")
ok("conforme al artículo 321 del Código Federal de Procedimientos Civiles, de aplicación supletoria en materia "
   "agraria" in _o52q and any("321" in a for a in _e52q.avisos),
   "fuera de la Ciudad de México: el 321 del CFPC, con su aviso")
_t52c, _, _, _ = componer3("amparo_directo", {"responsable": _TUA, "materia": "agraria", "tribunal": CDMX,
                                              "ciudad": "Ciudad de México"},
                           procesal=dict(_PRO52), ficha={"tipo": "amparo_directo", "materia": "agraria"},
                           regla="cnpcf_personal", nombre="r6_agr_cdmx")
_o52c = entre(_t52c, "Por cuanto hace", "\n")
ok("artículo 227, fracción I, del Código Nacional de Procedimientos Civiles y Familiares" in _o52c
   and "321" not in _o52c and "surtió efectos el mismo día" in _o52c,
   "en la Ciudad de México, con la regla del Código Nacional: el 227, fr. I, y nunca el 321")
_t52s, _, _, _ = componer3("amparo_directo", {"responsable": _TUA, "materia": "agraria", "tribunal": "", "ciudad": ""},
                           procesal=dict(_PRO52),
                           ficha={"tipo": "amparo_directo", "materia": "agraria",
                                  "sede": {"tribunal": CDMX, "ciudad": "Ciudad de México"}},
                           nombre="r6_agr_sede_ficha")
ok("321" not in entre(_t52s, "Por cuanto hace", "\n"),
   "la sede de la ficha (Ciudad de México) llega a `surtimiento_nacional`: no se cita el 321")
_c52 = types.SimpleNamespace(regla=types.SimpleNamespace(clave="cnpcf_personal"))
ok(dg._forma_de_omision({"tipo": "amparo_directo", "fuentes": {"forma_notificacion": "omision"}}, _c52) == "omision",
   "la personal del Código Nacional puesta por omisión tampoco se afirma como forma (F1)")
# LA PERSONAL DEL CÓDIGO FEDERAL ELEGIDA A PROPÓSITO (integración, 3-oct-2026).
_t52f, _e52f, _, _ = componer3("amparo_directo", {"responsable": _TUA, "materia": "agraria", "tribunal": CDMX,
                                                  "ciudad": "Ciudad de México"},
                               procesal=dict(_PRO52), ficha={"tipo": "amparo_directo", "materia": "agraria"},
                               regla="cfpc_personal", nombre="r6_agr_cdmx_cfpc")
_o52f = entre(_t52f, "Por cuanto hace", "\n")
ok("surtió efectos al día hábil siguiente, conforme al artículo 321 del Código Federal de Procedimientos Civiles, "
   "de aplicación supletoria en materia agraria" in _o52f and "Código Nacional" not in _o52f
   and not any(a.startswith("EL SURTIMIENTO SE FUNDÓ EN EL ARTÍCULO 321") for a in _e52f.avisos),
   "en la Ciudad de México, la del Código Federal elegida a propósito: el 321, sin el aviso del transitorio"
   + (f" — {_o52f[:300]}" if "321" not in _o52f else ""))
_t52g, _e52g, _, _ = componer3("amparo_directo", {"responsable": _TUA, "materia": "agraria"},
                               procesal=dict(_PRO52), ficha={"tipo": "amparo_directo", "materia": "agraria"},
                               regla="cfpc_personal", nombre="r6_agr_qro_cfpc")
ok(entre(_t52g, "Por cuanto hace", "\n").count("artículo 321") == 1
   and any(a.startswith("EL SURTIMIENTO SE FUNDÓ EN EL ARTÍCULO 321") for a in _e52g.avisos),
   "fuera de ella (lo que el desplegable propone): el 321 una vez y el aviso del transitorio, como la genérica")
_c52f = types.SimpleNamespace(regla=types.SimpleNamespace(clave="cfpc_personal"))
ok(dg._forma_de_omision({"tipo": "amparo_directo", "fuentes": {"forma_notificacion": "omision"}}, _c52f) == "omision",
   "elegir la regla del Código Federal no hace constar la forma (F1)")

print("\n53 · LA RESPONSABLE DEL FORMULARIO NO LE GANA AL ÓRGANO EN LA QUEJA NI EN LA RF (integración)")
# page.tsx manda en `responsable` la ordenadora leída del auto de admisión
# también en la queja y la revisión fiscal, que desde C3 y C4 no tienen el
# renglón del órgano en la carátula (`tipos_asunto.responsable_es_el_organo`).
_ORD53 = "Segunda Sala Civil del Tribunal Superior de Justicia del Estado de Querétaro"
_JQ53 = "Juzgado Quinto de Distrito en el Estado de Querétaro"
_t53a, _e53a, _, _ = componer3("queja", {"responsable": _ORD53},
                               procesal=dict(_PR44, organo_acto=_JQ53, juzgado=_JQ53, responsable=_JQ53),
                               ficha=_FI44, nombre="r6_q_ord_con_org")
_cons53a = considerandos(_t53a)
ok(_JQ53 in _cons53a and "Segunda Sala Civil" not in _t53a,
   "Q con el órgano en la ficha y la ordenadora en el formulario: la competencia y la procedencia nombran el "
   "juzgado; la ordenadora no aparece")
_t53b, _e53b, _, _ = componer3("queja", {"responsable": _ORD53}, procesal=dict(_PR44), ficha=_FI44,
                               nombre="r6_q_ord_sin_org")
ok("Segunda Sala Civil" not in _t53b and dg.HUECO in entre(_t53b, "Competencia.", "\n"),
   "Q sin órgano en la ficha: la ordenadora no ocupa su sitio en la competencia (hueco)"
   + (f" — {entre(_t53b, 'Competencia.', chr(10))[:300]}" if "Segunda Sala Civil" in _t53b else ""))
_t53c, _, _, _ = componer3("queja", {"responsable": _ORD53, "organo_recurrido": _JQ53}, ficha=_FI44,
                           nombre="r6_q_ord_sin_pro")
ok("Segunda Sala Civil" not in considerandos(_t53c) and _JQ53 in considerandos(_t53c),
   "Q sin lo procesal (sin compositor): manda el órgano que leyó la ficha de partes, no la ordenadora")
_pe53 = dg.prompt_estructura({"tipo_asunto": "queja", "quejoso": "Juan Pérez López", "responsable": _ORD53})
ok("órgano que dictó el auto recurrido" not in _pe53,
   "la hoja de datos del prompt no da la ordenadora como el órgano de la queja")
_t53r, _, _, _ = componer3("revision_fiscal", {"quejoso": "Administración Desconcentrada Jurídica de Querétaro «1»",
                                               "responsable": "Administración Desconcentrada Jurídica de Querétaro «1»"},
                           regla="lfpca_boletin", procesal={"sala": "Sala Regional del Centro II del Tribunal "
                                                                     "Federal de Justicia Administrativa"},
                           nombre="r6_rf_ord")
ok("Sala Regional del Centro II" in considerandos(_t53r),
   "RF: la Sala de la ficha manda sobre lo que llegue en `responsable`")

print("\n54 · REVISIÓN ADVERSARIAL (3-oct-2026): EL DOCUMENTO (bandera)")
# AR 72/2025 (fichas4): el Legislativo responsable y sólo el acto de aplicación en
# la ficha. El punto se lo atribuía al Legislativo; ahora remite y avisa.
_A54 = ["el pago de los derechos registrales correspondientes a la inscripción de una transmisión de propiedad en "
        "ejecución de fideicomiso y extinción parcial de éste sobre un inmueble"]
_L54 = ["Legislatura y Gobernador, ambos del Estado de Querétaro"]
_t54a, _e54a, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "concede"},
                               procesal={"juicio_amparo": "950/2024", "juzgado": _JZ47, "clase_recurrida": "sentencia",
                                         "actos": _A54, "autoridades": _L54},
                               ficha={"tipo": "amparo_revision"}, nombre="r7_ar72_norma")
_res54a = _t54a.split("R E S U E L V E")[1]
ok("consistente en el pago" not in _res54a
   and "ampara y protege a Juan Pérez López, en términos del último considerando de la resolución recurrida" in _res54a
   and any(a.startswith("LA DEMANDA RECLAMA UNA NORMA QUE LA FICHA NO TRAE ENTRE LOS ACTOS") for a in _e54a.avisos),
   "AR 72: el punto no nombra el acto de aplicación como si fuera todo lo reclamado; remite y avisa"
   + (f" — {_res54a[:400]}" if "consistente en el pago" in _res54a else ""))
# C1 QUITÓ LA EXISTENCIA DEL AR: el aviso de la perífrasis ya no la nombra (en el AD, sí).
_PERI = [{"titulo": "Trámite", "texto": "Se reclamó el acto reclamado precisado en antecedentes."}]
_t54b, _e54b, _, _ = componer3("amparo_revision", {"organo_recurrido": _JZ47, "resolvio_a_quo": "niega"},
                               resultandos=_PERI, nombre="r7_ar_perifrasis")
_t54c, _e54c, _, _ = componer3("amparo_directo", {"responsable": SALA_QRO}, resultandos=_PERI,
                               nombre="r7_ad_perifrasis")
_pb = [a for a in _e54b.avisos if a.startswith("UN RESULTANDO EVADE EL DATO")]
_pc = [a for a in _e54c.avisos if a.startswith("UN RESULTANDO EVADE EL DATO")]
ok(_pb and not any("existencia" in a for a in _pb) and _pc and any("considerando de existencia" in a for a in _pc),
   "la perífrasis en un resultando: en el AR el aviso no nombra la existencia; en el AD, sí")
# EL 64 DE LA LFPCA SE AVISA CADA VEZ QUE SE CITA.
ok(any(a.startswith("EL CONSIDERANDO DE RELACIONADOS CITA EL ARTÍCULO 64 DE LA LFPCA") for a in _e50.avisos)
   and not any("ARTÍCULO 64 DE LA LFPCA" in a for a in _e54c.avisos),
   "AD↔RF: el considerando cita el 64 y un aviso pide comprobar que impugnen la misma sentencia; sin él, nada")
# LA HOJA DE DATOS DE LA QUEJA DE LA FRACCIÓN II CONSERVA EL ÓRGANO (la Sala).
_pe54 = dg.prompt_estructura({"tipo_asunto": "queja", "quejoso": "Juan Pérez López", "responsable": _ORD53,
                              "procesal": {"fraccion_97": "II"}})
ok("órgano que dictó el auto recurrido: " + _ORD53 in _pe54,
   "queja de la fracción II: la hoja de datos da la Sala responsable como el órgano que dictó el auto")
# AGRARIO EN LA CIUDAD DE MÉXICO CON LA DEL CÓDIGO NACIONAL: el aviso del transitorio
# sale en el documento aunque nadie la haya cambiado (la propuso la pantalla).
_t54d, _e54d, _, _ = componer3("amparo_directo", {"responsable": _TUA, "materia": "agraria", "tribunal": CDMX,
                                                  "ciudad": "Ciudad de México"},
                               procesal=dict(_PRO52), ficha={"tipo": "amparo_directo", "materia": "agraria"},
                               regla="cnpcf_personal", nombre="r7_agr_cdmx_aviso")
ok(sum(1 for a in _e54d.avisos if a == f0.aviso_cnpcf_agrario("cnpcf_personal")) == 1
   and "pues conforme al artículo 227, fracción I" in entre(_t54d, "Por cuanto hace", "\n"),
   "agrario en la Ciudad de México con la regla del Código Nacional: el aviso del transitorio, una vez, y el "
   "párrafo dice lo que dice el 227-I")
# LA REVISIÓN FISCAL CON LA ACTORA ADHESIVA: un renglón, con su «(ACTORA)».
_t54e, _, _, _ = componer3("revision_fiscal", {"quejoso": "Titular de la Unidad Jurídica del Instituto de Ejemplo",
                                               "adherente": "María Pérez López", "tercero": "María Pérez López"},
                           regla="lfpca_boletin", procesal={"sala": "Sala Regional del Centro II del Tribunal "
                                                                    "Federal de Justicia Administrativa"},
                           nombre="r7_rf_adh_actora")
_car54 = _t54e.split("MAGISTRAD")[0]
ok("RECURRENTE ADHESIVO: MARÍA PÉREZ LÓPEZ (ACTORA)." in _car54 and "PARTE ACTORA" not in _car54,
   f"RF: la actora que se adhiere, una vez y con su carácter: {_car54.splitlines()[:6]}")

bandera(False)
print()
if FALLOS:
    print(f"FALLAN {len(FALLOS)}: " + " · ".join(FALLOS))
    raise SystemExit(1)
print("RESULTADO: TODAS LAS COMPROBACIONES PASAN")
