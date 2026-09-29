# -*- coding: utf-8 -*-
"""LA TESIS CITADA POR SU REGISTRO Y EL ARTÍCULO CITADO EN PLURAL — AR 631/2025, 28-sep-2026.

David, sobre el proyecto de la rama de coherencia (f4f88b7): «en este último
proyecto ni tesis, ni citas de artículos. Está mal». Tres anuncios «Sirve de
apoyo el criterio de registro N:» —la forma que enseña el propio prompt— salieron
como prosa, sin rubro ni texto; y de «los artículos 2284 y 2294», «los artículos
14 y 17» y «el artículo 49 de la legislación procesal civil» sólo bajaron al pie
el 2284 y el 14. Todo sin modelo: los párrafos son los del proyecto malo.
"""
import os
import re
import sys

os.environ.pop("ESTUDIO_PROMPT", None)

import docx  # noqa: E402

import documento_generado as dg  # noqa: E402
import fase6_estudio as f6  # noqa: E402
import fase6_rag as fr_  # noqa: E402

fallos = []
# (Con getattr, para que contra el código de antes FALLE comprobación a
# comprobación en vez de reventar al importar.)
SEP = getattr(dg, "SEP_NOTA", "\u2029")


def ok(c, nota):
    print(f"  {'OK ' if c else 'MAL'} {nota}")
    if not c:
        fallos.append(nota)


def _largo(semilla, n=100):
    return " ".join([semilla] * n)


TESIS = [
    {"registro": "2026918", "tipo": "Jurisprudencia", "instancia": "Primera Sala", "obligatoria": True,
     "rubro": "COSA JUZGADA Y SUS EFECTOS DIRECTO Y REFLEJO. DIFERENCIAS Y REQUISITOS PARA SU ACTUALIZACIÓN.",
     "localizacion": "[J]; 11a. Época; 1a. Sala; Gaceta S.J.F.; Libro 28, Agosto de 2023; Tomo II; Pág. 1157",
     "texto": "Hechos: " + _largo("identidad de cosa, causa y personas", 20)
              + "\n\nCriterio jurídico: efecto directo y efecto reflejo."},
    {"registro": "168958", "tipo": "Jurisprudencia", "instancia": "Pleno", "obligatoria": True,
     "rubro": "COSA JUZGADA. SUS LÍMITES OBJETIVOS Y SUBJETIVOS.",
     "localizacion": "[J]; 9a. Época; Pleno; S.J.F. y su Gaceta; Tomo XXVIII, Septiembre de 2008; Pág. 590",
     "texto": _largo("los límites subjetivos alcanzan a los causahabientes", 15)},
    {"registro": "171925", "tipo": "Jurisprudencia", "instancia": "Segunda Sala", "obligatoria": True,
     "rubro": "REVISIÓN EN AMPARO. RUBRO DE PRUEBA SOBRE LA REASUNCIÓN DE JURISDICCIÓN DEL TRIBUNAL REVISOR.",
     "localizacion": "ficha de prueba", "texto": _largo("reasume jurisdicción y estudia los conceptos", 15)},
    {"registro": "188480", "tipo": "Tesis aislada", "instancia": "Tribunales Colegiados de Circuito",
     "rubro": "SUSTITUCIÓN PROCESAL. RUBRO DE PRUEBA SOBRE LA RELACIÓN SUSTANCIAL PACTADA.",
     "localizacion": "ficha de prueba", "texto": _largo("la sustitución no afecta la relación sustancial", 15)},
    {"registro": "170307", "tipo": "Jurisprudencia", "instancia": "Tribunales Colegiados de Circuito",
     "rubro": "FUNDAMENTACIÓN Y MOTIVACIÓN. RUBRO DE PRUEBA SOBRE LA FALTA Y LA INDEBIDA SATISFACCIÓN.",
     "localizacion": "ficha de prueba", "texto": _largo("falta e indebida fundamentación", 15)},
]
CPC, CC = "Código de Procedimientos Civiles del Estado de Querétaro", "Código Civil del Estado de Querétaro"
CPEUM = "Constitución Política de los Estados Unidos Mexicanos"
NORMAS = [
    {"articulo": "49", "cuerpo_legal": CPC, "texto": "Artículo 49. Si durante la tramitación de un procedimiento se transfiere el derecho controvertido, quien transmitió el mismo dejará de ser parte."},
    # EL SEÑUELO: el 49 del Código CIVIL. La perífrasis procesal nunca puede casar con él.
    {"articulo": "49", "cuerpo_legal": CC, "texto": "Artículo 49. SEÑUELO DEL CÓDIGO CIVIL, que nunca debe transcribirse por el procesal."},
    {"articulo": "2284", "cuerpo_legal": CC, "texto": "Artículo 2284. Hay arrendamiento cuando las dos partes contratantes se obligan recíprocamente."},
    {"articulo": "2294", "cuerpo_legal": CC, "texto": "Artículo 2294. Si durante la vigencia del contrato de arrendamiento se verificare la transmisión de la propiedad del predio arrendado, el arrendamiento subsistirá."},
    {"articulo": "14", "cuerpo_legal": CPEUM, "texto": "Art. 14.- A ninguna ley se dará efecto retroactivo en perjuicio de persona alguna."},
    {"articulo": "17", "cuerpo_legal": CPEUM, "texto": "Art. 17.- Ninguna persona podrá hacerse justicia por sí misma, ni ejercer violencia para reclamar su derecho."},
]

# ── LOS PÁRRAFOS REALES DE LA SOLUCIÓN DEL PROYECTO MALO (entrega 1) ──────────
P_14_17 = ("Lo anterior, porque los artículos 14 y 17 de la Constitución Política de los Estados Unidos "
           "Mexicanos protegen la certeza derivada de las resoluciones firmes y la plena ejecución de las "
           "sentencias, de modo que la cosa juzgada impide reabrir lo decidido, pero no prohíbe que durante la "
           "ejecución se determine, conforme a la ley, quién está legitimado para recibir sus efectos.")
P_ANTES_2026918 = ("En ese sentido, la Primera Sala de la Suprema Corte de Justicia de la Nación ha distinguido entre "
                   "los efectos directo y reflejo de la cosa juzgada, y ha señalado que el primero exige identidad de "
                   "cosa, causa y personas, mientras que el segundo opera cuando, aun sin concurrir todos esos "
                   "elementos, lo decidido en el primer proceso resulta determinante para evitar sentencias contradictorias.")
A_2026918 = "Sirve de apoyo el criterio de registro 2026918:"
P_TRAS_2026918 = ("El precedente que dio origen a ese criterio examinó una resolución firme sobre la improcedencia de una "
                  "indemnización por error judicial y su influencia en un proceso posterior, lo que evidencia que la cosa "
                  "juzgada protege lo efectivamente decidido, sin extenderse a cuestiones distintas ni impedir que se "
                  "analicen consecuencias jurídicas sobrevinientes.")
A_168958 = "Sirve de apoyo el criterio de registro 168958:"
P_49 = ("La regla que la recurrente identifica como artículo 49 de la legislación procesal civil del Estado de "
        "Querétaro permite que, durante el procedimiento, el causahabiente sustituya a quien transmitió el derecho "
        "controvertido, salvo oposición justificada. Esa disposición no autoriza a modificar la sentencia, pero sí "
        "confirma que la identidad de quien intervino originalmente puede variar cuando existe una transmisión "
        "jurídicamente relevante del derecho vinculado con la ejecución.")
P_2284_2294 = ("En ese punto, asiste razón a la parte recurrente, porque los artículos 2284 y 2294 del Código Civil para "
               "el Estado de Querétaro vinculan el arrendamiento con el uso temporal del inmueble a cambio de una renta y "
               "disponen que, transmitida la propiedad del predio arrendado, el contrato subsiste en sus términos, mientras "
               "que el arrendatario debe pagar al nuevo propietario desde que se le notifique el título correspondiente.")
P_188480 = ("El criterio con registro 188480, invocado en el asunto para sostener que la sustitución procesal no afecta "
            "la relación sustancial pactada, tampoco conduce a una conclusión diversa. Su principio rector delimita la "
            "sustitución: ésta no puede transformar el contrato ni las obligaciones juzgadas.")
P_INNECESARIOS = ("Por la suficiencia de lo anterior, es innecesario examinar autónomamente la alegación relativa al "
                  "consentimiento de la parte quejosa respecto de la consideración de la Sala responsable, así como los "
                  "planteamientos sobre igualdad jurídica, fundamentación, motivación, congruencia y exhaustividad, y los "
                  "criterios con registros 187528, 2015679, 170307, 162826, 164590, 177529 y 2023791, pues aun cuando se "
                  "estimaran fundados no aportarían un beneficio adicional.")
P_93 = ("El artículo 93 de la Ley de Amparo obliga a que, cuando prosperan los agravios contra una sentencia que "
        "concedió la protección constitucional, el órgano revisor reasuma jurisdicción para ocuparse de los conceptos "
        "de violación cuyo estudio fue omitido; sin embargo, éstos no obran en el expediente del recurso.")
A_171925 = "Sirve de apoyo el criterio de registro 171925:"
ESTUDIO = [P_14_17, P_ANTES_2026918, A_2026918, P_TRAS_2026918, A_168958, P_49, P_2284_2294, P_188480,
           P_INNECESARIOS, P_93, A_171925]


def componer(estudio, tesis=TESIS, normas=NORMAS):
    d = docx.Document()
    notas = []
    n = dg._escribir_estudio(d, list(estudio), tesis, notas, normas)
    return n, d, [p.text for p in d.paragraphs], notas


def llamadas(p):
    return len(p._p.findall(".//" + dg.qn("w:footnoteReference")))


def lineas_de_notas(notas):
    return [ln for x in notas for ln in x.split(SEP)]


print("1 · LA TESIS CITADA POR SU REGISTRO LLEVA SU RUBRO Y SU TEXTO")
n, d, cuerpo, notas = componer(ESTUDIO)
todo = "\n".join(cuerpo)
ok(n == 3, f"las tres citas del proyecto malo se componen desde el acervo ({n})")
for reg in ("2026918", "168958", "171925"):
    t = next(x for x in TESIS if x["registro"] == reg)
    ok(t["rubro"].rstrip(".") in todo, f"{reg}: el rubro está en el cuerpo")
    ok(any(f"registro digital {reg}" in c and c.endswith("siguiente:") for c in cuerpo),
       f"{reg}: el anuncio lo compone el documento, con su registro digital")
    ok(any(f"registro digital {reg}. Texto:" in x for x in notas), f"{reg}: su texto (más de 80 palabras) baja al pie")
ok(not any(c.strip() == A_2026918 for c in cuerpo), "el anuncio pelado no queda como prosa")
ok(P_TRAS_2026918 in cuerpo, "la regla dicha por el estudio después de la cita se queda (no es eco)")
ok(P_ANTES_2026918 in cuerpo, "y la de antes también")

print("\n2 · LO QUE NO ES UN ANUNCIO SIGUE SIENDO PROSA")
ok(any(c.startswith("El criterio con registro 188480, invocado") for c in cuerpo)
   and "SUSTITUCIÓN PROCESAL. RUBRO DE PRUEBA" not in todo,
   "«El criterio con registro 188480, invocado…» no se convierte en cita")
ok(any(c.startswith("Por la suficiencia de lo anterior") and "170307" in c for c in cuerpo)
   and "FUNDAMENTACIÓN Y MOTIVACIÓN. RUBRO DE PRUEBA" not in todo,
   "la lista de registros innecesarios tampoco")
ok(not any(c.startswith("Sirve de apoyo") and ("188480" in c or "170307" in c) for c in cuerpo),
   "ninguna tesis contestada ni innecesaria sale como «Sirve de apoyo»")

print("\n3 · EL ANUNCIO QUE CONTESTA, LA LISTA Y LA PROSA DE DELANTE")
_, _, c3, _ = componer(["No resulta aplicable el criterio de registro 188480, porque exige una cesión de la "
                        "acción que en el caso no se alegó ni era necesaria para ejecutar la sentencia."])
ok(any(c.startswith("No resulta aplicable la tesis aislada") and "registro digital 188480" in c for c in c3)
   and not any(c.startswith("Sirve de apoyo") for c in c3), f"el que contesta conserva su arranque: {c3[:1]}")
ok(any(c.startswith("Ello, porque exige") for c in c3), "y la media frase de detrás recupera su sujeto")
_, _, c4, n4 = componer(["Sirven de apoyo los criterios de registros 2026918 y 168958:"])
ok(sum(1 for c in c4 if "registro digital" in c) == 2 and any(c.startswith("También sirve de apoyo") for c in c4)
   and len([x for x in n4 if "Texto:" in x]) == 2, f"una lista de registros da un bloque por tesis: {c4}")
_, _, c7, _ = componer(["La recurrente invoca el criterio de registro 188480 para sostener que la sustitución "
                        "procesal no puede afectar la relación sustancial pactada entre las partes originales."])
ok(len(c7) == 1 and c7[0].startswith("La recurrente invoca el criterio de registro 188480 para sostener"),
   "el arranque que sólo atribuye y sigue contando lo que pretendía la parte es prosa")
_, _, c8, _ = componer(["Sirve de apoyo el criterio de registro 168958, que extiende los límites subjetivos de la "
                        "cosa juzgada a los causahabientes de quienes litigaron."])
ok(any("registro digital 168958" in c for c in c8) and any(c.startswith("Ese criterio extiende") for c in c8),
   f"la media frase tras el registro recupera su sujeto: {c8[-1:]}")
_, _, c9, _ = componer([A_168958, "Los límites subjetivos alcanzan a los causahabientes de quienes litigaron, "
                        "como se dijo al examinar la ejecución de la sentencia definitiva.", A_168958])
ok(sum(1 for c in c9 if c.startswith("«COSA JUZGADA. SUS LÍMITES")) == 1
   and any(c.endswith("registro digital 168958, ya citada.") for c in c9)
   and not any(c.strip() == A_168958 for c in c9),
   f"la segunda vez se nombra como ya citada, sin anuncio colgando: {c9[-1:]}")
_pr = ("Los límites subjetivos de la cosa juzgada alcanzan a los causahabientes de quienes litigaron, porque su "
       "vinculación jurídica explica la sujeción a lo resuelto.")
_, _, c5, _ = componer([_pr + " " + A_168958])
ok(c5[:1] == [_pr] and "registro digital 168958" in c5[1], f"la prosa que precede al anuncio no se pierde: {c5[:2]}")
_, _, c6, _ = componer([_pr + " Sirve de apoyo la jurisprudencia de registro 168958, de rubro «COSA JUZGADA. "
                        "SUS LÍMITES OBJETIVOS Y SUBJETIVOS.»"])
ok(c6[:1] == [_pr], "…tampoco en la cita con rubro (el camino de siempre)")

print("\n4 · LOS ARTÍCULOS: PLURAL, PERÍFRASIS Y «DEL MISMO ORDENAMIENTO»")
lns = lineas_de_notas(notas)
for a in ("14", "17", "49", "2284", "2294"):
    ok(any(ln.startswith(f"«Artículo {a}.") for ln in lns), f"art. {a} transcrito al pie")
ok(not any("SEÑUELO" in ln for ln in lns), "el 49 de la legislación PROCESAL no se transcribe con el del Código Civil")
ok(any("«Artículo 2284." in x and "«Artículo 2294." in x for x in notas),
   "«los artículos 2284 y 2294» van juntos en una nota")
ok(any("«Artículo 14." in x and "«Artículo 17." in x for x in notas), "«los artículos 14 y 17», también")
ok(all(llamadas(p) <= 1 for p in d.paragraphs), "y nunca dos llamadas en un párrafo (las notas 3 y 4 no se leen «34»)")
_, _, _, n7 = componer(["En ese sentido, el artículo 2284 del Código Civil del Estado de Querétaro define el arrendamiento "
                        "como el contrato por el que una parte concede el uso o goce temporal de una cosa, mientras que el "
                        "artículo 2294 del mismo ordenamiento dispone que, transmitida la propiedad, el contrato subsiste."])
ok(any("«Artículo 2294." in ln for ln in lineas_de_notas(n7)), "«el artículo 2294 del mismo ordenamiento» hereda la ley")
_, _, _, n8 = componer(["El artículo 49 del Código Procesal Civil para el Estado de Querétaro permite la sustitución "
                        "del causahabiente durante el procedimiento, salvo oposición justificada de la contraria."])
ok(any(ln.startswith("«Artículo 49. Si durante") for ln in lineas_de_notas(n8)),
   "«Código Procesal Civil» es el Código de Procedimientos Civiles")

print("\n5 · LA NOTA DE VARIOS ARTÍCULOS EN EL .docx")
xml = dg._xml_notas(["«Artículo 2284. A» — CC" + SEP + "«Artículo 2294. B» — CC",
                     "Pleno, ficha. Texto: Hechos: uno.\n\nCriterio jurídico: dos."]).decode("utf8")
_n1 = re.search(r'<w:footnote w:id="1">(.*?)</w:footnote>', xml, re.S).group(1)
_n2 = re.search(r'<w:footnote w:id="2">(.*?)</w:footnote>', xml, re.S).group(1)
ok(_n1.count("<w:p>") == 2 and _n1.count("<w:footnoteRef/>") == 1, "dos renglones, una sola llamada")
ok(_n2.count("<w:p>") == 1, "la tesis con saltos de línea sigue siendo un párrafo (como antes)")

print("\n6 · EL MATERIAL SE COMPLETA CON LA LEY CORRECTA (preceptos_fuera)")
AGRAVIO = ("Alega que el órgano de amparo transgrede los artículos 14 y 16 de la Constitución Política de los Estados "
           "Unidos Mexicanos, así como los artículos 49, 57 y 279 del Código Procesal Civil para el Estado de Querétaro "
           "y 2284 y 2294 del Código Civil local, porque no expone de manera completa las razones.")
_c = f6.citas_de_articulos(AGRAVIO)
_nums = [a for a, _ in _c]
ok(all(x in _nums for x in ("14", "16", "49", "57", "279", "2284", "2294")),
   f"el agravio rinde sus siete preceptos, no sólo el 14 y el 16: {_nums}")
ok(all("procedimientos civiles" in c.lower() for a, c in _c if a in ("49", "57", "279")),
   "el 49, el 57 y el 279, del Código de Procedimientos Civiles")
ok(all(fr_.fuero_de(c) == "estatal" for a, c in _c if a in ("2284", "2294")),
   "«del Código Civil local» es ley del estado (no se busca en el silo federal)")
m = f6.Material()
m.normas = [{"cuerpo_legal": CPC, "articulo": "656"}, {"cuerpo_legal": CC, "articulo": "1"},
            {"cuerpo_legal": "Código Civil Federal", "articulo": "1"}]
_, pares_ag = f6.preceptos_fuera(AGRAVIO, m)
ok({(CC.lower(), "2284"), (CC.lower(), "2294"), (CPC.lower(), "49")} <= pares_ag
   and not any("federal" in c for c, a in pares_ag if a in ("2284", "2294")),
   f"el agravio pide el 2284 y el 2294 al Código Civil de Querétaro, nunca al federal: {sorted(pares_ag)}")
m.normas = [{"cuerpo_legal": CPC, "articulo": "656"}, {"cuerpo_legal": CC, "articulo": "2284"}]
_, pares = f6.preceptos_fuera(P_49, m)
ok(pares == {(CPC.lower(), "49")}, f"la perífrasis procesal se pide al CPC aunque el CC esté en el material: {sorted(pares)}")

print("\n7 · EL PROMPT v4 PIDE LAS DOS LLAVES QUE EL DOCUMENTO LEE")
C1 = [f6.Criterio("¿La sustitución alteró la cosa juzgada?", "fundado", "r", "principal")]
for forma in ("estandar", "moderna"):
    _m = f6.Material(tipo_asunto="amparo_revision", materia="civil", formato=forma, n_planteamientos=1, variante="v4")
    p4 = f6.prompt_estudio("ACTO", "CONC", C1, _m, es_recurso=True, guion="APARTADO 1 · A1.a A1.b")
    ok(getattr(f6, "_TECNICA_EJEMPLO_NUEVO", "\0") in p4
       and "  Y ahí se detiene el párrafo. NO ESCRIBAS" not in p4,
       f"{forma}: detrás del registro, el rubro entre comillas copiado del material")
    ok("el tipo, el órgano, el registro, el rubro— se encarga" not in p4
       and "su registro digital y su rubro entre comillas" in p4,
       f"{forma}: ya no dice que el registro y el rubro los pone el documento")

print()
if fallos:
    print(f"FALLAN {len(fallos)}: " + " · ".join(fallos))
    sys.exit(1)
print("TODAS LAS COMPROBACIONES PASAN")
