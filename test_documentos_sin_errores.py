#!/usr/bin/env python3
"""El análisis de documentos sin los errores del 29-oct-2026.

  1. late mientras el modelo piensa y al abrir la continuación (flujo_con_latido);
  2. lo que pasa del tope de hojas se lee en parte, en todos los planes, con la
     primera línea del servidor y las hojas siguientes (lectura_por_plan);
  3. un análisis vacío no se da por bueno: se repite con el motor de todos y,
     si también calla, se devuelve la consulta;
  4. si el navegador se va antes de la primera letra, la consulta vuelve;
  5. lo que el archivo es por dentro y el Word hecho de fotos (formato_documento);
  6. el chat sabe qué adjuntos no se leyeron (documento_sin_leer).

La ruta REAL de main.py con motores de mentira y un Supabase en memoria: sin
red y sin gasto. Sin arranque: el lifespan borra las cachés de Gemini de
producción, así que nada de `with TestClient(main.app)`; httpx habla con la app
por ASGI, que tampoco lo ejecuta.

    .venv/bin/python test_documentos_sin_errores.py
"""
import asyncio
import io
import json
import os
import sys
import types
import uuid
import zipfile

import dotenv
import fitz
import httpx

dotenv.load_dotenv = lambda *a, **k: False
for _k in ("OPENAI_API_KEY", "OPENROUTER_API_KEY", "DEEPSEEK_API_KEY", "GEMINI_API_KEY",
           "MISTRAL_API_KEY", "ELEVENLABS_API_KEY", "AZURE_DOCINT_KEY", "AZURE_SPEECH_KEY"):
    os.environ[_k] = "falsa-para-pruebas"
os.environ.update(QDRANT_URL="http://127.0.0.1:9", QDRANT_API_KEY="",
                  SUPABASE_URL="", SUPABASE_SERVICE_KEY="")

sys.path.insert(0, ".")
import main  # noqa: E402

assert main.supabase_admin is None, "main.py cargó un Supabase de verdad: la prueba no sigue"
import documento_sin_leer as dsl  # noqa: E402
import flujo_con_latido as fl  # noqa: E402
import formato_documento as fd  # noqa: E402
import lectura_por_plan as lp  # noqa: E402

FALLOS = []


def ok(cond, que):
    print(f"   {'PASA ' if cond else 'FALLA'}  {que}")
    if not cond:
        FALLOS.append(que)


# ═════════════════════════════════════════════════════════════════════════
print("\n1 · flujo_con_latido: el latido sale mientras el flujo calla")


async def _p1():
    async def abrir_lento():
        await asyncio.sleep(0.32)

        async def gen():
            yield "a"
            await asyncio.sleep(0.16)
            yield "b"
        return gen()

    vistos = [x async for x in fl.con_latido(abrir_lento, cada=0.05)]
    lat_antes = vistos.index("a")
    ok(lat_antes >= 4 and all(x is fl.LATIDO for x in vistos[:lat_antes]),
       f"antes del primer trozo, sólo latidos ({lat_antes})")
    ok([x for x in vistos if x is not fl.LATIDO] == ["a", "b"], "los trozos llegan todos y en orden")
    ok(any(x is fl.LATIDO for x in vistos[lat_antes:vistos.index("b")]), "y también late a mitad del flujo")

    async def abrir_falla():
        await asyncio.sleep(0.12)
        raise RuntimeError("cuota agotada")
    try:
        async for _ in fl.con_latido(abrir_falla, cada=0.05):
            pass
        ok(False, "el error al abrir sube en el consumidor")
    except RuntimeError as e:
        ok(str(e) == "cuota agotada", "el error al abrir sube en el consumidor, tal cual")

    cancelado = asyncio.Event()

    async def abrir_eterno():
        try:
            await asyncio.sleep(30)
        except asyncio.CancelledError:
            cancelado.set()
            raise
    g = fl.con_latido(abrir_eterno, cada=0.05)
    primero = await g.__anext__()
    await g.aclose()
    await asyncio.sleep(0.05)
    ok(primero is fl.LATIDO and cancelado.is_set(), "si el consumidor se va, la petición al modelo se cancela")

asyncio.run(_p1())


# ═════════════════════════════════════════════════════════════════════════
print("\n2 · lectura_por_plan: las hojas que se leen y lo que se dice")
ok(lp.tramo(201, 100) == (1, 100), "Pro con 201 hojas: las primeras 100")
ok(lp.tramo(80, 100) is None, "un documento que cabe se lee entero")
ok(lp.tramo(201, 100, 101) == (101, 200), "desde la 101: hasta 100 hojas más")
ok(lp.tramo(150, 100, 101) == (101, 150), "desde la 101 de 150: hasta el final")
try:
    lp.tramo(150, 100, 151)
    ok(False, "una hoja de inicio fuera del documento es un error claro")
except ValueError as e:
    ok("150 hojas" in str(e), "una hoja de inicio fuera del documento es un error claro")
linea = lp.primera_linea("pro_monthly", 100, 201, 1, 100)
ok(linea == ("*Leí las primeras 100 de las 201 hojas de su documento: su plan Pro lee hasta 100 por "
             "documento. Para el resto, adjunte las hojas 101 a 201 por separado o pase a Platinum "
             "(hasta 600).*\n\n"), f"la línea de Pro dice qué se leyó y qué hacer: {linea!r}")
ok("pase a" not in lp.primera_linea("platinum_monthly", 600, 750, 1, 600)
   and "las hojas 601 a 750" in lp.primera_linea("platinum_monthly", 600, 750, 1, 600),
   "Platinum no recibe una invitación a sí mismo")
ok(lp.primera_linea("gratuito", 20, 34, 1, 20) ==
   "*Leí las primeras 20 de las 34 hojas de su documento (el plan gratuito lee hasta 20).*\n\n",
   "el gratuito conserva la frase del 25-oct")
ok("Leí las hojas 101 a 200 de las 201" in lp.primera_linea("pro_annual", 100, 201, 101, 200),
   "un tramo pedido se nombra por sus hojas")
av = lp.aviso("basico_monthly", 50, 80, 1, 50)
ok(av == {"leidas": 50, "total": 80, "plan": "Pro", "plan_actual": "Básico", "tope": 50, "desde": 1,
          "hasta": 50, "siguiente": 51}, f"el aviso de Básico con 80 hojas ({av})")
ok([lp.aviso("gratuito", 20, n, 1, 20)["plan"] for n in (34, 150, 700)] == ["Pro", "Platinum", None],
   "el gratuito recibe la misma invitación que su tarjeta (Pro, Platinum, ninguna)")
ok(lp.desde_valido("101") == 101 and lp.desde_valido("x") is None and lp.desde_valido("0") is None,
   "desde_hoja raro se ignora")
ok(lp.nombre_del_plan("ultra_secretarios") == "Ultra Secretarios" and lp.tope_del_plan("raro") == 30,
   "nombres y topes de planes fuera de la tabla")


# ═════════════════════════════════════════════════════════════════════════
print("\n3 · formato_documento: el archivo por dentro")


def _docx(texto="", imagenes=0):
    import docx
    d = docx.Document()
    if texto:
        d.add_paragraph(texto)
    for i in range(imagenes):
        # Una «foto» con ruido: una imagen lisa pesa menos de 2 KB y el lector
        # la tomaría por un icono (que no es una hoja).
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 400 + i, 300), False)
        pix.set_rect(pix.irect, (250, 250, 250))
        for y in range(0, 300, 3):
            for x in range((y * 7 + i) % 5, 400, 9):
                pix.set_pixel(x, y, ((x * 31 + y * 17) % 256, (x * 13) % 256, (y * 29) % 256))
        d.add_picture(io.BytesIO(pix.tobytes("png")))
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _pdf(n, texto="Hoja {i}. "):
    """`n` hojas con texto seleccionable que pasa `_texto_nativo_sirve` (si no,
    la ruta lo mandaría al OCR, que en esta prueba está prohibido)."""
    d = fitz.open()
    parrafo = ("PRIMER AGRAVIO. La sentencia recurrida valoró indebidamente la prueba testimonial "
               "ofrecida por la parte actora, pues omitió considerar las contradicciones en que "
               "incurrieron los declarantes respecto de la fecha, el lugar y las circunstancias del "
               "acto reclamado. SEGUNDO AGRAVIO. El juzgador de origen carecía de competencia por razón "
               "de territorio, ya que el domicilio de la demandada se ubica en Apizaco, Tlaxcala, y no "
               "en la capital del estado como erróneamente se asentó en el auto admisorio. TERCERO. "
               "Se violaron las formalidades esenciales del procedimiento al negarse el desahogo de la "
               "pericial en materia de grafoscopía ofrecida oportunamente por la recurrente.")
    def letras(k):
        return "".join(chr(97 + (k // 26 ** j) % 26) for j in range(3))
    for i in range(n):
        p = d.new_page()
        # Treinta palabras propias por hoja: un párrafo repetido cien veces
        # parece un sello y `_texto_nativo_sirve` lo mandaría al OCR.
        propias = " ".join("vocablo" + letras(i * 30 + k) for k in range(30))
        p.insert_textbox(fitz.Rect(50, 50, 550, 800), texto.format(i=i + 1) + parrafo + " " + propias,
                         fontsize=9)
    return d.tobytes()


ok(fd.formato_real(_docx("hola")) == "docx", "un .docx de verdad")
ok(fd.formato_real(_pdf(1)) == "pdf", "un PDF aunque se llame .docx")
ok(fd.formato_real(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1" + b"\0" * 600) == "doc", "un .doc viejo (OLE2)")
RTF = (rb"{\rtf1\ansi\deff0{\fonttbl{\f0 Times New Roman;}}{\colortbl;\red0\green0\blue0;}"
       rb"{\*\generator Riched20;}\f0 CONTRATO DE ARRENDAMIENTO\par Cl\'e1usula cuarta: el pago\tab mensual.\par"
       rb" Se\u241?ora Pe\u241?a}")
ok(fd.formato_real(RTF) == "rtf", "un RTF con nombre de Word")
t_rtf = fd.texto_rtf(RTF)
ok(t_rtf == "CONTRATO DE ARRENDAMIENTO\nCláusula cuarta: el pago\tmensual.\nSeñora Peña", f"su texto, limpio: {t_rtf!r}")
HTML = b"<html><head><style>p{}</style></head><body><p>CONTRATO</p><p>Cl&aacute;usula primera.</p></body></html>"
ok(fd.formato_real(HTML) == "html" and "Cláusula primera." in fd.texto_html(HTML), "un HTML guardado como .doc")
con_fotos = _docx("", imagenes=3)
imgs = fd.imagenes_docx(con_fotos)
ok(len(imgs) == 3, f"las imágenes del Word, en orden ({len(imgs)})")
pdf_imgs = fd.pdf_de_imagenes(imgs)
with fitz.open(stream=pdf_imgs, filetype="pdf") as _d:
    ok(_d.page_count == 3, "un PDF con una imagen por hoja, para el OCR")
ras = fd.pdf_rasterizado(_pdf(2))
with fitz.open(stream=ras, filetype="pdf") as _d:
    ok(_d.page_count == 2 and not _d[0].get_text().strip(), "la copia en imágenes de un PDF (sin capa de texto)")


# ═════════════════════════════════════════════════════════════════════════
print("\n4 · documento_sin_leer: el chat sabe qué adjunto no se leyó")
M = lambda r, c: types.SimpleNamespace(role=r, content=c)  # noqa: E731
conv = [M("user", "📄 **Documento adjunto:** agravios.pdf\n\nAnaliza los agravios"),
        M("assistant", "⏱️ **Se perdió el contacto con el servidor.**"),
        M("user", "📄 **Documento adjunto:** demanda.pdf\n\nResume\n\n<!-- DOCUMENTO_INICIO -->\nTEXTO\n<!-- DOCUMENTO_FIN -->"),
        M("user", "dame el proyecto valorando todo")]
ok(dsl.adjuntos_sin_texto(conv) == ["agravios.pdf"], "sólo el que no trae su texto")
nota = dsl.nota_documentos_sin_leer(conv)
ok("«agravios.pdf»" in nota and "NO está en la conversación" in nota and "vuelva a adjuntar" in nota,
   "la instrucción lo nombra y pide que lo vuelva a adjuntar")
ok(dsl.nota_documentos_sin_leer(conv[2:]) == "", "sin adjuntos sin texto, nada")
reintento = conv + [M("user", "📄 **Documento adjunto:** agravios.pdf\n\notra vez\n\n<!-- DOCUMENTO_INICIO -->x<!-- DOCUMENTO_FIN -->")]
ok(dsl.adjuntos_sin_texto(reintento) == [], "si después se volvió a adjuntar y entró, ya no cuenta")


# ═════════════════════════════════════════════════════════════════════════
# La ruta real, con motores de mentira y un Supabase en memoria.

def _trozo(texto=None, fin=None, nativo=None):
    delta = types.SimpleNamespace(content=texto)
    return types.SimpleNamespace(choices=[types.SimpleNamespace(
        delta=delta, finish_reason=fin, model_extra={"native_finish_reason": nativo} if nativo else {})])


class _Motor:
    """`guion`: lista de (espera antes de abrir, [trozos]) por llamada."""

    def __init__(self, guion):
        self.guion = list(guion)
        self.llamadas = []
        self.chat = types.SimpleNamespace(completions=types.SimpleNamespace(create=self._create))

    async def _create(self, **kw):
        self.llamadas.append(kw)
        espera, trozos = self.guion.pop(0)
        await asyncio.sleep(espera)

        async def gen():
            for t in trozos:
                await asyncio.sleep(0.005)
                yield t
        return gen()


def _token(uid: str) -> str:
    """Un token con forma de JWT que dice de quién es. Desde el 29-oct-2026 el
    `user_id` sin token se atiende como anónimo (puerta_sesion.py): para que la
    ruta lea el plan, la petición tiene que traer el token de esa cuenta."""
    import base64
    import time
    carga = base64.urlsafe_b64encode(json.dumps({"sub": uid, "exp": time.time() + 3600}).encode())
    return f"e30.{carga.decode().rstrip('=')}.firma"


class _Sb:
    def __init__(self, plan="pro_monthly"):
        self.plan, self.rpcs = plan, []
        self.auth = types.SimpleNamespace(get_user=self._get_user)

    @staticmethod
    def _get_user(tok):
        import base64
        carga = tok.split(".")[1]
        uid = json.loads(base64.urlsafe_b64decode(carga + "=" * (-len(carga) % 4)))["sub"]
        return types.SimpleNamespace(user=types.SimpleNamespace(id=uid, email="abogada@ejemplo.mx"))

    def table(self, nombre):
        sb = self

        class _Q:
            def select(self, *a):
                return self

            def eq(self, *a):
                return self

            def limit(self, *a):
                return self

            def execute(self):
                if nombre == "user_profiles":
                    return types.SimpleNamespace(data=[{"subscription_type": sb.plan, "email": "abogada@ejemplo.mx",
                                                         "estado": None}])
                return types.SimpleNamespace(data=[])
        return _Q()

    def rpc(self, nombre, params):
        sb = self

        class _R:
            def execute(self):
                sb.rpcs.append(nombre)
                if nombre == "consume_query":
                    return types.SimpleNamespace(data={"allowed": True, "used": 1, "limit": 140})
                if nombre == "devolver_consulta":
                    return types.SimpleNamespace(data={"devuelta": True, "used": 0, "limit": 140})
                return types.SimpleNamespace(data={})
        return _R()


async def _analizar(sb, motor, archivo, campos=None, nombre="agravios.pdf"):
    main.supabase_admin = sb
    main._via_documento = lambda modelo, esfuerzo=None: (motor, {})
    transporte = httpx.ASGITransport(app=main.app)
    uid = str(uuid.uuid4())
    async with httpx.AsyncClient(transport=transporte, base_url="http://prueba", timeout=60) as c:
        r = await c.post("/analyze-document",
                         headers={"Authorization": f"Bearer {_token(uid)}"},
                         files={"file": (nombre, archivo, "application/pdf")},
                         data={"prompt": "¿Qué agravios hace valer?", "usar_acervo": "0",
                               "user_id": uid, **(campos or {})})
    for _ in range(100):
        if not main._DEVOLUCIONES:
            break
        await asyncio.sleep(0.01)
    lineas = r.text.split("\n")
    eventos = [json.loads(l[6:]) for l in lineas if l.startswith("data: ")]
    return r, lineas, eventos


main.DOCUMENTO_LATIDO_S = 0.05
PLATINUM = main.DOCUMENT_MODEL_PLATINUM


async def _ocr_prohibido(subida):
    raise AssertionError("OCR en una prueba sin red")
main._extract_text_from_upload = _ocr_prohibido
main.get_gemini_client = lambda *a, **k: (_ for _ in ()).throw(AssertionError("Gemini en una prueba sin red"))

print("\n5 · la ruta late mientras el modelo piensa, también al continuar")
motor = _Motor([(0.4, [_trozo("## Los agravios\n\n"), _trozo("Primero, la prueba."),
                       _trozo(None, "content_filter", "RECITATION")]),
                (0.3, [_trozo(" Segundo, la competencia."), _trozo(None, "stop")])])
r, lineas, ev = asyncio.run(_analizar(_Sb("platinum_monthly"), motor, _pdf(2)))
i_redacta = next(i for i, l in enumerate(lineas) if "Redactando" in l)
i_token = next(i for i, l in enumerate(lineas) if '"token"' in l)
ok(sum(1 for l in lineas[i_redacta:i_token] if l == ": latido") >= 4,
   "entre «Redactando el análisis…» y el primer token, latidos (antes: silencio)")
i_cont = next(i for i, l in enumerate(lineas) if "Segundo, la competencia" in l)
i_ult = max(i for i, l in enumerate(lineas[:i_cont]) if "Primero, la prueba" in l)
ok(sum(1 for l in lineas[i_ult:i_cont] if l == ": latido") >= 3, "y mientras abre la continuación por RECITATION")
ok("".join(e.get("token", "") for e in ev).startswith("## Los agravios") and any(e.get("done") for e in ev),
   "el análisis sale entero y termina con done")

print("\n6 · lectura parcial para los planes de pago (antes: 413)")
motor = _Motor([(0.0, [_trozo("Análisis de las hojas leídas."), _trozo(None, "stop")])])
sb = _Sb("pro_monthly")
r, lineas, ev = asyncio.run(_analizar(sb, motor, _pdf(120)))
lim = next((e["limite_plan"] for e in ev if "limite_plan" in e), None)
ok(r.status_code == 200 and lim and lim["leidas"] == 100 and lim["total"] == 120 and lim["plan"] == "Platinum"
   and lim["siguiente"] == 101, f"Pro con 120 hojas: se leen 100 y se ofrecen las siguientes ({lim})")
texto = "".join(e.get("token", "") for e in ev)
ok(texto.startswith("*Leí las primeras 100 de las 120 hojas de su documento: su plan Pro lee hasta 100"),
   "la respuesta empieza con la línea del servidor")
usuario = motor.llamadas[0]["messages"][-1]["content"]
ok("sólo recibes las primeras 100" in usuario and "Hoja 100." in usuario and "Hoja 101." not in usuario,
   "el modelo recibe sólo las 100 hojas, y lo sabe")
ok("consume_query" in sb.rpcs and "devolver_consulta" not in sb.rpcs, "se cobra como un análisis: lo es")
motor = _Motor([(0.0, [_trozo("Las hojas siguientes."), _trozo(None, "stop")])])
r, lineas, ev = asyncio.run(_analizar(_Sb("pro_monthly"), motor, _pdf(120), {"desde_hoja": "101"}))
usuario = motor.llamadas[0]["messages"][-1]["content"]
lim = next((e["limite_plan"] for e in ev if "limite_plan" in e), None)
ok("Hoja 101." in usuario and "Hoja 120." in usuario and "Hoja 100." not in usuario
   and lim and lim["desde"] == 101 and lim["hasta"] == 120 and lim["siguiente"] is None,
   "con desde_hoja=101 se leen las hojas 101 a 120")
r, _, _ = asyncio.run(_analizar(_Sb("pro_monthly"), _Motor([]), _pdf(5), {"desde_hoja": "9"}))
ok(r.status_code == 400 and "5 hojas" in r.json().get("detail", ""), "una hoja de inicio que no existe: 400 claro")
os.environ["DOCUMENTO_LECTURA_PARCIAL"] = "0"
r, _, _ = asyncio.run(_analizar(_Sb("pro_monthly"), _Motor([]), _pdf(120)))
ok(r.status_code == 413, "DOCUMENTO_LECTURA_PARCIAL=0 devuelve el 413 de antes, sin desplegar")
os.environ.pop("DOCUMENTO_LECTURA_PARCIAL")

print("\n7 · un análisis vacío no se da por bueno")
motor = _Motor([(0.0, [_trozo(None, "stop", "STOP")]),
                (0.0, [_trozo("Ahora sí, el análisis."), _trozo(None, "stop")])])
sb = _Sb("platinum_monthly")
r, _, ev = asyncio.run(_analizar(sb, motor, _pdf(2)))
texto = "".join(e.get("token", "") for e in ev)
ok(len(motor.llamadas) == 2 and motor.llamadas[0]["model"] == PLATINUM
   and motor.llamadas[1]["model"] == main.DOCUMENT_MODEL and texto == "Ahora sí, el análisis.",
   "si Platinum calla, se repite UNA vez con el motor de todos")
motor = _Motor([(0.0, [_trozo(None, "stop", "STOP")]), (0.0, [_trozo(None, "stop")])])
sb = _Sb("platinum_monthly")
r, _, ev = asyncio.run(_analizar(sb, motor, _pdf(2)))
err = next((e["error"] for e in ev if "error" in e), "")
ok(not any(e.get("token") for e in ev) and "No se te descontó la consulta" in err
   and "devolver_consulta" in sb.rpcs, f"si también calla: error, sin tokens, y la consulta vuelve ({err[:60]}…)")

print("\n8 · el navegador se va antes de la primera letra: la consulta vuelve")


async def _p8():
    sb = _Sb("pro_monthly")
    main.supabase_admin = sb
    main._via_documento = lambda modelo, esfuerzo=None: (_Motor([(5.0, [_trozo("tarde")])]), {})
    from starlette.datastructures import Headers, UploadFile
    archivo = UploadFile(file=io.BytesIO(_pdf(2)), filename="agravios.pdf",
                         headers=Headers({"content-type": "application/pdf"}))
    pedido = types.SimpleNamespace(headers={})
    resp = await main.analyze_document(
        http_request=pedido, file=archivo, prompt="Analiza", user_id=str(uuid.uuid4()), fuentes_web=None,
        estado=None, usar_acervo="0", fuentes=None, historial=None, historial_archivo=None, peticion_id=None,
        conversation_id=None, expediente_id=None, mensaje_visible=None, desde_hoja=None)
    it = resp.body_iterator.__aiter__()
    vistos = []
    while len(vistos) < 8:
        x = await it.__anext__()
        vistos.append(x)
        if x == ": latido\n\n" and any("Redactando" in v for v in vistos):
            break
    await it.aclose()            # el navegador cerró: Starlette cierra el flujo
    for _ in range(100):
        if "devolver_consulta" in sb.rpcs:
            break
        await asyncio.sleep(0.02)
    ok("devolver_consulta" in sb.rpcs, "cerrado a medio pensar, sin una letra: se devuelve la consulta")

asyncio.run(_p8())

print("\n9 · el Word hecho de fotos se lee con OCR y el .docx que es RTF, como RTF")


async def _ocr_falso(subida):
    datos = await subida.read()
    with fitz.open(stream=datos, filetype="pdf") as d:
        return "\n\n".join(f"[[PÁGINA {i + 1}]]\nTexto reconocido de la imagen {i + 1}." for i in range(d.page_count))

main._extract_text_from_upload = _ocr_falso
motor = _Motor([(0.0, [_trozo("Analizado."), _trozo(None, "stop")])])
r, _, ev = asyncio.run(_analizar(_Sb("pro_monthly"), motor, con_fotos, nombre="imagenescontrato.docx"))
doc = next((e["documento"] for e in ev if "documento" in e), {})
ok("Texto reconocido de la imagen 3." in (doc.get("texto") or "") and any(e.get("done") for e in ev),
   "las 3 imágenes del Word pasan por el OCR y el análisis sale")
motor = _Motor([(0.0, [_trozo("Analizado."), _trozo(None, "stop")])])
r, _, ev = asyncio.run(_analizar(_Sb("pro_monthly"), motor, RTF, nombre="CONTRATO ARENDAMIENTO IVAN.docx"))
doc = next((e["documento"] for e in ev if "documento" in e), {})
ok("Cláusula cuarta" in (doc.get("texto") or ""), "el .docx que por dentro es RTF se lee (antes: «File is not a zip file»)")
r, _, ev = asyncio.run(_analizar(_Sb("pro_monthly"), _Motor([]), _docx(""), nombre="vacio.docx"))
err = next((e["error"] for e in ev if "error" in e), "")
ok("Guardar como → PDF" in err and "No se le descontó la consulta." in err,
   f"un Word vacío dice qué hacer y que no se cobró ({err[:70]}…)")
main._extract_text_from_upload = _ocr_prohibido

print()
if FALLOS:
    print("FALLOS:\n  · " + "\n  · ".join(FALLOS))
    sys.exit(1)
print("El análisis de documentos late, lee en parte lo que pasa del tope, no cobra lo que no entregó "
      "y lee lo que el archivo es por dentro.")
