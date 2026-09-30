# -*- coding: utf-8 -*-
"""El lector de DOCX por flujo (30-sep-2026): mismo texto que python-docx, sin
cargar el documento entero, y las tres rutas que reciben un DOCX lo usan.

    .venv/bin/python test_lectura_docx.py
"""
import asyncio, copy, io, os, subprocess, sys, zipfile
from pathlib import Path
sys.path.insert(0, ".")
import docx
from docx.oxml.ns import qn
from lxml import etree
import lectura_docx as ld

FALLOS = []


def ok(c, q):
    print(f"   {'PASA ' if c else 'FALLA'}  {q}")
    if not c:
        FALLOS.append(q)


def a_bytes(d) -> bytes:
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def viejo(contenido: bytes, sep="\n\n") -> str:
    d = docx.Document(io.BytesIO(contenido))
    return sep.join(p.text for p in d.paragraphs if p.text.strip())


print("\n1 · EL MISMO TEXTO QUE python-docx")
d = docx.Document()
d.add_paragraph("PRIMERO. Se admite la demanda.")
p = d.add_paragraph("Con tabulador:")
p.add_run().add_tab()
p.add_run("después del tab")
p = d.add_paragraph("Línea uno")
p.add_run().add_break()                                   # salto de línea → "\n"
p.add_run("línea dos")
p = d.add_paragraph("Antes del salto de página")
from docx.enum.text import WD_BREAK
p.add_run().add_break(WD_BREAK.PAGE)                      # de página → nada
p.add_run(" y después")
d.add_paragraph("   ")                                    # vacío: fuera
d.add_paragraph("Artículo 1°. «Comillas», acentos: áéíóú ñ.")
# Un hipervínculo, como lo escribe Word.
p = d.add_paragraph("Véase ")
h = etree.SubElement(p._p, qn("w:hyperlink"))
r = etree.SubElement(h, qn("w:r"))
t = etree.SubElement(r, qn("w:t"))
t.text = "la tesis 2a./J. 5/2020"
# Paradas de tabulación en las propiedades del párrafo: NO son tabuladores.
p = d.add_paragraph("Con paradas de tabulación")
p.paragraph_format.tab_stops.add_tab_stop(docx.shared.Inches(1))
b = a_bytes(d)
ok(ld.texto_docx(b) == viejo(b), "párrafos, tab, saltos, hipervínculo y paradas: idéntico a python-docx")
ok(ld.texto_docx(b, "\n") == viejo(b, "\n"), "con el separador de una línea del taller, también")
ok("línea uno" in ld.texto_docx(b).lower() and "Línea uno\nlínea dos" in ld.texto_docx(b), "el salto de línea es \\n")
ok("Antes del salto de página y después" in ld.texto_docx(b), "el salto de página no mete nada")
ok("\t" in ld.texto_docx(b) and ld.texto_docx(b).count("\t") == 1, "sólo el tab de la corrida, no la parada")

print("\n2 · LO QUE python-docx DEJABA FUERA")
d = docx.Document()
d.add_paragraph("Antes de la tabla.")
tb = d.add_table(rows=2, cols=2)
tb.cell(0, 0).text = "Autoridad"
tb.cell(0, 1).text = "Acto reclamado"
tb.cell(1, 0).text = "Juez Quinto"
tb.cell(1, 1).text = "Auto de 7 de octubre"
d.add_paragraph("Después de la tabla.")
b = a_bytes(d)
tx = ld.texto_docx(b)
ok("Acto reclamado" in tx and "Auto de 7 de octubre" in tx and "Acto reclamado" not in viejo(b),
   "el texto de las tablas entra (python-docx lo perdía)")
ok(tx.index("Antes de la tabla.") < tx.index("Autoridad") < tx.index("Después de la tabla."), "y en su lugar")

# Un cuadro de texto escrito dos veces, como lo guarda Word: mc:Choice y mc:Fallback.
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
d = docx.Document()
p = d.add_paragraph("Párrafo con cuadro de texto.")
r = etree.SubElement(p._p, qn("w:r"))
alt = etree.SubElement(r, f"{{{MC}}}AlternateContent")
for rama in ("Choice", "Fallback"):
    c = etree.SubElement(alt, f"{{{MC}}}{rama}")
    caja = etree.SubElement(c, qn("w:txbxContent"))
    pp = etree.SubElement(caja, qn("w:p"))
    rr = etree.SubElement(pp, qn("w:r"))
    tt = etree.SubElement(rr, qn("w:t"))
    tt.text = "TEXTO DEL CUADRO"
b = a_bytes(d)
tx = ld.texto_docx(b)
ok(tx.count("TEXTO DEL CUADRO") == 1, "el cuadro de texto sale UNA vez (el mc:Fallback se salta)")
ok("Párrafo con cuadro de texto." in tx and "TEXTO DEL CUADROPárrafo" not in tx,
   "y no se pega al párrafo que lo contiene")

# El texto borrado en control de cambios (w:delText) no entra.
d = docx.Document()
p = d.add_paragraph("Quedó ")
dl = etree.SubElement(p._p, qn("w:del"))
rr = etree.SubElement(dl, qn("w:r"))
etree.SubElement(rr, qn("w:delText")).text = "BORRADO "
rr = etree.SubElement(p._p, qn("w:r"))
etree.SubElement(rr, qn("w:t")).text = "firme."
ok(ld.texto_docx(a_bytes(d)) == "Quedó firme.", "lo tachado en control de cambios no se lee")

print("\n3 · ROBUSTEZ")
d = docx.Document()
for i in range(50):
    d.add_paragraph(f"Párrafo {i} " + "x" * 100)
b = a_bytes(d)
ok(len(ld.texto_docx(b, tope=1000)) < 1300, "el tope detiene la lectura (defensa ante un ZIP que se infla)")
# El documento principal con otro nombre, declarado en _rels/.rels.
zin = zipfile.ZipFile(io.BytesIO(a_bytes(docx.Document())))
d2 = docx.Document()
d2.add_paragraph("Documento con otro nombre de parte.")
src = zipfile.ZipFile(io.BytesIO(a_bytes(d2)))
out = io.BytesIO()
with zipfile.ZipFile(out, "w") as z:
    for n in src.namelist():
        datos = src.read(n)
        if n == "word/document.xml":
            n = "word/document2.xml"
        if n == "_rels/.rels":
            datos = datos.replace(b"word/document.xml", b"word/document2.xml")
        if n == "[Content_Types].xml":
            datos = datos.replace(b"/word/document.xml", b"/word/document2.xml")
        z.writestr(n, datos)
ok(ld.texto_docx(out.getvalue()) == "Documento con otro nombre de parte.", "la parte principal se busca en _rels/.rels")
try:
    ld.texto_docx(b"esto no es un zip")
    ok(False, "un archivo que no es DOCX lanza")
except Exception:
    ok(True, "un archivo que no es DOCX lanza (quien llama decide)")
ok(asyncio.run(ld.leer_docx(a_bytes(d2))) == "Documento con otro nombre de parte.", "leer_docx: en un hilo, mismo texto")
try:
    asyncio.run(ld.leer_docx(b"nada"))
    ok(False, "leer_docx lanza si ni el repliegue puede")
except Exception:
    ok(True, "leer_docx lanza si ni el repliegue puede (las rutas ya lo convierten en su 400)")
ok(ld.devolver_memoria() is None, "devolver_memoria nunca lanza (en macOS no hay malloc_trim)")
# Una entidad externa no se resuelve (XXE).
x = zipfile.ZipFile(io.BytesIO(a_bytes(docx.Document())))
out = io.BytesIO()
with zipfile.ZipFile(out, "w") as z:
    for n in x.namelist():
        datos = x.read(n)
        if n == "word/document.xml":
            datos = (b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY e SYSTEM "file:///etc/passwd">]>'
                     b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     b'<w:body><w:p><w:r><w:t>A&e;B</w:t></w:r></w:p></w:body></w:document>')
        z.writestr(n, datos)
try:
    r = ld.texto_docx(out.getvalue())
    ok("root:" not in r, f"una entidad externa no se resuelve (salió {r!r})")
except Exception:
    ok(True, "una entidad externa no se resuelve (el XML se rechaza)")

print("\n4 · LA MEMORIA: un DOCX como el del 25-sep (cada palabra, una corrida con formato)")
SCRIPT = r'''
import io, resource, sys, docx
sys.path.insert(0, ".")
import lectura_docx as ld
b = open(sys.argv[2], "rb").read()
antes = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
if sys.argv[1] == "viejo":
    d = docx.Document(io.BytesIO(b)); t = "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
else:
    t = ld.texto_docx(b)
despues = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
print(len(t), despues - antes)
'''
ruta = Path(os.environ.get("TMPDIR", "/tmp")) / "docx_pesado_prueba.docx"
if not ruta.exists():
    d = docx.Document()
    for i in range(2500):
        p = d.add_paragraph()
        for j in range(60):
            r = p.add_run(f"palabra{j} ")
            r.bold = j % 2 == 0
            r.font.size = docx.shared.Pt(11 + j % 3)
            r.font.name = "Arial"
    d.save(ruta)
res = {}
for modo in ("viejo", "nuevo"):
    o = subprocess.run([sys.executable, "-c", SCRIPT, modo, str(ruta)], capture_output=True, text=True, cwd=".")
    n, delta = o.stdout.split()
    # ru_maxrss: bytes en macOS, KB en Linux.
    mb = int(delta) / (1024 * 1024 if sys.platform == "darwin" else 1024)
    res[modo] = (int(n), mb)
print(f"      {ruta.stat().st_size / 1e6:.1f} MB en disco · python-docx +{res['viejo'][1]:.0f} MB · "
      f"por flujo +{res['nuevo'][1]:.0f} MB")
ok(res["viejo"][0] == res["nuevo"][0], "el mismo número de caracteres por los dos caminos")
ok(res["nuevo"][1] < res["viejo"][1] / 5, "por flujo usa menos de la quinta parte de memoria")

print("\n5 · DÓNDE SE ENGANCHA")
F = Path("main.py").read_text(encoding="utf-8")
AD = F[F.index("async def analyze_document("):]
AD = AD[:AD.index("\n@app.")]
ok("await lectura_docx.leer_docx(content" in AD and "DocxDocument(io.BytesIO(content))" not in AD,
   "/analyze-document lee el DOCX por flujo, sin python-docx en la ruta")
ok("await asyncio.to_thread(_leer_pdf_nativo, content)" in AD, "/analyze-document: el PDF con texto se lee en un hilo")
ok("await asyncio.to_thread(_leer_doc_ole, content)" in AD, "/analyze-document: el .doc también")
ET = F[F.index("async def extract_text_from_document("):]
ET = ET[:ET.index("\n@app.")]
ok("await lectura_docx.leer_docx(content" in ET and "Document(io.BytesIO(content))" not in ET,
   "/extract-text: DOCX por flujo")
ok("await _extract_text_from_upload(_SubidaDeBytes(content, filename))" in ET
   and "gemini_client.models.generate_content" not in ET,
   "/extract-text: el PDF por el lector del taller (texto nativo, OCR sin bloquear), no Gemini síncrono")
ok("await asyncio.to_thread(_leer_doc_ole_extract, content)" in ET, "/extract-text: el .doc en un hilo")
UP = F[F.index("async def _extract_text_from_upload("):]
UP = UP[:UP.index("\n@app.")]
ok("await lectura_docx.leer_docx(content, \"\\n\")" in UP and "_Docx(io.BytesIO(content))" not in UP,
   "el lector del taller: DOCX por flujo, con su separador de una línea")

print()
if FALLOS:
    print("FALLAS:\n  " + "\n  ".join(FALLOS))
    sys.exit(1)
print("RESULTADO: TODAS LAS COMPROBACIONES PASAN")
