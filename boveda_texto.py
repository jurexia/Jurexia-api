"""LA BÓVEDA · EL TEXTO DE CADA PÁGINA.   (oct-2026)

Todo lo de la Bóveda cita a documento Y página: el chat, la tabla de revisión
y el visor que abre el PDF ahí. Por eso aquí nada se aplana en un solo texto:
cada documento es una lista de páginas.

PDF.   PyMuPDF lee la capa de texto gratis. Sólo las páginas SIN texto útil
       van a Azure (`prebuilt-read`, 1.50 USD por mil páginas), y van solas:
       Azure cobra por página ENVIADA, y el parámetro `pages` le manda sólo
       las escaneadas. La lección de FIREL (21-sep): una página escaneada con
       el sello de la firma electrónica trae ~93 caracteres de texto y pasaba
       por «con texto». Aquí una página es escaneada si casi no tiene texto
       (menos de TEXTO_MINIMO) Y una imagen le cubre buena parte de la hoja.

WORD.  No tiene páginas. Se parte en «páginas» de ~CUARTILLA caracteres por
       párrafos enteros, para que la cita diga dónde está y el visor la
       encuentre. Se avisa en la interfaz que es una página aproximada.

FRAGMENTOS. Cada fragmento vive dentro de UNA página (nunca cruza), con un
       poco de traslape, para que la página citada sea siempre exacta.
"""
from __future__ import annotations

import asyncio
import io
import os
import re
import unicodedata
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Optional

TEXTO_MINIMO = int(os.getenv("BOVEDA_TEXTO_MINIMO", "200"))     # caracteres
IMAGEN_MINIMA = float(os.getenv("BOVEDA_IMAGEN_MINIMA", "0.35"))  # fracción de la hoja
CUARTILLA = 3000
FRAGMENTO = int(os.getenv("BOVEDA_FRAGMENTO", "1400"))
TRASLAPE = 200

MIME_PDF = "application/pdf"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

AZURE_KEY = os.getenv("AZURE_DOCINT_KEY", "")
AZURE_ENDPOINT = os.getenv("AZURE_DOCINT_ENDPOINT", "").rstrip("/")
AZURE_PRECIO_PAGINA = 0.0015  # USD, S0 Read Pages


@dataclass
class Inspeccion:
    """Lo que se sabe de un archivo SIN pagar nada: cuántas páginas y cuáles
    irán a OCR. Es lo que se cotiza."""
    paginas: int
    paginas_ocr: list[int] = field(default_factory=list)   # 1-based
    textos: list[str] = field(default_factory=list)         # capa de texto por página
    aproximadas: bool = False                               # Word: páginas de cuartilla

    @property
    def n_ocr(self) -> int:
        return len(self.paginas_ocr)


def tipo_de(nombre: str, mime: Optional[str] = None) -> str:
    n = (nombre or "").lower()
    if (mime or "") == MIME_PDF or n.endswith(".pdf"):
        return "pdf"
    if (mime or "") == MIME_DOCX or n.endswith(".docx"):
        return "docx"
    return ""


def _limpio(t: str) -> str:
    t = (t or "").replace("\x00", "")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _cobertura_imagen(pagina) -> float:
    """Fracción de la hoja cubierta por imágenes (la mayor, no la suma: un
    escaneo es UNA imagen de página completa)."""
    area = abs(pagina.rect) or 1.0
    mayor = 0.0
    try:
        for info in pagina.get_image_info():
            b = info.get("bbox")
            if b:
                import fitz
                r = fitz.Rect(b) & pagina.rect
                mayor = max(mayor, abs(r) / area)
    except Exception:
        return 0.0
    return mayor


def inspeccionar_pdf(datos: bytes) -> Inspeccion:
    import fitz
    doc = fitz.open(stream=datos, filetype="pdf")
    try:
        textos, ocr = [], []
        for i, p in enumerate(doc, 1):
            t = _limpio(p.get_text("text"))
            textos.append(t)
            if len(t) < TEXTO_MINIMO and (_cobertura_imagen(p) >= IMAGEN_MINIMA or len(t) < 20):
                # Sin texto y sin imagen es una hoja en blanco: no se paga OCR por ella.
                if _cobertura_imagen(p) > 0.02:
                    ocr.append(i)
        return Inspeccion(paginas=doc.page_count, paginas_ocr=ocr, textos=textos)
    finally:
        doc.close()


def inspeccionar_docx(datos: bytes) -> Inspeccion:
    import docx
    d = docx.Document(io.BytesIO(datos))
    parrafos = [p.text.strip() for p in d.paragraphs if p.text and p.text.strip()]
    # Las tablas también son contrato (cuadros de pagos, anexos).
    for tabla in d.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text and c.text.strip()]
            if celdas:
                parrafos.append(" | ".join(dict.fromkeys(celdas)))
    paginas, actual = [], ""
    for p in parrafos:
        if actual and len(actual) + len(p) + 1 > CUARTILLA:
            paginas.append(actual)
            actual = ""
        actual = f"{actual}\n{p}" if actual else p
    if actual:
        paginas.append(actual)
    paginas = paginas or [""]
    return Inspeccion(paginas=len(paginas), textos=paginas, aproximadas=True)


def inspeccionar(datos: bytes, nombre: str, mime: Optional[str] = None) -> Inspeccion:
    t = tipo_de(nombre, mime)
    if t == "pdf":
        return inspeccionar_pdf(datos)
    if t == "docx":
        return inspeccionar_docx(datos)
    raise ValueError("Sólo se aceptan PDF y Word (.docx).")


def _rangos(paginas: list[int]) -> str:
    """[1,2,3,7,9,10] → '1-3,7,9-10' (el formato del parámetro `pages` de Azure)."""
    if not paginas:
        return ""
    s = sorted(set(paginas))
    out, ini, ant = [], s[0], s[0]
    for n in s[1:]:
        if n == ant + 1:
            ant = n
            continue
        out.append(f"{ini}-{ant}" if ini != ant else f"{ini}")
        ini = ant = n
    out.append(f"{ini}-{ant}" if ini != ant else f"{ini}")
    return ",".join(out)


async def ocr_azure(datos: bytes, paginas: list[int], plazo_s: float = 600) -> dict[int, str]:
    """OCR de SÓLO esas páginas. {número de página: texto}."""
    if not paginas:
        return {}
    if not (AZURE_KEY and AZURE_ENDPOINT):
        raise RuntimeError("Faltan AZURE_DOCINT_KEY y AZURE_DOCINT_ENDPOINT.")
    import httpx
    url = (f"{AZURE_ENDPOINT}/documentintelligence/documentModels/"
           f"prebuilt-read:analyze?api-version=2024-11-30&pages={_rangos(paginas)}")
    cab = {"Ocp-Apim-Subscription-Key": AZURE_KEY}
    async with httpx.AsyncClient(timeout=120) as cli:
        for intento in range(4):
            r = await cli.post(url, content=datos, headers={**cab, "Content-Type": MIME_PDF})
            if r.status_code == 429 and intento < 3:   # Azure S0: 15 peticiones por segundo
                await asyncio.sleep(2 ** intento)
                continue
            r.raise_for_status()
            break
        operacion = r.headers.get("Operation-Location")
        espera, gastado, j = 1.0, 0.0, {}
        while gastado < plazo_s:
            await asyncio.sleep(espera)
            gastado += espera
            espera = min(espera * 1.5, 5.0)
            s = await cli.get(operacion, headers=cab)
            if s.status_code == 429:
                continue
            j = s.json()
            if j.get("status") in ("succeeded", "failed"):
                break
        if j.get("status") != "succeeded":
            raise RuntimeError(f"Azure no terminó: {j.get('status') or 'plazo agotado'}")
    salida: dict[int, str] = {}
    for pg in j.get("analyzeResult", {}).get("pages", []):
        n = int(pg.get("pageNumber") or 0)
        salida[n] = _limpio("\n".join(l.get("content", "") for l in pg.get("lines", [])))
    return salida


# ── Fragmentos ──────────────────────────────────────────────────────────────
_CORTE = re.compile(r"(?<=[.;:!?])\s+|\n+")


def fragmentar_pagina(texto: str, tam: int = FRAGMENTO, traslape: int = TRASLAPE) -> list[str]:
    """Trozos de ~tam caracteres que cortan en fin de oración cuando se puede."""
    texto = (texto or "").strip()
    if not texto:
        return []
    if len(texto) <= tam:
        return [texto]
    trozos, ini = [], 0
    while ini < len(texto):
        fin = min(len(texto), ini + tam)
        if fin < len(texto):
            ventana = texto[ini + tam // 2: fin]
            cortes = [m.end() for m in _CORTE.finditer(ventana)]
            if cortes:
                fin = ini + tam // 2 + cortes[-1]
        trozos.append(texto[ini:fin].strip())
        if fin >= len(texto):
            break
        ini = max(fin - traslape, ini + 1)
    return [t for t in trozos if t]


@dataclass
class Fragmento:
    pagina: int
    orden: int
    texto: str


def fragmentar(paginas: list[str]) -> list[Fragmento]:
    out = []
    for i, t in enumerate(paginas, 1):
        for j, trozo in enumerate(fragmentar_pagina(t)):
            out.append(Fragmento(pagina=i, orden=j, texto=trozo))
    return out


# ── Verificar una cita literal contra el texto ─────────────────────────────
def plano(t: str) -> str:
    """Minúsculas, sin acentos, sin comillas tipográficas y un solo espacio."""
    t = unicodedata.normalize("NFD", t or "")
    t = "".join(c for c in t if unicodedata.category(c) != "Mn").lower()
    t = t.replace("«", '"').replace("»", '"').replace("“", '"').replace("”", '"').replace("’", "'")
    t = re.sub(r"[\"'`]", "", t)
    t = re.sub(r"-\s*\n\s*", "", t)          # palabra partida al final de renglón
    return re.sub(r"\s+", " ", t).strip()


@dataclass
class Verificacion:
    verificada: bool
    pagina: Optional[int]      # dónde está DE VERDAD (puede corregir la del modelo)
    similitud: float           # 1.0 = literal


_NUMERO = re.compile(r"\d+(?:[.,]\d+)*")


def _numeros(t: str) -> list[str]:
    return [n.replace(",", "").replace(".", "") for n in _NUMERO.findall(t)]


def _cifras_iguales(aguja: str, tramo: str) -> bool:
    """Cada número de la cita está en el tramo. En una cita jurídica una cifra
    ES la cita: «30%» por «15%» es otra cláusula aunque se parezca al 95%."""
    disponibles = _numeros(tramo)
    for n in _numeros(aguja):
        if n in disponibles:
            disponibles.remove(n)
        else:
            return False
    return True


def _mejor_parecido(aguja: str, pajar: str) -> tuple[float, str]:
    """El mejor parecido de `aguja` con algún tramo de `pajar` de su mismo
    largo, y ese tramo."""
    if not aguja or not pajar:
        return 0.0, ""
    if aguja in pajar:
        return 1.0, aguja
    n = len(aguja)
    sm = SequenceMatcher(None, autojunk=False)
    sm.set_seq2(aguja)
    mejor, mejor_tramo = 0.0, ""
    # Anclas: dónde aparecen las primeras y las últimas palabras de la cita.
    palabras = aguja.split()
    anclas = set()
    for clave in (" ".join(palabras[:3]), " ".join(palabras[-3:])):
        if len(clave) < 8:
            continue
        for m in re.finditer(re.escape(clave), pajar):
            anclas.add(max(0, m.start() - n))
            anclas.add(m.start())
    if not anclas:
        paso = max(1, n // 3)
        anclas = set(range(0, max(1, len(pajar) - n + 1), paso))
        if len(anclas) > 400:            # páginas enormes: no vale la pena
            return 0.0, ""
    for a in anclas:
        tramo = pajar[a:a + n + n // 5]
        sm.set_seq1(tramo)
        # Cota barata antes de la cara: quick_ratio·(l1+l2)/2 acota lo que coincide.
        if sm.quick_ratio() * (len(tramo) + n) / 2 / n <= mejor:
            continue
        # Qué parte de la CITA aparece, en orden, en el tramo (no ratio(), que
        # castiga el largo del tramo y deja una cita literal en 0.91).
        cubierto = sum(b.size for b in sm.get_matching_blocks()) / n
        if cubierto > mejor:
            mejor, mejor_tramo = cubierto, tramo
        if mejor >= 0.98:
            break
    return mejor, mejor_tramo


def verificar_cita(cita: str, paginas: list[str], pagina_dicha: Optional[int] = None,
                   umbral: float = 0.9) -> Verificacion:
    """¿La cita está en el documento? Busca primero en la página que dijo el
    modelo, luego en las vecinas y luego en todas. Una cita de menos de 12
    caracteres no prueba nada y no se verifica."""
    aguja = plano(cita)
    if len(aguja) < 12 or not paginas:
        return Verificacion(False, pagina_dicha, 0.0)
    orden = list(range(1, len(paginas) + 1))
    if pagina_dicha and 1 <= pagina_dicha <= len(paginas):
        cerca = [pagina_dicha, pagina_dicha + 1, pagina_dicha - 1]
        orden = [p for p in cerca if 1 <= p <= len(paginas)] + [p for p in orden if p not in cerca]
    planos = {}
    mejor = Verificacion(False, pagina_dicha, 0.0)
    for p in orden:
        pl = planos.setdefault(p, plano(paginas[p - 1]))
        if aguja in pl:
            return Verificacion(True, p, 1.0)
    # Una cita que cruza de una página a la siguiente.
    for p in range(1, len(paginas)):
        junto = planos.get(p, plano(paginas[p - 1])) + " " + planos.get(p + 1, plano(paginas[p]))
        if aguja in junto:
            return Verificacion(True, p, 1.0)
    for p in orden:
        s, tramo = _mejor_parecido(aguja, planos[p])
        if s > mejor.similitud:
            mejor = Verificacion(s >= umbral and _cifras_iguales(aguja, tramo), p, round(s, 3))
        if s >= 0.98:
            break
    return mejor
