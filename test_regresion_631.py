# -*- coding: utf-8 -*-
"""EL RETROCESO DEL 28-SEP-2026 (AR 631/2025), SIN LLAMAR A NINGÚN MODELO.

David generó en la cuenta de soporte un amparo en revisión que la tercera
interesada interpuso contra una sentencia que CONCEDIÓ el amparo. Dictó
«fundado» (revocar) para todo el asunto. El proyecto salió así:

  1. LA RAZÓN ARGUMENTABA LO CONTRARIO DEL SENTIDO. El cuadro de la razón
     global traía la del motor para «infundado»; al pedir la de «fundado», esa
     razón viajó como «la base del secretario —desarróllala, no la discutas—» y
     el modelo la desarrolló: «la sustitución sí alteró la cosa juzgada…».
  2. DIECINUEVE ARGUMENTOS «FUNDADOS» QUE EL PLANIFICADOR HABÍA DESESTIMADO. El
     planificador escribió la RAZÓN «fondo_desestimado» en el campo de la
     calificación; `reparar` la trató como ilegible y les puso la del problema.
  3. REVOCÓ PARA VOLVER A CONCEDER. Lo que hizo el juzgado se leyó contando
     verbos en el PDF entero —la sentencia narraba un amparo ANTERIOR que «negó
     el amparo»— y salió «niega»; con el recurso fundado, la rama fue
     «revoca_fondo_concede» y el resolutivo amparó a la recurrente contra un
     juzgado cuyo acto se había sobreseído. Su propio resolutivo decía «ampara y
     protege»: revocar esa concesión es NEGAR.

Todo es SINTÉTICO: nombres, números y fechas de prueba.

    .venv/bin/python test_regresion_631.py
"""
import ast
import copy
import datetime as _dt
import os
import sys
import tempfile
import types

os.environ.pop("ESTUDIO_PROMPT", None)
os.environ.pop("PLAN_ESTUDIO", None)

import fase5_propuesta as f5
import fase6_estudio as f6
import fase_rama as fr
import fases123_pipeline as f123
import modos_decision as md
import plan_estudio as pe
import tipos_asunto as ta

FALLOS = []


def ok(cond, que):
    print(f"   {'PASA ' if cond else 'FALLA'}  {que}")
    if not cond:
        FALLOS.append(que)


AQUI = os.path.dirname(os.path.abspath(__file__))

# ═══ EL ASUNTO DE PRUEBA ═══════════════════════════════════════════════════
ACTO = ("El Juzgado de Distrito consideró que la sustitución de la parte ejecutante alteró la cosa "
        "juzgada, porque la adquirente del inmueble no era causahabiente del contrato de "
        "arrendamiento que sustentó la acción personal de rescisión, y concedió el amparo. "
        "Consideró innecesario estudiar los restantes conceptos de violación.")
ESCRITO = (
    "AGRAVIOS\n"
    "PRIMERO. La sentencia recurrida vulnera los principios de exhaustividad y congruencia, porque "
    "la interlocutoria sólo determinó quién cuenta con legitimación para continuar la ejecución sin "
    "modificar las prestaciones de la sentencia definitiva. Además, la sentencia de amparo omite "
    "ponderar que el artículo 49 del Código Procesal Civil local permite que el causahabiente "
    "sustituya a quien transmitió el derecho controvertido. También sostiene que se vulnera el "
    "principio de igualdad jurídica porque no se aplican de manera uniforme las normas procesales "
    "a las dos partes. Asimismo afirma que la sentencia carece de fundamentación y motivación "
    "porque no aplica los artículos 2284 y 2294 del Código Civil local al caso concreto. Finalmente "
    "invoca el principio de privilegio del fondo sobre la forma para sostener que la interlocutoria "
    "atiende al fondo de la controversia incidental.\n\n"
    "SEGUNDO. Solicita que, como consecuencia de los argumentos expuestos en el primero, se tenga "
    "por reproducido su contenido para controvertir el considerando octavo de la sentencia recurrida.")
_o = [ESCRITO.index(x) for x in ("PRIMERO.", "SEGUNDO.")] + [len(ESCRITO)]
CONTEO = {"estado": "contado", "n": 2, "tramos": [[_o[i], _o[i + 1]] for i in range(2)]}
PREG1 = "¿La sustitución de la parte ejecutante alteró la cosa juzgada?"
PREG2 = "¿La sentencia recurrida fue exhaustiva y debidamente fundada al resolver sobre la sustitución?"
RAZON1 = ("La interlocutoria sólo determinó quién cuenta con legitimación para continuar la "
          "ejecución, sin modificar las prestaciones de la sentencia definitiva; el artículo 49 del "
          "Código Procesal Civil local permite al causahabiente sustituir a quien transmitió el "
          "derecho controvertido.")
RAZON2 = "Dado el sentido del estudio del problema principal, queda sin materia el análisis de este planteamiento."


def fases(**kw):
    f = f123.Fases123(
        resumen_acto=ACTO,
        resumen_conceptos="Primero: la interlocutoria no alteró la cosa juzgada.\nSegundo: reproduce el primero.",
        problemas=[{"pregunta": PREG1, "cubre": [1], "clase": "fondo", "jerarquia": "principal",
                    "combate": "La recurrente sostiene que la interlocutoria sólo determinó quién "
                               "ejecuta, sin modificar las prestaciones.",
                    "resolvio": "El Juzgado de Distrito determinó que la sustitución alteró la cosa "
                                "juzgada y concedió el amparo."},
                   {"pregunta": PREG2, "cubre": [2], "clase": "fondo", "jerarquia": "accesorio",
                    "depende_de": 1}],
        fuentes=[ACTO, ESCRITO], conteo=dict(CONTEO))
    for k, v in kw.items():
        setattr(f, k, v)
    return f


def crit(s1="fundado", s2="innecesario"):
    return [f6.Criterio(problema=PREG1, sentido=s1, razonamiento=RAZON1, jerarquia="principal"),
            f6.Criterio(problema=PREG2, sentido=s2, razonamiento=RAZON2, jerarquia="accesorio")]


def material():
    return f6.Material(tesis=[], normas=[
        {"cuerpo_legal": "Código Procesal Civil local", "articulo": "49",
         "texto": "El causahabiente puede sustituir a quien transmitió el derecho controvertido."}],
        tipo_asunto="amparo_revision", materia="civil", variante="v4")


SEGS = [pe.normalizar_segmento_piso(s) for s in [
    {"id": "A1.a", "concepto": 1, "texto": "La interlocutoria sólo determinó quién ejecuta.",
     "cita": "la interlocutoria sólo determinó quién cuenta con legitimación para continuar la ejecución",
     "anclas": []},
    {"id": "A1.b", "concepto": 1, "texto": "El artículo 49 permite la sustitución del causahabiente.",
     "cita": "el artículo 49 del Código Procesal Civil local permite que el causahabiente sustituya",
     "anclas": ["art 49 cpc"]},
    {"id": "A1.c", "concepto": 1, "texto": "Igualdad jurídica.",
     "cita": "se vulnera el principio de igualdad jurídica porque no se aplican de manera uniforme",
     "anclas": []},
    {"id": "A1.d", "concepto": 1, "texto": "Falta de fundamentación: artículos 2284 y 2294.",
     "cita": "no aplica los artículos 2284 y 2294 del Código Civil local al caso concreto",
     "anclas": ["art 2284 cc", "art 2294 cc"]},
    {"id": "A1.e", "concepto": 1, "texto": "Privilegio del fondo sobre la forma.",
     "cita": "invoca el principio de privilegio del fondo sobre la forma para sostener",
     "anclas": []},
    {"id": "A2.a", "concepto": 2, "texto": "Reproduce el primero contra el considerando octavo.",
     "cita": "se tenga por reproducido su contenido para controvertir el considerando octavo",
     "anclas": []},
]]


def _s(sid, **kw):
    base = {"id": sid, "problema_id": 1, "vicio": "fondo", "ataca": "P1", "reitera": None,
            "dato": None, "etiqueta": "fondo_desestimado", "razon": "fondo_desestimado",
            "trat": "desarrolla", "diferencia": "norma", "pendiente": None}
    base.update(kw)
    return base


PLAN_631 = {
    "proposiciones": [
        {"id": "P1", "dice": "La sustitución alteró la cosa juzgada", "caracter": "toral",
         "relacion": "necesaria", "fuente": "recurrida",
         "cita": "la sustitución de la parte ejecutante alteró la cosa juzgada",
         "vinculada_por_ejecutoria": False}],
    "segmentos": [
        # Lo que devolvió el planificador en el 631: UNO fundado, y los demás
        # con la RAZÓN «fondo_desestimado» escrita en el campo de la etiqueta.
        _s("A1.a", etiqueta="fundado", razon="fundado", trat="aplica", diferencia=None),
        _s("A1.b"), _s("A1.c", vicio="fondo"), _s("A1.d", vicio="forma"), _s("A1.e"),
        _s("A2.a", problema_id=2, etiqueta="innecesario", razon="cae_con_principal",
           trat="no_se_estudia", diferencia=None, ataca=None)],
    "premisas": [],
    "unidades": [{"id": "U1", "problemas": [1], "segmentos": ["A1.a"], "premisa": None, "objecion": None},
                 {"id": "U2", "problemas": [1], "segmentos": ["A1.b", "A1.c", "A1.e"], "premisa": None,
                  "objecion": None},
                 {"id": "U3", "problemas": [1], "segmentos": ["A1.d"], "premisa": None, "objecion": None}],
    "propuestas": [], "avisos_al_secretario": [],
    "orden": {"criterio": "promovente", "por_que": "el orden del escrito"},
}


# LA TRADUCCIÓN SOLA (577c700, plan-5). Desde plan-6 lo que en un problema que
# prospera no lo funda es innecesario por suficiencia —la jerarquía lo
# resuelve después de traducir—; para ver la traducción sin ella, cada uno
# alega aquí una omisión, que nunca depende del principal. (Hasta la revisión
# adversarial del 28-sep-2026 pedían una «consecuencia distinta»; pero un
# infundado que la pide no da nada más y ahora es innecesario igual.)
PLAN_631_C = copy.deepcopy(PLAN_631)
for _x in PLAN_631_C["segmentos"]:
    if _x["id"] in ("A1.b", "A1.c", "A1.d", "A1.e"):
        _x["vicio"] = "omision"


def _norm(d):
    p = pe.normalizar(copy.deepcopy(d))
    p["tipo_asunto"] = "amparo_revision"
    return p


def _rep(d, c=None):
    return pe.reparar(_norm(d), c or crit(), SEGS, fases(), material(), "", {})


def seg(p, sid):
    return next(s for s in p["segmentos"] if s["id"] == sid)


# ═══════════════════════════════════════════════════════════════════════════
print("\n1 · EL PLAN: LA RAZÓN ESCRITA COMO CALIFICACIÓN NO HEREDA LA DEL PROBLEMA")
ok(pe.etiqueta_fuera("fondo_desestimado", "fundado") != "",
   "el caso reproduce la trampa: «fondo_desestimado» no es una calificación del catálogo")
_r, _av = _rep(PLAN_631_C)
_et = {s["id"]: s["etiqueta"] for s in _r["segmentos"]}
ok(_et["A1.a"] == "fundado", "el argumento que el planificador fundó sigue fundado")
ok(all(_et[x] == "infundado" for x in ("A1.b", "A1.c", "A1.d", "A1.e")),
   f"los desestimados salen «infundado», la calificación que su razón implica, NO «fundado»: {_et}")
ok(sum(1 for x in ("A1.a", "A1.b", "A1.c", "A1.d", "A1.e") if _et[x] == "fundado") == 1,
   "en el problema que prospera, UN argumento lo funda: los demás no heredan su sentido")
ok(all(seg(_r, x)["razon"] == "fondo_desestimado" for x in ("A1.b", "A1.c", "A1.d", "A1.e")),
   "la razón del planificador se queda (antes se «ajustaba» a «fundado»)")
ok(not any("lleva la calificación del problema" in a for a in _av),
   "ningún argumento «lleva la calificación del problema»")
ok(any("leídas como la calificación que esa razón implica" in a and "A1.b" in a
       for a in _r["avisos_al_secretario"]),
   "el panel dice qué calificaciones se leyeron de su razón")
ok(_et["A2.a"] == "innecesario", "el accesorio que no se estudia conserva su sentido")
_f = pe.validar(_r, crit(), SEGS, fases(), material(), "", suplencia={})
ok(not any("etiqueta" in x or "problema 1" in x for x in _f),
   f"V0 no rechaza las calificaciones del plan reparado: {[x for x in _f if 'etiqueta' in x][:3]}")
# plan-6: SIN esa consecuencia distinta, lo que el planificador desestimó dentro
# del problema que prospera por A1.a es innecesario por suficiencia —nunca
# «fundado» heredado (el 631), ni un infundado que el estudio desarrolle—.
_r6, _ = _rep(PLAN_631)
_et6 = {s["id"]: s["etiqueta"] for s in _r6["segmentos"]}
ok(_et6["A1.a"] == "fundado"
   and all(_et6[x] == "innecesario" and seg(_r6, x)["razon"] == "innecesario_por_suficiencia"
           and seg(_r6, x)["con"] == "A1.a" and seg(_r6, x)["trat"] == "no_se_estudia"
           for x in ("A1.b", "A1.c", "A1.d", "A1.e")),
   f"plan-6: los desestimados, innecesarios por suficiencia (decide A1.a), ninguno fundado: {_et6}")
ok(any("leídas como la calificación que esa razón implica" in a and "A1.b" in a
       for a in _r6["avisos_al_secretario"]),
   "…y la traducción de su razón se sigue diciendo: se tradujo antes de resolverlos")

print("\n2 · EL PLAN: TODOS DESESTIMADOS DENTRO DE UN PROBLEMA FUNDADO → LO FUNDA UNO, NO TODOS")
_d = copy.deepcopy(PLAN_631_C)
_d["segmentos"][0].update(etiqueta="fondo_desestimado", razon="fondo_desestimado",
                          trat="desarrolla", diferencia="norma")
_r2, _av2 = _rep(_d)
_et2 = {s["id"]: s["etiqueta"] for s in _r2["segmentos"]}
ok(_et2["A1.a"] == "fundado" and all(_et2[x] == "infundado" for x in ("A1.b", "A1.c", "A1.d", "A1.e")),
   f"lo funda el primero que ataca la proposición toral; los demás conservan lo suyo: {_et2}")
ok(_r2["propuestas"] and _r2["propuestas"][0]["a"] == "infundado" and _r2["propuestas"][0]["de"] == "fundado",
   "lo que opinaba el planificador va a PROPUESTA del problema (sólo la aplica el secretario)")
ok(any("problema 1" in a and "A1.a" in a and "conservan" in a for a in _r2["avisos_al_secretario"]),
   "y el panel dice quién lo funda y que los demás conservan su calificación")
ok(pe.problemas_sin_quien_los_funde(_r2["segmentos"], pe.problemas_del_criterio(crit(), fases())) == [],
   "el sentido del secretario tiene quien lo funde")

print("\n3 · EL PLAN: LA ETIQUETA VACÍA SE LEE DE SU RAZÓN")
_d = copy.deepcopy(PLAN_631_C)
_d["segmentos"][2].update(etiqueta="", razon="no_combate(P1)")
_r3, _ = _rep(_d)
ok(seg(_r3, "A1.c")["etiqueta"] == "inoperante",
   f"vacía con razón «no_combate» → «inoperante», no «fundado»: {seg(_r3, 'A1.c')['etiqueta']}")
ok(pe.PLAN_VERSION == "plan-6", "la versión del plan sube: ni un plan-4 guardado con la herencia ni un plan-5 "
                                 "con pendientes de razón se reutilizan")
_pr = pe.prompt_plan(tipo_asunto="amparo_revision", probs=pe.problemas_del_criterio(crit(), fases()),
                     segs=SEGS, resumen_acto=ACTO, tramos=[("escrito", ESCRITO)],
                     indice=pe.indice_material(material(), fases()))
ok("etiqueta (la calificación del segmento): fundado | esencialmente_fundado" in _pr
   and "nunca en «etiqueta»" in _pr,
   "el prompt del planificador trae el catálogo de las etiquetas y dice que la razón no va ahí")

# ═══════════════════════════════════════════════════════════════════════════
print("\n4 · LA RAZÓN DE LA OTRA VÍA NO ES LA DEL SECRETARIO")
R_MOTOR = ("La adquisición del inmueble no acredita la transmisión de los derechos litigiosos; la "
           "sustitución alteró los límites subjetivos de la cosa juzgada.")
R_ALT = ("La interlocutoria sólo sustituyó a quien ejecuta, sin alterar la cosa, la causa ni las "
         "prestaciones de la sentencia firme.")
GLOB = {"sentido": "infundado", "razon": R_MOTOR,
        "alternativa": {"sentido": "fundado", "razon": R_ALT}}
_t, _a = md.razon_de_la_otra_via(R_MOTOR, "fundado", GLOB)
ok(_t == R_ALT and "VÍA CONTRARIA" in _a,
   "el eco del motor para «infundado» pedido como base de «fundado» → la de la vía contraria ya escrita")
_t, _a = md.razon_de_la_otra_via("  " + R_MOTOR.replace(" ", "  ") + " ", "fundado", GLOB)
ok(_t == R_ALT, "los espacios no esconden el eco")
_t, _a = md.razon_de_la_otra_via(R_MOTOR, "fundado", {"sentido": "infundado", "razon": R_MOTOR})
ok(_t == "" and "no se usó" in _a, "sin vía contraria escrita, la razón se queda vacía y se dice")
_t, _a = md.razon_de_la_otra_via(R_MOTOR, "inoperante", GLOB)
ok(_t == R_MOTOR and _a == "", "la misma vía (infundado → inoperante) no es eco: se respeta")
_t, _a = md.razon_de_la_otra_via("Mi razón: la interlocutoria no modificó las prestaciones.", "fundado", GLOB)
ok(_a == "" and _t.startswith("Mi razón"), "lo que escribió el secretario se respeta siempre")
_t, _a = md.razon_de_la_otra_via(R_MOTOR, "fundado", None)
ok(_t == R_MOTOR and _a == "", "sin propuesta global no hay nada que comparar")

print("\n5 · LA RAZÓN SABE QUIÉN GANA SI PROSPERA")
_p = f5.prompt_razon(PREG1, "fundado", f6.Material(), ACTO, "", True, "amparo_revision",
                     combate=fases().problemas[0]["combate"], resolvio=fases().problemas[0]["resolvio"])
ok("QUÉ QUIERE DECIR ESA CALIFICACIÓN EN ESTE PLANTEAMIENTO" in _p
   and "prosperan los agravios: cae lo que resolvió" in _p
   and "sólo determinó quién ejecuta" in _p and "concedió el amparo" in _p,
   "«fundado» en una revisión: prosperan los agravios y cae lo resuelto, con lo que sostienen y lo resuelto")
_p = f5.prompt_razon(PREG1, "infundado", f6.Material(), ACTO, "", True, "amparo_revision",
                     combate="x sostiene y", resolvio="el juez resolvió z")
ok("no prosperan los agravios: subsiste lo que resolvió" in _p, "«infundado»: subsiste lo resuelto")
_p = f5.prompt_razon(PREG2, "innecesario", f6.Material(), ACTO, "", True, "amparo_revision",
                     combate="x", resolvio="z")
ok("QUÉ QUIERE DECIR" not in _p, "lo que no se estudia no lleva dirección")
_p = f5.prompt_razon(PREG1, "fundado", f6.Material(), ACTO, "", True, "amparo_revision",
                     directriz="Una base cualquiera del secretario.")
ok("La calificación de arriba es la decisión" in _p,
   "con directriz: la base se lleva hasta la calificación, no a la contraria")
ok("QUÉ QUIERE DECIR" not in f5.prompt_razon(PREG1, "fundado", f6.Material(), ACTO, "", True,
                                              "amparo_revision"),
   "sin datos del problema, el prompt queda como antes")

print("\n6 · LAS PUERTAS DEL SERVIDOR USAN LA REGLA (por AST)")
_src = open(os.path.join(AQUI, "main.py"), encoding="utf8").read()
_arbol = ast.parse(_src)


def _fn(nombre):
    return next(n for n in ast.walk(_arbol)
                if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name == nombre)


def _llama(fn, atributo):
    return sum(1 for n in ast.walk(fn) if isinstance(n, ast.Attribute) and n.attr == atributo)


_raz = _fn("taller_razonar")
ok(_llama(_raz, "razon_de_la_otra_via") >= 1, "/taller/razonar limpia la directriz que es eco de la otra vía")
_kws = {k.arg for n in ast.walk(_raz) if isinstance(n, ast.Call) and getattr(n.func, "attr", "") == "razonar"
        for k in n.keywords}
ok({"combate", "resolvio"} <= _kws, "/taller/razonar le pasa al modelo qué sostienen y qué se resolvió")
ok(_llama(_fn("_taller_armar_criterio"), "razon_de_la_otra_via") >= 2,
   "el criterio común limpia el eco en el modo global y por problema (las cuatro puertas lo usan)")

# ═══════════════════════════════════════════════════════════════════════════
print("\n7 · LO QUE HIZO EL JUZGADO LO DICE SU RESOLUTIVO")
RESOL = ("La Justicia de la Unión ampara y protege a Unión Ejemplo, A.C., en contra del acto "
         "atribuido a la Sala Civil Uno del Tribunal Superior de Justicia, por los motivos "
         "expresados en el considerando séptimo de esta sentencia y para los efectos precisados en "
         "el último considerando.")
SENTENCIA = (
    "ANTECEDENTES. Inconforme, la demandada promovió amparo directo y el Tribunal Colegiado negó el "
    "amparo solicitado. Después, en otro juicio, el Tribunal negó la protección. Más tarde el "
    "Colegiado negó el amparo por segunda vez. "
    "CONSIDERANDO SÉPTIMO. Es fundado el concepto de violación y procede conceder el amparo. "
    "Por lo expuesto, se R E S U E L V E: ÚNICO. " + RESOL + " Notifíquese.")
ok(fr.resolvio_a_quo(SENTENCIA) == "niega",
   "el caso reproduce la trampa: contar verbos en todo el texto da «niega» (un amparo anterior)")
ok(fr.resolvio_segun_resolutivos(SENTENCIA) == "concede", "sus puntos resolutivos dicen «concede»")
ok(fr.resolvio_a_quo(SENTENCIA, resolutivo=fr.resolutivo_recurrida(SENTENCIA)) == "concede",
   "con el resolutivo delante, `resolvio_a_quo` dice «concede»")
_mixta = ("se R E S U E L V E: PRIMERO. Se sobresee en el juicio respecto del acto del Juzgado "
          "Quinto. SEGUNDO. " + RESOL + " Notifíquese.")
ok(fr.resolvio_segun_resolutivos(_mixta) == "sobresee_concede", "sobresee un acto y ampara por otro → mixta")
ok(fr.que_dice_el_resolutivo("La Justicia de la Unión no ampara ni protege a X contra el acto de Y.")
   == "niega", "«no ampara ni protege» → niega")
ok(fr.que_dice_el_resolutivo("ampara y protege a X; no ampara ni protege a Z") == "",
   "ampara a uno y niega a otro: no se adivina")
ok(fr.que_dice_el_resolutivo("Se estima que el juez negó el amparo.") == "",
   "la prosa en pasado no es un resolutivo")
_fx = types.SimpleNamespace(resolutivo_recurrida=RESOL, resolvio_a_quo="niega",
                            antecedentes="El juzgado negó el amparo.")
ok(fr.que_hizo_el_juzgado(_fx, "") == "concede", "el resolutivo manda sobre la lectura guardada del PDF")
_fx = types.SimpleNamespace(resolutivo_recurrida="", resolvio_a_quo="",
                            antecedentes="Seguido el juicio, el juzgado concedió el amparo a la quejosa.")
ok(fr.que_hizo_el_juzgado(_fx, "") == "concede",
   "los antecedentes son una CADENA: ya no se leen letra por letra")
_fx = types.SimpleNamespace(resolutivo_recurrida="", resolvio_a_quo="", antecedentes="")
ok(fr.que_hizo_el_juzgado(_fx, "Concedió el amparo porque la sustitución alteró la cosa juzgada.")
   == "concede", "sin papel, lo declarado por el motor")
ok(ta.rama_revision("concede", "fundado") == "revoca_fondo_niega",
   "juzgado que concede + recurso fundado → se revoca y se NIEGA")
_pts = ta.puntos_revoca_concesion(fr.resolutivo_recurrida(SENTENCIA))
ok(_pts == ["PRIMERO. Se revoca la sentencia recurrida.",
            "SEGUNDO. La Justicia de la Unión no ampara ni protege a Unión Ejemplo, A.C., en contra del "
            "acto atribuido a la Sala Civil Uno del Tribunal Superior de Justicia, por los motivos y "
            "fundamentos expuestos en el último considerando de esta ejecutoria."],
   f"revocar la concesión es negar ESO, con el sujeto y el acto del juzgado: {_pts}")
ok(ta.puntos_revoca_concesion("Se sobresee en el juicio.") == []
   and ta.puntos_revoca_concesion("La Justicia de la Unión no ampara ni protege a X contra el acto de Y.") == [],
   "si el resolutivo no es un «ampara y protege» limpio, va la fórmula genérica")

print("\n8 · EL DOCUMENTO: REVOCA Y NIEGA A QUIEN SE AMPARÓ, AUNQUE LA SESIÓN GUARDARA «niega»")
try:
    import documento_generado as dg
    import fase0_oportunidad as f0
    from docx import Document

    _est = dg.Estructura(apertura="V.", visto="para resolver.",
                         resultandos=[{"titulo": "Presentación de la demanda de amparo indirecto",
                                       "texto": "Unión Ejemplo, A.C. promovió amparo contra la Sala "
                                                "Civil Uno y el Juzgado Quinto."}],
                         competencia="", existencia="", procedencia="")
    _c = f0.computar(_dt.date(2026, 3, 2), _dt.date(2026, 3, 9), plazo=10)
    _ruta = os.path.join(tempfile.mkdtemp(prefix="regresion631_"), "ar.docx")
    dg.componer({"tipo_asunto": "amparo_revision", "numero": "1/2026",
                 "encabezado": "AMPARO EN REVISIÓN 1/2026",
                 # Lo que tecleó el formulario en el 631: la RECURRENTE en «quejoso».
                 "quejoso": "Inmobiliaria Ejemplo, S.A. de C.V.",
                 "recurrente": "Inmobiliaria Ejemplo, S.A. de C.V.",
                 "responsable": "Juzgado Séptimo de Distrito", "magistrado": "M", "secretario": "S",
                 "tribunal": "Tercer Tribunal Colegiado", "ciudad": "Querétaro",
                 "resolvio_a_quo": "niega",               # lo que guardó la sesión
                 "resolutivo_recurrida": fr.resolutivo_recurrida(SENTENCIA)},
                _est, _c, f0.fecha_en_letra, _ruta,
                estudio=["Resolución recurrida",
                         "El Juzgado de Distrito precisó que la protección se concedía para que la "
                         "Sala Civil Uno dejara insubsistente la resolución reclamada.",
                         "Solución",
                         "Es fundado el primer agravio: la interlocutoria sólo determinó quién "
                         "ejecuta, sin modificar las prestaciones; por tanto, se revoca la sentencia "
                         "recurrida y procede negar el amparo."],
                calificaciones=["fundado"], tipo_asunto="amparo_revision")
    _ps = [q.text for q in Document(_ruta).paragraphs if q.text.strip()]
    _res = " ".join(_ps[next(i for i, t in enumerate(_ps) if "R E S U E L V E" in t or t.strip() == "RESUELVE"):])
    ok("Se revoca la sentencia recurrida" in _res, "PRIMERO: se revoca")
    ok("no ampara ni protege a Unión Ejemplo, A.C." in _res and "Sala Civil Uno" in _res,
       "SEGUNDO: NIEGA, a quien amparó el juzgado y contra el acto por el que lo amparó")
    ok("ampara y protege a Inmobiliaria" not in _res and "Juzgado Quinto" not in _res,
       "no ampara a la recurrente ni contra el juzgado sobreseído")
    ok(any("SE TOMÓ DE SU PUNTO RESOLUTIVO" in a for a in _est.avisos)
       and any("revoca_fondo_niega" in a for a in _est.avisos),
       "el aviso dice que manda el resolutivo y la rama es «revoca_fondo_niega»")
    import ensamblar_adelanto as ens
    _cg = ens.revisar_congruencia(_ruta, ["fundado"], "amparo_revision")
    ok(not any("INCONGRUENCIA QUE INVALIDA" in a for a in _cg),
       f"narrar que el juzgado concedió «para que … dejara insubsistente» no son efectos de "
       f"este proyecto: sin la alarma de incongruencia {[a[:60] for a in _cg]}")
    ok(bool(ens._RX_EFECTOS.search("la autoridad responsable deberá dejar insubsistente la resolución"))
       and bool(ens._RX_EFECTOS.search("la Sala dejará insubsistente la resolución"))
       and not ens._RX_EFECTOS.search("se concedía para que la Sala dejara insubsistente la resolución"),
       "«deberá dejar» y «dejará» son órdenes; «dejara» es la narración de otra sentencia")
    # LA ALARMA VERDADERA SIGUE: efectos escritos en la Solución y un resolutivo que niega.
    _est2 = dg.Estructura(apertura="V.", visto="para resolver.",
                          resultandos=[{"titulo": "Presentación de la demanda de amparo indirecto",
                                        "texto": "Unión Ejemplo, A.C. promovió amparo contra la Sala "
                                                 "Civil Uno."}],
                          competencia="", existencia="", procedencia="")
    _ruta2 = os.path.join(os.path.dirname(_ruta), "ar2.docx")
    dg.componer({"tipo_asunto": "amparo_revision", "numero": "2/2026",
                 "encabezado": "AMPARO EN REVISIÓN 2/2026", "quejoso": "Unión Ejemplo, A.C.",
                 "responsable": "Juzgado Séptimo de Distrito", "magistrado": "M", "secretario": "S",
                 "tribunal": "Tercer Tribunal Colegiado", "ciudad": "Querétaro",
                 "resolvio_a_quo": "concede",
                 "resolutivo_recurrida": fr.resolutivo_recurrida(SENTENCIA)},
                _est2, _c, f0.fecha_en_letra, _ruta2,
                # Desde el 28-sep-2026 el segundo punto sale de lo que concluye
                # el estudio (art. 93, fr. VI) y, si calla, va con hueco: para que
                # el resolutivo NIEGUE, el estudio tiene que decirlo.
                estudio=["Solución",
                         "Es fundado el agravio y la autoridad responsable deberá dejar insubsistente "
                         "la resolución reclamada; por tanto, procede negar el amparo."],
                calificaciones=["fundado"], tipo_asunto="amparo_revision")
    ok(any("INCONGRUENCIA QUE INVALIDA" in a
           for a in ens.revisar_congruencia(_ruta2, ["fundado"], "amparo_revision")),
       "la alarma verdadera sigue: efectos en la Solución y un resolutivo que niega")
except Exception as ex:  # pragma: no cover
    ok(False, f"el documento no se pudo componer: {type(ex).__name__}: {ex}")

# ═══════════════════════════════════════════════════════════════════════════
# plan-6 (28-sep-2026). David: «esta ventana entorpeció demasiado el diálogo
# jurídico coherente […] el redactor debería de entender en automático cómo
# funciona la decisión judicial cuando el tema principal está resuelto y de
# ella dependen muchos argumentos secundarios». La Decisión 6 (una razón por
# argumento) se retira: dentro de cada problema decide el que ataca la
# proposición toral y los demás se resuelven por dependencia de él.
print("\n9 · plan-6, FUNDADO: DECIDE UNO; LOS DEL MISMO TEMA CON EL PRINCIPAL; LO DEMÁS, INNECESARIO")
ESCRITO6 = (
    "AGRAVIOS\n"
    "PRIMERO. La sentencia recurrida vulnera los principios de exhaustividad y congruencia porque no "
    "atendió lo que se planteó sobre la ejecución de la sentencia definitiva. La interlocutoria sólo "
    "determinó quién cuenta con legitimación para continuar la ejecución sin modificar las prestaciones "
    "de la sentencia definitiva. El artículo 49 del Código Procesal Civil local permite que el "
    "causahabiente sustituya a quien transmitió el derecho controvertido. Los artículos 2284 y 2294 del "
    "Código Civil local transmiten al adquirente los derechos del arrendador sobre el inmueble "
    "arrendado. La recurrente adquirió el inmueble mediante escritura pública de fecha anterior a la "
    "sustitución que se reclama. La sentencia carece de fundamentación porque no cita el precepto que "
    "sustenta la cosa juzgada. La sentencia carece de motivación porque no expone las razones "
    "particulares del caso concreto. Se vulnera el principio de igualdad jurídica porque no se aplican "
    "de manera uniforme las normas procesales. La falta de fundamentación y motivación es una violación "
    "formal distinta de la indebida fundamentación y motivación. Invoca la jurisprudencia sobre los "
    "argumentos que deben examinarse cuando se alega la ausencia de fundamentación. REVISIÓN ADHESIVA. "
    "La quejosa adherente sostiene que el juicio de origen se tramitó con violaciones al procedimiento "
    "que la dejaron sin defensa.\n\n"
    "SEGUNDO. Solicita que, como consecuencia de los argumentos expuestos en el primero, se tenga por "
    "reproducido su contenido para controvertir el considerando octavo de la sentencia recurrida.")
_o6 = [ESCRITO6.index(x) for x in ("PRIMERO.", "SEGUNDO.")] + [len(ESCRITO6)]
CONTEO6 = {"estado": "contado", "n": 2, "tramos": [[_o6[i], _o6[i + 1]] for i in range(2)]}


def fases6():
    return fases(fuentes=[ACTO, ESCRITO6], conteo=dict(CONTEO6))


_CITAS6 = {
    "A1.a": ("forma", "vulnera los principios de exhaustividad y congruencia porque no atendió", []),
    "A1.b": ("fondo", "sólo determinó quién cuenta con legitimación para continuar la ejecución", []),
    "A1.c": ("fondo", "el artículo 49 del Código Procesal Civil local permite que el causahabiente",
             ["art 49 cpc"]),
    "A1.d": ("fondo", "Los artículos 2284 y 2294 del Código Civil local transmiten al adquirente",
             ["art 2284 cc", "art 2294 cc"]),
    "A1.e": ("fondo", "La recurrente adquirió el inmueble mediante escritura pública de fecha anterior", []),
    "A1.f": ("forma", "La sentencia carece de fundamentación porque no cita el precepto", []),
    "A1.g": ("forma", "La sentencia carece de motivación porque no expone las razones particulares", []),
    "A1.h": ("forma", "Se vulnera el principio de igualdad jurídica porque no se aplican", []),
    "A1.i": ("forma", "La falta de fundamentación y motivación es una violación formal distinta", []),
    "A1.j": ("forma", "Invoca la jurisprudencia sobre los argumentos que deben examinarse", []),
    "AD1.a": ("procesal", "el juicio de origen se tramitó con violaciones al procedimiento que la dejaron", []),
}
SEGS6 = [pe.normalizar_segmento_piso({"id": k, "concepto": 1, "texto": v[1], "cita": v[1], "anclas": v[2]})
         for k, v in _CITAS6.items()] + [SEGS[-1]]           # A2.a, el del segundo agravio
_FORMA6 = ("A1.a", "A1.f", "A1.g", "A1.h", "A1.i", "A1.j")


def _s6(sid, **kw):
    base = {"id": sid, "problema_id": 1, "vicio": _CITAS6[sid][0], "ataca": "P1", "reitera": None,
            "dato": None, "etiqueta": "fundado", "razon": "fundado", "trat": "aplica", "diferencia": None,
            "pendiente": None}
    base.update(kw)
    return base


PLAN6 = {
    "proposiciones": [
        {"id": "P1", "dice": "La sustitución alteró la cosa juzgada", "caracter": "toral", "relacion": "suficiente",
         "fuente": "recurrida", "cita": "la sustitución de la parte ejecutante alteró la cosa juzgada",
         "vinculada_por_ejecutoria": False},
        {"id": "P2", "dice": "La adquirente no era causahabiente", "caracter": "accesoria",
         "relacion": "dependiente_de:P1", "fuente": "recurrida",
         "cita": "la adquirente del inmueble no era causahabiente del contrato de arrendamiento",
         "vinculada_por_ejecutoria": False},
        {"id": "P3", "dice": "La acción era personal", "caracter": "accesoria", "relacion": "dependiente_de:P2",
         "fuente": "recurrida", "cita": "que sustentó la acción personal de rescisión",
         "vinculada_por_ejecutoria": False},
        {"id": "P5", "dice": "Era innecesario estudiar los demás conceptos", "caracter": "accesoria",
         "relacion": "dependiente_de:P1", "fuente": "recurrida",
         "cita": "Consideró innecesario estudiar los restantes conceptos de violación",
         "vinculada_por_ejecutoria": False}],
    "segmentos": [
        # Lo que devolvía el planificador del 631 con la Decisión 6: la forma
        # «fundada» y pendiente de una razón del secretario, uno por uno.
        _s6("A1.a", trat="desarrolla", diferencia="norma", pendiente="razon"),
        _s6("A1.b"),
        _s6("A1.c", ataca="P2", dato={"texto": "el artículo 49", "fuente": "escrito",
                                      "cita": "el artículo 49 del Código Procesal Civil local permite que el "
                                              "causahabiente sustituya"}),
        _s6("A1.d", ataca="P2", etiqueta="esencialmente_fundado", razon="esencialmente_fundado"),
        _s6("A1.e", ataca="P3", diferencia="hecho"),
        _s6("A1.f", trat="desarrolla", diferencia="norma", pendiente="razon"),
        _s6("A1.g", trat="desarrolla", diferencia="norma", pendiente="razon"),
        _s6("A1.h", etiqueta="inoperante", razon="generico", trat="residual"),
        _s6("A1.i", trat="desarrolla", diferencia="norma", pendiente="razon"),
        _s6("A1.j", trat="desarrolla", diferencia="precedente", pendiente="razon"),
        # Un suplido y la violación procesal del adhesivo: nunca dependen.
        {"id": "S1.a", "problema_id": 1, "vicio": "fondo", "ataca": "P1", "reitera": None, "dato": None,
         "etiqueta": "fundado", "razon": "fundado", "trat": "desarrolla", "diferencia": "hecho", "pendiente": None},
        _s6("AD1.a", etiqueta="infundado", razon="fondo_desestimado", trat="desarrolla", diferencia="procesal"),
        {"id": "A2.a", "problema_id": 2, "vicio": "fondo", "ataca": "P5", "reitera": None, "dato": None,
         "etiqueta": "innecesario", "razon": "cae_con_principal", "trat": "no_se_estudia", "diferencia": None,
         "pendiente": None}],
    "premisas": [{"id": "M1", "responde_a": ["P1", "P2", "P3"], "fuentes": {"tesis": [], "normas": ["art. 49"]},
                  "anclas": ["causahabiente", "legitimación"], "rastro": "razon",
                  "rastro_cita": "el artículo 49 del Código Procesal Civil local permite al causahabiente "
                                 "sustituir a quien transmitió el derecho controvertido"}],
    # Partido como en el 631: la misma premisa en tres unidades.
    "unidades": [{"id": "U1", "problemas": [1], "segmentos": ["A1.b"], "premisa": "M1",
                  "objecion": {"de": "el juzgado", "anclas": ["cosa juzgada"]}},
                 {"id": "U2", "problemas": [1], "segmentos": ["A1.c", "A1.d"], "premisa": "M1", "objecion": None},
                 {"id": "U3", "problemas": [1], "segmentos": ["A1.e"], "premisa": "M1", "objecion": None},
                 {"id": "U4", "problemas": [1], "segmentos": ["A1.a", "A1.f", "A1.g", "A1.i", "A1.j"],
                  "premisa": None, "objecion": None},
                 {"id": "U5", "problemas": [1], "segmentos": ["S1.a"], "premisa": None, "objecion": None},
                 {"id": "U6", "problemas": [1], "segmentos": ["AD1.a"], "premisa": None, "objecion": None}],
    "propuestas": [],
    "avisos_al_secretario": ["A1.a, A1.f, A1.g, A1.i y A1.j tienen pendiente la razón específica del secretario"],
    "orden": {"criterio": "prelacion", "por_que": "el principal primero"},
}
_r9, _av9 = pe.reparar(_norm(PLAN6), crit(), SEGS6, fases6(), material(), "", {})
_s9 = {s["id"]: s for s in _r9["segmentos"]}
_j9 = next(j for j in _r9["jerarquia"] if j["problema"] == 1)
ok(_j9["decide"] == "A1.b" and _s9["A1.b"]["dependencia"] == "decide" and _j9["ataca"] == "P1",
   "decide A1.b, el de fondo que ataca la proposición toral (P1), no el agravio de forma que va antes")
ok(all(_s9[x]["dependencia"] == "con_el_principal" and _s9[x]["con"] == "A1.b"
       and _s9[x]["depende_de"] == "P1" and _s9[x]["etiqueta"] == "fundado" for x in ("A1.c", "A1.d", "A1.e")),
   "con el principal: el art. 49 (P2), los arts. 2284 y 2294 (P2) y la escritura (P3 → P2 → P1), con la "
   "calificación del grupo (el «esencialmente fundado» ya no se acota por argumento)")
ok(all(_s9[x]["etiqueta"] == "innecesario" and _s9[x]["razon"] == "innecesario_por_suficiencia"
       and _s9[x]["trat"] == "no_se_estudia" and _s9[x]["con"] == "A1.b" for x in _FORMA6)
   and not any(x in (u.get("segmentos") or []) for u in _r9["unidades"] for x in _FORMA6),
   "los seis de forma, innecesarios por suficiencia (decide A1.b) y fuera de las unidades")
ok(_s9["S1.a"]["dependencia"] == "autonomo" and _s9["AD1.a"]["dependencia"] == "autonomo"
   and _s9["S1.a"]["etiqueta"] == "fundado" and _s9["AD1.a"]["etiqueta"] == "infundado",
   "el suplido y la violación procesal del adhesivo no se absorben ni se declaran innecesarios")
ok(not any(s.get("pendiente") == "razon" for s in _r9["segmentos"]), "ningún pendiente de razón")
ok(_s9["A2.a"].get("dependencia") is None and _s9["A2.a"]["etiqueta"] == "innecesario",
   "el problema que no se estudia, como hoy")
_f9 = pe.validar(_r9, crit(), SEGS6, fases6(), material(), "", suplencia={})
ok(_f9 == [], f"V0 limpio: {_f9[:3]}")
ok(sum(1 for u in _r9["unidades"] if u.get("premisa") == "M1") == 3
   and not any("juntaba" in a and "U2" in a for a in _r9["avisos_al_secretario"]),
   "el tema no se parte por proposición: P1, P2 y P3 son una cadena (la premisa M1 no se reparte más)")
ok(any(a.startswith("Resueltos por consecuencia del principal (problema 1, decide A1.b")
       and "innecesarios por suficiencia A1.a" in a for a in _r9["avisos_al_secretario"])
   and not any("pendiente" in a and "razón" in a for a in _r9["avisos_al_secretario"]),
   "el panel lo dice en una línea, como información; lo que el planificador escribió de «pendientes» sale")
_rev = types.SimpleNamespace(encargo=types.SimpleNamespace(tipo_asunto="amparo_revision",
                                                           resolvio_declarado="Concedió el amparo."),
                             fases=fases6())
ok(pe.concede_de(_rev, crit()) is False, "la revisión que revoca una concesión NIEGA: no hay efectos")
_g9 = pe.vista(_r9, "estandar", concede=pe.concede_de(_rev, crit()))
print("      " + _g9.replace("\n", "\n      ")[:2600])
ok("JERARQUÍA DEL PROBLEMA 1 (fundado): DECIDE A1.b (ataca P1) · CON EL PRINCIPAL: A1.c, A1.d, A1.e · "
   "INNECESARIOS POR SUFICIENCIA: A1.a, A1.f, A1.g, A1.h, A1.i, A1.j · AUTÓNOMOS: S1.a, AD1.a" in _g9,
   "el guion trae la JERARQUÍA del problema, como datos")
ok("PENDIENTE DE RAZÓN" not in _g9 and "UNIDADES QUE PROSPERAN" not in _g9 and "ADVERTENCIAS" not in _g9,
   "sin PENDIENTE DE RAZÓN, sin nada a ADVERTENCIAS y sin «de ellas salen los efectos» en una revisión que niega")
ok("APLICA A1.c" not in _g9 and "con el principal: A1.c (P2) · dato (escrito): «el artículo 49" in _g9
   and _g9.count("EXPONE M1") == 1,
   "los que van con el principal no tienen renglón propio: su dato, dentro del renglón del que decide")
ok("INNECESARIOS POR SUFICIENCIA A1.a, A1.f, A1.g, A1.h, A1.i, A1.j · etiqueta: innecesario" in _g9
   and "NO SE ESTUDIA A1.f" not in _g9,
   "los innecesarios, un renglón por grupo al cerrar el apartado")
ok("UNIDADES QUE PROSPERAN" in pe.vista(_r9, "estandar", concede=True),
   "(control: si el asunto concediera, sí se dicen las unidades de las que salen los efectos)")
_b9 = pe.bloque(_g9)
ok("JERARQUÍA DEL PROBLEMA: cómo se sigue la suerte" in _b9 and "PENDIENTE DE RAZÓN" not in _b9,
   "el bloque describe la jerarquía; ya no describe el pendiente de razón")
_p9 = f6.prompt_estudio("ACTO", "CONC", crit(), f6.Material(tipo_asunto="amparo_revision", variante="v4",
                                                            formato="estandar"), guion=_g9)
ok("JERARQUÍA del guion lo resuelva por" in _p9 and "PENDIENTE DE RAZÓN" not in _p9,
   "el estudio (v4) acepta la jerarquía como criterio y no pone nada primero en ADVERTENCIAS")
ok(pe.resolver_por_dependencia(_r9) == pe.resolver_por_dependencia(pe.resolver_por_dependencia(_r9)),
   "resolver por dependencia es idempotente sobre el plan reparado")
_props9 = {p["id"]: p for p in _r9["proposiciones"]}
ok(pe.depende_del_principal({"id": "A1.z", "vicio": "procesal", "ataca": "P3"}, _props9, "amparo_directo") is None
   and pe.depende_del_principal({"id": "A1.z", "vicio": "procesal", "ataca": "P3"}, _props9,
                                "amparo_revision") == "P1"
   and pe.depende_del_principal({"id": "A1.z", "vicio": "omision", "ataca": "P1"}, _props9) is None
   and pe.raiz_toral({"P1": {"relacion": "dependiente_de:P2", "caracter": "toral"},
                      "P2": {"relacion": "dependiente_de:P1", "caracter": "toral"}}, "P1") is None,
   "la procesal del amparo directo (arts. 74-V, 174 y 189) y la omisión nunca dependen; un ciclo no tiene raíz")

print("\n10 · plan-6, INFUNDADO: LOS 13 DE LA CAPTURA DE DAVID, SIN CAJA Y SIN PEDIR NADA")
# El plan-4 «infundado» guardado en producción (sesión 462, clave 83a4…),
# reducido a su forma: 24 argumentos del primer agravio, P2 y P3 dependientes
# de la toral P1, trece «pendiente: razon» —la caja que vio David— y las
# premisas ya retiradas. Más una omisión, que no depende de nadie. Desde la
# revisión adversarial (28-sep-2026), por consecuencia del principal sólo va lo
# de fondo; la forma, la procesal y el inoperante sin caída declarada se
# desarrollan con su diferencia: ninguno queda pendiente ni sin respuesta.
_PEND = ("A1.d", "A1.g", "A1.j", "A1.m", "A1.n", "A1.o", "A1.p", "A1.q", "A1.r", "A1.s", "A1.t", "A1.v", "A1.w")
_ATACA = {"A1.c": "P3", "A1.d": "P3", "A1.f": "P3", "A1.g": "P3", "A1.p": "P3", "A1.x": "P2"}
_segs10 = []
for _c in "abcdefghijklmnopqrstuvwx":
    _sid = f"A1.{_c}"
    _p = _sid in _PEND
    _segs10.append({"id": _sid, "problema_id": 1, "concepto": 1,
                    "vicio": "forma" if _c in "lmpqrst" else ("procesal" if _c in "vw" else "fondo"),
                    "ataca": _ATACA.get(_sid, "P1"), "reitera": None, "dato": None,
                    "etiqueta": "inoperante" if _sid in ("A1.n", "A1.o") else "infundado",
                    "razon": "" if _p else "fondo_desestimado",
                    "trat": "residual" if _sid == "A1.k" else "desarrolla",
                    "diferencia": ("precedente" if _c in "jorstvw" else "norma") if _p else None,
                    "pendiente": "razon" if _p else None, "sin_premisa": not _p and _sid != "A1.k"})
_segs10.append({"id": "A1.y", "problema_id": 1, "concepto": 1, "vicio": "omision", "ataca": "P1", "reitera": None,
                "dato": None, "etiqueta": "infundado", "razon": "", "trat": "desarrolla", "diferencia": "hecho",
                "pendiente": "razon"})
# Un inoperante que el planificador SÍ declaró caído por derivar de lo
# desestimado, nombrando la proposición (lo único que cae así desde la
# revisión adversarial del 28-sep-2026).
_segs10.append({"id": "A1.z", "problema_id": 1, "concepto": 1, "vicio": "fondo", "ataca": "P2", "reitera": None,
                "dato": None, "etiqueta": "inoperante", "razon": "deriva_de_desestimado", "razon_p": "P2",
                "trat": "residual", "diferencia": None, "pendiente": None})
_segs10.append({"id": "A2.a", "problema_id": 2, "concepto": 2, "vicio": "omision", "ataca": "P4", "reitera": None,
                "dato": None, "etiqueta": "infundado", "razon": "omision_inexistente", "trat": "remite",
                "diferencia": None, "pendiente": None})
PLAN4_GUARDADO = {
    "version": "plan-4", "tipo_asunto": "amparo_revision",
    "problemas": [{"id": 1, "pregunta": PREG1, "sentido": "infundado", "clase": "fondo", "jerarquia": "principal",
                   "grupo": ""},
                  {"id": 2, "pregunta": PREG2, "sentido": "infundado", "clase": "fondo", "jerarquia": "accesorio",
                   "grupo": ""}],
    "proposiciones": [{"id": "P1", "dice": "x", "caracter": "toral", "relacion": "suficiente"},
                      {"id": "P2", "dice": "x", "caracter": "accesoria", "relacion": "dependiente_de:P1"},
                      {"id": "P3", "dice": "x", "caracter": "accesoria", "relacion": "dependiente_de:P1"},
                      {"id": "P4", "dice": "x", "caracter": "toral", "relacion": "suficiente"}],
    "segmentos": _segs10,
    "premisas": [{"id": "M3", "responde_a": ["P4"], "fuentes": {"tesis": [], "normas": []}, "anclas": [],
                  "rastro": "razon", "rastro_cita": "x"}],
    "unidades": [{"id": "U1", "problemas": [1], "segmentos": ["A1.a", "A1.b", "A1.e", "A1.h", "A1.i", "A1.u"],
                  "premisa": None}]
                + [{"id": f"U{i + 2}", "problemas": [1], "segmentos": [x["id"]], "premisa": None}
                   for i, x in enumerate(s_ for s_ in _segs10 if s_["id"] not in
                                         ("A1.a", "A1.b", "A1.e", "A1.h", "A1.i", "A1.u", "A1.k", "A2.a"))]
                + [{"id": "U40", "problemas": [2], "segmentos": ["A2.a"], "premisa": "M3"}],
    "propuestas": [], "avisos_al_secretario": [], "orden": {"criterio": "prelacion", "por_que": ""},
}
_r10 = pe.resolver_por_dependencia(PLAN4_GUARDADO)
_s10 = {s["id"]: s for s in _r10["segmentos"]}
_con10 = [x for x in _PEND if _s10[x]["dependencia"] == "con_el_principal"]
_aut10 = [x for x in _PEND if _s10[x]["dependencia"] == "autonomo"]
# REVISIÓN ADVERSARIAL (28-sep-2026): «por las mismas razones» sólo lo de fondo
# —la forma, la igualdad o la vía no se desestiman con la premisa de la cosa
# juzgada—, y nada cae por derivar sin que el planificador lo haya declarado.
ok(_con10 == ["A1.d", "A1.g", "A1.j"]
   and _aut10 == ["A1.m", "A1.n", "A1.o", "A1.p", "A1.q", "A1.r", "A1.s", "A1.t", "A1.v", "A1.w"],
   f"los 13: los de fondo, infundados por las mismas razones ({_con10}); la forma, la procesal y los "
   f"inoperantes sin caída declarada, autónomos ({_aut10})")
ok(all(_s10[x]["etiqueta"] == "infundado" and _s10[x]["razon"] == "fondo_desestimado"
       and _s10[x]["con"] == "A1.a" and _s10[x]["depende_de"] == "P1" for x in _con10)
   and _s10["A1.l"]["dependencia"] == "autonomo" and _s10["A1.b"]["dependencia"] == "con_el_principal",
   "con el principal: infundados de fondo, fondo_desestimado, siguen a A1.a (P1; los de P3, por su cadena); "
   "el de forma que no estaba pendiente (A1.l) tampoco va con él")
ok(all(_s10[x]["trat"] == "desarrolla" and _s10[x].get("diferencia") and _s10[x].get("pendiente") is None
       and _s10[x]["etiqueta"] == ("inoperante" if x in ("A1.n", "A1.o") else "infundado") for x in _aut10),
   "los autónomos se desarrollan con su diferencia, sin pendiente y con la calificación que les dio el planificador")
ok(_s10["A1.z"]["dependencia"] == "deriva" and _s10["A1.z"]["razon"] == "deriva_de_desestimado"
   and _s10["A1.z"]["razon_p"] == "P1" and _s10["A1.z"]["trat"] == "residual",
   "el inoperante que el planificador declaró derivado de lo desestimado cae con P1")
ok(not any(s["etiqueta"] == "innecesario" or s.get("razon") == "innecesario_por_suficiencia"
           for s in _r10["segmentos"]), "ninguno «innecesario»: en un problema que no prospera todo se contesta")
ok(_s10["A1.y"]["dependencia"] == "autonomo" and _s10["A1.y"]["trat"] == "desarrolla"
   and _s10["A1.y"].get("pendiente") is None and any("A1.y" in u["segmentos"] for u in _r10["unidades"]),
   "la omisión sigue autónoma: se desarrolla con su diferencia, sin pendiente y sin pedir nada")
ok(not any(s.get("pendiente") == "razon" for s in _r10["segmentos"])
   and pe.resolver_por_dependencia(_r10) == _r10,
   "el plan-4 guardado queda sin pendientes, y resolver otra vez no cambia nada (idempotente)")
_g10 = pe.vista(_r10, "estandar")
ok("PENDIENTE DE RAZÓN" not in _g10 and "CAEN POR DERIVAR A1.z" in _g10
   and "JERARQUÍA DEL PROBLEMA 1 (infundado): DECIDE A1.a" in _g10 and "DESARROLLA A1.y" in _g10
   and "DESARROLLA A1.q" in _g10 and "CON EL PRINCIPAL, POR LAS MISMAS RAZONES: A1.b" in _g10,
   "el guion: la jerarquía en su dirección, los que caen en un renglón, y la omisión y la forma con el suyo; "
   "sin PENDIENTE DE RAZÓN")

print("\n11 · REVISIÓN ADVERSARIAL DEL PLAN-6 (28-sep-2026): LA DEPENDENCIA NO OMITE NI MEZCLA")


def _sg(sid, **kw):
    b = {"id": sid, "problema_id": 1, "concepto": 1, "vicio": "fondo", "ataca": "P1", "reitera": None,
         "dato": None, "etiqueta": "infundado", "razon": "fondo_desestimado", "trat": "aplica",
         "diferencia": None, "pendiente": None}
    b.update(kw)
    return b


def _pl(segs, props, sentido="infundado", tipo="amparo_directo", premisas=None, unidades=None):
    return {"tipo_asunto": tipo, "problemas": [{"id": 1, "pregunta": "q", "sentido": sentido, "clase": "fondo"}],
            "proposiciones": props, "segmentos": segs, "premisas": premisas or [],
            "unidades": unidades or [{"id": "U1", "problemas": [1], "segmentos": [s["id"] for s in segs],
                                      "premisa": None}],
            "propuestas": [], "avisos_al_secretario": [], "orden": {"criterio": "prelacion", "por_que": ""}}


def _P(i, car, rel):
    return {"id": i, "dice": "x " + i, "caracter": car, "relacion": rel}


def _dep(plan):
    return {s["id"]: s for s in pe.resolver_por_dependencia(plan)["segmentos"]}


_T = [_P("P1", "toral", "suficiente"), _P("P2", "accesoria", "dependiente_de:P1")]
# (a) FONDO CON FONDO: en un problema que no prospera, la forma y una procesal
# fuera del amparo directo no se desestiman con la premisa de fondo del
# principal (el 83a4: falta de motivación, vía incorrecta).
_x = _dep(_pl([_sg("A1.a"), _sg("A1.b", vicio="forma"), _sg("A1.c", vicio="procesal")], _T[:1],
              tipo="amparo_revision"))
ok(_x["A1.a"]["dependencia"] == "decide" and _x["A1.b"]["dependencia"] == "autonomo"
   and _x["A1.c"]["dependencia"] == "autonomo",
   "no prospera: la forma y la procesal no van «con el principal»; se contestan con su razón")
# (b) PARCIALMENTE FUNDADO: no cabe la suficiencia; lo infundado se desestima.
_x = _dep(_pl([_sg("A1.a", etiqueta="fundado", razon="fundado"), _sg("A1.b", ataca="P2", etiqueta="fundado",
                                                                      razon="fundado"),
               _sg("A1.c", ataca="P3", trat="desarrolla", diferencia="hecho"), _sg("A1.d")],
              _T + [_P("P3", "toral", "suficiente")], sentido="parcialmente_fundado"))
ok(not any(s["etiqueta"] == "innecesario" for s in _x.values())
   and _x["A1.b"]["dependencia"] == "con_el_principal" and _x["A1.b"]["etiqueta"] == "fundado"
   and _x["A1.c"]["dependencia"] == "autonomo" and _x["A1.d"]["dependencia"] == "autonomo"
   and _x["A1.d"]["etiqueta"] == "infundado",
   "parcialmente fundado: nada innecesario, la parte infundada se contesta y el grupo lleva «fundado», "
   "no «parcialmente_fundado»")
# (c) EN EL AMPARO DIRECTO DECIDE EL FONDO (art. 189): la forma contra la toral
# suficiente no le gana a un fondo fundado contra otra toral.
_x = _dep(_pl([_sg("A1.a", vicio="forma", etiqueta="fundado", razon="fundado"),
               _sg("A1.b", ataca="P2", etiqueta="fundado", razon="fundado", trat="desarrolla", diferencia="norma")],
              [_P("P1", "toral", "suficiente"), _P("P2", "toral", "necesaria")], sentido="fundado"))
ok(_x["A1.b"]["dependencia"] == "decide" and _x["A1.a"]["dependencia"] == "innecesario",
   "amparo directo: decide el fondo fundado y la forma es la innecesaria, no al revés")
# …y en un recurso, donde la forma puede decidir, el fondo que prospera nunca
# se absorbe en ella ni queda innecesario.
_x = _dep(_pl([_sg("A1.a", vicio="forma", etiqueta="fundado", razon="fundado"),
               _sg("A1.b", ataca="P2", etiqueta="fundado", razon="fundado")], _T, sentido="fundado",
              tipo="amparo_revision"))
ok(_x["A1.a"]["dependencia"] == "decide" and _x["A1.b"]["dependencia"] == "autonomo"
   and _x["A1.b"]["etiqueta"] == "fundado",
   "recurso: con un principal de forma, el fondo fundado sigue autónomo (ni innecesario ni absorbido)")
# (d) NADA CAE POR DERIVAR SIN QUE EL PLANIFICADOR LO DECLARE.
_x = _dep(_pl([_sg("A1.a"), _sg("A1.b", ataca="P2", etiqueta="inoperante", razon="", pendiente="razon",
                                trat="desarrolla", diferencia="precedente")], _T, tipo="amparo_revision"))
ok(_x["A1.b"]["dependencia"] == "autonomo" and _x["A1.b"]["razon"] != "deriva_de_desestimado"
   and _x["A1.b"]["trat"] == "desarrolla",
   "el inoperante pendiente sin caída declarada sigue autónomo: el código no elige la razón de la caída")
# (e) LA PREMISA COMÚN NO UNE LA SUERTE DE UNA TORAL INDEPENDIENTE (las costas
# junto a la cosa juzgada) en un problema que no prospera.
_x = _dep(_pl([_sg("A1.a"), _sg("A1.b", ataca="P5", etiqueta="inoperante", razon="", pendiente="razon",
                                trat="desarrolla", diferencia="norma"), _sg("A1.c", ataca="P5")],
              [_P("P1", "toral", "suficiente"), _P("P5", "toral", "suficiente")],
              premisas=[{"id": "M1", "responde_a": ["P1", "P5"], "fuentes": {"tesis": [], "normas": []},
                         "anclas": [], "rastro": "razon", "rastro_cita": "x"}]))
ok(_x["A1.b"]["dependencia"] == "autonomo" and _x["A1.c"]["dependencia"] == "autonomo",
   "no prospera: contra P5, independiente de P1, ni «por las mismas razones» ni «cae por derivar» de P1")
# (f) SIN RAÍZ TORAL NO SE FABRICA UNA CAÍDA QUE V0 RECHACE.
_x = _dep(_pl([_sg("A1.a", ataca="P1"), _sg("A1.b", ataca="P1", etiqueta="inoperante",
                                            razon="deriva_de_desestimado", razon_p="P1", trat="residual")],
              [_P("P1", "accesoria", "necesaria")]))
ok(_x["A1.b"]["dependencia"] == "autonomo",
   "si lo que ataca el que decide no llega a una toral, nada se le une ni cae por derivar de él")
# (g) «FUNDADO PERO INSUFICIENTE» CONTRA EL TEMA QUE CAE (el e34c): va con el
# principal; el infundado que pide otra consecuencia sigue autónomo (el 642).
_x = _dep(_pl([_sg("A1.a", etiqueta="fundado", razon="fundado"),
               _sg("A1.b", etiqueta="fundado_insuficiente", razon="fundado_insuficiente", diferencia="consecuencia"),
               _sg("A1.c", ataca="P2", diferencia="consecuencia")], _T, sentido="fundado"))
ok(_x["A1.b"]["dependencia"] == "con_el_principal" and _x["A1.b"]["etiqueta"] == "fundado"
   and _x["A1.c"]["dependencia"] == "autonomo" and _x["A1.c"]["etiqueta"] == "infundado",
   "en un problema fundado, el «fundado pero insuficiente» contra P1 va con el principal; el infundado que "
   "pide otra consecuencia se contesta")
_pg = pe.resolver_por_dependencia(_pl([_sg("A1.a", etiqueta="fundado", razon="fundado"),
                                       _sg("A1.b", etiqueta="fundado_insuficiente", razon="fundado_insuficiente",
                                           diferencia="consecuencia")], _T, sentido="fundado"))
ok(pe.resolver_por_dependencia(_pg) == _pg, "…y resolver otra vez no lo devuelve a autónomo (idempotente)")
# (h) UNA CALIFICACIÓN POR GRUPO: la del que decide, también en un problema
# «esencialmente fundado».
_x = _dep(_pl([_sg("A1.a", etiqueta="fundado", razon="fundado"),
               _sg("A1.b", etiqueta="esencialmente_fundado", razon="esencialmente_fundado"),
               _sg("A1.c", ataca="P2", etiqueta="fundado", razon="fundado")], _T, sentido="esencialmente_fundado"))
ok({(_x[k]["etiqueta"], _x[k]["razon"]) for k in ("A1.a", "A1.b", "A1.c")} == {("fundado", "fundado")},
   "el grupo lleva la calificación del que decide: ni dos calificaciones ni la unidad partida por razón")
# (i) EL GUION: lo que conserva «desarrolla» con su diferencia tiene renglón; la
# objeción del que decide, una; con concesión, los datos de los innecesarios.
_pv = _pl([_sg("A1.a", trat="desarrolla", diferencia="hecho"),
           _sg("A1.b", ataca="P2", trat="desarrolla", diferencia="norma",
               dato={"texto": "art", "cita": "el artículo 2294 impone", "fuente": "escrito"}),
           _sg("A1.c", ataca="P2", dato={"texto": "d", "cita": "la escritura pública", "fuente": "escrito"})],
          _T, tipo="amparo_revision",
          unidades=[{"id": "U1", "problemas": [1], "segmentos": ["A1.a"], "premisa": None,
                     "objecion": {"de": "recurrente", "anclas": ["uno"]}},
                    {"id": "U2", "problemas": [1], "segmentos": ["A1.b"], "premisa": None,
                     "objecion": {"de": "recurrente", "anclas": ["dos"]}},
                    {"id": "U3", "problemas": [1], "segmentos": ["A1.c"], "premisa": None,
                     "objecion": {"de": "recurrente", "anclas": ["tres"]}}])
_gv = pe.vista(pe.resolver_por_dependencia(_pv), "estandar")
_ren = next(x for x in _gv.splitlines() if "decide el problema 1" in x)
ok("DESARROLLA A1.b" in _gv and "diferencia: norma" in _gv and "sigue a A1.a" in _gv
   and "con el principal, por las mismas razones: A1.c (P2)" in _ren,
   "el que trae su diferencia lleva renglón y dice a quién sigue; el que no, entra al del que decide")
ok(_ren.count("objeción aquí") == 1,
   f"el renglón del que decide lleva UNA objeción, no la de cada unidad absorbida ({_ren.count('objeción aquí')})")
_pc = _pl([_sg("A1.a", etiqueta="fundado", razon="fundado"),
           _sg("A1.b", vicio="forma", dato={"texto": "p", "cita": "la pericial de fecha 3 de marzo", "fuente": "escrito"})],
          _T, sentido="fundado")
_gc = pe.vista(pe.resolver_por_dependencia(_pc), "estandar", concede=True)
ok(f"{pe._RX_A_EFECTOS}: A1.b (P1) · dato (escrito): «la pericial de fecha 3 de marzo»" in _gc
   and pe._RX_A_EFECTOS not in pe.vista(pe.resolver_por_dependencia(_pc), "estandar", concede=False)
   and "EFECTOS, con su dato" in pe.bloque(_gc),
   "con concesión, el renglón de los innecesarios trae su dato para los EFECTOS (0ad0379), y el bloque lo describe")
# (j) LA GUARDA PROCESAL DEL AMPARO DIRECTO (arts. 74-V, 174 y 189): dentro de
# un problema procesal fundado, ninguna procesal queda «innecesaria».
_fp = fases()
_fp.problemas[0]["clase"] = "procesal"
_mp = material()
_mp.tipo_asunto = "amparo_directo"
_PLP = copy.deepcopy(PLAN_631)
_PLP["segmentos"] = [
    _s("A1.a", vicio="procesal", etiqueta="fundado", razon="fundado", trat="aplica", diferencia=None),
    _s("A1.b", vicio="procesal", etiqueta="innecesario", razon="innecesario_por_suficiencia",
       trat="no_se_estudia", diferencia=None),
    _s("A1.c", vicio="procesal", etiqueta="sin_materia", razon="sin_materia", trat="no_se_estudia", diferencia=None),
    _s("A1.d", vicio="procesal", etiqueta="fundado", razon="fundado", trat="aplica", diferencia=None),
    _s("A1.e", vicio="procesal", etiqueta="fundado", razon="fundado", trat="aplica", diferencia=None),
    _s("A2.a", problema_id=2, etiqueta="innecesario", razon="cae_con_principal", trat="no_se_estudia",
       diferencia=None, ataca=None)]
_PLP["unidades"] = [{"id": "U1", "problemas": [1], "segmentos": ["A1.a", "A1.d", "A1.e"], "premisa": None,
                     "objecion": None}]
_np = pe.normalizar(copy.deepcopy(_PLP))
_np["tipo_asunto"] = "amparo_directo"
_rp, _ = pe.reparar(_np, crit(), SEGS, _fp, _mp, "", {})
_sp = {s["id"]: s for s in _rp["segmentos"]}
ok(all(_sp[x]["etiqueta"] == "fundado" and _sp[x]["trat"] not in ("no_se_estudia", "")
       for x in ("A1.b", "A1.c")),
   "reparar: la procesal que el planificador dejó «innecesaria» o «sin materia» vuelve al estudio")
ok(not [f for f in pe.validar(_rp, crit(), SEGS, _fp, _mp, "", suplencia={}) if "procesal" in f],
   "…y V0 no tiene nada que decir de las procesales")
_mal = copy.deepcopy(_rp)
next(s for s in _mal["segmentos"] if s["id"] == "A1.b").update(
    etiqueta="innecesario", razon="innecesario_por_suficiencia", trat="no_se_estudia")
ok(any("A1.b" in f and "violación procesal" in f
       for f in pe.validar(_mal, crit(), SEGS, _fp, _mp, "", suplencia={})),
   "V0 rechaza una procesal «innecesaria por suficiencia» aunque ningún fondo prospere")
# (k) EL PROMPT DEL ESTUDIO Y EL BLOQUE CON JERARQUÍA.
_p9n = " ".join(_p9.split())
ok("Salvo lo que la JERARQUÍA del guion resuelve por consecuencia de su principal" in _p9n
   and "se contesta dentro de la respuesta de su grupo, con su dato y su marca" in _p9n
   and "reciba su respuesta —la de su grupo si la JERARQUÍA del guion lo resuelve" in _p9n
   and "recibe su respuesta —la de su grupo, si la JERARQUÍA lo resuelve" in _p9n,
   "v4 con JERARQUÍA: las cuatro reglas de «una respuesta por argumento» dicen su salvedad")
ok("UNIDADES QUE PROSPERAN" not in " ".join(_b9.split()) and "UNIDADES QUE PROSPERAN" not in _p9n
   and "UNIDADES QUE PROSPERAN" in " ".join(pe.bloque(pe.vista(_r9, "estandar", concede=True)).split()),
   "sin concesión, ni el bloque habla de las unidades de las que salen los efectos (con los espacios "
   "normalizados)")
ok("manda la JERARQUÍA" in _g9, "el porqué del orden del planificador cede ante la jerarquía")

# ═══════════════════════════════════════════════════════════════════════════
# SPEC B (28-sep-2026). El fondo del 631 negó sin estudiar los conceptos que el
# juzgado no estudió (art. 93, fr. VI), revocó «la sentencia recurrida» entera
# (con el sobreseimiento del Juez Quinto, que nadie impugnó), avisó «se
# concede» en un proyecto que niega, arrastró la rama del adelanto y cambió los
# papeles (la recurrente como «quejosa y recurrente»).
print("\n12 · SPEC B: REVOCAR LA CONCESIÓN REASUME JURISDICCIÓN; RAMA Y PAPELES DEL 631")
import redactor_adelanto as ra
import documento_generado as dg
import ensamblar_adelanto as ens
import fase0_oportunidad as f0
from docx import Document

RECURRIDA_B = (
    "JUZGADO SÉPTIMO DE DISTRITO EN EL ESTADO DE EJEMPLO\nAMPARO INDIRECTO 950/2024\n"
    "CONSIDERANDO SÉPTIMO. Es fundado el segundo concepto de violación; resulta innecesario el "
    "estudio de los restantes conceptos de violación.\n"
    "Por lo expuesto, se R E S U E L V E: PRIMERO. Se sobresee en el juicio respecto del acto del "
    "Juzgado Quinto. SEGUNDO. " + RESOL + " Notifíquese.")
DEMANDA_B = ("DEMANDA DE AMPARO INDIRECTO\nCONCEPTOS DE VIOLACIÓN\n"
             "PRIMERO. La resolución reclamada altera la cosa juzgada. " + "Argumento. " * 70 + "\n"
             "SEGUNDO. La interlocutoria no valoró las pruebas del incidente. " + "Argumento. " * 40
             + "\nSUSPENSIÓN\nSe pide la suspensión.")
RECURRENTE_B = "Inmobiliaria Ejemplo, S.A. de C.V."


def fases_b(**kw):
    f = fases(fuentes=[RECURRIDA_B, ESCRITO6], conteo=dict(CONTEO6))
    f.resolutivo_recurrida = fr.resolutivo_recurrida("se R E S U E L V E: ÚNICO. " + RESOL + " Notifíquese.")
    f.resolvio_a_quo = "niega"               # el recuento que guardó la sesión del 631
    for k, v in kw.items():
        setattr(f, k, v)
    return f


def encargo_b(**kw):
    e = ra.Encargo(numero="631/2025", encabezado="AMPARO EN REVISIÓN 631/2025",
                   quejoso=RECURRENTE_B,           # lo que guardó el formulario (desde la admisión)
                   magistrado="M", secretario="S", notificacion=_dt.date(2025, 4, 25),
                   presentacion=_dt.date(2025, 5, 15), tipo_asunto="amparo_revision",
                   responsable="Magistrada de la Sala Civil Uno", tribunal="Tercer Tribunal Colegiado",
                   ciudad="Querétaro")
    e.es_recurso = True
    for k, v in kw.items():
        setattr(e, k, v)
    return e


def r_b(**kw):
    return types.SimpleNamespace(encargo=encargo_b(), fases=fases_b(**kw), partes=None, avisos=[])


# (a) EL FLUJO LOS PIDE: la misma función para la pantalla, la tarjeta y el estudio.
_info_b = ra.info_de_rama(r_b())
_co_b = fr.conceptos_omitidos(_info_b, "fundado", fases_b())
ok(_info_b["que_hizo"] == "concede" and _info_b["quien_recurre"] == "tercero"
   and _co_b and _co_b["hacen_falta"] and not _co_b["tenemos"],
   "rama con conceptos omitidos → los pide (concedió según su resolutivo; recurre la tercera)")
ok(fr.conceptos_omitidos(_info_b, "infundado", fases_b()) is None, "si el recurso no prospera, no")
_mb = material()
ra._formato_al_material(r_b(), _mb, None, crit())
ok(isinstance(_mb.reasuncion, dict) and _mb.reasuncion.get("reasuncion") == "concesion"
   and _mb.reasuncion.get("tenemos") is False and _mb.reasuncion.get("sobresee_ademas"),
   "al resolver, el material sabe que se reasume jurisdicción, que faltan y que la recurrida sobreseyó")
_pb = f6.prompt_estudio(ACTO, "c", crit(), _mb, es_recurso=True, rama="revoca_fondo_niega")
ok("REVOCAR NO ES NEGAR" in _pb and "FALTAN ESOS CONCEPTOS" in _pb and "93, fracción VI" in _pb,
   "el estudio recibe la técnica de la fracción VI y la orden de no concluir sin los conceptos")
_mb2 = material()
ra._formato_al_material(r_b(autos=DEMANDA_B), _mb2, None, crit())
_pb2 = f6.prompt_estudio(ACTO, "c", crit(), _mb2, es_recurso=True, rama="revoca_fondo_niega")
ok(_mb2.reasuncion["donde"] == "constancias" and "LOS CONCEPTOS DE VIOLACIÓN NO ESTUDIADOS, tal como constan" in _pb2
   and "La interlocutoria no valoró las pruebas del incidente" in _pb2,
   "si la demanda está entre las constancias, se toman de ahí antes de pedirlos")


def componer_b(estudio, reasuncion, est=None, calif=("fundado",), partes=None):
    _e = encargo_b()
    datos = ra._datos_estructura(_e, "", acto=RECURRIDA_B, partes=partes, fases=fases_b())
    datos.update({"resolvio_a_quo": "niega", "resolutivo_recurrida": fases_b().resolutivo_recurrida})
    if reasuncion is not None:
        datos["reasuncion"] = reasuncion
    est = est or dg.Estructura(apertura="V.", visto="para resolver.",
                               resultandos=[{"titulo": "Presentación de la demanda de amparo indirecto",
                                             "texto": "Unión Ejemplo, A.C. reclamó la resolución del toca civil "
                                                      "2338/2024. El Juzgado Séptimo de Distrito registró la "
                                                      "demanda con el número 950/2024."}],
                               competencia="", existencia="", procedencia="")
    _c = f0.computar(_dt.date(2025, 4, 28), _dt.date(2025, 5, 15), plazo=10)
    ruta = os.path.join(tempfile.mkdtemp(prefix="regresion631b_"), "ar.docx")
    dg.componer(datos, est, _c, f0.fecha_en_letra, ruta, estudio=estudio,
                calificaciones=list(calif), tipo_asunto="amparo_revision")
    ps = [q.text for q in Document(ruta).paragraphs if q.text.strip()]
    i = next(k for k, t in enumerate(ps) if "R E S U E L V E" in t or t.strip() in ("RESUELVE", "Resuelve"))
    return ps, " ".join(ps[i:]), est


# (b) SIN LOS CONCEPTOS, EL PROYECTO NO AFIRMA «NO AMPARA».
_reas_sin = {"reasuncion": "concesion", "hacen_falta": True, "tenemos": False, "donde": "",
             "quien_recurre": "tercero", "sobresee_ademas": True}
_ps, _res, _est = componer_b(["Solución", "Es fundado el primer agravio: la interlocutoria sólo cambió quién "
                              "ejecuta. Se revoca y el tribunal reasume jurisdicción; el estudio de los conceptos "
                              "no estudiados queda pendiente."], _reas_sin)
ok("no ampara ni protege" not in _res and "*********" in _res and "Unión Ejemplo, A.C." in _res,
   f"sin los conceptos, el punto del amparo va con hueco, no «no ampara»: {_res[:260]}")
ok(any(a.startswith("SE REVOCA UNA CONCESIÓN Y LOS CONCEPTOS") for a in _est.avisos),
   "y el aviso dice que faltan y que no se firme un «no ampara» sin ese estudio")
# (c) CON ELLOS Y TODOS DESESTIMADOS: «no ampara», en la materia de la revisión, sobreseimiento firme.
_reas_con = dict(_reas_sin, tenemos=True, donde="constancias")
_ps2, _res2, _est2 = componer_b(["Solución", "Es fundado el primer agravio. El sobreseimiento respecto del Juzgado "
                                 "Quinto no fue impugnado y queda firme.", "Estudio de los conceptos de violación no "
                                 "estudiados", "Los conceptos descansan en la tesis ya desestimada y son inoperantes; "
                                 "por tanto, procede revocar la sentencia recurrida y negar el amparo."], _reas_con)
ok("PRIMERO. Queda firme el sobreseimiento decretado en la sentencia recurrida." in _res2
   and "SEGUNDO. En la materia de la revisión, se revoca la sentencia recurrida." in _res2
   and "TERCERO. La Justicia de la Unión no ampara ni protege a Unión Ejemplo, A.C." in _res2,
   f"con ellos y desestimados: firme, en la materia de la revisión, y no ampara a quien amparó el juzgado: {_res2[:420]}")
ok(any(a.startswith("SE REVOCA LA CONCESIÓN Y SE REASUME JURISDICCIÓN") and "NEGAR" in a for a in _est2.avisos)
   and any(a.startswith("LOS CONCEPTOS DE VIOLACIÓN NO ESTUDIADOS") for a in _est2.avisos),
   "el aviso dice de dónde sale la negativa y de dónde salieron los conceptos")
# (d) NINGÚN AVISO DE «SE CONCEDE» EN EL REVOCA-Y-NIEGA.
_txt2 = " ".join(_ps2)
_rev = f6.revisar(_txt2, crit(), material(), rama="revoca_fondo_niega")
ok(not any(a.startswith("Se concede y los efectos") for a in _rev)
   and f6._efectos_de_reposicion(_txt2, crit(), True, rama="revoca_fondo_niega") == ""
   and ens.formula_resolutivo(["fundado", "innecesario"], rama="revoca_fondo_niega",
                              sentido_amparo=fr.sentido_en_plenitud(_txt2))[1] == "",
   "ni «Se concede y los efectos van en prosa», ni «EFECTOS INCOMPLETOS…», ni «El resolutivo concede…»")
# (e) LOS PAPELES: carátula, competencia, existencia y legitimación.
_cab = " ".join(_ps2[:8]).upper()
ok("QUEJOSA: UNIÓN EJEMPLO, A.C." in _cab and "TERCERA INTERESADA Y RECURRENTE: INMOBILIARIA EJEMPLO" in _cab
   and "QUEJOSA Y RECURRENTE" not in _cab,
   f"la carátula separa a la quejosa de la recurrente, con su carácter: {_cab[:200]}")
# SIN ÓRGANO RECURRIDO EN EL RUBRO: el corpus de la revisión no lo lleva (0 de
# 478, SCR/caratulas_revision.md). El juzgado se nombra en la competencia.
ok(not any(t.upper().startswith("ÓRGANO RECURRIDO") for t in _ps2[:8])
   and "MAGISTRADA" not in _cab,
   f"la carátula no rotula al juzgado ni a la Magistrada: {_cab[:260]}")
_comp = next((t for t in _ps2 if "es competente" in t), "")
ok("por el Juzgado Séptimo de Distrito" in _comp and "Magistrada" not in _comp,
   f"competencia: dictada por el juzgado de distrito: {_comp[180:330]}")
_exi = next((t for t in _ps2 if "existencia del acto reclamado" in t.lower()), "")
ok(not _exi or ("950/2024" in _exi and "2338/2024" not in _exi),
   f"existencia: el número del amparo, no el del toca: {_exi[:260]}")
_legb = next((t for t in _ps2 if "legitimad" in t), "")
ok("Inmobiliaria Ejemplo" in _legb and "5o., fracción III" in _legb and "6º" not in _legb,
   f"legitimación: la de la tercera interesada, no la del quejoso: {_legb[:260]}")
# (f) LA RAMA DEL ADELANTO NO SE QUEDA: la misma estructura, primero sin criterio y después con él.
_, _, _est3 = componer_b(["Solución", "Texto del adelanto."], None, calif=())
ok(any("RESOLUTIVO DE REVISIÓN, rama «confirma_concede»" in a for a in _est3.avisos),
   "el adelanto (sin criterio) anota su rama provisional")
_, _, _est3 = componer_b(["Solución", "Es fundado el primer agravio; procede revocar la sentencia "
                          "recurrida y negar el amparo."], _reas_con, est=_est3)
ok(not any("«confirma_concede»" in a for a in _est3.avisos)
   and any("«revoca_fondo_niega»" in a for a in _est3.avisos),
   "al componer el proyecto con la misma estructura, sólo queda la rama del proyecto")

# ═══════════════════════════════════════════════════════════════════════════
# REVISIÓN ADVERSARIAL (28-sep-2026): el documento con la quejosa recurrente y
# con los conceptos «por confirmar».
print("\n13 · REVISIÓN DEL 28-sep: LA QUEJOSA QUE GANA Y LOS CONCEPTOS POR CONFIRMAR")
# (g) RECURRE LA QUEJOSA CONTRA SU CONCESIÓN Y GANA: no se le niega el amparo.
_reas_q = {"reasuncion": "", "quien_recurre": "quejoso"}
_, _res_q, _est_q = componer_b(["Solución", "Es fundado el agravio de la quejosa: los efectos de la "
                                "concesión no restituyen el derecho violado. Procede revocar la sentencia "
                                "recurrida y conceder el amparo para los efectos que se precisan."], _reas_q)
ok("no ampara" not in _res_q and "La Justicia de la Unión ampara y protege a Unión Ejemplo, A.C." in _res_q
   and "PRIMERO. Se revoca la sentencia recurrida." in _res_q,
   f"la quejosa que gana su recurso sigue amparada (art. 93, fr. V): {_res_q[:320]}")
ok(any(a.startswith("RECURRE LA QUEJOSA CONTRA UNA CONCESIÓN") for a in _est_q.avisos)
   and any("«revoca_fondo_concede»" in a for a in _est_q.avisos),
   "y el aviso dice por qué y que compruebe revocar o modificar")
_, _res_qm, _ = componer_b(["Solución", "Es fundado el agravio; procede modificar la sentencia recurrida "
                            "para ampliar el alcance de la concesión."], _reas_q)
ok("PRIMERO. Se modifica la sentencia recurrida." in _res_qm and "no ampara" not in _res_qm,
   "si el estudio dice modificar, se modifica")
# (h) «POR CONFIRMAR»: la recurrida no dice que quedaran conceptos sin estudiar;
#     si el estudio concluye, el punto sale de ahí y no con hueco.
_reas_pc = dict(_reas_sin, hacen_falta="por_confirmar")
_, _res_pc, _ = componer_b(["Solución", "Es fundado el primer agravio. El juzgado examinó y desestimó los "
                            "demás conceptos y la quejosa no los combatió: no hay nada que reasumir; por "
                            "tanto, procede revocar la sentencia recurrida y negar el amparo."], _reas_pc)
ok("no ampara ni protege a Unión Ejemplo, A.C." in _res_pc and "*********" not in _res_pc,
   f"por confirmar y el estudio concluye: el punto sale del estudio, sin hueco: {_res_pc[:300]}")
_, _res_pc2, _est_pc2 = componer_b(["Solución", "Es fundado el primer agravio. Se revoca y el tribunal "
                                    "reasume jurisdicción."], _reas_pc)
ok("*********" in _res_pc2 and any("NO DICE QUE EL JUZGADO DEJARA CONCEPTOS" in a for a in _est_pc2.avisos),
   "por confirmar y el estudio no concluye: hueco, con el aviso que dice qué comprobar")

print()
if FALLOS:
    print(f"FALLAN {len(FALLOS)}: " + " · ".join(FALLOS))
    sys.exit(1)
print("RESULTADO: TODAS LAS COMPROBACIONES PASAN")
