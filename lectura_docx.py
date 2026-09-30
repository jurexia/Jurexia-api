# -*- coding: utf-8 -*-
"""LEER UN DOCX SIN CARGARLO ENTERO, Y SIN CONGELAR AL TRABAJADOR (30-sep-2026).

POR QUÉ. El 25-sep a las 15:59 (CDMX) Render mató una instancia por falta de
memoria: 2,139 de 2,147 MB. Medido minuto a minuto con su API, la memoria subió
en escalones y nunca bajó:
  · +626 MB al leer «GUIA PARA EL USO CLINICO DE LA SANGRE.docx» (9 MB);
  · +73 MB con otra guía de 2.3 MB, y
  · el resto, leyes enteras extraídas en texto plano (887,433 caracteres).
Las tres rutas que reciben un DOCX lo abrían con python-docx
(`Document(io.BytesIO(...))`): construye el árbol XML COMPLETO del documento
—en un DOCX convertido de PDF, cada palabra es una corrida con su formato: millones
de nodos— y carga también las imágenes. Y lo hacía de forma síncrona dentro de una
ruta async: con el de 9 MB, once segundos en los que ese trabajador no atendió a
nadie más. Ésa era la única «cola» real que se encontró en una semana de registros.

LO QUE HACE. Lee `word/document.xml` directo del ZIP y lo recorre como flujo
(iterparse), soltando cada párrafo ya leído: la memoria no depende del tamaño del
documento. Da el mismo texto que `paragraph.text` de python-docx —corridas,
hipervínculos, tabuladores, saltos de línea— y además el de las tablas y los
cuadros de texto, que python-docx dejaba fuera de `doc.paragraphs`. De un cuadro de
texto escrito dos veces (mc:Choice y su mc:Fallback) se lee sólo el primero.

`leer_docx` lo corre en un hilo, cae a python-docx si el flujo falla, y tras un
archivo grande le pide al proceso que devuelva al sistema la memoria que liberó.
"""
from __future__ import annotations

import asyncio
import gc
import io
import zipfile

from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_P, _R, _TBL = W + "p", W + "r", W + "tbl"
_T, _TAB, _PTAB, _BR, _CR, _NOBREAK = W + "t", W + "tab", W + "ptab", W + "br", W + "cr", W + "noBreakHyphen"
_TIPO = W + "type"
_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_REL_DOCUMENTO = "/officeDocument"

# Más de lo que cualquier ruta usa (la de mayor tope, 1,000,000): pasado esto se
# deja de leer. Protege de un ZIP que se descomprime en gigas de texto.
TOPE_CHARS = 3_000_000
# Desde este tamaño de archivo se le pide al proceso que devuelva la memoria.
UMBRAL_DEVOLVER = 1_000_000

_OPCIONES_XML = dict(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)


def _parte_principal(z: zipfile.ZipFile) -> str:
    """La ruta del documento principal según `_rels/.rels`; casi siempre
    `word/document.xml`, pero hay generadores que lo llaman distinto."""
    nombres = set(z.namelist())
    try:
        with z.open("_rels/.rels") as f:
            for _, el in etree.iterparse(f, events=("end",), **_OPCIONES_XML):
                if str(el.tag).endswith("Relationship") and str(el.get("Type") or "").endswith(_REL_DOCUMENTO):
                    destino = str(el.get("Target") or "").lstrip("/")
                    if destino in nombres:
                        return destino
    except Exception:
        pass
    return "word/document.xml"


def _soltar(el) -> None:
    """Vacía un elemento ya leído y borra a sus hermanos anteriores: lo leído no
    se queda en memoria esperando al final del documento."""
    el.clear()
    padre = el.getparent()
    if padre is not None:
        while el.getprevious() is not None:
            del padre[0]


def texto_docx(contenido: bytes, separador: str = "\n\n", tope: int = TOPE_CHARS) -> str:
    """El texto de un DOCX, párrafo por párrafo, sin cargar el documento entero.
    Lanza si el archivo no es un DOCX legible (quien llama decide qué hacer)."""
    parrafos: list[str] = []
    pila: list[list[str]] = []      # un búfer por párrafo abierto (los hay anidados)
    saltar = 0                      # >0 dentro de un mc:Fallback
    total = 0
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        with z.open(_parte_principal(z)) as f:
            for ev, el in etree.iterparse(f, events=("start", "end"), **_OPCIONES_XML):
                tag = el.tag
                if tag == _FALLBACK:
                    saltar += 1 if ev == "start" else -1
                    continue
                if saltar:
                    continue
                if ev == "start":
                    if tag == _P:
                        pila.append([])
                    continue
                if tag == _T:
                    if pila:
                        pila[-1].append(el.text or "")
                elif tag in (_TAB, _PTAB):
                    # Sólo el de una corrida: el `w:tab` de `w:pPr/w:tabs` es una
                    # parada de tabulación, no un carácter.
                    padre = el.getparent()
                    if pila and padre is not None and padre.tag == _R:
                        pila[-1].append("\t")
                elif tag == _BR:
                    # Como python-docx: salto de línea "\n"; de página o columna, nada.
                    if pila and (el.get(_TIPO) or "textWrapping") == "textWrapping":
                        pila[-1].append("\n")
                elif tag == _CR:
                    if pila:
                        pila[-1].append("\n")
                elif tag == _NOBREAK:
                    if pila:
                        pila[-1].append("-")
                elif tag == _P:
                    texto = "".join(pila.pop()) if pila else ""
                    if texto.strip():
                        parrafos.append(texto)
                        total += len(texto)
                    _soltar(el)
                    if total >= tope:
                        print(f"   ✂️ DOCX: lectura detenida en {total:,} caracteres (tope {tope:,})")
                        break
                elif tag == _TBL:
                    _soltar(el)
    return separador.join(parrafos)


def _texto_python_docx(contenido: bytes, separador: str) -> str:
    """El camino de antes, de repliegue."""
    from docx import Document
    d = Document(io.BytesIO(contenido))
    return separador.join(p.text for p in d.paragraphs if p.text.strip())


def devolver_memoria() -> None:
    """Que el proceso devuelva al sistema la memoria que ya liberó. Python y
    glibc se quedan con lo liberado para reusarlo, y la instancia se ve llena
    aunque no lo esté: eso es la escalera que nunca baja. En Linux, malloc_trim
    lo suelta; en otro sistema no hace nada."""
    try:
        gc.collect()
    except Exception:
        pass
    try:
        import ctypes
        ctypes.CDLL("libc.so.6").malloc_trim(0)
    except Exception:
        pass


async def leer_docx(contenido: bytes, separador: str = "\n\n") -> str:
    """El texto de un DOCX sin bloquear al trabajador. Si el flujo falla, se
    intenta con python-docx (también en un hilo); si eso falla, lanza."""
    def _leer() -> str:
        try:
            return texto_docx(contenido, separador)
        except Exception as e:
            print(f"   ⚠️ DOCX: el lector por flujo falló ({type(e).__name__}); se intenta con python-docx")
            return _texto_python_docx(contenido, separador)

    try:
        return await asyncio.to_thread(_leer)
    finally:
        if len(contenido or b"") >= UMBRAL_DEVOLVER:
            await asyncio.to_thread(devolver_memoria)
