"""La Bóveda sin red: texto por página, citas, cobro y exportaciones — oct-2026.

    .venv/bin/python test_boveda.py

Lo que importa:
  · que sólo las páginas SIN texto vayan a OCR (cada una cuesta 0.0015 USD) y
    que una hoja en blanco no se pague;
  · que un fragmento nunca cruce de página (la cita dice la página exacta);
  · que una cita literal se verifique aunque cambien acentos, comillas o
    renglones, y que una inventada NO;
  · que la regla de cobro redondee hacia arriba y nunca cobre por nada;
  · que la tabla exportada lleve las citas.
No toca Supabase, Qdrant, Azure ni OpenAI.
"""
import io
import sys

import fitz

import boveda_tabla as btab
import boveda_texto as bt

FALLOS = []


def ok(cond, que):
    print(f"   {'PASA ' if cond else 'FALLA'}  {que}")
    if not cond:
        FALLOS.append(que)


def pdf_mixto() -> bytes:
    """3 páginas: con texto, escaneada (sólo imagen) y en blanco."""
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_textbox(fitz.Rect(72, 72, 540, 760),
                      "CONTRATO DE ARRENDAMIENTO. " + "Las partes convienen en lo siguiente. " * 30, fontsize=10)
    fuente = fitz.open()
    q = fuente.new_page()
    q.insert_textbox(fitz.Rect(72, 72, 540, 760), "Cláusula escaneada. " * 40, fontsize=11)
    pix = q.get_pixmap(dpi=80)
    p2 = doc.new_page()
    p2.insert_image(p2.rect, pixmap=pix)
    doc.new_page()   # en blanco
    datos = doc.tobytes()
    doc.close()
    return datos


print("\n── Inspección: qué se paga ──")
ins = bt.inspeccionar(pdf_mixto(), "x.pdf")
ok(ins.paginas == 3, "cuenta 3 páginas")
ok(ins.paginas_ocr == [2], f"sólo la escaneada va a OCR (dio {ins.paginas_ocr})")
ok(len(ins.textos[0]) > 500, "la página con texto se lee gratis")

print("\n── Word: páginas aproximadas ──")
import docx  # noqa: E402
d = docx.Document()
for i in range(40):
    d.add_paragraph(f"Párrafo {i}: " + "texto del convenio " * 20)
b = io.BytesIO()
d.save(b)
iw = bt.inspeccionar(b.getvalue(), "convenio.docx")
ok(iw.aproximadas and iw.paginas >= 4, f"un Word largo se parte en cuartillas ({iw.paginas})")
ok(all(len(t) <= bt.CUARTILLA + 600 for t in iw.textos), "ninguna «página» se pasa mucho de la cuartilla")
ok(bt.tipo_de("a.DOCX") == "docx" and bt.tipo_de("b.Pdf") == "pdf" and bt.tipo_de("c.doc") == "", "tipos admitidos")

print("\n── Fragmentos ──")
frags = bt.fragmentar(["a " * 2000, "", "Una página corta."])
ok({f.pagina for f in frags} == {1, 3}, "una página vacía no da fragmentos")
ok(all(len(f.texto) <= bt.FRAGMENTO + 10 for f in frags), "ningún fragmento rebasa el tamaño")
ok(bt._rangos([1, 2, 3, 7, 9, 10]) == "1-3,7,9-10", "rangos para el parámetro pages de Azure")

print("\n── Verificar citas ──")
paginas = [
    "PRIMERA. OBJETO. El arrendador concede el uso del inmueble.",
    "CUARTA. PENA CONVENCIONAL. En caso de incumplimiento, la parte incumplida pagará el 15% (por "
    "ciento) del monto total del contrato.",
    "DÉCIMA. JURISDICCIÓN. Las partes se someten a los tribunales de Querétaro.",
]
v = bt.verificar_cita("la parte incumplida pagará el 15% (por ciento) del monto total", paginas, 2)
ok(v.verificada and v.pagina == 2, "cita literal en la página dicha")
v = bt.verificar_cita("Las partes se someten a los tribunales de Queretaro", paginas, 1)
ok(v.verificada and v.pagina == 3, "corrige la página y tolera acentos")
v = bt.verificar_cita("«la parte incumplida pagará el 15%\n(por ciento)»", paginas, 2)
ok(v.verificada, "tolera comillas y saltos de renglón")
v = bt.verificar_cita("la parte incumplida pagará el 30% del monto total del contrato", paginas, 2)
ok(not v.verificada, f"una cifra cambiada NO se verifica (similitud {v.similitud})")
v = bt.verificar_cita("el arrendatario renuncia a todo derecho de preferencia", paginas, 1)
ok(not v.verificada, "una cita inventada NO se verifica")
ok(not bt.verificar_cita("15%", paginas, 2).verificada, "una cita de 3 caracteres no prueba nada")
v = bt.verificar_cita("incumplimiento, la parte incumplida pagarа el 15% (por ciento) del monto total del contrato",
                      paginas, 2)
ok(v.verificada, f"una letra rara de OCR no tumba la cita ({v.similitud})")

print("\n── La respuesta del modelo ──")
cols = btab.COLUMNAS_SUGERIDAS[:3]
crudo = {"respuestas": [
    {"columna_id": "partes", "encontrado": True, "valor": "A; B", "cita": "«PRIMERA. OBJETO.»", "pagina": "1"},
    {"columna_id": "fecha", "encontrado": False, "valor": "No se menciona", "cita": "", "pagina": None},
]}
cel = btab._normalizar(crudo, cols)
ok(cel[0].pagina == 1 and cel[0].cita == "PRIMERA. OBJETO.", "página en texto y comillas sobrantes")
ok(not cel[1].encontrado, "lo no encontrado queda no encontrado")
ok(cel[2].error == "sin respuesta del modelo", "la columna que el modelo omitió se marca")
ctx = btab._contexto("c.pdf", paginas, [3, 1])
ok(ctx.index("PÁGINA 1") < ctx.index("PÁGINA 3") and "PÁGINA 2" not in ctx, "sólo las páginas elegidas, en orden")

print("\n── La regla de cobro ──")
import boveda  # noqa: E402
ok(boveda.consultas_ingesta(0, 0) == 0, "nada no cuesta")
ok(boveda.consultas_ingesta(0, 1) == 1, "una página ya cuesta una consulta")
ok(boveda.consultas_ingesta(400, 0) == 40, "400 escaneadas = 40")
ok(boveda.consultas_ingesta(0, 750) == 8, "750 con texto = 8 (hacia arriba)")
ok(boveda.consultas_ingesta(10, 100) == 2, "10 escaneadas + 100 con texto = 2, sin redondear de más")
ok(boveda.consultas_celdas(300) == 30 and boveda.consultas_celdas(1) == 1, "celdas: 300 → 30, 1 → 1")

print("\n── Exportar ──")
columnas = [{"id": "pena", "titulo": "Pena convencional"}]
filas = [{"id": "d1", "nombre": "contrato.pdf"}]
celdas = {("d1", "pena"): {"estado": "lista", "respuesta": {
    "valor": "15% del monto", "encontrado": True, "cita": "pagará el 15%", "pagina": 2, "verificada": True}}}
x = boveda._xlsx("Revisión", columnas, filas, celdas)
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(io.BytesIO(x))
ok(wb["Revisión"]["B2"].value == "15% del monto", "Excel: la respuesta en su celda")
ok("pagará el 15%" in (wb["Revisión"]["B2"].comment.text or ""), "Excel: la cita como comentario")
ok(wb["Citas"]["D2"].value == "pagará el 15%" and wb["Citas"]["E2"].value == 2, "Excel: hoja de citas con página")
w = boveda._docx("Revisión", "Carpeta X", columnas, filas, celdas)
dd = docx.Document(io.BytesIO(w))
texto = "\n".join(c.text for r in dd.tables[0].rows for c in r.cells)
ok("15% del monto" in texto and "pagará el 15%" in texto and "(p. 2)" in texto, "Word: respuesta, cita y página")

print()
if FALLOS:
    print(f"✗ {len(FALLOS)} fallas")
    sys.exit(1)
print("✓ todo pasa")
