"""LOS EFECTOS DE LA CONCESIÓN: APERTURA CON EL 77, DE TRES A CINCO, EN INFINITIVO.

David, 27-sep-2026: «cuando concede los efectos no están bien configurados,
son muy extensos y el diálogo jurídico no parece congruente: inicia con
"Efectos.1. Deje insubsistente…" cuando debería decir: "Efectos. Con fundamento
en el artículo … de la Ley de Amparo, la autoridad responsable deberá:…" con
efectos más reducidos pero que comprendan el objetivo de la concesión.
Regularmente los efectos son 3 a máximo 5 en función de la complejidad».

  1 · `componer_efectos`: numera, pasa a infinitivo el verbo de cada orden
      (y sus coordinados, nunca los de una subordinada), quita la introducción
      que el modelo escribía y avisa si pasan de cinco o se alargan.
  2 · el prompt de la v4: de tres a cinco, breves, en infinitivo y sin
      introducción; y ningún bloque de efectos en la queja ni en la revisión
      fiscal, que no conceden amparo.
  3 · el documento entero: «SÉPTIMO. Efectos. Con fundamento en el artículo 77
      de la Ley de Amparo, la autoridad responsable deberá:» y cada orden en su
      párrafo; y en la revisión que concede, los efectos ya no se pierden.
"""
import datetime as _dt
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import documento_generado as dg
import fase6_estudio as f6
import tipos_asunto as ta

FALLAS = []


def ok(cond, que):
    print(("   PASA   " if cond else "   FALLA  ") + que)
    if not cond:
        FALLAS.append(que)


print("\n1 · COMPONER LAS ÓRDENES")
ok(ta.APERTURA_EFECTOS == ("Con fundamento en el artículo 77 de la Ley de Amparo, la "
                           "autoridad responsable deberá:"),
   "la apertura funda en el 77 (efectos de la concesión), no en el 93 (reglas de la revisión)")

# Los efectos del ADC 93/2026 v8, tal como los escribió el estudio.
V8 = ["1. Deje insubsistente la sentencia definitiva de veintiocho de noviembre de dos mil "
      "veinticinco dictada en el expediente 765/25-09-01-8-ST.",
      "2. Deje sin efectos la sentencia interlocutoria de dieciocho de septiembre y el acuerdo "
      "de trece de agosto, en la parte en que decretaron la preclusión.",
      "3. Admita la ampliación de demanda presentada el catorce de agosto y provea lo conducente "
      "respecto de los actos que en ella se hagan valer.",
      "4. Corra traslado a la autoridad demandada con la ampliación admitida y continúe el "
      "procedimiento.",
      "5. Cerrada nuevamente la instrucción, dicte una sentencia de fondo con plenitud de "
      "jurisdicción, en la que examine los argumentos de la ampliación."]
o, a = dg.componer_efectos(V8)
ok([x[:3] for x in o] == ["1. ", "2. ", "3. ", "4. ", "5. "], "numeradas del 1 al 5")
ok(o[0].startswith("1. Dejar insubsistente") and o[1].startswith("2. Dejar sin efectos")
   and o[2].startswith("3. Admitir") and o[3].startswith("4. Correr traslado"),
   "el primer verbo de cada orden, en infinitivo")
ok("y proveer lo conducente" in o[2] and "que en ella se hagan valer" in o[2]
   and "y continuar el procedimiento" in o[3], "los coordinados también («y provea» → «y proveer»)")
ok(o[4].startswith("5. Cerrada nuevamente la instrucción, dictar una sentencia")
   and "en la que examine" in o[4],
   "tras la frase de enlace, infinitivo; dentro de «en la que», el subjuntivo se queda")
ok(a == [], "cinco órdenes breves: sin aviso")

# ADC 810/2025 (engrose real): sujeto + «deberá», incisos que cuelgan de «en la que».
o, a = dg.componer_efectos([
    "La Justicia de la Unión ampara y protege a Luis, para los efectos siguientes:",
    "1. La autoridad responsable, Segunda Sala Civil del Tribunal Superior de Justicia, "
    "deberá dejar insubsistente la sentencia reclamada.",
    "2. Acto seguido, deberá emitir una nueva sentencia en la que, con plena jurisdicción:",
    "a) Reitere las consideraciones que no fueron materia de la concesión.",
    "b) Analice y resuelva los agravios fundados.",
    "3. Emita otra en la que deberá analizar la prueba pericial;"])
ok(o[0] == "1. Dejar insubsistente la sentencia reclamada.",
   "el sujeto y el «deberá» sobran: los pone la apertura")
ok(o[1] == "2. Acto seguido, emitir una nueva sentencia en la que, con plena jurisdicción:",
   "la frase de enlace se queda; la orden que abre incisos cierra en dos puntos")
ok(o[2].startswith("a) Reitere") and o[3].startswith("b) Analice"),
   "los incisos de segundo nivel conservan su subjuntivo y su letra")
ok(o[4] == "3. Emitir otra en la que deberá analizar la prueba pericial.",
   "un «deberá» DENTRO de la orden no se toca; el punto y coma final pasa a punto")
ok(len(o) == 5 and not o[0].lower().startswith("la justicia"), "la introducción del modelo sobra")

# ADA 767/2025 (engrose real): sin números, «realice lo siguiente:» y un párrafo por orden.
o, a = dg.componer_efectos([
    "La protección constitucional se concede para el efecto de que el Tribunal Unitario "
    "Agrario del Distrito 42 realice lo siguiente:",
    "Deje insubsistente el auto de fecha trece de septiembre.",
    "En su lugar, emita un nuevo proveído en el que determine que no operó la caducidad.",
    "Hecho lo anterior, verifique si se encuentra integrada la litis y provea lo conducente."])
ok(o == ["1. Dejar insubsistente el auto de fecha trece de septiembre.",
         "2. En su lugar, emitir un nuevo proveído en el que determine que no operó la caducidad.",
         "3. Hecho lo anterior, verificar si se encuentra integrada la litis y proveer lo conducente."],
   "sin marcas: una orden por párrafo tras la introducción")

o, a = dg.componer_efectos(["La concesión es para que la Sala deje insubsistente la sentencia y dicte otra."])
ok(o == [], "un párrafo de prosa suelto no se convierte: se escribe como vino y se avisa")

o, a = dg.componer_efectos([f"{i}. Deje insubsistente el acto {i}." for i in range(1, 7)])
ok(len([x for x in o if x[0].isdigit()]) == 6 and any("SON 6 ÓRDENES" in x for x in a),
   "seis órdenes: se avisa (lo habitual es de tres a cinco)")
o, a = dg.componer_efectos(["1. Examine " + "la prueba pericial en materia contable " * 12 + "."])
ok(any("PALABRAS" in x for x in a), "una orden de más de sesenta palabras: se avisa")
o, a = dg.componer_efectos(["1. Cierre de la instrucción conforme a derecho."])
ok(o == ["1. Cierre de la instrucción conforme a derecho."],
   "un sustantivo que parece verbo («Cierre de…») no se convierte")

print("\n2 · EL PROMPT DE LA v4")
_C = [f6.Criterio(problema="¿La Sala debió admitir la ampliación?", sentido="fundado",
                  razonamiento="Porque el acuerdo la dejó a salvo.", jerarquia="principal")]


def bloque(tipo):
    return f6._bloque_criterio(_C, "administrativa", "", tipo, "estandar", [], variante="v2")


_ad = " ".join(bloque("amparo_directo").split())
ok("EFECTOS DE LA CONCESIÓN" in _ad and "DE TRES A CINCO ÓRDENES" in _ad
   and "nunca más de cinco" in _ad, "amparo directo: de tres a cinco órdenes, nunca más de cinco")
ok("EMPIEZA CON EL VERBO EN INFINITIVO" in _ad and "SIN frase de introducción" in _ad
   and "el artículo 77 de la Ley de Amparo" in _ad,
   "en infinitivo y sin introducción, porque el documento abre con el 77")
ok("no repite las razones del estudio" in _ad, "dice qué hacer, no por qué")
for t in ("queja", "revision_fiscal"):
    ok("EFECTOS DE LA CONCESIÓN" not in bloque(t), f"{t}: no se piden efectos (no hay amparo que conceder)")
_rv = " ".join(bloque("amparo_revision").split())
ok("EFECTOS DE LA CONCESIÓN" in _rv and "SÓLO SI, al resolver esta revisión, ESTE tribunal concede" in _rv,
   "revisión: sólo si este tribunal concede o modifica los efectos")

print("\n3 · EL DOCUMENTO ENTERO")
try:
    import fase0_oportunidad as _f0
    from docx import Document

    _est = dg.Estructura(apertura="V.", visto="para resolver.",
                         resultandos=[{"titulo": "Presentación de la demanda",
                                       "texto": "Se presentó la demanda."}],
                         competencia="", existencia="", procedencia="")
    _c = _f0.computar(_dt.date(2026, 3, 2), _dt.date(2026, 3, 9), plazo=15)
    _salida = "/tmp/_efectos_ad.docx"
    dg.componer({"tipo_asunto": "amparo_directo", "numero": "93/2026",
                 "encabezado": "AMPARO DIRECTO ADMINISTRATIVO 93/2026", "quejoso": "Juan Pérez",
                 "responsable": "Sala Regional", "magistrado": "M", "secretario": "S",
                 "tribunal": "Tercer Tribunal Colegiado", "ciudad": "Querétaro"},
                _est, _c, _f0.fecha_en_letra, _salida,
                estudio=["Es fundado el concepto de violación.", "EFECTOS DE LA CONCESIÓN"] + V8,
                calificaciones=["fundado"], tipo_asunto="amparo_directo")
    _ps = [q.text for q in Document(_salida).paragraphs if q.text.strip()]
    _k = [i for i, t in enumerate(_ps) if "Efectos." in t and "Con fundamento en el artículo 77" in t]
    ok(bool(_k), "el considerando abre «… Efectos. Con fundamento en el artículo 77 de la Ley de "
                 "Amparo, la autoridad responsable deberá:»")
    if _k:
        k = _k[0]
        ok(_ps[k].rstrip().endswith("deberá:") and "1." not in _ps[k],
           "la primera orden ya no va pegada al rótulo")
        ok(_ps[k + 1].startswith("1. Dejar insubsistente") and _ps[k + 5].startswith("5. Cerrada"),
           "cada orden en su párrafo, numerada y en infinitivo")
    ok(not any("Efectos.1." in t or "Efectos. 1." in t for t in _ps), "nunca «Efectos.1. Deje…»")
    # LA REVISIÓN QUE CONCEDE: los efectos cierran el último considerando, que
    # es a donde remite el resolutivo; antes se perdían.
    _est_rv = dg.Estructura(apertura="V.", visto="para resolver.",
                            resultandos=[{"titulo": "Presentación de la demanda de amparo indirecto",
                                          "texto": "Juan Pérez promovió amparo contra la Directora "
                                                   "de Ingresos."}],
                            competencia="", existencia="", procedencia="")
    _c10 = _f0.computar(_dt.date(2026, 3, 2), _dt.date(2026, 3, 9), plazo=10)
    dg.componer({"tipo_asunto": "amparo_revision", "numero": "12/2026",
                 "encabezado": "AMPARO EN REVISIÓN 12/2026", "quejoso": "Juan Pérez",
                 "responsable": "Juzgado Segundo de Distrito", "magistrado": "M", "secretario": "S",
                 "tribunal": "Tercer Tribunal Colegiado", "ciudad": "Querétaro",
                 "resolvio_a_quo": "sobresee", "recurrente": "Juan Pérez"},
                _est_rv, _c10, _f0.fecha_en_letra, "/tmp/_efectos_rv.docx",
                estudio=["Son fundados los agravios: el juez sobreseyó indebidamente y este "
                         "tribunal reasume jurisdicción.",
                         "El concepto de violación es fundado; por tanto, procede conceder el amparo.",
                         "EFECTOS DE LA CONCESIÓN",
                         "1. Deje insubsistente la resolución de la Directora de Ingresos.",
                         "2. Emita otra fundada y motivada en la que determine la base del impuesto."],
                calificaciones=["fundado"], tipo_asunto="amparo_revision")
    _pr = [q.text for q in Document("/tmp/_efectos_rv.docx").paragraphs if q.text.strip()]
    _kr = [i for i, t in enumerate(_pr) if t.strip() == "Efectos de la concesión"]
    ok(bool(_kr) and _pr[_kr[0] + 1] == ta.APERTURA_EFECTOS
       and _pr[_kr[0] + 2].startswith("1. Dejar insubsistente"),
       "revisión que concede: los efectos van al final del último considerando, con su apertura")
    ok(any("al resultar fundados los agravios planteados" in t for t in _pr),
       "y el cierre concuerda («fundados los agravios planteados», no «fundados lo planteado»)")
except Exception as ex:  # pragma: no cover
    ok(False, f"el documento entero no se pudo componer: {type(ex).__name__}: {ex}")

print()
if FALLAS:
    print(f"FALLAN {len(FALLAS)}: " + " · ".join(FALLAS))
    sys.exit(1)
print("RESULTADO: TODAS LAS COMPROBACIONES PASAN")
